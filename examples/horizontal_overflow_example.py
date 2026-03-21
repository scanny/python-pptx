#!/usr/bin/env python3
"""Example demonstrating horizontal overflow detection in python-pptx.

This example shows how to detect when text will overflow horizontally,
which occurs when:
1. Word wrap is enabled: Individual words are too wide to fit
2. Word wrap is disabled: Entire lines exceed the available width
"""

import sys

from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])

shape = slide.shapes.placeholders[1]
text_frame = shape.text_frame

# Font path varies by platform - adjust for your system
if sys.platform == "darwin":
    font_file = "/System/Library/Fonts/Geneva.ttf"
elif sys.platform == "win32":
    font_file = "C:\\Windows\\Fonts\\arial.ttf"
else:
    font_file = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"

print("=" * 70)
print("Horizontal Overflow Detection Examples")
print("=" * 70)
print()

# Example 1: Checking for horizontal overflow (word wrap enabled)
print("Example 1: Long URL that won't fit")
print("-" * 70)
text_frame.word_wrap = True
text_frame.text = "Visit https://very-long-subdomain.example.com/path/to/resource for more info"

# Quick check
if text_frame.will_overflow(direction='horizontal', font_size=18, font_file=font_file):
    print("⚠️  Text will overflow horizontally!")

# Detailed information
info = text_frame.overflow_info(direction='horizontal', font_size=18, font_file=font_file)
print(f"Widest element: {info.widest_element}")
print(f"Required width: {info.required_width.inches:.2f} inches")
print(f"Available width: {info.available_width.inches:.2f} inches")
print(f"Overflow: {info.overflow_width_percentage:.1f}%")
print(f"Will fit at: {info.fits_at_font_size_horizontal}pt font size")
print()

# Example 2: Checking both vertical and horizontal overflow
print("Example 2: Checking both dimensions")
print("-" * 70)
text_frame.text = "Short text"

# Check both directions at once
if text_frame.will_overflow(direction='both', font_size=18, font_file=font_file):
    print("Text overflows in at least one direction")
else:
    print("✓ Text fits perfectly in both dimensions")
print()

# Example 3: Word wrap disabled
print("Example 3: Long line with word wrap disabled")
print("-" * 70)
text_frame.word_wrap = False
text_frame.text = "This is a very long line that will not wrap and may overflow horizontally"

info = text_frame.overflow_info(direction='horizontal', font_size=14, font_file=font_file)
if info.will_overflow_horizontally:
    print(f"⚠️  Line is too wide!")
    print(f"Widest line: '{info.widest_element[:50]}...'")
    print(f"Overflow: {info.overflow_width_percentage:.1f}%")
print()

# Example 4: Backward compatibility - vertical overflow still works
print("Example 4: Backward compatibility (vertical overflow)")
print("-" * 70)
text_frame.word_wrap = True
text_frame.text = " ".join(["Lorem ipsum"] * 100)  # Lots of text

# Default direction is 'vertical' for backward compatibility
if text_frame.will_overflow(font_size=24, font_file=font_file):
    print("⚠️  Text will overflow vertically (backward compatible check)")

info = text_frame.overflow_info(font_size=24, font_file=font_file)
print(f"Estimated lines: {info.estimated_lines}")
print(f"Vertical overflow: {info.overflow_percentage:.1f}%")
print()

# Example 5: Comprehensive check
print("Example 5: Comprehensive overflow analysis")
print("-" * 70)
text_frame.word_wrap = True
text_frame.text = "Check https://example.com/very/long/url/path and " + " ".join(["text"] * 50)

# Check both directions
info = text_frame.overflow_info(direction='both', font_size=20, font_file=font_file)

print(f"Overall overflow status: {info.will_overflow}")
print()
print("Horizontal:")
print(f"  Overflow: {info.will_overflow_horizontally}")
if info.will_overflow_horizontally:
    print(f"  Widest element: '{info.widest_element}'")
    print(f"  Percentage: {info.overflow_width_percentage:.1f}%")
print()
print("Vertical:")
if info.overflow_percentage and info.overflow_percentage > 0:
    print(f"  Overflow: True")
    print(f"  Percentage: {info.overflow_percentage:.1f}%")
    print(f"  Lines: {info.estimated_lines}")
else:
    print(f"  Overflow: False")

print()
print("=" * 70)
print("Examples completed!")
print("=" * 70)
