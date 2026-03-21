"""Example: Text Overflow Detection

This example demonstrates how to detect when text will overflow a text frame
before applying changes to the presentation.
"""

from pptx import Presentation
from pptx.util import Inches, Pt

# Create a presentation with a text box
prs = Presentation()
blank_slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_slide_layout)

# Add a text box
textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3))
text_frame = textbox.text_frame
text_frame.word_wrap = True

# Add some content
text_frame.text = "Project Overview"
for i in range(1, 11):
    p = text_frame.add_paragraph()
    p.text = f"Key point #{i}: Important information about the project"

# Example 1: Simple overflow check
if text_frame.will_overflow(font_family="Arial", font_size=18):
    print("Warning: Content will be cut off at 18pt!")
else:
    print("All content fits at 18pt")

# Example 2: Get detailed overflow information
info = text_frame.overflow_info(font_family="Arial", font_size=18)

if info.will_overflow:
    print(f"\nOverflow detected:")
    print(f"  Required height: {info.required_height.inches:.2f} inches")
    print(f"  Available height: {info.available_height.inches:.2f} inches")
    print(f"  Overflow: {info.overflow_percentage:.1f}%")
    print(f"  Recommendation: Use {info.fits_at_font_size}pt font to fit")
else:
    print("\nAll content fits!")

# Example 3: Find the best font size
for size in [20, 18, 16, 14, 12]:
    if not text_frame.will_overflow(font_family="Arial", font_size=size):
        print(f"\nLargest fitting font size: {size}pt")
        break

# Example 4: Validate before applying font size
desired_size = 18
if text_frame.will_overflow(font_family="Arial", font_size=desired_size):
    print(f"\nWarning: Cannot use {desired_size}pt - would overflow!")
    info = text_frame.overflow_info(font_family="Arial", font_size=desired_size)
    print(f"Using recommended size: {info.fits_at_font_size}pt instead")
    if info.fits_at_font_size is not None:
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(info.fits_at_font_size)
else:
    # Apply the desired size
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(desired_size)

prs.save("overflow_example.pptx")
