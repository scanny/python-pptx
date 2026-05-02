"""Gherkin step implementations for shape-animation round-trip (issue #102)."""

from __future__ import annotations

import io

from behave import then, when

from pptx import Presentation
from pptx.enum.animation import MSO_ANIMATION_TRIGGER, MSO_ANIMATION_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches


# when ============================================================


@when("I add a rectangle shape")
def when_add_rectangle_shape(context):
    context.shape = context.slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
    )
    context.shape_id = context.shape.shape_id


@when("I set shape.animation to {name} with delay {delay:d}")
def when_set_shape_animation(context, name, delay):
    effect_type = getattr(MSO_ANIMATION_TYPE, name)
    context.shape.set_animation(effect_type, delay=delay)


@when("I clear shape.animation")
def when_clear_shape_animation(context):
    context.shape.set_animation(None)


@when("I save and reopen the presentation preserving the shape")
def when_save_and_reopen_preserving_shape(context):
    buf = io.BytesIO()
    context.prs.save(buf)
    buf.seek(0)
    context.prs = Presentation(buf)
    context.slide = context.prs.slides[0]
    context.shape = next(
        s for s in context.slide.shapes if s.shape_id == context.shape_id
    )


# then ============================================================


@then("shape.animation is None")
def then_shape_animation_is_None(context):
    actual = context.shape.animation
    assert actual is None, "shape.animation is %s (expected None)" % actual


@then("shape.animation.type is {name}")
def then_shape_animation_type_is(context, name):
    expected = getattr(MSO_ANIMATION_TYPE, name)
    actual = context.shape.animation.type
    assert actual is expected, "shape.animation.type is %s (expected %s)" % (
        actual,
        expected,
    )


@then("shape.animation.delay is {delay:d}")
def then_shape_animation_delay_is(context, delay):
    actual = context.shape.animation.delay
    assert actual == delay, "shape.animation.delay is %s (expected %d)" % (actual, delay)


@then("shape.animation.trigger is {name}")
def then_shape_animation_trigger_is(context, name):
    expected = getattr(MSO_ANIMATION_TRIGGER, name)
    actual = context.shape.animation.trigger
    assert actual is expected, "shape.animation.trigger is %s (expected %s)" % (
        actual,
        expected,
    )
