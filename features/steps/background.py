"""Gherkin step implementations for slide background-related features."""

from __future__ import annotations

from behave import given, then, when
from helpers import test_pptx
from lxml import etree

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

# given ===================================================


@given("a _Background object having {type} background as background")
def given_a_Background_object_having_type_background(context, type):
    sld_idx = {"no": 0, "a fill": 1, "a style reference": None}[type]
    prs = Presentation(test_pptx("sld-background"))
    # -- retain the parent presentation so the slide objects stay alive --
    context.prs = prs
    slide = prs.slide_masters[0] if sld_idx is None else prs.slides[sld_idx]
    context.slide = slide
    context.background = slide.background


@given("a Slide object with no explicit background")
def given_a_Slide_object_with_no_explicit_background(context):
    prs = Presentation(test_pptx("sld-background"))
    context.dst_prs = prs
    context.dst_slide = prs.slides[0]  # -- fixture slide 0 has no explicit bg --
    assert context.dst_slide.background.bg_element is None


@given("a Slide object with an explicit fill background")
def given_a_Slide_object_with_an_explicit_fill_background(context):
    prs = Presentation(test_pptx("sld-background"))
    context.dst_prs = prs
    context.dst_slide = prs.slides[1]  # -- fixture slide 1 has an explicit fill --
    assert context.dst_slide.background.bg_element is not None


@given("a source Slide object with an explicit fill background")
def given_a_source_Slide_with_an_explicit_fill_background(context):
    prs = Presentation(test_pptx("sld-background"))
    context.src_prs = prs
    context.src_slide = prs.slides[1]
    assert context.src_slide.background.bg_element is not None


@given("a source Slide object with no explicit background")
def given_a_source_Slide_with_no_explicit_background(context):
    prs = Presentation(test_pptx("sld-background"))
    context.src_prs = prs
    context.src_slide = prs.slides[0]
    assert context.src_slide.background.bg_element is None


# when ====================================================


@when("I call slide.copy_background_from(source)")
def when_I_call_slide_copy_background_from(context):
    context.dst_slide.copy_background_from(context.src_slide)


# then ====================================================


@then("background.fill is a FillFormat object")
def then_background_fill_is_a_Fill_object(context):
    cls_name = context.background.fill.__class__.__name__
    assert cls_name == "FillFormat", "background.fill is a %s object" % cls_name


@then("background.bg_element is the p:bg element")
def then_background_bg_element_is_p_bg(context):
    bg = context.background.bg_element
    assert bg is not None, "background.bg_element is None"
    assert bg.tag == qn("p:bg"), "expected p:bg, got %s" % bg.tag


@then("background.bg_element is None")
def then_background_bg_element_is_None(context):
    assert context.background.bg_element is None, (
        "expected None, got %r" % context.background.bg_element
    )


@then("the destination slide has an explicit p:bg element")
def then_destination_has_explicit_bg(context):
    bg = context.dst_slide.background.bg_element
    assert bg is not None, "destination slide has no explicit p:bg"


@then("the destination slide has no explicit p:bg element")
def then_destination_has_no_explicit_bg(context):
    bg = context.dst_slide.background.bg_element
    assert bg is None, "destination slide still has a p:bg element"


@then("the destination p:bg subtree matches the source")
def then_destination_bg_matches_source(context):
    dst_bg = context.dst_slide.background.bg_element
    src_bg = context.src_slide.background.bg_element
    assert dst_bg is not None
    assert src_bg is not None
    dst_c14n = etree.tostring(dst_bg, method="c14n")
    src_c14n = etree.tostring(src_bg, method="c14n")
    assert dst_c14n == src_c14n, (
        "destination p:bg differs from source p:bg\nDST: %s\nSRC: %s"
        % (dst_c14n, src_c14n)
    )


# --- Slide.effective_background (issue #809) ---------------------------------


@given("a slide with its own explicit fill background")
def given_a_slide_with_its_own_explicit_fill_background(context):
    # -- slide 1 in the fixture carries a solid red p:bg --
    prs = Presentation(test_pptx("sld-background"))
    context.prs = prs
    context.slide = prs.slides[1]
    # -- sanity: fixture precondition --
    assert context.slide.background.bg_element is not None


@given("a slide inheriting a layout-level fill background")
def given_a_slide_inheriting_a_layout_level_fill_bg(context):
    # -- build a fresh deck with a layout-level solid color and a slide
    # -- that uses that layout without its own override. --
    prs = Presentation()
    layout = prs.slide_layouts[0]
    layout.background.fill.solid()
    layout.background.fill.fore_color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    slide = prs.slides.add_slide(layout)
    context.prs = prs
    context.slide = slide
    # -- sanity: precondition --
    assert slide._element.cSld.bg is None
    assert layout._element.cSld.bg is not None


@given("a slide with no explicit background and no layout background")
def given_a_slide_with_no_explicit_bg_and_no_layout_bg(context):
    # -- default template: neither slide nor layout carries a p:bg;
    # -- the master does (a p:bgRef). --
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    context.prs = prs
    context.slide = slide
    assert slide._element.cSld.bg is None
    assert slide.slide_layout._element.cSld.bg is None


@when("I read slide.effective_background.fill.fore_color.rgb")
def when_I_read_slide_effective_background_fill_fore_color_rgb(context):
    eff = context.slide.effective_background
    assert eff is not None
    assert eff.fill is not None
    # -- reading .rgb is the usage the #809 reporter asked about --
    context.read_rgb = eff.fill.fore_color.rgb


@then('slide.effective_background.source is "slide"')
def then_effective_background_source_is_slide(context):
    eff = context.slide.effective_background
    assert eff is not None, "effective_background is None"
    assert eff.source == "slide", 'source is %r, expected "slide"' % eff.source


@then('slide.effective_background.source is "layout"')
def then_effective_background_source_is_layout(context):
    eff = context.slide.effective_background
    assert eff is not None, "effective_background is None"
    assert eff.source == "layout", 'source is %r, expected "layout"' % eff.source


@then('slide.effective_background.source is "master"')
def then_effective_background_source_is_master(context):
    eff = context.slide.effective_background
    assert eff is not None, "effective_background is None"
    assert eff.source == "master", 'source is %r, expected "master"' % eff.source


@then("slide.effective_background.fill reads the slide color")
def then_effective_background_fill_reads_slide_color(context):
    # -- fixture slide 1 uses FF0000 --
    rgb = context.slide.effective_background.fill.fore_color.rgb
    assert rgb == RGBColor(0xFF, 0x00, 0x00), "expected FF0000, got %r" % rgb


@then("slide.effective_background.fill reads the layout color")
def then_effective_background_fill_reads_layout_color(context):
    rgb = context.slide.effective_background.fill.fore_color.rgb
    assert rgb == RGBColor(0xAB, 0xCD, 0xEF), "expected ABCDEF, got %r" % rgb


@then("the slide still has no explicit p:bg element")
def then_slide_still_has_no_explicit_p_bg_element(context):
    assert context.slide._element.cSld.bg is None, (
        "slide was mutated to carry a p:bg element"
    )
