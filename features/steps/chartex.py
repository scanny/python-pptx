"""Gherkin step implementations for Office 2016+ extended (chartex) chart passthrough."""

from __future__ import annotations

import io
import zipfile

from behave import given, then, when
from helpers import test_pptx

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE

_CHARTEX_URI = "http://schemas.microsoft.com/office/drawing/2014/chartex"


# given ===================================================


@given("a presentation containing a chartex chart")
def given_a_presentation_containing_a_chartex_chart(context):
    context.prs = Presentation(test_pptx("cht-chartex"))
    context.slide = context.prs.slides[0]


@given("the {which} chartex graphic-frame shape from a presentation")
def given_the_nth_chartex_graphic_frame_shape(context, which):
    prs = Presentation(test_pptx("cht-chartex"))
    # -- slide has two shapes: index 0 is the mc:AlternateContent-wrapped chartex,
    # -- index 1 is the directly-placed chartex graphicFrame
    idx = {"AlternateContent-wrapped": 0, "direct": 1}[which]
    context.shape = prs.slides[0].shapes[idx]


# when ====================================================


@when("I save the presentation and reload it")
def when_I_save_the_presentation_and_reload_it(context):
    buf = io.BytesIO()
    context.prs.save(buf)
    context.saved_bytes = buf.getvalue()
    buf.seek(0)
    context.prs = Presentation(buf)
    context.slide = context.prs.slides[0]


# then ====================================================


@then("slide.shapes contains the chartex graphic-frame shapes")
def then_slide_shapes_contains_the_chartex_graphic_frame_shapes(context):
    shapes = list(context.slide.shapes)
    chartex_shapes = [s for s in shapes if getattr(s, "has_chartex", False)]
    assert (
        len(chartex_shapes) == 2
    ), "expected 2 chartex graphic-frame shapes; got %d (total shapes: %d)" % (
        len(chartex_shapes),
        len(shapes),
    )


@then("shape.has_chartex is {value}")
def then_shape_has_chartex_is_value(context, value):
    expected = {"True": True, "False": False}[value]
    actual = context.shape.has_chartex
    assert actual is expected, "shape.has_chartex is %s" % actual


@then("shape.chart_type is XL_CHART_TYPE.UNSUPPORTED_CHARTEX")
def then_shape_chart_type_is_UNSUPPORTED_CHARTEX(context):
    actual = context.shape.chart_type
    assert actual is XL_CHART_TYPE.UNSUPPORTED_CHARTEX, "shape.chart_type is %s" % actual


@then('shape.chartex_type == "{expected}"')
def then_shape_chartex_type_equals(context, expected):
    actual = context.shape.chartex_type
    assert actual == expected, "shape.chartex_type is %r" % actual


@then("shape.shape_type == MSO_SHAPE_TYPE.CHART")
def then_shape_shape_type_is_chart(context):
    # -- existing step in shape.py asserts via parameterised string match;
    # -- we duplicate here to avoid coupling feature file ordering to unrelated steps.
    actual = context.shape.shape_type
    assert actual is MSO_SHAPE_TYPE.CHART, "shape.shape_type is %s" % actual


@then("the reloaded slide still exposes the chartex graphic-frame shapes")
def then_reloaded_slide_still_exposes_chartex_shapes(context):
    shapes = list(context.slide.shapes)
    chartex_shapes = [s for s in shapes if s.has_chartex]
    assert len(chartex_shapes) == 2, "got %d chartex shapes after reload" % len(chartex_shapes)
    for shape in chartex_shapes:
        assert shape.chart_type is XL_CHART_TYPE.UNSUPPORTED_CHARTEX


@then("the saved package still contains the chartex part and its AlternateContent wrapper")
def then_saved_package_still_contains_chartex(context):
    with zipfile.ZipFile(io.BytesIO(context.saved_bytes)) as z:
        assert "ppt/charts/chartEx1.xml" in z.namelist(), "chartex part missing from saved package"
        slide_xml = z.read("ppt/slides/slide1.xml").decode()
    assert "AlternateContent" in slide_xml, "mc:AlternateContent wrapper missing from saved slide"
    assert _CHARTEX_URI in slide_xml, "chartex graphicData URI missing from saved slide"
