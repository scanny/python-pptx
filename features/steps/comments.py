"""Gherkin step implementations for legacy-comments acceptance tests."""

from __future__ import annotations

import datetime as dt

from behave import given, then, when
from helpers import saved_pptx_path

from pptx import Presentation

# -- given -----------------------------------------------------------------


@given("a presentation with one slide")
def given_a_presentation_with_one_slide(context):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5])
    context.prs = prs
    context.slide = prs.slides[0]


# -- when ------------------------------------------------------------------


@when("I add a comment authored by 'Alice' to the slide")
def when_add_comment_authored_by_alice(context):
    comments = context.slide.comments
    alice = comments._authors.add_author("Alice", "A.")
    context.alice = alice
    comments.add_comment(
        alice,
        "Please revise this slide.",
        position=(100, 200),
        datetime_value=dt.datetime(2025, 5, 2, 12, 0, 0),
    )


# -- then ------------------------------------------------------------------


@then("the reloaded slide reports it has a comments part")
def then_reloaded_slide_has_comments(context):
    context.prs2 = Presentation(saved_pptx_path)
    context.slide2 = context.prs2.slides[0]
    assert context.slide2.has_comments is True, "expected has_comments==True"


@then("the reloaded slide has one comment with the expected text and author")
def then_reloaded_comment_matches(context):
    comments = list(context.slide2.comments)
    assert len(comments) == 1, "expected 1 comment, got %d" % len(comments)
    c = comments[0]
    assert c.text == "Please revise this slide.", "unexpected text: %r" % c.text
    assert c.author is not None, "comment has no author"
    assert c.author.name == "Alice"
    assert c.datetime == dt.datetime(2025, 5, 2, 12, 0, 0)


@then("the slide reports it has no comments part")
def then_slide_has_no_comments_part(context):
    assert context.slide.has_comments is False, "expected has_comments==False"


@then("the slide's Comments collection has length zero")
def then_comments_collection_is_empty(context):
    assert len(context.slide.comments) == 0
