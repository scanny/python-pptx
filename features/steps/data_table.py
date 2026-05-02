"""Gherkin step implementations for chart data-table features (issue #373)."""

from __future__ import annotations

from behave import given, then, when

from pptx import Presentation
from pptx.chart.chart import _DataTable
from pptx.chart.data import CategoryChartData
from pptx.dml.chtfmt import ChartFormat
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def _make_chart(with_data_table: bool):
    """Build a fresh in-memory chart and optionally add a data-table."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_data = CategoryChartData()
    chart_data.categories = ["East", "West", "Mid-West"]
    chart_data.add_series("Q1", (19.2, 21.4, 16.7))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    ).chart
    if with_data_table:
        chart.has_data_table = True
    return chart


# given ===================================================


@given("a chart {having_or_not} a data-table")
def given_a_chart_having_or_not_a_data_table(context, having_or_not):
    with_dt = {"having": True, "not having": False}[having_or_not]
    context.chart = _make_chart(with_dt)


# when ====================================================


@when("I assign {value} to chart.has_data_table")
def when_I_assign_value_to_chart_has_data_table(context, value):
    new_value = {"True": True, "False": False}[value]
    context.chart.has_data_table = new_value


@when("I assign {value} to chart.data_table.{flag}")
def when_I_assign_value_to_chart_data_table_flag(context, value, flag):
    new_value = {"True": True, "False": False, "None": None}[value]
    setattr(context.chart.data_table, flag, new_value)


# then ====================================================


@then("chart.has_data_table is {value}")
def then_chart_has_data_table_is_value(context, value):
    expected = {"True": True, "False": False}[value]
    assert context.chart.has_data_table is expected, "got %s" % context.chart.has_data_table


@then("chart.data_table is {description}")
def then_chart_data_table_is(context, description):
    data_table = context.chart.data_table
    if description == "None":
        assert data_table is None, "got %r" % data_table
    elif description == "a _DataTable":
        assert isinstance(data_table, _DataTable), "got %r" % data_table
    else:
        raise ValueError("unrecognized description: %r" % description)


@then("chart.data_table.show_{flag} is {value}")
def then_chart_data_table_show_flag_is_value(context, flag, value):
    expected = {"True": True, "False": False}[value]
    actual = getattr(context.chart.data_table, "show_%s" % flag)
    assert actual is expected, "got %r" % actual


@then("chart.data_table.format is a ChartFormat object")
def then_chart_data_table_format_is_a_ChartFormat_object(context):
    fmt = context.chart.data_table.format
    assert isinstance(fmt, ChartFormat), "got %r" % fmt
