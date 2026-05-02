"""Gherkin step implementations for extended-properties feature (issue #131)."""

from __future__ import annotations

import re
import zipfile

from behave import given, then
from behave.runner import Context
from helpers import saved_pptx_path

from pptx import Presentation

# given ===================================================


@given("I have a presentation with three slides added")
def step_given_three_slide_presentation(context: Context):
    prs = Presentation()
    layout = prs.slide_layouts[0]
    for _ in range(3):
        prs.slides.add_slide(layout)
    context.prs = prs


# then ====================================================


@then("the <Slides> element in docProps/app.xml equals 3")
def step_then_slides_element_is_three(context: Context):
    with zipfile.ZipFile(saved_pptx_path) as pkg:
        app_xml = pkg.read("docProps/app.xml").decode("utf-8")
    match = re.search(r"<Slides>(\d+)</Slides>", app_xml)
    assert match is not None, "no <Slides> element found in docProps/app.xml"
    assert match.group(1) == "3", (
        "expected <Slides>3</Slides>, got <Slides>%s</Slides>" % match.group(1)
    )
