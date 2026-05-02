"""Gherkin step implementations for header/footer/slide-number feature."""

from __future__ import annotations

from behave import given, then, when

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

# given ===================================================


@given("a SlideMaster with no p:hf element")
def given_a_SlideMaster_with_no_hf_element(context):
    prs = Presentation()
    context.slide_master = prs.slide_masters[0]
    # ---defensive: ensure master has no existing p:hf so the scenarios start
    # ---from a well-known baseline---
    sldMaster = context.slide_master._element
    hf = sldMaster.hf
    if hf is not None:
        sldMaster.remove(hf)


@given("a SlideLayout with no p:hf element")
def given_a_SlideLayout_with_no_hf_element(context):
    prs = Presentation()
    context.slide_layout = prs.slide_layouts[0]
    sldLayout = context.slide_layout._element
    hf = sldLayout.hf
    if hf is not None:
        sldLayout.remove(hf)


@given("a paragraph")
def given_a_paragraph(context):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(0.5))
    context.paragraph = shape.text_frame.paragraphs[0]


# when ====================================================


@when("I set slide_master.header_footer.{flag} to False")
def when_I_set_slide_master_header_footer_flag_to_False(context, flag: str):
    setattr(context.slide_master.header_footer, flag, False)


@when('I call paragraph.add_field("slidenum", "#")')
def when_I_call_paragraph_add_field_slidenum(context):
    context.field = context.paragraph.add_field("slidenum", "#")


# then ====================================================


@then("slide_master.header_footer.{flag} is {value}")
def then_slide_master_header_footer_flag_is(context, flag: str, value: str):
    expected = {"True": True, "False": False}[value]
    actual = getattr(context.slide_master.header_footer, flag)
    assert actual is expected, (
        f"slide_master.header_footer.{flag} is {actual!r}, expected {expected!r}"
    )


@then("slide_layout.header_footer.{flag} is {value}")
def then_slide_layout_header_footer_flag_is(context, flag: str, value: str):
    expected = {"True": True, "False": False}[value]
    actual = getattr(context.slide_layout.header_footer, flag)
    assert actual is expected, (
        f"slide_layout.header_footer.{flag} is {actual!r}, expected {expected!r}"
    )


@then("the p:sldMaster element has a p:hf child")
def then_sldMaster_has_hf_child(context):
    sldMaster = context.slide_master._element
    assert sldMaster.hf is not None, "expected p:sldMaster to have a p:hf child"


@then("paragraph has one a:fld child")
def then_paragraph_has_one_fld_child(context):
    p = context.paragraph._p
    flds = p.findall(qn("a:fld"))
    assert len(flds) == 1, "expected one a:fld child, got %d" % len(flds)


@then('the a:fld type is "{fld_type}"')
def then_the_fld_type_is(context, fld_type: str):
    assert context.field.field_type == fld_type, (
        f"a:fld/@type is {context.field.field_type!r}, expected {fld_type!r}"
    )


@then('the a:fld text is "{fld_text}"')
def then_the_fld_text_is(context, fld_text: str):
    assert context.field.text == fld_text, (
        f"a:fld text is {context.field.text!r}, expected {fld_text!r}"
    )
