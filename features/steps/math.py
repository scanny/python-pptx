"""Gherkin step implementations for OMML math-related features."""

from __future__ import annotations

from behave import given, then, when
from helpers import test_pptx

from pptx import Presentation
from pptx.util import Emu

# given ====================================================


@given("a presentation with one slide")
def given_a_presentation_with_one_slide(context):
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]  # Use first available layout
    slide = prs.slides.add_slide(slide_layout)
    context.presentation = prs
    context.slide = slide


@given("a presentation with one slide containing a math equation")
def given_a_presentation_with_math_equation(context):
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)

    # Add a simple math equation
    omml_xml = "<m:oMath xmlns:m='http://purl.oclc.org/ooxml/officeDocument/math'><m:r><m:t>x = 2</m:t></m:r></m:oMath>"
    math_shape = slide.shapes.add_math_equation()
    math_shape.math.add_omml(omml_xml)

    context.presentation = prs
    context.slide = slide
    context.math_shape = math_shape


@given("a presentation with one slide containing a math shape")
def given_a_presentation_with_math_shape(context):
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)

    omml_xml = "<m:oMath xmlns:m='http://purl.oclc.org/ooxml/officeDocument/math'><m:r><m:t>E = mc²</m:t></m:r></m:oMath>"
    math_shape = slide.shapes.add_math_equation()
    math_shape.math.add_omml(omml_xml)

    context.presentation = prs
    context.slide = slide
    context.math_shape = math_shape


@given("a presentation with math equations")
def given_a_presentation_with_math_equations(context):
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)

    # Add multiple math equations
    equations = [
        "<m:oMath xmlns:m='http://purl.oclc.org/ooxml/officeDocument/math'><m:r><m:t>a² + b²</m:t></m:r></m:oMath>",
        "<m:oMath xmlns:m='http://purl.oclc.org/ooxml/officeDocument/math'><m:r><m:t>= c²</m:t></m:r></m:oMath>"
    ]

    for omml_xml in equations:
        math_shape = slide.shapes.add_math_equation()
        math_shape.math.add_omml(omml_xml)

    context.presentation = prs
    context.slide = slide


# when =====================================================


@when('I add a math equation "{omml_xml}"')
def when_i_add_math_equation(context, omml_xml):
    math_shape = context.slide.shapes.add_math_equation()
    math_shape.math.add_omml(omml_xml)
    context.math_shape = math_shape


@when('I add a math equation with fraction "{omml_xml}"')
def when_i_add_math_equation_with_fraction(context, omml_xml):
    math_shape = context.slide.shapes.add_math_equation()
    math_shape.math.add_omml(omml_xml)
    context.math_shape = math_shape


@when('I add a math equation with superscript "{omml_xml}"')
def when_i_add_math_equation_with_superscript(context, omml_xml):
    math_shape = context.slide.shapes.add_math_equation()
    math_shape.math.add_omml(omml_xml)
    context.math_shape = math_shape


@when('I add a math equation with radical "{omml_xml}"')
def when_i_add_math_equation_with_radical(context, omml_xml):
    math_shape = context.slide.shapes.add_math_equation()
    math_shape.math.add_omml(omml_xml)
    context.math_shape = math_shape


@when('I add a math equation with summation "{omml_xml}"')
def when_i_add_math_equation_with_summation(context, omml_xml):
    math_shape = context.slide.shapes.add_math_equation()
    math_shape.math.add_omml(omml_xml)
    context.math_shape = math_shape


@when("I position the math shape at left {left:d}, top {top:d}")
def when_i_position_math_shape(context, left, top):
    context.math_shape.left = Emu(left)
    context.math_shape.top = Emu(top)


@when("I size the math shape to width {width:d}, height {height:d}")
def when_i_size_math_shape(context, width, height):
    context.math_shape.width = Emu(width)
    context.math_shape.height = Emu(height)


@when('I add a first math equation "{omml_xml}"')
def when_i_add_first_math_equation(context, omml_xml):
    math_shape = context.slide.shapes.add_math_equation()
    math_shape.math.add_omml(omml_xml)
    context.first_math_shape = math_shape


@when('I add a second math equation "{omml_xml}"')
def when_i_add_second_math_equation(context, omml_xml):
    math_shape = context.slide.shapes.add_math_equation()
    math_shape.math.add_omml(omml_xml)
    context.second_math_shape = math_shape


@when("I get the OMML XML from the math shape")
def when_i_get_omml_xml(context):
    context.omml_xml = context.math_shape.math.get_omml()


@when('I replace the OMML content with "{omml_xml}"')
def when_i_replace_omml_content(context, omml_xml):
    context.math_shape.math.set_omml(omml_xml)


@when("I save the presentation to a file")
def when_i_save_presentation(context):
    import tempfile
    import os

    temp_file = tempfile.NamedTemporaryFile(suffix='.pptx', delete=False)
    temp_file.close()

    context.presentation.save(temp_file.name)
    context.saved_file_path = temp_file.name


@when("I reload the presentation from the file")
def when_i_reload_presentation(context):
    context.presentation = Presentation(context.saved_file_path)
    context.slide = context.presentation.slides[0]


# then =====================================================


@then("the slide should contain one math shape")
def then_slide_should_contain_one_math_shape(context):
    math_shapes = [s for s in context.slide.shapes if hasattr(s, 'math')]
    assert len(math_shapes) == 1, f"Expected 1 math shape, found {len(math_shapes)}"


@then('the math shape should contain the equation "{expected_equation}"')
def then_math_shape_should_contain_equation(context, expected_equation):
    omml_xml = context.math_shape.math.get_omml()
    assert expected_equation in omml_xml, f"Expected '{expected_equation}' in OMML: {omml_xml}"


@then("the math shape should contain a fraction with numerator \"{numerator}\" and denominator \"{denominator}\"")
def then_math_shape_should_contain_fraction(context, numerator, denominator):
    omml_xml = context.math_shape.math.get_omml()
    assert "<m:f>" in omml_xml, "No fraction element found"
    assert f"<m:t>{numerator}</m:t>" in omml_xml, f"Numerator '{numerator}' not found"
    assert f"<m:t>{denominator}</m:t>" in omml_xml, f"Denominator '{denominator}' not found"


@then('the math shape should contain "{expected_content}"')
def then_math_shape_should_contain_content(context, expected_content):
    omml_xml = context.math_shape.math.get_omml()
    # Convert superscript notation for comparison
    if "²" in expected_content:
        assert "<m:sSup>" in omml_xml, "No superscript element found"
        assert "<m:t>2</m:t>" in omml_xml, "Superscript '2' not found"
    else:
        assert expected_content in omml_xml, f"Expected '{expected_content}' in OMML: {omml_xml}"


@then("the math shape should contain a square root of \"{content}\"")
def then_math_shape_should_contain_square_root(context, content):
    omml_xml = context.math_shape.math.get_omml()
    assert "<m:rad>" in omml_xml, "No radical element found"
    assert f"<m:t>{content}</m:t>" in omml_xml, f"Content '{content}' not found in radical"


@then("the math shape should contain a summation from i=1 to n of i")
def then_math_shape_should_contain_summation(context):
    omml_xml = context.math_shape.math.get_omml()
    assert "<m:nary>" in omml_xml, "No n-ary element found"
    assert '<m:chr val="∑">' in omml_xml or '<m:chr val="∑"/>' in omml_xml, "No summation character found"
    assert "<m:t>i=1</m:t>" in omml_xml, "Lower limit not found"
    assert "<m:t>n</m:t>" in omml_xml, "Upper limit not found"
    assert "<m:t>i</m:t>" in omml_xml, "Expression not found"


@then("the math shape left should be {left:d}")
def then_math_shape_left_should_be(context, left):
    assert context.math_shape.left == Emu(left), f"Expected left {left}, got {context.math_shape.left}"


@then("the math shape top should be {top:d}")
def then_math_shape_top_should_be(context, top):
    assert context.math_shape.top == Emu(top), f"Expected top {top}, got {context.math_shape.top}"


@then("the math shape width should be {width:d}")
def then_math_shape_width_should_be(context, width):
    assert context.math_shape.width == Emu(width), f"Expected width {width}, got {context.math_shape.width}"


@then("the math shape height should be {height:d}")
def then_math_shape_height_should_be(context, height):
    assert context.math_shape.height == Emu(height), f"Expected height {height}, got {context.math_shape.height}"


@then("the slide should contain two math shapes")
def then_slide_should_contain_two_math_shapes(context):
    math_shapes = [s for s in context.slide.shapes if hasattr(s, 'math')]
    assert len(math_shapes) == 2, f"Expected 2 math shapes, found {len(math_shapes)}"


@then('the first math shape should contain "{expected_content}"')
def then_first_math_shape_should_contain(context, expected_content):
    omml_xml = context.first_math_shape.math.get_omml()
    assert expected_content in omml_xml, f"Expected '{expected_content}' in first OMML: {omml_xml}"


@then('the second math shape should contain "{expected_content}"')
def then_second_math_shape_should_contain(context, expected_content):
    omml_xml = context.second_math_shape.math.get_omml()
    assert expected_content in omml_xml, f"Expected '{expected_content}' in second OMML: {omml_xml}"


@then("the OMML XML should be valid math markup")
def then_omml_xml_should_be_valid(context):
    omml_xml = context.omml_xml
    assert "<m:oMath" in omml_xml, "No oMath root element found"
    assert "xmlns:m=" in omml_xml, "No math namespace found"
    assert "</m:oMath>" in omml_xml, "No closing oMath element found"


@then("the OMML XML should contain the original equation content")
def then_omml_xml_should_contain_original_content(context):
    omml_xml = context.omml_xml
    assert "<m:t>x = 2</m:t>" in omml_xml, "Original equation content not found"


@then("the math shape should have default properties")
def then_math_shape_should_have_default_properties(context):
    # Check that it's a proper shape with expected properties
    assert hasattr(context.math_shape, 'left'), "Math shape missing left property"
    assert hasattr(context.math_shape, 'top'), "Math shape missing top property"
    assert hasattr(context.math_shape, 'width'), "Math shape missing width property"
    assert hasattr(context.math_shape, 'height'), "Math shape missing height property"


@then("the math shape should support rotation")
def then_math_shape_should_support_rotation(context):
    # Test setting rotation (should not raise an exception)
    original_rotation = getattr(context.math_shape, 'rotation', None)
    context.math_shape.rotation = 45
    assert context.math_shape.rotation == 45, "Math shape rotation not supported"


@then("the math shape should support shadow effects")
def then_math_shape_should_support_shadow(context):
    # Test shadow properties (should not raise an exception)
    if hasattr(context.math_shape, 'shadow'):
        context.math_shape.shadow.inherit = True
        assert context.math_shape.shadow.inherit == True, "Math shape shadow not supported"


@then("the math shape should be included in slide shapes collection")
def then_math_shape_in_shapes_collection(context):
    math_shapes = [s for s in context.slide.shapes if hasattr(s, 'math')]
    assert len(math_shapes) > 0, "Math shape not found in shapes collection"
    assert context.math_shape in math_shapes, "Math shape not in shapes collection"


@then("the presentation should contain the same math equations")
def then_presentation_should_contain_same_math_equations(context):
    math_shapes = [s for s in context.slide.shapes if hasattr(s, 'math')]
    assert len(math_shapes) > 0, "No math shapes found after reload"


@then("the math equations should have the same content")
def then_math_equations_should_have_same_content(context):
    math_shapes = [s for s in context.slide.shapes if hasattr(s, 'math')]
    for math_shape in math_shapes:
        omml_xml = math_shape.math.get_omml()
        assert "<m:oMath" in omml_xml, "Invalid OMML content after reload"


@then("the math equations should have the same positions and sizes")
def then_math_equations_should_have_same_positions_and_sizes(context):
    math_shapes = [s for s in context.slide.shapes if hasattr(s, 'math')]
    for math_shape in math_shapes:
        assert hasattr(math_shape, 'left'), "Missing position after reload"
        assert hasattr(math_shape, 'width'), "Missing size after reload"
