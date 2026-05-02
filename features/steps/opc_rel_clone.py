"""Gherkin step implementations for the cross-part rel cloning helper."""

from __future__ import annotations

from behave import given, then, when
from helpers import test_image

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import PartRelationshipCloner
from pptx.oxml.ns import qn
from pptx.util import Inches

# given ===================================================


@given("a slide containing a picture")
def given_a_slide_containing_a_picture(context):
    prs = Presentation()
    context.prs = prs
    slide_layout = prs.slide_layouts[6]  # blank layout
    context.src_slide = prs.slides.add_slide(slide_layout)
    context.src_picture = context.src_slide.shapes.add_picture(
        test_image("monty-truth.png"), Inches(1), Inches(1)
    )


@given("a second slide with no picture")
def given_a_second_slide_with_no_picture(context):
    slide_layout = context.prs.slide_layouts[6]
    context.tgt_slide = context.prs.slides.add_slide(slide_layout)


# when ====================================================


@when("I use PartRelationshipCloner to copy the picture element to the second slide's part")
def when_I_use_PartRelationshipCloner(context):
    src_part = context.src_slide.part
    tgt_part = context.tgt_slide.part
    src_pic_element = context.src_picture._element  # pyright: ignore[reportPrivateUsage]

    # -- capture pre-clone state for later assertions --
    old_embed_rId = src_pic_element.find(qn("p:blipFill") + "/" + qn("a:blip")).get(qn("r:embed"))
    context.old_embed_rId = old_embed_rId
    context.old_image_part = src_part.related_part(old_embed_rId)

    # -- exercise the helper --
    context.new_pic_element = PartRelationshipCloner.clone(src_part, tgt_part, src_pic_element)


# then ====================================================


@then("the returned element has a fresh r:embed value")
def then_the_returned_element_has_a_fresh_r_embed_value(context):
    new_blip = context.new_pic_element.find(qn("p:blipFill") + "/" + qn("a:blip"))
    assert new_blip is not None, "no <a:blip> found in cloned element"
    new_rId = new_blip.get(qn("r:embed"))
    assert new_rId is not None
    assert new_rId != ""
    # -- the rewritten rId must resolve on the target part --
    tgt_part = context.tgt_slide.part
    assert new_rId in tgt_part.rels, "new rId %r not registered on tgt_part" % new_rId


@then("the second slide's part is related to a new image part with the same blob")
def then_the_second_slide_part_has_a_new_image_part(context):
    new_blip = context.new_pic_element.find(qn("p:blipFill") + "/" + qn("a:blip"))
    new_rId = new_blip.get(qn("r:embed"))
    tgt_part = context.tgt_slide.part
    new_rel = tgt_part.rels[new_rId]
    assert new_rel.reltype == RT.IMAGE, "unexpected reltype %r" % new_rel.reltype
    # -- same image data (both sides of the same package → reuse is legitimate) --
    assert new_rel.target_part.blob == context.old_image_part.blob


@then("the original slide's picture still points at its original image part")
def then_the_original_slide_picture_is_unchanged(context):
    src_part = context.src_slide.part
    src_pic_element = context.src_picture._element  # pyright: ignore[reportPrivateUsage]
    old_rId = src_pic_element.find(qn("p:blipFill") + "/" + qn("a:blip")).get(qn("r:embed"))
    # -- source element was not mutated --
    assert old_rId == context.old_embed_rId
    # -- source part still resolves that rId to its original image part --
    assert src_part.related_part(old_rId) is context.old_image_part
