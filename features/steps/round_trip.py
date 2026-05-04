"""Step implementations for save/reload round-trip scenarios."""

from __future__ import annotations

import io

from behave import given, then, when
from helpers import test_image

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_DISPLAY_BLANKS_AS,
    XL_ERROR_BAR_INCLUDE,
    XL_ERROR_BAR_TYPE,
)
from pptx.enum.dml import MSO_LINE
from pptx.enum.text import MSO_STRIKE
from pptx.util import Emu, Inches, Pt


def _build_prs_with_blank_slide():
    prs = Presentation()
    # -- layout[6] is Blank — zero placeholders, so newly-added shapes
    # -- start at slide.shapes[0] without a title-placeholder offset --
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return prs, slide


def _round_trip(context):
    buf = io.BytesIO()
    context.prs.save(buf)
    buf.seek(0)
    context.prs = Presentation(buf)


@when("I round-trip the presentation")
def when_round_trip_the_presentation(context):
    _round_trip(context)


# -- Shape-level round-trips --------------------------------------------


@given("a fresh presentation with a flipped autoshape")
def given_fresh_prs_with_flipped_autoshape(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    shape.flip_horizontal = True
    shape.flip_vertical = True
    context.prs = prs


@then("the reloaded autoshape flip_horizontal is True")
def then_reloaded_flip_horizontal_True(context):
    shape = context.prs.slides[0].shapes[0]
    assert shape.flip_horizontal is True, shape.flip_horizontal


@then("the reloaded autoshape flip_vertical is True")
def then_reloaded_flip_vertical_True(context):
    shape = context.prs.slides[0].shapes[0]
    assert shape.flip_vertical is True, shape.flip_vertical


@given("a fresh presentation with a shape carrying alt_text and title")
def given_fresh_prs_with_alt_text_and_title(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    shape.alt_text = "description of chart"
    shape.title = "Quarterly Revenue"
    context.prs = prs


@then('the reloaded shape.alt_text is "description of chart"')
def then_reloaded_alt_text(context):
    shape = context.prs.slides[0].shapes[0]
    assert shape.alt_text == "description of chart", repr(shape.alt_text)


@then('the reloaded shape.title is "Quarterly Revenue"')
def then_reloaded_title(context):
    shape = context.prs.slides[0].shapes[0]
    assert shape.title == "Quarterly Revenue", repr(shape.title)


@given("a fresh presentation with a hidden shape")
def given_fresh_prs_with_hidden_shape(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    shape.is_hidden = True
    context.prs = prs


@then("the reloaded shape.is_hidden is True")
def then_reloaded_shape_is_hidden(context):
    shape = context.prs.slides[0].shapes[0]
    assert shape.is_hidden is True, shape.is_hidden


@given('a fresh presentation with a shape renamed to "LogoBanner"')
def given_fresh_prs_with_renamed_shape(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    shape.name = "LogoBanner"
    context.prs = prs
    context.expected_shape_id = shape.shape_id


@then('slide.shapes.get_by_name("LogoBanner") is the renamed shape')
def then_get_by_name_locates_shape(context):
    shapes = context.prs.slides[0].shapes
    got = shapes.get_by_name("LogoBanner")
    assert got is not None, "expected a shape, got None"
    assert got.name == "LogoBanner", got.name


@then('slide.shapes.get_by_name("DoesNotExist") is None')
def then_get_by_name_miss(context):
    shapes = context.prs.slides[0].shapes
    assert shapes.get_by_name("DoesNotExist") is None


@given('a fresh presentation with three shapes all named "Tag"')
def given_fresh_prs_three_tag_shapes(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    for i in range(3):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1 + i), Inches(1), Inches(0.5), Inches(0.5)
        )
        shape.name = "Tag"
    context.prs = prs


@then('slide.shapes.find_all_by_name("Tag") returns 3 shapes')
def then_find_all_by_name_Tag_3(context):
    shapes = context.prs.slides[0].shapes
    hits = list(shapes.find_all_by_name("Tag"))
    assert len(hits) == 3, "expected 3, got %d" % len(hits)


@then('slide.shapes.find_all_by_name("Nope") returns 0 shapes')
def then_find_all_by_name_Nope_0(context):
    shapes = context.prs.slides[0].shapes
    hits = list(shapes.find_all_by_name("Nope"))
    assert len(hits) == 0, "expected 0, got %d" % len(hits)


@given("a fresh presentation with a group of two shapes plus a lone rectangle")
def given_fresh_prs_with_group_and_rect(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    # -- two autoshapes then group them --
    a = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1), Inches(1), Inches(1), Inches(1))
    b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.5), Inches(1), Inches(1), Inches(1))
    # -- group them via shapes.add_group_shape --
    slide.shapes.add_group_shape([a, b])
    # -- one lone rectangle outside the group --
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(1), Inches(1), Inches(1))
    context.prs = prs


@then("the reloaded slide.shapes.iter_leaf_shapes yields 3 leaves")
def then_reloaded_iter_leaf_shapes_3(context):
    leaves = list(context.prs.slides[0].shapes.iter_leaf_shapes())
    assert len(leaves) == 3, "expected 3, got %d: %r" % (
        len(leaves),
        [s.shape_type for s in leaves],
    )


@given("a fresh presentation with three rectangles in z-order A B C")
def given_fresh_prs_three_rects(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    for letter in ("A", "B", "C"):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1)
        )
        shape.name = letter
    context.prs = prs


@when("I call bring_to_front on A")
def when_bring_to_front_A(context):
    shapes = context.prs.slides[0].shapes
    for shape in shapes:
        if shape.name == "A":
            shape.bring_to_front()
            return
    raise AssertionError("no shape named 'A'")


@then("the reloaded shapes z-order is B C A")
def then_reloaded_z_order_B_C_A(context):
    shapes = context.prs.slides[0].shapes
    names = [sh.name for sh in shapes]
    assert names == ["B", "C", "A"], names


# -- Text-frame round-trips --------------------------------------------


@given("a fresh presentation with a text_frame rotated 45 degrees")
def given_fresh_prs_with_rotated_text_frame(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.rotation = 45.0
    context.prs = prs


@then("the reloaded text_frame.rotation == 45.0")
def then_reloaded_text_frame_rotation(context):
    tf = context.prs.slides[0].shapes[0].text_frame
    assert tf.rotation == 45.0, tf.rotation


@given("a fresh presentation with autofit font_scale 75% and line_space_reduction 10%")
def given_fresh_prs_with_font_scale(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "sample"
    # -- font_scale / line_space_reduction require an a:normAutofit present --
    tb.text_frame.font_scale = 75.0
    tb.text_frame.line_space_reduction = 10.0
    context.prs = prs


@then("the reloaded text_frame.font_scale == 75.0")
def then_reloaded_font_scale(context):
    tf = context.prs.slides[0].shapes[0].text_frame
    assert tf.font_scale == 75.0, tf.font_scale


@then("the reloaded text_frame.line_space_reduction == 10.0")
def then_reloaded_line_space_reduction(context):
    tf = context.prs.slides[0].shapes[0].text_frame
    assert tf.line_space_reduction == 10.0, tf.line_space_reduction


@given('a fresh presentation with a text_frame containing "Hello world" runs')
def given_fresh_prs_text_frame_hello_world(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "Hello world"
    context.prs = prs


@when('I call text_frame.replace_text("world", "Python")')
def when_call_text_frame_replace_text(context):
    tf = context.prs.slides[0].shapes[0].text_frame
    tf.replace_text("world", "Python")


@then('the reloaded text_frame.text == "Hello Python"')
def then_reloaded_text_frame_text_hello_python(context):
    tf = context.prs.slides[0].shapes[0].text_frame
    assert tf.text == "Hello Python", repr(tf.text)


# -- Paragraph / Run delete -------------------------------------------


@given("a fresh presentation with a textbox of three paragraphs")
def given_fresh_prs_with_three_paragraphs(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    tf = tb.text_frame
    tf.text = "first"
    tf.add_paragraph().text = "middle"
    tf.add_paragraph().text = "third"
    context.prs = prs


@when("I call paragraph.delete on the middle paragraph")
def when_delete_middle_paragraph(context):
    tf = context.prs.slides[0].shapes[0].text_frame
    tf.paragraphs[1].delete()


@then("the reloaded textbox has 2 paragraphs")
def then_reloaded_textbox_has_2_paragraphs(context):
    tf = context.prs.slides[0].shapes[0].text_frame
    assert len(tf.paragraphs) == 2, len(tf.paragraphs)


@then('the reloaded paragraph texts are "first" and "third"')
def then_reloaded_paragraph_texts_first_third(context):
    tf = context.prs.slides[0].shapes[0].text_frame
    texts = [p.text for p in tf.paragraphs]
    assert texts == ["first", "third"], texts


@given('a fresh presentation with a textbox of one paragraph containing "only"')
def given_fresh_prs_with_one_paragraph_only(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "only"
    context.prs = prs


@when("I call paragraph.delete on that paragraph")
def when_call_paragraph_delete_on_that_paragraph(context):
    tf = context.prs.slides[0].shapes[0].text_frame
    tf.paragraphs[0].delete()


@then("the reloaded textbox has 1 paragraph")
def then_reloaded_textbox_has_1_paragraph(context):
    tf = context.prs.slides[0].shapes[0].text_frame
    assert len(tf.paragraphs) == 1, len(tf.paragraphs)


@then("the reloaded first paragraph has no runs")
def then_reloaded_first_paragraph_has_no_runs(context):
    tf = context.prs.slides[0].shapes[0].text_frame
    assert len(tf.paragraphs[0].runs) == 0, len(tf.paragraphs[0].runs)


@given('a fresh presentation with a paragraph of three runs "A" "B" "C"')
def given_fresh_prs_with_three_runs_ABC(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    p = tb.text_frame.paragraphs[0]
    for letter in ("A", "B", "C"):
        p.add_run().text = letter
    context.prs = prs


@when("I call run.delete on the middle run")
def when_run_delete_middle_run(context):
    p = context.prs.slides[0].shapes[0].text_frame.paragraphs[0]
    p.runs[1].delete()


@then("the reloaded paragraph has 2 runs")
def then_reloaded_paragraph_has_2_runs(context):
    p = context.prs.slides[0].shapes[0].text_frame.paragraphs[0]
    assert len(p.runs) == 2, len(p.runs)


@then('the reloaded run texts are "A" and "C"')
def then_reloaded_run_texts_A_C(context):
    p = context.prs.slides[0].shapes[0].text_frame.paragraphs[0]
    texts = [r.text for r in p.runs]
    assert texts == ["A", "C"], texts


@given("a fresh presentation with an empty paragraph")
def given_fresh_prs_empty_paragraph(context):
    prs, slide = _build_prs_with_blank_slide()
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    context.prs = prs


@when("I call paragraph.write_rich with plain bold and sized parts")
def when_paragraph_write_rich_parts(context):
    p = context.prs.slides[0].shapes[0].text_frame.paragraphs[0]
    p.write_rich("plain ", ("bold", {"bold": True}), (" big", {"size": Pt(24)}))


@then("the reloaded paragraph has 3 runs with the expected formatting")
def then_reloaded_paragraph_write_rich(context):
    p = context.prs.slides[0].shapes[0].text_frame.paragraphs[0]
    runs = p.runs
    assert len(runs) == 3, "expected 3 runs, got %d" % len(runs)
    assert runs[0].text == "plain ", repr(runs[0].text)
    assert runs[1].text == "bold", repr(runs[1].text)
    assert runs[1].font.bold is True, runs[1].font.bold
    assert runs[2].text == " big", repr(runs[2].text)
    assert runs[2].font.size == Pt(24), runs[2].font.size


# -- Font round-trips -----------------------------------------------


def _first_run(context):
    tb = context.prs.slides[0].shapes[0]
    return tb.text_frame.paragraphs[0].runs[0]


@given("a fresh presentation with a run whose font.baseline is 30000")
def given_fresh_prs_font_baseline(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "x"
    tb.text_frame.paragraphs[0].runs[0].font.baseline = 30000
    context.prs = prs


@then("the reloaded font.baseline is 30000")
def then_reloaded_font_baseline_30000(context):
    run = _first_run(context)
    assert run.font.baseline == 30000, run.font.baseline


@then("the reloaded font.superscript is True")
def then_reloaded_font_superscript_True(context):
    run = _first_run(context)
    assert run.font.superscript is True, run.font.superscript


@then("the reloaded font.subscript is False")
def then_reloaded_font_subscript_False(context):
    run = _first_run(context)
    assert run.font.subscript is False, run.font.subscript


@given("a fresh presentation with a run whose font.subscript is True")
def given_fresh_prs_font_subscript(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "x"
    tb.text_frame.paragraphs[0].runs[0].font.subscript = True
    context.prs = prs


@then("the reloaded font.subscript is True")
def then_reloaded_font_subscript_True(context):
    run = _first_run(context)
    assert run.font.subscript is True, run.font.subscript


@then("the reloaded font.baseline is -25000")
def then_reloaded_font_baseline_minus_25000(context):
    run = _first_run(context)
    assert run.font.baseline == -25000, run.font.baseline


@given("a fresh presentation with a run carrying a text shadow")
def given_fresh_prs_run_text_shadow(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "shadow"
    run = tb.text_frame.paragraphs[0].runs[0]
    run.font.shadow.blur_radius = 50800
    run.font.shadow.distance = 38100
    run.font.shadow.direction = 45.0
    context.prs = prs


@then("the reloaded font.shadow.blur_radius is 50800")
def then_reloaded_font_shadow_blur_radius(context):
    run = _first_run(context)
    assert run.font.shadow.blur_radius == 50800, run.font.shadow.blur_radius


@then("the reloaded font.shadow.distance is 38100")
def then_reloaded_font_shadow_distance(context):
    run = _first_run(context)
    assert run.font.shadow.distance == 38100, run.font.shadow.distance


@then("the reloaded font.shadow.direction is 45.0")
def then_reloaded_font_shadow_direction(context):
    run = _first_run(context)
    assert run.font.shadow.direction == 45.0, run.font.shadow.direction


@given("a fresh presentation with a run carrying a glow effect")
def given_fresh_prs_run_glow(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "glow"
    run = tb.text_frame.paragraphs[0].runs[0]
    run.font.effect_format.glow.size = 63500
    context.prs = prs


@then("the reloaded font.effect_format.glow.size is 63500")
def then_reloaded_font_effect_glow_size(context):
    run = _first_run(context)
    assert run.font.effect_format.glow.size == 63500, run.font.effect_format.glow.size


@given("a fresh presentation with a run whose font.strikethrough is SINGLE_LINE")
def given_fresh_prs_font_strike(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "strike"
    run = tb.text_frame.paragraphs[0].runs[0]
    run.font.strikethrough = MSO_STRIKE.SINGLE_LINE
    context.prs = prs


@then("the reloaded font.strikethrough is MSO_STRIKE.SINGLE_LINE")
def then_reloaded_font_strikethrough_single(context):
    run = _first_run(context)
    assert run.font.strikethrough == MSO_STRIKE.SINGLE_LINE, run.font.strikethrough


@given("a fresh presentation with a run whose Latin EA and CS fonts are set")
def given_fresh_prs_font_latin_ea_cs(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "x"
    run = tb.text_frame.paragraphs[0].runs[0]
    run.font.name = "Calibri"
    run.font.name_ea = "MS Gothic"
    run.font.name_cs = "Arial"
    context.prs = prs


@then('the reloaded font.name is "Calibri"')
def then_reloaded_font_name_calibri(context):
    run = _first_run(context)
    assert run.font.name == "Calibri", run.font.name


@then('the reloaded font.name_ea is "MS Gothic"')
def then_reloaded_font_name_ea(context):
    run = _first_run(context)
    assert run.font.name_ea == "MS Gothic", run.font.name_ea


@then('the reloaded font.name_cs is "Arial"')
def then_reloaded_font_name_cs(context):
    run = _first_run(context)
    assert run.font.name_cs == "Arial", run.font.name_cs


# -- Slide-level round-trips ---------------------------------------


@given("a fresh presentation with a hidden slide")
def given_fresh_prs_with_hidden_slide(context):
    prs, slide = _build_prs_with_blank_slide()
    slide.is_hidden = True
    context.prs = prs


@then("the reloaded slide.is_hidden is True")
def then_reloaded_slide_is_hidden_true(context):
    assert context.prs.slides[0].is_hidden is True


@given("a fresh presentation with show_master_shapes False on slide 0")
def given_fresh_prs_show_master_shapes_False(context):
    prs, slide = _build_prs_with_blank_slide()
    slide.show_master_shapes = False
    context.prs = prs


@then("the reloaded slide.show_master_shapes is False")
def then_reloaded_slide_show_master_shapes_false(context):
    assert context.prs.slides[0].show_master_shapes is False


@given('a fresh presentation with slide 0 named "Title slide"')
def given_fresh_prs_slide_named_Title_slide(context):
    prs, slide = _build_prs_with_blank_slide()
    slide.name = "Title slide"
    context.prs = prs


@then('the reloaded slide.name is "Title slide"')
def then_reloaded_slide_name_Title_slide(context):
    assert context.prs.slides[0].name == "Title slide", context.prs.slides[0].name


# -- Chart round-trips ---------------------------------------------


def _build_simple_bar_chart(slide):
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
    chart_data.add_series("Sales", (10, 20, 30, 40))
    chart_data.add_series("Cost", (5, 15, 25, 35))
    return slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    ).chart


@given("a fresh presentation with a chart whose display_blanks_as is GAPS")
def given_fresh_prs_chart_display_blanks_GAPS(context):
    prs, slide = _build_prs_with_blank_slide()
    chart = _build_simple_bar_chart(slide)
    chart.display_blanks_as = XL_DISPLAY_BLANKS_AS.GAPS
    context.prs = prs


@then("the reloaded chart.display_blanks_as is XL_DISPLAY_BLANKS_AS.GAPS")
def then_reloaded_chart_display_blanks_GAPS(context):
    chart = context.prs.slides[0].shapes[0].chart
    assert chart.display_blanks_as == XL_DISPLAY_BLANKS_AS.GAPS, chart.display_blanks_as


@given("a fresh presentation with a chart whose tick_label_skip is 3")
def given_fresh_prs_chart_tick_label_skip(context):
    prs, slide = _build_prs_with_blank_slide()
    chart = _build_simple_bar_chart(slide)
    # -- in a bar chart the category axis is the perpendicular axis --
    chart.category_axis.tick_label_skip = 3
    chart.category_axis.tick_mark_skip = 2
    context.prs = prs


@then("the reloaded category_axis.tick_label_skip == 3")
def then_reloaded_category_axis_tick_label_skip_3(context):
    chart = context.prs.slides[0].shapes[0].chart
    assert chart.category_axis.tick_label_skip == 3, chart.category_axis.tick_label_skip


@then("the reloaded category_axis.tick_mark_skip == 2")
def then_reloaded_category_axis_tick_mark_skip_2(context):
    chart = context.prs.slides[0].shapes[0].chart
    assert chart.category_axis.tick_mark_skip == 2, chart.category_axis.tick_mark_skip


@given("a fresh presentation with a chart whose tick-label rotation is -45")
def given_fresh_prs_chart_tick_labels_rotation(context):
    prs, slide = _build_prs_with_blank_slide()
    chart = _build_simple_bar_chart(slide)
    chart.value_axis.tick_labels.rotation = -45
    context.prs = prs


@then("the reloaded value_axis.tick_labels.rotation == 315.0")
def then_reloaded_value_axis_tick_labels_rotation(context):
    chart = context.prs.slides[0].shapes[0].chart
    # -- rotation normalizes to [0, 360); -45 round-trips as 315 --
    assert chart.value_axis.tick_labels.rotation == 315.0, chart.value_axis.tick_labels.rotation


@given("a fresh presentation with a chart whose first series has error bars")
def given_fresh_prs_chart_with_error_bars(context):
    prs, slide = _build_prs_with_blank_slide()
    chart = _build_simple_bar_chart(slide)
    series = chart.plots[0].series[0]
    series.set_error_bars(XL_ERROR_BAR_TYPE.FIXED_VALUE, 1.0, XL_ERROR_BAR_INCLUDE.BOTH)
    context.prs = prs


@then("the reloaded first series has error bars")
def then_reloaded_first_series_has_error_bars(context):
    chart = context.prs.slides[0].shapes[0].chart
    series = chart.plots[0].series[0]
    assert series.has_error_bars is True, series.has_error_bars


# -- Table cell/style round-trips ----------------------------------


def _build_simple_table(slide, rows=2, cols=2):
    graphic_frame = slide.shapes.add_table(rows, cols, Inches(1), Inches(1), Inches(6), Inches(2))
    return graphic_frame.table


@given("a fresh presentation with a cell border_left red 1.5pt dashed")
def given_fresh_prs_cell_border_left(context):
    prs, slide = _build_prs_with_blank_slide()
    table = _build_simple_table(slide)
    cell = table.cell(0, 0)
    cell.border_left.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    cell.border_left.width = Pt(1.5)
    cell.border_left.dash_style = MSO_LINE.DASH
    context.prs = prs


@then("the reloaded cell.border_left.color.rgb is FF0000")
def then_reloaded_cell_border_left_color(context):
    table = context.prs.slides[0].shapes[0].table
    cell = table.cell(0, 0)
    assert cell.border_left.color.rgb == RGBColor(0xFF, 0x00, 0x00), cell.border_left.color.rgb


@then("the reloaded cell.border_left.width.pt == 1.5")
def then_reloaded_cell_border_left_width(context):
    table = context.prs.slides[0].shapes[0].table
    cell = table.cell(0, 0)
    assert cell.border_left.width.pt == 1.5, cell.border_left.width.pt


@given("a fresh presentation with a table whose style_id is set")
def given_fresh_prs_table_style_id(context):
    prs, slide = _build_prs_with_blank_slide()
    table = _build_simple_table(slide)
    table.style_id = "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"
    context.prs = prs


@then("the reloaded table.style_id matches the expected GUID")
def then_reloaded_table_style_id(context):
    table = context.prs.slides[0].shapes[0].table
    assert table.style_id == "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}", table.style_id


@given("a fresh presentation with a 2x2 table")
def given_fresh_prs_2x2_table(context):
    prs, slide = _build_prs_with_blank_slide()
    _build_simple_table(slide, rows=2, cols=2)
    context.prs = prs


@when("I call table.add_row and table.add_column")
def when_table_add_row_and_add_column(context):
    table = context.prs.slides[0].shapes[0].table
    table.add_row()
    table.add_column()


@then("the reloaded table is 3 rows by 3 columns")
def then_reloaded_table_3x3(context):
    table = context.prs.slides[0].shapes[0].table
    assert len(table.rows) == 3, "rows=%d" % len(table.rows)
    assert len(table.columns) == 3, "cols=%d" % len(table.columns)


# -- ActionSetting / hover round-trips ---------------------------


@given("a fresh presentation with a shape carrying a hover screen_tip")
def given_fresh_prs_shape_hover_screen_tip(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    # -- hover needs an actionable hyperlink or target to survive meaningfully,
    # -- but the screen_tip itself still round-trips. Give it an address
    # -- so PowerPoint would actually render the tooltip.
    shape.hover_action.hyperlink.address = "https://example.com/hover"
    shape.hover_action.screen_tip = "Hover me"
    context.prs = prs


@then('the reloaded hover_action.screen_tip is "Hover me"')
def then_reloaded_hover_screen_tip(context):
    shape = context.prs.slides[0].shapes[0]
    assert shape.hover_action.screen_tip == "Hover me", shape.hover_action.screen_tip


@given("a fresh presentation with a shape carrying a hover hyperlink")
def given_fresh_prs_shape_hover_hyperlink(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    shape.hover_action.hyperlink.address = "https://example.com/hover"
    context.prs = prs


@then('the reloaded hover_action.hyperlink.address is "https://example.com/hover"')
def then_reloaded_hover_hyperlink(context):
    shape = context.prs.slides[0].shapes[0]
    assert (
        shape.hover_action.hyperlink.address == "https://example.com/hover"
    ), shape.hover_action.hyperlink.address


# -- Custom document-part props on shapes ------------------------


@given("a fresh presentation with a shape carrying three custom_props")
def given_fresh_prs_shape_three_custom_props(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    shape.custom_props["k1"] = "v1"
    shape.custom_props["k2"] = "v2"
    shape.custom_props["k3"] = "v3"
    context.prs = prs


@then("the reloaded shape.custom_props has 3 entries")
def then_reloaded_custom_props_has_3_entries(context):
    shape = context.prs.slides[0].shapes[0]
    assert len(shape.custom_props) == 3, len(shape.custom_props)


@then("the reloaded shape.custom_props gets k1 v1 k2 v2 k3 v3")
def then_reloaded_custom_props_values(context):
    shape = context.prs.slides[0].shapes[0]
    assert shape.custom_props["k1"] == "v1", shape.custom_props["k1"]
    assert shape.custom_props["k2"] == "v2", shape.custom_props["k2"]
    assert shape.custom_props["k3"] == "v3", shape.custom_props["k3"]


# -- Color / fill round-trips ------------------------------------


@given("a fresh presentation with a solid-fill shape whose alpha is 0.5")
def given_fresh_prs_shape_alpha_half(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xFF, 0x00, 0x00)
    shape.fill.fore_color.alpha = 0.5
    context.prs = prs


@then("the reloaded fill.fore_color.alpha == 0.5")
def then_reloaded_fill_alpha_half(context):
    shape = context.prs.slides[0].shapes[0]
    alpha = shape.fill.fore_color.alpha
    assert abs(alpha - 0.5) < 1e-6, alpha


# -- Picture / connector additions ---------------------------------


@given("a fresh presentation with a picture_link to an external URL")
def given_fresh_prs_picture_link(context):
    prs, slide = _build_prs_with_blank_slide()
    slide.shapes.add_picture_link(
        "https://example.com/logo.png", Inches(1), Inches(1), Inches(2), Inches(2)
    )
    context.prs = prs


@then("the reloaded picture has an external rel to the URL")
def then_reloaded_picture_has_external_rel(context):
    slide = context.prs.slides[0]
    pic = slide.shapes[0]
    # -- the picture should be present and reference an external rel --
    assert pic.shape_type is not None
    # -- check for external rel on the slide part matching the URL --
    rels = slide.part.rels
    found = False
    for rId, rel in rels.items():
        if rel.is_external and "example.com/logo.png" in rel.target_ref:
            found = True
            break
    assert found, "expected external relationship to example.com/logo.png"


# -- Trendline round-trips ----------------------------------------


@given("a fresh presentation with a series carrying a LINEAR trendline")
def given_fresh_prs_series_trendline(context):
    from pptx.enum.chart import XL_TRENDLINE_TYPE

    prs, slide = _build_prs_with_blank_slide()
    chart = _build_simple_bar_chart(slide)
    series = chart.plots[0].series[0]
    series.add_trendline(XL_TRENDLINE_TYPE.LINEAR, display_equation=True)
    context.prs = prs


@when("I call trendline.delete on the first trendline")
def when_trendline_delete(context):
    chart = context.prs.slides[0].shapes[0].chart
    chart.plots[0].series[0].trendlines[0].delete()


@then("the reloaded series has 1 trendline")
def then_reloaded_series_1_trendline(context):
    chart = context.prs.slides[0].shapes[0].chart
    trendlines = chart.plots[0].series[0].trendlines
    assert len(trendlines) == 1, len(trendlines)


@then("the reloaded series has 0 trendlines")
def then_reloaded_series_0_trendlines(context):
    chart = context.prs.slides[0].shapes[0].chart
    trendlines = chart.plots[0].series[0].trendlines
    assert len(trendlines) == 0, len(trendlines)


@then("the reloaded trendline type is LINEAR")
def then_reloaded_trendline_type_LINEAR(context):
    from pptx.enum.chart import XL_TRENDLINE_TYPE

    chart = context.prs.slides[0].shapes[0].chart
    tl = chart.plots[0].series[0].trendlines[0]
    assert tl.trendline_type == XL_TRENDLINE_TYPE.LINEAR, tl.trendline_type


@then("the reloaded trendline displays the equation")
def then_reloaded_trendline_displays_equation(context):
    chart = context.prs.slides[0].shapes[0].chart
    tl = chart.plots[0].series[0].trendlines[0]
    assert tl.display_equation is True, tl.display_equation


# -- Chart.replace_data_preserve_formulas --------------------------


@given("a fresh presentation with a chart and a tweaked embedded workbook")
def given_fresh_prs_chart_for_replace_preserve(context):
    prs, slide = _build_prs_with_blank_slide()
    chart = _build_simple_bar_chart(slide)
    context.prs = prs
    context.initial_categories = list(chart.plots[0].categories)


@when("I call chart.replace_data_preserve_formulas with new category data")
def when_replace_data_preserve_formulas(context):
    chart = context.prs.slides[0].shapes[0].chart
    cd = CategoryChartData()
    cd.categories = ["A1", "A2", "A3", "A4"]
    cd.add_series("Sales", (100, 200, 300, 400))
    cd.add_series("Cost", (50, 60, 70, 80))
    chart.replace_data_preserve_formulas(cd)


@then("the reloaded chart has the replacement category values")
def then_reloaded_chart_has_replacement_values(context):
    chart = context.prs.slides[0].shapes[0].chart
    cats = list(chart.plots[0].categories)
    assert cats == ["A1", "A2", "A3", "A4"], cats


# -- Chart secondary value axis --------------------------------


@given("a fresh presentation with a combo chart carrying a LINE plot")
def given_fresh_prs_combo_chart_line_plot(context):
    prs, slide = _build_prs_with_blank_slide()
    chart = _build_simple_bar_chart(slide)
    line_data = CategoryChartData()
    line_data.categories = ["Q1", "Q2", "Q3", "Q4"]
    line_data.add_series("Trend", (15, 25, 35, 45))
    chart.add_plot(XL_CHART_TYPE.LINE, line_data)
    context.prs = prs


@then("the reloaded chart has 2 plots")
def then_reloaded_chart_has_2_plots(context):
    chart = context.prs.slides[0].shapes[0].chart
    assert len(chart.plots) == 2, len(chart.plots)


# -- LineEndFormat round-trip -------------------------------------


@given("a fresh presentation with a connector whose end_arrow.type is TRIANGLE")
def given_fresh_prs_connector_arrow(context):
    from pptx.enum.dml import MSO_LINE_END_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR

    prs, slide = _build_prs_with_blank_slide()
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Emu(0), Emu(0), Emu(914400), Emu(914400)
    )
    conn.line.end_arrow.type = MSO_LINE_END_TYPE.TRIANGLE
    context.prs = prs


@then("the reloaded connector end_arrow.type is MSO_LINE_END_TYPE.TRIANGLE")
def then_reloaded_connector_end_arrow(context):
    from pptx.enum.dml import MSO_LINE_END_TYPE

    conn = context.prs.slides[0].shapes[0]
    assert conn.line.end_arrow.type == MSO_LINE_END_TYPE.TRIANGLE, conn.line.end_arrow.type


# -- Paragraph bullets round-trip ---------------------------------


@given("a fresh presentation with a paragraph having a bullet character")
def given_fresh_prs_paragraph_bullet_char(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "bullet line"
    p = tb.text_frame.paragraphs[0]
    p.bullet.character("•")
    context.prs = prs


@then('the reloaded paragraph bullet character is "•"')
def then_reloaded_paragraph_bullet_character(context):
    p = context.prs.slides[0].shapes[0].text_frame.paragraphs[0]
    # -- bullet proxy exposes __repr__/.character via element access; read the
    # -- underlying buChar @char attribute for robustness --
    from pptx.oxml.ns import qn

    pPr = p._pPr  # type: ignore[attr-defined]
    buChar = pPr.find(qn("a:buChar"))
    assert buChar is not None, "expected an a:buChar element"
    assert buChar.get("char") == "•", buChar.get("char")


@given("a fresh presentation with an auto-numbered paragraph")
def given_fresh_prs_paragraph_auto_number(context):
    from pptx.enum.text import PP_AUTO_NUMBER

    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "item"
    p = tb.text_frame.paragraphs[0]
    p.bullet.auto_number(PP_AUTO_NUMBER.ARABIC_PERIOD)
    context.prs = prs


@then("the reloaded paragraph uses auto-numbering")
def then_reloaded_paragraph_auto_number(context):
    from pptx.oxml.ns import qn

    p = context.prs.slides[0].shapes[0].text_frame.paragraphs[0]
    pPr = p._pPr  # type: ignore[attr-defined]
    buAutoNum = pPr.find(qn("a:buAutoNum"))
    assert buAutoNum is not None, "expected an a:buAutoNum element"


# -- Sections round-trip ------------------------------------------


@given("a fresh presentation with two sections covering three slides")
def given_fresh_prs_two_sections(context):
    prs = Presentation()
    layout = prs.slide_layouts[6]
    prs.slides.add_slide(layout)
    prs.slides.add_slide(layout)
    prs.slides.add_slide(layout)
    prs.sections.add_section("Intro", [prs.slides[0]])
    prs.sections.add_section("Main", [prs.slides[1], prs.slides[2]])
    context.prs = prs


@then("the reloaded presentation has 2 sections")
def then_reloaded_prs_has_2_sections(context):
    assert len(context.prs.sections) == 2, len(context.prs.sections)


@then('the reloaded section names are "Intro" and "Main"')
def then_reloaded_section_names(context):
    names = [s.name for s in context.prs.sections]
    assert names == ["Intro", "Main"], names


# -- Slide.tags round-trip ---------------------------------------


@given("a fresh presentation with a slide carrying two custom tags")
def given_fresh_prs_slide_two_tags(context):
    prs, slide = _build_prs_with_blank_slide()
    slide.tags["project"] = "atlas"
    slide.tags["status"] = "draft"
    context.prs = prs


@then("the reloaded slide has 2 tags")
def then_reloaded_slide_has_2_tags(context):
    tags = context.prs.slides[0].tags
    assert len(tags) == 2, len(tags)


@then('the reloaded slide tag "project" is "atlas"')
def then_reloaded_slide_tag_project(context):
    tags = context.prs.slides[0].tags
    assert tags["project"] == "atlas", tags["project"]


@then('the reloaded slide tag "status" is "draft"')
def then_reloaded_slide_tag_status(context):
    tags = context.prs.slides[0].tags
    assert tags["status"] == "draft", tags["status"]


# -- slide_layout setter round-trip ------------------------------


@given("a fresh presentation with a slide on layout 1 carrying a textbox")
def given_fresh_prs_slide_layout_1_textbox(context):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1)).text_frame.text = "marker"
    context.prs = prs


@when("I re-point the slide to layout 5")
def when_repoint_slide_to_layout_5(context):
    prs = context.prs
    slide = prs.slides[0]
    slide.slide_layout = prs.slide_layouts[5]


@then("the reloaded slide's layout index is 5")
def then_reloaded_slide_layout_index_5(context):
    prs = context.prs
    slide = prs.slides[0]
    layouts = list(prs.slide_layouts)
    idx = layouts.index(slide.slide_layout)
    assert idx == 5, idx


@then("the reloaded slide still has the textbox")
def then_reloaded_slide_still_has_textbox(context):
    slide = context.prs.slides[0]
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert any(t == "marker" for t in texts), texts


# -- LineFormat dash_style round-trip ----------------------------


@given("a fresh presentation with a shape having a dashed line")
def given_fresh_prs_shape_dashed_line(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    shape.line.dash_style = MSO_LINE.DASH
    context.prs = prs


@then("the reloaded shape.line.dash_style is MSO_LINE.DASH")
def then_reloaded_shape_line_dash(context):
    shape = context.prs.slides[0].shapes[0]
    assert shape.line.dash_style == MSO_LINE.DASH, shape.line.dash_style


# -- Notes slide round-trip --------------------------------------


@given('a fresh presentation with notes "Speaker reminder" on slide 0')
def given_fresh_prs_notes(context):
    prs, slide = _build_prs_with_blank_slide()
    slide.notes_slide.notes_text_frame.text = "Speaker reminder"
    context.prs = prs


@then('the reloaded slide.notes_slide.notes_text_frame.text is "Speaker reminder"')
def then_reloaded_notes_text(context):
    slide = context.prs.slides[0]
    assert slide.has_notes_slide is True, slide.has_notes_slide
    text = slide.notes_slide.notes_text_frame.text
    assert text == "Speaker reminder", repr(text)


# -- Font highlight_color round-trip ------------------------------


@given("a fresh presentation with a run carrying a yellow highlight")
def given_fresh_prs_highlight(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "lit"
    run = tb.text_frame.paragraphs[0].runs[0]
    run.font.highlight_color.rgb = RGBColor(0xFF, 0xFF, 0x00)
    context.prs = prs


@then("the reloaded font.highlight_color.rgb is FFFF00")
def then_reloaded_highlight_yellow(context):
    run = _first_run(context)
    assert run.font.highlight_color.rgb == RGBColor(0xFF, 0xFF, 0x00), run.font.highlight_color.rgb


# -- Picture crop round-trip --------------------------------------


@given("a fresh presentation with a cropped picture")
def given_fresh_prs_cropped_picture(context):
    prs, slide = _build_prs_with_blank_slide()
    pic = slide.shapes.add_picture(
        test_image("python-icon.jpeg"),
        Inches(1),
        Inches(1),
        Inches(2),
        Inches(2),
    )
    pic.crop_left = 0.25
    pic.crop_right = 0.25
    context.prs = prs


@then("the reloaded picture crop_left is 0.25")
def then_reloaded_picture_crop_left(context):
    pic = context.prs.slides[0].shapes[0]
    assert abs(pic.crop_left - 0.25) < 1e-6, pic.crop_left


@then("the reloaded picture crop_right is 0.25")
def then_reloaded_picture_crop_right(context):
    pic = context.prs.slides[0].shapes[0]
    assert abs(pic.crop_right - 0.25) < 1e-6, pic.crop_right


# -- Table merge round-trip -------------------------------------


@given("a fresh presentation with a 3x3 table merging the top row")
def given_fresh_prs_table_merged_row(context):
    prs, slide = _build_prs_with_blank_slide()
    table = _build_simple_table(slide, rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2))
    context.prs = prs


@then("the reloaded cell(0, 0) is_merge_origin is True")
def then_reloaded_cell00_merge_origin(context):
    table = context.prs.slides[0].shapes[0].table
    cell = table.cell(0, 0)
    assert cell.is_merge_origin is True, cell.is_merge_origin


@then("the reloaded cell(0, 0) span_width is 3")
def then_reloaded_cell00_span_width_3(context):
    table = context.prs.slides[0].shapes[0].table
    cell = table.cell(0, 0)
    assert cell.span_width == 3, cell.span_width


# -- Slides.duplicate round-trip -------------------------------


@given('a fresh presentation with a slide carrying a textbox "original"')
def given_fresh_prs_textbox_original(context):
    prs, slide = _build_prs_with_blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "original"
    context.prs = prs


@when("I call slides.duplicate on the first slide")
def when_slides_duplicate_first(context):
    context.prs.slides.duplicate(context.prs.slides[0])


@then("the reloaded presentation has 2 slides")
def then_reloaded_prs_2_slides(context):
    assert len(context.prs.slides) == 2, len(context.prs.slides)


@then('both reloaded slides have a textbox with text "original"')
def then_reloaded_both_slides_textbox_original(context):
    for s in context.prs.slides:
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        assert any(t == "original" for t in texts), (s, texts)


# -- Animation round-trip ----------------------------------------


@given("a fresh presentation with a shape carrying an APPEAR entrance animation")
def given_fresh_prs_shape_animation(context):
    from pptx.enum.shapes import MSO_SHAPE

    prs, slide = _build_prs_with_blank_slide()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    # -- APPEAR is the simplest entrance effect --
    from pptx.enum.animation import MSO_ANIMATION_TYPE

    shape.set_animation(MSO_ANIMATION_TYPE.APPEAR)
    context.prs = prs


@then("the reloaded slide.has_animations is True")
def then_reloaded_slide_has_animations(context):
    slide = context.prs.slides[0]
    assert slide.has_animations is True, slide.has_animations
