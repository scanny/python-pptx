"""Gherkin step implementations for presentation-section features."""

from __future__ import annotations

from behave import given, then, when
from behave.runner import Context
from helpers import test_pptx

from pptx import Presentation
from pptx.oxml.ns import qn

# given ===================================================


@given("a presentation with three slides")
def given_a_presentation_with_three_slides(context: Context):
    context.prs = Presentation(test_pptx("sld-slides"))


# when ====================================================


@when('I add a section named "{name}" with slides 1 and 2')
def when_add_section_with_two_slides(context: Context, name: str):
    slides = [context.prs.slides[0], context.prs.slides[1]]
    context.prs.sections.add_section(name, slides=slides)


@when('I add a section named "{name}" with slide 3')
def when_add_section_with_slide_three(context: Context, name: str):
    context.prs.sections.add_section(name, slides=[context.prs.slides[2]])


@when('I add three sections "{a}", "{b}", "{c}"')
def when_add_three_sections(context: Context, a: str, b: str, c: str):
    context.prs.sections.add_section(a)
    context.prs.sections.add_section(b)
    context.prs.sections.add_section(c)


@when('I move the "{mover}" section before the "{anchor}" section')
def when_move_section_before(context: Context, mover: str, anchor: str):
    section = context.prs.sections.get_by_name(mover)
    other = context.prs.sections.get_by_name(anchor)
    assert section is not None and other is not None
    section.move_before(other)


@when('I move the "{mover}" section after the "{anchor}" section')
def when_move_section_after(context: Context, mover: str, anchor: str):
    section = context.prs.sections.get_by_name(mover)
    other = context.prs.sections.get_by_name(anchor)
    assert section is not None and other is not None
    section.move_after(other)


@when('I move slide {n:d} into the "{target_name}" section')
def when_move_slide_into_section(context: Context, n: int, target_name: str):
    target = context.prs.sections.get_by_name(target_name)
    assert target is not None
    target.move_slide(context.prs.slides[n - 1])


@when('I rename the first section to "{new_name}"')
def when_rename_first_section(context: Context, new_name: str):
    context.prs.sections[0].name = new_name


@when("I remove the first section")
def when_remove_first_section(context: Context):
    section = context.prs.sections[0]
    context.prs.sections.remove(section)


# then ====================================================


@then("Presentation.sections has {count:d} section")
@then("Presentation.sections has {count:d} sections")
def then_sections_has_count(context: Context, count: int):
    actual = len(context.prs.sections)
    assert actual == count, "expected %d sections, got %d" % (count, actual)


@then('the section is named "{name}"')
def then_section_is_named(context: Context, name: str):
    actual = context.prs.sections[0].name
    assert actual == name, "expected %r, got %r" % (name, actual)


@then("the section has two assigned slides")
def then_section_has_two_assigned_slides(context: Context):
    section = context.prs.sections[0]
    slides = section.slides
    assert len(slides) == 2, "expected 2 assigned slides, got %d" % len(slides)
    # -- they should be the first two slides of the deck --
    assert slides[0] == context.prs.slides[0]
    assert slides[1] == context.prs.slides[1]


@then("the presentation.xml has a p14:sectionLst under p:extLst")
def then_prs_xml_has_sectionLst(context: Context):
    prs_element = context.prs._element  # pyright: ignore[reportPrivateUsage]
    extLst = prs_element.find(qn("p:extLst"))
    assert extLst is not None, "no p:extLst present on presentation"
    sectionLst = extLst.find("./{0}/{1}".format(qn("p:ext"), qn("p14:sectionLst")))
    assert sectionLst is not None, "no p14:sectionLst under p:extLst/p:ext"


@then("the presentation has no p:extLst element")
def then_prs_has_no_extLst(context: Context):
    prs_element = context.prs._element  # pyright: ignore[reportPrivateUsage]
    assert prs_element.find(qn("p:extLst")) is None


@then('the section names are "{a}", "{b}", "{c}"')
def then_section_names_are(context: Context, a: str, b: str, c: str):
    actual = [s.name for s in context.prs.sections]
    assert actual == [a, b, c], "expected %r, got %r" % ([a, b, c], actual)


@then('find_containing for slide {n:d} returns the "{expected_name}" section')
def then_find_containing(context: Context, n: int, expected_name: str):
    slide = context.prs.slides[n - 1]
    section = context.prs.sections.find_containing(slide)
    assert section is not None, "find_containing returned None"
    assert section.name == expected_name, "expected %r, got %r" % (
        expected_name,
        section.name,
    )


@then('adding slide {n:d} to the "{target_name}" section raises a ValueError mentioning "{phrase}"')
def then_add_slide_raises(context: Context, n: int, target_name: str, phrase: str):
    target = context.prs.sections.get_by_name(target_name)
    assert target is not None
    try:
        target.add_slide(context.prs.slides[n - 1])
    except ValueError as exc:
        assert phrase in str(exc), "expected %r in %r" % (phrase, str(exc))
        return
    raise AssertionError("add_slide did not raise ValueError")


@then('the "{section_name}" section contains slides {a:d}')
def then_section_contains_one_slide(context: Context, section_name: str, a: int):
    section = context.prs.sections.get_by_name(section_name)
    assert section is not None
    expected = (context.prs.slides[a - 1],)
    assert section.slides == expected, "expected %r, got %r" % (expected, section.slides)


@then('the "{section_name}" section contains slides {a:d} and {b:d}')
def then_section_contains_two_slides(context: Context, section_name: str, a: int, b: int):
    section = context.prs.sections.get_by_name(section_name)
    assert section is not None
    expected = (context.prs.slides[a - 1], context.prs.slides[b - 1])
    assert section.slides == expected, "expected %r, got %r" % (expected, section.slides)
