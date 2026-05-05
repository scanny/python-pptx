"""Gherkin step implementations for `SlideShapes.add_picture_from` (CLO-12)."""

from __future__ import annotations

from behave import then, when

from pptx.util import Inches

# when ====================================================


@when("I add the picture_from the source onto the second slide")
def when_add_picture_from(context):
    context.clone = context.tgt_slide.shapes.add_picture_from(context.src_shape)


@when("I add_picture_from the source at (3in, 4in)")
def when_add_picture_from_at_pos(context):
    context.clone = context.tgt_slide.shapes.add_picture_from(
        context.src_shape, left=Inches(3), top=Inches(4)
    )


@when("I add_picture_from the source with size (2in, 1in)")
def when_add_picture_from_with_size(context):
    context.clone = context.tgt_slide.shapes.add_picture_from(
        context.src_shape, width=Inches(2), height=Inches(1)
    )


@when("I add_picture_from the source onto the target presentation's slide")
def when_add_picture_from_cross_prs(context):
    context.clone = context.tgt_slide.shapes.add_picture_from(context.src_shape)


# then ====================================================


@then("the clone's size is (2in, 1in)")
def then_clone_size(context):
    assert context.clone.width == Inches(2), "width is %r" % context.clone.width
    assert context.clone.height == Inches(1), "height is %r" % context.clone.height
