"""Gherkin step implementations for `BaseShape.clone_onto`."""

from __future__ import annotations

from behave import given, then, when
from helpers import test_image

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR_TYPE, MSO_SHAPE
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture
from pptx.util import Inches

# given ===================================================


@given("a slide with an auto-shape")
def given_a_slide_with_an_autoshape(context):
    prs = Presentation()
    context.prs = prs
    context.src_slide = prs.slides.add_slide(prs.slide_layouts[5])
    rect = context.src_slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
    )
    rect.text_frame.text = "Greetings"
    context.src_shape = rect


@given("a slide with a picture")
def given_a_slide_with_a_picture(context):
    prs = Presentation()
    context.prs = prs
    context.src_slide = prs.slides.add_slide(prs.slide_layouts[5])
    context.src_shape = context.src_slide.shapes.add_picture(
        test_image("python-icon.jpeg"), Inches(1), Inches(1)
    )


@given("a slide with a connector")
def given_a_slide_with_a_connector(context):
    prs = Presentation()
    context.prs = prs
    context.src_slide = prs.slides.add_slide(prs.slide_layouts[5])
    context.src_shape = context.src_slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT, Inches(1), Inches(1), Inches(3), Inches(1)
    )


@given("a slide with a table")
def given_a_slide_with_a_table(context):
    prs = Presentation()
    context.prs = prs
    context.src_slide = prs.slides.add_slide(prs.slide_layouts[5])
    context.src_shape = context.src_slide.shapes.add_table(
        2, 2, Inches(1), Inches(1), Inches(3), Inches(1)
    )


@given("a slide with a group shape")
def given_a_slide_with_a_group(context):
    prs = Presentation()
    context.prs = prs
    context.src_slide = prs.slides.add_slide(prs.slide_layouts[5])
    context.src_shape = context.src_slide.shapes.add_group_shape()


@given("a second slide in the same presentation")
def given_a_second_slide_same_prs(context):
    context.tgt_slide = context.prs.slides.add_slide(context.prs.slide_layouts[5])


@given("a source presentation with a picture on slide 1")
def given_source_prs_with_pic(context):
    prs = Presentation()
    context.src_prs = prs
    context.src_slide = prs.slides.add_slide(prs.slide_layouts[5])
    context.src_shape = context.src_slide.shapes.add_picture(
        test_image("python-icon.jpeg"), Inches(1), Inches(1)
    )


@given("a target presentation with one empty slide")
def given_target_prs(context):
    prs = Presentation()
    context.tgt_prs = prs
    context.tgt_slide = prs.slides.add_slide(prs.slide_layouts[5])


# when ====================================================


@when("I clone the auto-shape onto the second slide")
def when_clone_autoshape(context):
    context.clone = context.src_shape.clone_onto(context.tgt_slide.shapes)


@when("I clone the picture onto the second slide")
def when_clone_picture(context):
    context.clone = context.src_shape.clone_onto(context.tgt_slide.shapes)


@when("I clone the picture onto the target presentation's slide")
def when_clone_picture_cross_prs(context):
    context.clone = context.src_shape.clone_onto(context.tgt_slide.shapes)


@when("I clone the connector onto the second slide")
def when_clone_connector(context):
    context.clone = context.src_shape.clone_onto(context.tgt_slide.shapes)


@when("I clone the table onto the second slide")
def when_clone_table(context):
    context.clone = context.src_shape.clone_onto(context.tgt_slide.shapes)


@when("I clone the group onto the second slide")
def when_clone_group(context):
    context.clone = context.src_shape.clone_onto(context.tgt_slide.shapes)


@when("I clone the auto-shape onto the second slide at (3in, 4in)")
def when_clone_with_pos(context):
    context.clone = context.src_shape.clone_onto(
        context.tgt_slide.shapes, left=Inches(3), top=Inches(4)
    )


# then ====================================================


@then("the clone is a Shape with a fresh shape id")
def then_clone_is_shape_fresh_id(context):
    from pptx.shapes.autoshape import Shape

    assert isinstance(context.clone, Shape), "clone is %r" % type(context.clone)
    assert context.clone.shape_id != context.src_shape.shape_id


@then("the clone has the same text as the source")
def then_clone_same_text(context):
    assert context.clone.text_frame.text == context.src_shape.text_frame.text


@then("the clone is topmost in the second slide's z-order")
def then_clone_topmost(context):
    assert context.tgt_slide.shapes[-1]._element is context.clone._element


@then("the clone is a Picture")
def then_clone_is_picture(context):
    assert isinstance(context.clone, Picture), "clone is %r" % type(context.clone)


@then("the clone's image bytes match the source")
def then_clone_image_bytes(context):
    assert context.clone.image is not None
    assert context.clone.image.blob == context.src_shape.image.blob


@then("the clone's image part lives in the target presentation's package")
def then_clone_image_in_tgt_pkg(context):
    assert context.clone.part.package is context.tgt_prs.part.package


@then("the clone is a Connector with the same begin and end points")
def then_clone_connector(context):
    from pptx.shapes.connector import Connector

    assert isinstance(context.clone, Connector)
    assert context.clone.begin_x == context.src_shape.begin_x
    assert context.clone.end_x == context.src_shape.end_x
    assert context.clone.begin_y == context.src_shape.begin_y
    assert context.clone.end_y == context.src_shape.end_y


@then("the clone is a GraphicFrame that has a table")
def then_clone_gf_table(context):
    assert isinstance(context.clone, GraphicFrame)
    assert context.clone.has_table


@then("the clone is a GroupShape with a fresh shape id")
def then_clone_group_fresh_id(context):
    assert isinstance(context.clone, GroupShape)
    assert context.clone.shape_id != context.src_shape.shape_id


@then("the clone's position is (3in, 4in)")
def then_clone_position(context):
    assert context.clone.left == Inches(3)
    assert context.clone.top == Inches(4)
