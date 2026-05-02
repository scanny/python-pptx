"""Gherkin step implementations for issue #94 view-settings feature."""

from __future__ import annotations

from behave import given, then, when
from behave.runner import Context
from helpers import saved_pptx_path

from pptx import Presentation
from pptx.enum.presentation import PP_VIEW_TYPE

# given ===================================================


@given("I open a freshly-created presentation")
def step_given_fresh_presentation(context: Context):
    context.prs = Presentation()


# when ====================================================


@when("I set view_type to OUTLINE and slide_view_zoom to 1.5")
def step_when_set_view_type_and_zoom(context: Context):
    context.prs.view_props.view_type = PP_VIEW_TYPE.OUTLINE
    context.prs.view_props.slide_view_zoom = 1.5


@when("I set show_comments to False and first_slide_num to 7")
def step_when_set_show_comments_and_first_slide_num(context: Context):
    context.prs.view_props.show_comments = False
    context.prs.first_slide_num = 7


# then ====================================================


@then("view_props.view_type reports PP_VIEW_TYPE.SLIDE_THUMBNAIL")
def step_then_view_type_is_slide_thumbnail(context: Context):
    actual = context.prs.view_props.view_type
    assert actual == PP_VIEW_TYPE.SLIDE_THUMBNAIL, (
        "expected PP_VIEW_TYPE.SLIDE_THUMBNAIL, got %r" % actual
    )


@then("opening the saved presentation shows view_type OUTLINE")
def step_then_reopened_view_type_outline(context: Context):
    prs = Presentation(saved_pptx_path)
    assert prs.view_props.view_type == PP_VIEW_TYPE.OUTLINE, (
        "expected OUTLINE, got %r" % prs.view_props.view_type
    )


@then("opening the saved presentation shows slide_view_zoom 1.5")
def step_then_reopened_slide_view_zoom(context: Context):
    prs = Presentation(saved_pptx_path)
    actual = prs.view_props.slide_view_zoom
    assert abs(actual - 1.5) < 1e-9, "expected 1.5, got %r" % actual


@then("opening the saved presentation shows show_comments False")
def step_then_reopened_show_comments_false(context: Context):
    prs = Presentation(saved_pptx_path)
    assert prs.view_props.show_comments is False, (
        "expected False, got %r" % prs.view_props.show_comments
    )


@then("opening the saved presentation shows first_slide_num 7")
def step_then_reopened_first_slide_num_7(context: Context):
    prs = Presentation(saved_pptx_path)
    assert prs.first_slide_num == 7, "expected 7, got %r" % prs.first_slide_num
