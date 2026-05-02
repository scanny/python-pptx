"""Unit-test suite for `pptx.oxml.shapes.shared` module."""

from __future__ import annotations

import pytest

from pptx.oxml.ns import qn
from pptx.oxml.shapes.shared import CT_AlternateContent

from ...unitutil.cxml import element


class DescribeCT_AlternateContent(object):
    """Unit-test suite for `pptx.oxml.shapes.shared.CT_AlternateContent` objects."""

    def it_parses_to_the_custom_element_class(self):
        ac = element("mc:AlternateContent")
        assert isinstance(ac, CT_AlternateContent)

    def it_provides_access_to_its_Choice_children(self):
        ac = element(
            "mc:AlternateContent/(mc:Choice{Requires=a14},mc:Choice{Requires=p14"
            "},mc:Fallback)"
        )
        assert isinstance(ac, CT_AlternateContent)
        choices = ac.choices
        assert len(choices) == 2
        assert all(c.tag == qn("mc:Choice") for c in choices)

    def it_returns_empty_choices_list_when_none_present(self):
        ac = element("mc:AlternateContent/mc:Fallback")
        assert isinstance(ac, CT_AlternateContent)
        assert ac.choices == []

    def it_provides_access_to_its_Fallback_child(self):
        ac = element("mc:AlternateContent/(mc:Choice,mc:Fallback)")
        assert isinstance(ac, CT_AlternateContent)
        fallback = ac.fallback
        assert fallback is not None
        assert fallback.tag == qn("mc:Fallback")

    def it_returns_None_when_no_Fallback_child(self):
        ac = element("mc:AlternateContent/mc:Choice")
        assert isinstance(ac, CT_AlternateContent)
        assert ac.fallback is None

    @pytest.fixture(
        params=[
            # -- no Choice: nothing yielded even if Fallback contains shapes --
            ("mc:AlternateContent/mc:Fallback/p:sp", 0),
            # -- empty Choice: nothing yielded --
            ("mc:AlternateContent/mc:Choice", 0),
            # -- single shape in Choice --
            ("mc:AlternateContent/mc:Choice/p:sp", 1),
            # -- multiple shapes in Choice in document order --
            ("mc:AlternateContent/mc:Choice/(p:sp,p:sp,p:sp)", 3),
            # -- nested AlternateContent flattens by preferring inner Choice --
            (
                "mc:AlternateContent/mc:Choice/(p:sp,mc:AlternateContent/(mc:Cho"
                "ice/p:sp,mc:Fallback/p:sp))",
                2,
            ),
            # -- non-shape children are skipped --
            ("mc:AlternateContent/mc:Choice/(p:extLst,p:sp)", 1),
            # -- mc:Fallback is ignored even if it has shapes --
            ("mc:AlternateContent/(mc:Choice/p:sp,mc:Fallback/(p:sp,p:sp))", 1),
            # -- only first Choice is walked --
            ("mc:AlternateContent/(mc:Choice/p:sp,mc:Choice/(p:sp,p:sp))", 1),
        ]
    )
    def iter_choice_fixture(self, request):
        cxml, expected_count = request.param
        return element(cxml), expected_count

    def it_iterates_shape_elements_from_the_first_Choice(self, iter_choice_fixture):
        ac, expected_count = iter_choice_fixture
        shape_tags = (qn("p:sp"), qn("p:grpSp"), qn("p:graphicFrame"))
        shape_elms = list(ac.iter_choice_shape_elms(shape_tags))
        assert len(shape_elms) == expected_count
        assert all(e.tag in shape_tags for e in shape_elms)
