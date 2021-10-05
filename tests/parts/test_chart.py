# encoding: utf-8

"""Unit-test suite for `pptx.parts.chart` module."""

import pytest

from pptx.chart.chart import Chart
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE as XCT
from pptx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from pptx.opc.package import OpcPackage
from pptx.opc.packuri import PackURI
from pptx.oxml.chart.chart import CT_ChartSpace
from pptx.parts.chart import ChartPart, ChartWorkbook
from pptx.parts.embeddedpackage import EmbeddedXlsxPart

from ..unitutil.cxml import element, xml
from ..unitutil.mock import class_mock, instance_mock, method_mock, property_mock


class DescribeChartPart(object):
    """Unit-test suite for `pptx.parts.chart.ChartPart` objects."""

    def it_can_construct_from_chart_type_and_data(self, request):
        chart_data_ = instance_mock(request, ChartData, xlsx_blob=b"xlsx-blob")
        chart_data_.xml_bytes.return_value = b"chart-blob"
        package_ = instance_mock(request, OpcPackage)
        package_.next_partname.return_value = PackURI("/ppt/charts/chart42.xml")
        chart_part_ = instance_mock(request, ChartPart)
        # --- load() must have autospec turned off to work in Python 2.7 mock ---
        load_ = method_mock(
            request, ChartPart, "load", autospec=False, return_value=chart_part_
        )

        chart_part = ChartPart.new(XCT.RADAR, chart_data_, package_)

        package_.next_partname.assert_called_once_with("/ppt/charts/chart%d.xml")
        chart_data_.xml_bytes.assert_called_once_with(XCT.RADAR)
        load_.assert_called_once_with(
            "/ppt/charts/chart42.xml", CT.DML_CHART, package_, b"chart-blob"
        )
        chart_part_.chart_workbook.update_from_xlsx_blob.assert_called_once_with(
            b"xlsx-blob"
        )
        assert chart_part is chart_part_

    def it_provides_access_to_the_chart_object(self, request, chartSpace_):
        chart_ = instance_mock(request, Chart)
        Chart_ = class_mock(request, "pptx.parts.chart.Chart", return_value=chart_)
        chart_part = ChartPart(None, None, None, chartSpace_)

        chart = chart_part.chart

        Chart_.assert_called_once_with(chart_part._element, chart_part)
        assert chart is chart_

    def it_provides_access_to_the_chart_workbook(self, request, chartSpace_):
        chart_workbook_ = instance_mock(request, ChartWorkbook)
        ChartWorkbook_ = class_mock(
            request, "pptx.parts.chart.ChartWorkbook", return_value=chart_workbook_
        )
        chart_part = ChartPart(None, None, None, chartSpace_)

        chart_workbook = chart_part.chart_workbook

        ChartWorkbook_.assert_called_once_with(chartSpace_, chart_part)
        assert chart_workbook is chart_workbook_

    # fixture components ---------------------------------------------

    @pytest.fixture
    def chartSpace_(self, request):
        return instance_mock(request, CT_ChartSpace)


class DescribeChartWorkbook(object):
    """Unit-test suite for `pptx.parts.chart.ChartWorkbook` objects."""

    def it_can_get_the_chart_xlsx_part(self, chart_part_, xlsx_part_):
        chart_part_.related_part.return_value = xlsx_part_
        chart_workbook = ChartWorkbook(
            element("c:chartSpace/c:externalData{r:id=rId42}"), chart_part_
        )

        xlsx_part = chart_workbook.xlsx_part

        chart_part_.related_part.assert_called_once_with("rId42")
        assert xlsx_part is xlsx_part_

    def but_it_returns_None_when_the_chart_has_no_xlsx_part(self):
        chart_workbook = ChartWorkbook(element("c:chartSpace"), None)
        assert chart_workbook.xlsx_part is None

    @pytest.mark.parametrize(
        "chartSpace_cxml, expected_cxml",
        (
            (
                "c:chartSpace{r:a=b}",
                "c:chartSpace{r:a=b}/c:externalData{r:id=rId" "42}/c:autoUpdate{val=0}",
            ),
            (
                "c:chartSpace/c:externalData{r:id=rId66}",
                "c:chartSpace/c:externalData{r:id=rId42}",
            ),
        ),
    )
    def it_can_change_the_chart_xlsx_part(
        self, chart_part_, xlsx_part_, chartSpace_cxml, expected_cxml
    ):
        chart_part_.relate_to.return_value = "rId42"
        chart_data = ChartWorkbook(element(chartSpace_cxml), chart_part_)

        chart_data.xlsx_part = xlsx_part_

        chart_part_.relate_to.assert_called_once_with(xlsx_part_, RT.PACKAGE)
        assert chart_data._chartSpace.xml == xml(expected_cxml)

    def it_adds_an_xlsx_part_on_update_if_needed(
        self, request, chart_part_, package_, xlsx_part_, xlsx_part_prop_
    ):
        EmbeddedXlsxPart_ = class_mock(request, "pptx.parts.chart.EmbeddedXlsxPart")
        EmbeddedXlsxPart_.new.return_value = xlsx_part_
        chart_part_.package = package_
        xlsx_part_prop_.return_value = None
        chart_data = ChartWorkbook(element("c:chartSpace"), chart_part_)

        chart_data.update_from_xlsx_blob(b"xlsx-blob")

        EmbeddedXlsxPart_.new.assert_called_once_with(b"xlsx-blob", package_)
        xlsx_part_prop_.assert_called_with(xlsx_part_)

    def but_it_replaces_the_xlsx_blob_when_the_part_exists(
        self, xlsx_part_prop_, xlsx_part_
    ):
        xlsx_part_prop_.return_value = xlsx_part_
        chart_data = ChartWorkbook(None, None)
        chart_data.update_from_xlsx_blob(b"xlsx-blob")

        assert chart_data.xlsx_part.blob == b"xlsx-blob"

    # fixture components ---------------------------------------------

    @pytest.fixture
    def chart_part_(self, request, package_, xlsx_part_):
        return instance_mock(request, ChartPart)

    @pytest.fixture
    def package_(self, request):
        return instance_mock(request, OpcPackage)

    @pytest.fixture
    def xlsx_part_(self, request):
        return instance_mock(request, EmbeddedXlsxPart)

    @pytest.fixture
    def xlsx_part_prop_(self, request):
        return property_mock(request, ChartWorkbook, "xlsx_part")
