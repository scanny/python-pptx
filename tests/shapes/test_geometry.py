# pyright: reportPrivateUsage=false

"""Unit-test suite for `pptx.shapes.geometry` module."""

from __future__ import annotations

from typing import cast

import pytest

from pptx.oxml.shapes.autoshape import CT_Shape
from pptx.shapes.geometry import (
    ArcTo,
    Close,
    CubicBezierTo,
    LineTo,
    MoveTo,
    Path,
    PathGeometry,
    QuadBezierTo,
)
from pptx.util import Emu

from ..unitutil.cxml import element


class DescribeDrawingOps(object):
    """Unit-test suite for the drawing-operation value objects in `pptx.shapes.geometry`."""

    def it_exposes_MoveTo_coordinates(self):
        op = MoveTo(Emu(100), Emu(200))
        assert op.x == 100
        assert op.y == 200
        assert op.points == ((100, 200),)

    def it_exposes_LineTo_coordinates(self):
        op = LineTo(Emu(300), Emu(400))
        assert op.x == 300
        assert op.y == 400

    def it_holds_CubicBezierTo_control_points_and_endpoint(self):
        op = CubicBezierTo((Emu(1), Emu(2)), (Emu(3), Emu(4)), (Emu(5), Emu(6)))
        assert op.points == ((1, 2), (3, 4), (5, 6))

    def it_holds_QuadBezierTo_control_point_and_endpoint(self):
        op = QuadBezierTo((Emu(1), Emu(2)), (Emu(3), Emu(4)))
        assert op.points == ((1, 2), (3, 4))

    def it_stores_ArcTo_parameters(self):
        op = ArcTo(Emu(100), Emu(200), 0, 5400000)
        assert op.wR == 100
        assert op.hR == 200
        assert op.stAng == 0
        assert op.swAng == 5400000

    def it_represents_Close_with_no_points(self):
        op = Close()
        assert op.points == ()

    def it_compares_equal_for_same_type_and_points(self):
        assert MoveTo(Emu(1), Emu(2)) == MoveTo(Emu(1), Emu(2))
        assert MoveTo(Emu(1), Emu(2)) != LineTo(Emu(1), Emu(2))
        assert ArcTo(Emu(1), Emu(2), 0, 100) == ArcTo(Emu(1), Emu(2), 0, 100)
        assert ArcTo(Emu(1), Emu(2), 0, 100) != ArcTo(Emu(1), Emu(2), 0, 200)

    def it_is_hashable(self):
        hash(MoveTo(Emu(1), Emu(2)))
        hash(ArcTo(Emu(1), Emu(2), 0, 100))
        hash(Close())


class DescribePath(object):
    """Unit-test suite for `pptx.shapes.geometry.Path`."""

    def it_exposes_its_coord_system_dimensions(self):
        path = Path(Emu(914400), Emu(914400), (MoveTo(Emu(0), Emu(0)),))
        assert path.width == 914400
        assert path.height == 914400

    def it_is_a_sequence_of_drawing_ops(self):
        ops = (MoveTo(Emu(0), Emu(0)), LineTo(Emu(1), Emu(0)), Close())
        path = Path(Emu(100), Emu(100), ops)
        assert len(path) == 3
        assert path[0] == ops[0]
        assert list(path) == list(ops)


class DescribePathGeometry(object):
    """Unit-test suite for `pptx.shapes.geometry.PathGeometry`."""

    def it_is_a_sequence_of_Paths(self):
        p1 = Path(Emu(100), Emu(100), ())
        p2 = Path(Emu(200), Emu(200), ())
        pg = PathGeometry((p1, p2))
        assert len(pg) == 2
        assert pg[0] is p1
        assert list(pg) == [p1, p2]

    def it_returns_None_for_placeholder_shape_without_geometry(self):
        sp = cast(CT_Shape, element("p:sp/(p:nvSpPr/(p:cNvPr,p:cNvSpPr,p:nvPr),p:spPr)"))
        assert PathGeometry.from_shape(sp) is None

    def it_returns_None_for_dynamic_preset_with_adjustment(self):
        # -- roundRect is a dynamic preset (has adjustment) --
        xml = (
            '<p:sp xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
            ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:nvSpPr><p:cNvPr id="1" name="x"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            "<p:spPr>"
            '<a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></a:xfrm>'
            '<a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>'
            "</p:spPr></p:sp>"
        )
        sp = cast(CT_Shape, element_from_xml(xml))
        assert PathGeometry.from_shape(sp) is None

    def it_returns_None_for_static_preset_with_avLst_adjustments_set(self):
        # -- even a normally-static preset returns None if the avLst holds guides --
        xml = (
            '<p:sp xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
            ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:nvSpPr><p:cNvPr id="1" name="x"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            "<p:spPr>"
            '<a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst>'
            '<a:gd name="adj" fmla="val 50000"/>'
            "</a:avLst></a:prstGeom>"
            "</p:spPr></p:sp>"
        )
        sp = cast(CT_Shape, element_from_xml(xml))
        assert PathGeometry.from_shape(sp) is None

    def it_resolves_a_rect_preset_against_the_shape_dimensions(self):
        xml = _sp_xml("rect", 914400, 914400)
        sp = cast(CT_Shape, element_from_xml(xml))
        pg = PathGeometry.from_shape(sp)
        assert pg is not None
        assert len(pg) == 1
        path = pg[0]
        assert path.width == 914400
        assert path.height == 914400
        assert list(path) == [
            MoveTo(Emu(0), Emu(0)),
            LineTo(Emu(914400), Emu(0)),
            LineTo(Emu(914400), Emu(914400)),
            LineTo(Emu(0), Emu(914400)),
            Close(),
        ]

    def it_resolves_a_lineInv_preset(self):
        xml = _sp_xml("lineInv", 200, 100)
        sp = cast(CT_Shape, element_from_xml(xml))
        pg = PathGeometry.from_shape(sp)
        assert pg is not None
        assert list(pg[0]) == [MoveTo(Emu(0), Emu(100)), LineTo(Emu(200), Emu(0))]

    def it_scales_path_local_coords_to_shape_bbox(self):
        # -- flowChartProcess has path w=1 h=1 for a shape of any size --
        xml = _sp_xml("flowChartProcess", 914400, 457200)
        sp = cast(CT_Shape, element_from_xml(xml))
        pg = PathGeometry.from_shape(sp)
        assert pg is not None
        path = pg[0]
        assert list(path) == [
            MoveTo(Emu(0), Emu(0)),
            LineTo(Emu(914400), Emu(0)),
            LineTo(Emu(914400), Emu(457200)),
            LineTo(Emu(0), Emu(457200)),
            Close(),
        ]

    def it_returns_multiple_paths_for_a_multi_path_preset(self):
        xml = _sp_xml("chartPlus", 100, 100)
        sp = cast(CT_Shape, element_from_xml(xml))
        pg = PathGeometry.from_shape(sp)
        assert pg is not None
        assert len(pg) == 2  # -- chartPlus has one stroke and one fill path --

    def it_reads_custom_geometry_from_the_pathLst(self):
        xml = (
            '<p:sp xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
            ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:nvSpPr><p:cNvPr id="1" name="x"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            "<p:spPr>"
            '<a:xfrm><a:off x="0" y="0"/><a:ext cx="1000" cy="500"/></a:xfrm>'
            "<a:custGeom><a:avLst/><a:pathLst>"
            '<a:path w="1000" h="500">'
            '<a:moveTo><a:pt x="0" y="0"/></a:moveTo>'
            '<a:lnTo><a:pt x="1000" y="0"/></a:lnTo>'
            '<a:cubicBezTo>'
            '<a:pt x="1000" y="250"/>'
            '<a:pt x="500" y="500"/>'
            '<a:pt x="0" y="500"/>'
            '</a:cubicBezTo>'
            "<a:close/>"
            "</a:path>"
            "</a:pathLst></a:custGeom>"
            "</p:spPr></p:sp>"
        )
        sp = cast(CT_Shape, element_from_xml(xml))
        pg = PathGeometry.from_shape(sp)
        assert pg is not None
        assert len(pg) == 1
        path = pg[0]
        assert path.width == 1000
        assert path.height == 500
        assert list(path) == [
            MoveTo(Emu(0), Emu(0)),
            LineTo(Emu(1000), Emu(0)),
            CubicBezierTo((Emu(1000), Emu(250)), (Emu(500), Emu(500)), (Emu(0), Emu(500))),
            Close(),
        ]

    def it_reads_arcTo_and_quadBezTo_in_custom_geometry(self):
        xml = (
            '<p:sp xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
            ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:nvSpPr><p:cNvPr id="1" name="x"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            "<p:spPr>"
            '<a:xfrm><a:off x="0" y="0"/><a:ext cx="1000" cy="1000"/></a:xfrm>'
            "<a:custGeom><a:avLst/><a:pathLst>"
            '<a:path w="1000" h="1000">'
            '<a:moveTo><a:pt x="0" y="0"/></a:moveTo>'
            '<a:arcTo wR="500" hR="500" stAng="0" swAng="5400000"/>'
            '<a:quadBezTo><a:pt x="500" y="500"/><a:pt x="0" y="1000"/></a:quadBezTo>'
            "</a:path>"
            "</a:pathLst></a:custGeom>"
            "</p:spPr></p:sp>"
        )
        sp = cast(CT_Shape, element_from_xml(xml))
        pg = PathGeometry.from_shape(sp)
        assert pg is not None
        ops = list(pg[0])
        assert ops[0] == MoveTo(Emu(0), Emu(0))
        assert ops[1] == ArcTo(Emu(500), Emu(500), 0, 5400000)
        assert ops[2] == QuadBezierTo((Emu(500), Emu(500)), (Emu(0), Emu(1000)))

    def it_returns_empty_geometry_for_custom_geom_without_pathLst(self):
        xml = (
            '<p:sp xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
            ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:nvSpPr><p:cNvPr id="1" name="x"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            "<p:spPr>"
            '<a:xfrm><a:off x="0" y="0"/><a:ext cx="1000" cy="500"/></a:xfrm>'
            "<a:custGeom><a:avLst/></a:custGeom>"
            "</p:spPr></p:sp>"
        )
        sp = cast(CT_Shape, element_from_xml(xml))
        pg = PathGeometry.from_shape(sp)
        assert pg is not None
        assert len(pg) == 0

    @pytest.mark.parametrize("prst", ["rect", "lineInv", "flowChartProcess", "actionButtonBlank"])
    def it_resolves_each_static_preset_without_error(self, prst: str):
        xml = _sp_xml(prst, 914400, 914400)
        sp = cast(CT_Shape, element_from_xml(xml))
        pg = PathGeometry.from_shape(sp)
        assert pg is not None
        assert len(pg) >= 1


# -- helpers ---------------------------------------------------------------------


def _sp_xml(prst: str, cx: int, cy: int) -> str:
    """Return a `p:sp` XML string with a preset-geometry child."""
    return (
        '<p:sp xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<p:nvSpPr><p:cNvPr id="1" name="x"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        "<p:spPr>"
        f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>'
        "</p:spPr></p:sp>"
    )


def element_from_xml(xml: str):
    from pptx.oxml import parse_xml

    return parse_xml(xml)
