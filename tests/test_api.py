"""Unit-test suite for `pptx.api` module."""

from __future__ import annotations

import os

import pytest

import pptx
from pptx.api import Presentation, _default_pptx_path, _is_pptx_package
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.parts.presentation import PresentationPart

from .unitutil.mock import class_mock, instance_mock


class DescribePresentation(object):
    def it_opens_default_template_on_no_path_provided(self, call_fixture):
        Package_, path, prs_ = call_fixture
        prs = Presentation()
        Package_.open.assert_called_once_with(path, password=None)
        assert prs is prs_

    def it_can_open_a_password_protected_presentation(self, call_fixture):
        Package_, path, prs_ = call_fixture
        prs = Presentation(path, password="s3cret")
        Package_.open.assert_called_once_with(path, password="s3cret")
        assert prs is prs_

    @pytest.mark.parametrize(
        ("pptx_format", "template_filename"),
        [
            ("16x9", "default-16x9.pptx"),
            ("widescreen", "default-16x9.pptx"),
            ("4x3", "default.pptx"),
            ("standard", "default.pptx"),
            ("16X9", "default-16x9.pptx"),
        ],
    )
    def it_opens_preset_template_on_pptx_format(
        self, pptx_format, template_filename, Package_, prs_, prs_part_
    ):
        Package_.open.return_value.main_document_part = prs_part_
        prs_part_.content_type = CT.PML_PRESENTATION_MAIN
        prs_part_.presentation = prs_

        prs = Presentation(pptx_format=pptx_format)

        expected_path = os.path.join(
            os.path.split(pptx.__file__)[0], "templates", template_filename
        )
        Package_.open.assert_called_once_with(expected_path, password=None)
        assert prs is prs_

    def it_raises_on_unknown_pptx_format(self):
        with pytest.raises(ValueError, match="unknown pptx_format '21x9'"):
            Presentation(pptx_format="21x9")

    def it_raises_when_pptx_format_used_with_explicit_pptx(self):
        with pytest.raises(ValueError, match="pptx_format is only valid when opening the default"):
            Presentation("foobar.pptx", pptx_format="16x9")

    @pytest.mark.parametrize(
        "content_type",
        [
            CT.PML_PRESENTATION_MAIN,
            CT.PML_PRES_MACRO_MAIN,
            CT.PML_TEMPLATE_MAIN,
            CT.PML_SLIDESHOW_MAIN,
        ],
    )
    def it_opens_any_presentationml_main_content_type(
        self, content_type, Package_, prs_, prs_part_
    ):
        Package_.open.return_value.main_document_part = prs_part_
        prs_part_.content_type = content_type
        prs_part_.presentation = prs_

        prs = Presentation("foobar.pptx")

        assert prs is prs_

    def it_raises_on_non_presentationml_main_content_type(self, Package_, prs_part_):
        Package_.open.return_value.main_document_part = prs_part_
        prs_part_.content_type = "application/vnd.ms-word"

        with pytest.raises(ValueError, match="is not a PowerPoint file"):
            Presentation("foobar.docx")

    def it_ships_a_16x9_template_with_widescreen_slide_size(self):
        """Built-in widescreen template uses 13.333in x 7.5in (16:9)."""
        from pptx import Presentation as PresentationFn

        prs = PresentationFn(pptx_format="16x9")
        assert prs.slide_width == 12192000
        assert prs.slide_height == 6858000

    def it_ships_a_4x3_template_with_standard_slide_size(self):
        """Built-in default template remains 10in x 7.5in (4:3)."""
        from pptx import Presentation as PresentationFn

        prs = PresentationFn()
        assert prs.slide_width == 9144000
        assert prs.slide_height == 6858000

    # fixtures -------------------------------------------------------

    @pytest.fixture
    def call_fixture(self, Package_, prs_, prs_part_):
        path = os.path.abspath(
            os.path.join(os.path.split(pptx.__file__)[0], "templates", "default.pptx")
        )
        Package_.open.return_value.main_document_part = prs_part_
        prs_part_.content_type = CT.PML_PRESENTATION_MAIN
        prs_part_.presentation = prs_
        return Package_, path, prs_

    # fixture components ---------------------------------------------

    @pytest.fixture
    def Package_(self, request):
        return class_mock(request, "pptx.api.Package")

    @pytest.fixture
    def prs_(self, request):
        return instance_mock(request, Presentation)

    @pytest.fixture
    def prs_part_(self, request):
        return instance_mock(request, PresentationPart)


class Describe_default_pptx_path(object):
    def it_returns_4x3_template_path_by_default(self):
        path = _default_pptx_path()
        assert os.path.basename(path) == "default.pptx"
        assert os.path.isfile(path)

    @pytest.mark.parametrize("preset", ["4x3", "standard", "STANDARD"])
    def it_returns_4x3_template_for_4x3_aliases(self, preset):
        path = _default_pptx_path(preset)
        assert os.path.basename(path) == "default.pptx"

    @pytest.mark.parametrize("preset", ["16x9", "widescreen", "Widescreen"])
    def it_returns_widescreen_template_for_16x9_aliases(self, preset):
        path = _default_pptx_path(preset)
        assert os.path.basename(path) == "default-16x9.pptx"
        assert os.path.isfile(path)

    def it_raises_on_unknown_preset(self):
        with pytest.raises(ValueError, match="unknown pptx_format 'foo'"):
            _default_pptx_path("foo")


class Describe_is_pptx_package(object):
    @pytest.mark.parametrize(
        "content_type",
        [
            CT.PML_PRESENTATION_MAIN,
            CT.PML_PRES_MACRO_MAIN,
            CT.PML_TEMPLATE_MAIN,
            CT.PML_SLIDESHOW_MAIN,
        ],
    )
    def it_is_True_for_each_presentationml_main_content_type(self, content_type, prs_part_):
        prs_part_.content_type = content_type
        assert _is_pptx_package(prs_part_) is True

    def it_is_False_for_any_other_content_type(self, prs_part_):
        prs_part_.content_type = "application/vnd.ms-excel"
        assert _is_pptx_package(prs_part_) is False

    # fixture components ---------------------------------------------

    @pytest.fixture
    def prs_part_(self, request):
        return instance_mock(request, PresentationPart)
