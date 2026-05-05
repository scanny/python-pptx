"""Regression tests for audit item 2: wrap missing ``[Content_Types].xml``.

When a caller opens a file/stream that is a valid zip but lacks the
``[Content_Types].xml`` part required by the Open Packaging Conventions,
python-pptx previously surfaced a bare :class:`KeyError` originating from the
zip lookup. This module exercises the replacement behavior, which raises
:class:`pptx.exc.PackageNotFoundError` with a clear diagnostic message while
preserving the underlying :class:`KeyError` via ``__cause__``.

Tracks upstream issue scanny/python-pptx#3.
"""

from __future__ import annotations

import io
import os
import zipfile
from pathlib import Path

import pytest

from pptx import Presentation
from pptx.exc import PackageNotFoundError


def _make_zip_without_content_types() -> io.BytesIO:
    """Return an in-memory zip that is valid but carries no OPC parts."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("garbage.txt", "hello")
    buf.seek(0)
    return buf


class DescribeInvalidZipPackageHandling:
    """Verify PackageNotFoundError surfaces instead of KeyError."""

    def it_raises_PackageNotFoundError_for_stream_missing_content_types(self):
        buf = _make_zip_without_content_types()

        with pytest.raises(PackageNotFoundError) as excinfo:
            Presentation(buf)

        assert "[Content_Types].xml" in str(excinfo.value)
        assert "valid zip" in str(excinfo.value)

    def it_raises_PackageNotFoundError_for_path_missing_content_types(self, tmp_path: Path):
        pkg_path = os.fspath(tmp_path / "bogus.pptx")
        with zipfile.ZipFile(pkg_path, "w") as zf:
            zf.writestr("garbage.txt", "hello")

        with pytest.raises(PackageNotFoundError) as excinfo:
            Presentation(pkg_path)

        assert "[Content_Types].xml" in str(excinfo.value)

    def it_chains_the_original_KeyError_as_cause(self):
        buf = _make_zip_without_content_types()

        with pytest.raises(PackageNotFoundError) as excinfo:
            Presentation(buf)

        # -- the underlying `raise ... from exc` preserves the KeyError so
        # -- downstream diagnostics (tracebacks, logging) keep the full story.
        assert isinstance(excinfo.value.__cause__, KeyError)
        assert "[Content_Types].xml" in str(excinfo.value.__cause__)
