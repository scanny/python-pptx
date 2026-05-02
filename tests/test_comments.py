# pyright: reportPrivateUsage=false

"""Unit-test suite for `pptx.comments` high-level API module."""

from __future__ import annotations

import datetime as dt
import io

import pytest

from pptx import Presentation
from pptx.comments import Comment, CommentAuthor, CommentAuthors, Comments
from pptx.oxml.comments import CT_CommentAuthorList, CT_CommentList
from pptx.parts.comments import CommentAuthorsPart, CommentsPart


@pytest.fixture
def _restore_part_factory():
    """Guard against pollution from other tests that mutate `PartFactory.part_type_for`.

    In particular, `tests/opc/test_package.py::DescribePartFactory` overwrites
    `PartFactory.part_type_for[CT.PML_SLIDE]` with a Mock and does not restore it.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.comments import CommentAuthorsPart, CommentsPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_COMMENTS: CommentsPart,
        CT.PML_COMMENT_AUTHORS: CommentAuthorsPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


@pytest.fixture
def prs_with_slide(_restore_part_factory):
    prs = Presentation()
    slide_layout = prs.slide_layouts[5]
    prs.slides.add_slide(slide_layout)
    return prs


class DescribeCommentAuthor:
    def it_exposes_id_name_and_initials(self):
        cmAuthorLst = CT_CommentAuthorList.new()
        cmAuthor = cmAuthorLst.add_author("Alice", "A.")

        author = CommentAuthor(cmAuthor)

        assert author.id == 0
        assert author.name == "Alice"
        assert author.initials == "A."


class DescribeCommentAuthors:
    def _make(self) -> CommentAuthors:
        from unittest.mock import Mock

        part = Mock(spec=CommentAuthorsPart)
        part.author_list = CT_CommentAuthorList.new()
        return CommentAuthors(part)

    def it_has_zero_length_when_empty(self):
        authors = self._make()
        assert len(authors) == 0
        assert list(authors) == []

    def it_appends_new_authors(self):
        authors = self._make()

        alice = authors.add_author("Alice", "A.")
        bob = authors.add_author("Bob", "B.")

        assert len(authors) == 2
        assert alice.id == 0
        assert bob.id == 1
        assert [a.name for a in authors] == ["Alice", "Bob"]

    def it_supports_indexed_access(self):
        authors = self._make()
        authors.add_author("Alice", "A.")
        authors.add_author("Bob", "B.")

        assert authors[0].name == "Alice"
        assert authors[1].name == "Bob"

    def it_can_look_up_by_id(self):
        authors = self._make()
        authors.add_author("Alice", "A.")
        authors.add_author("Bob", "B.")

        found = authors.get_by_id(1)

        assert found is not None
        assert found.name == "Bob"

    def but_get_by_id_returns_None_for_unknown_id(self):
        authors = self._make()

        assert authors.get_by_id(99) is None

    def it_reuses_existing_author_in_get_or_add(self):
        authors = self._make()
        authors.add_author("Alice", "A.")

        again = authors.get_or_add("Alice", "A.")

        assert len(authors) == 1
        assert again.name == "Alice"

    def it_adds_new_author_in_get_or_add_when_not_present(self):
        authors = self._make()

        new_author = authors.get_or_add("Bob", "B.")

        assert len(authors) == 1
        assert new_author.name == "Bob"


class DescribeComments:
    def _make(self) -> tuple[Comments, CommentAuthors]:
        from unittest.mock import Mock

        authors_part = Mock(spec=CommentAuthorsPart)
        authors_part.author_list = CT_CommentAuthorList.new()
        authors = CommentAuthors(authors_part)

        comments_part = Mock(spec=CommentsPart)
        comments_part.comment_list = CT_CommentList.new()
        comments = Comments(Mock(name="slide"), comments_part, authors)
        return comments, authors

    def it_has_zero_length_initially(self):
        comments, _ = self._make()
        assert len(comments) == 0
        assert list(comments) == []

    def it_can_add_a_comment(self):
        comments, authors = self._make()
        alice = authors.add_author("Alice", "A.")
        when = dt.datetime(2025, 5, 1, 12, 0, 0)

        c = comments.add_comment(alice, "hi", position=(10, 20), datetime_value=when)

        assert isinstance(c, Comment)
        assert c.text == "hi"
        assert c.position == (10, 20)
        assert c.datetime == when
        assert c.author_id == alice.id
        assert c.idx == 1
        assert c.author is not None
        assert c.author.name == "Alice"

    def it_supports_iteration(self):
        comments, authors = self._make()
        alice = authors.add_author("Alice", "A.")

        comments.add_comment(alice, "one")
        comments.add_comment(alice, "two")

        texts = [c.text for c in comments]
        assert texts == ["one", "two"]

    def it_allows_setting_text_position_and_datetime(self):
        comments, authors = self._make()
        alice = authors.add_author("Alice", "A.")
        c = comments.add_comment(alice, "hi")

        c.text = "updated"
        c.position = (50, 60)
        c.datetime = dt.datetime(2030, 1, 1, 0, 0, 0)

        assert c.text == "updated"
        assert c.position == (50, 60)
        assert c.datetime == dt.datetime(2030, 1, 1, 0, 0, 0)


class DescribeSlide_comments_integration:
    def it_is_empty_when_the_slide_has_no_comments_part(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        assert slide.has_comments is False

    def it_creates_comments_part_on_first_access(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        comments = slide.comments

        assert isinstance(comments, Comments)
        assert slide.has_comments is True

    def it_caches_the_Comments_proxy(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        c1 = slide.comments
        c2 = slide.comments

        assert c1 is c2

    def it_writes_the_expected_xml_payloads_when_saved(self, prs_with_slide):
        slide = prs_with_slide.slides[0]
        comments = slide.comments
        alice = comments._authors.add_author("Alice", "A.")
        comments.add_comment(
            alice,
            "hello",
            position=(100, 200),
            datetime_value=dt.datetime(2025, 2, 3, 4, 5, 6),
        )

        buf = io.BytesIO()
        prs_with_slide.save(buf)
        buf.seek(0)

        import zipfile

        with zipfile.ZipFile(buf, "r") as z:
            author_xml = z.read("ppt/commentAuthors.xml").decode()
            comments_xml = z.read("ppt/comments/comment1.xml").decode()
            pres_rels = z.read("ppt/_rels/presentation.xml.rels").decode()
            slide_rels = z.read("ppt/slides/_rels/slide1.xml.rels").decode()

        assert 'name="Alice"' in author_xml
        assert 'initials="A."' in author_xml
        assert "<p:text>hello</p:text>" in comments_xml
        assert 'dt="2025-02-03T04:05:06"' in comments_xml
        assert '<p:pos x="100" y="200"/>' in comments_xml
        # -- commentAuthors rel hangs off presentation.xml per ECMA-376 --
        assert "commentAuthors" in pres_rels
        # -- comments rel hangs off the slide --
        assert "comments" in slide_rels

    def it_round_trips_through_save_and_reopen(self, prs_with_slide):
        slide = prs_with_slide.slides[0]
        comments = slide.comments
        alice = comments._authors.add_author("Alice", "A.")
        comments.add_comment(
            alice,
            "hello",
            position=(100, 200),
            datetime_value=dt.datetime(2025, 2, 3, 4, 5, 6),
        )

        buf = io.BytesIO()
        prs_with_slide.save(buf)
        buf.seek(0)

        prs2 = Presentation(buf)
        slide2 = prs2.slides[0]
        assert slide2.has_comments is True
        reloaded = list(slide2.comments)
        assert len(reloaded) == 1
        c = reloaded[0]
        assert c.text == "hello"
        assert c.position == (100, 200)
        assert c.datetime == dt.datetime(2025, 2, 3, 4, 5, 6)
        assert c.author is not None
        assert c.author.name == "Alice"

    def it_reuses_a_single_comment_authors_part_across_slides(self, prs_with_slide):
        prs = prs_with_slide
        prs.slides.add_slide(prs.slide_layouts[5])
        slide1, slide2 = prs.slides[0], prs.slides[1]

        slide1.comments  # triggers creation
        slide2.comments  # must reuse the same authors part

        package = prs.part.package
        # -- exactly one COMMENT_AUTHORS relationship on the presentation part --
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT

        count = sum(
            1
            for rel in package.presentation_part.rels.values()
            if rel.reltype == RT.COMMENT_AUTHORS
        )
        assert count == 1
