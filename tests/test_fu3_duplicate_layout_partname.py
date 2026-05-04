# pyright: reportPrivateUsage=false

"""Pin "no duplicate layout partname on save-reopen-save" invariant (FU-3).

FU-3 in ``TODO.md`` records a report from the Wave 12 #956 investigation that
opening a default-template deck and re-saving it emitted ``UserWarning:
Duplicate name: 'ppt/slideLayouts/slideLayout7.xml'`` from the stdlib zipfile
writer. The hypothesis was that two distinct layout parts end up with the
same partname when ``iter_parts()`` walks the reachable graph — typically the
symptom of a partname-collision elsewhere in the package (e.g. a cloner that
forgot to rename the copy, or a reachability walk that visits the same logical
part via two different rel-chains).

A directed retry with a fresh ``Presentation()`` — save, reopen, save — does
NOT reproduce the warning on current master, nor do any of the obvious
variations: adding a slide of every built-in layout, attaching notes, round-
tripping a chart, or triple save-reopen-save. The default template's 11
layouts (``slideLayout1.xml`` … ``slideLayout11.xml``) all emit as unique
zip entries on every save cycle.

The original report most likely referred to a deck that had already been
manipulated by the in-progress #956 clone / delete paths during that wave;
those intermediate code paths did not land on master. Without a concrete
repro input we cannot fix a bug we can't observe, so this test pins the
current correct behavior as a regression guard: if a future change to the
cloner or the reachability walk re-introduces a duplicate-partname emission
on the save-reopen-save cycle, these tests fail immediately.
"""

from __future__ import annotations

import io
import warnings
import zipfile

import pytest

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


class DescribeFU3NoDuplicatePartnameOnSave:
    """Pin invariant: save-reopen-save emits no ``UserWarning`` about dup zip names."""

    def it_emits_no_warning_on_default_deck_save_reopen_save(self):
        buf = io.BytesIO()
        Presentation().save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        buf2 = io.BytesIO()

        with warnings.catch_warnings():
            warnings.simplefilter("error")
            prs2.save(buf2)

    def it_emits_no_warning_on_triple_save_reopen_cycle(self):
        # Three save-reopen cycles — a stricter check that nothing drifts
        # across repeated round-trips (the FU-3 report phrased the symptom
        # as "on save-after-reopen", implying it may emerge after a cycle).
        buf = io.BytesIO()
        Presentation().save(buf)

        for _ in range(3):
            buf.seek(0)
            prs = Presentation(buf)
            buf = io.BytesIO()
            with warnings.catch_warnings():
                warnings.simplefilter("error")
                prs.save(buf)

    def it_emits_no_warning_when_every_builtin_layout_is_used(self):
        prs = Presentation()
        for layout in prs.slide_layouts:
            prs.slides.add_slide(layout)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        buf2 = io.BytesIO()

        with warnings.catch_warnings():
            warnings.simplefilter("error")
            prs2.save(buf2)

    def it_emits_no_warning_on_save_reopen_save_with_chart_and_notes(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        notes_tf = slide.notes_slide.notes_text_frame
        assert notes_tf is not None
        notes_tf.text = "hello"
        chart_data = CategoryChartData()
        chart_data.categories = ["Q1"]
        chart_data.add_series("A", (1,))  # pyright: ignore[reportUnknownMemberType]
        slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(4),
            Inches(3),
            chart_data,  # pyright: ignore[reportArgumentType]
        )

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        buf2 = io.BytesIO()

        with warnings.catch_warnings():
            warnings.simplefilter("error")
            prs2.save(buf2)


class DescribeFU3ZipEntriesAreUnique:
    """Pin invariant: saved zip contains no duplicate partnames."""

    @pytest.mark.parametrize("n_cycles", [1, 2, 3])
    def it_has_unique_zip_entry_names_after_n_save_reopen_cycles(self, n_cycles: int):
        buf = io.BytesIO()
        Presentation().save(buf)

        for _ in range(n_cycles):
            buf.seek(0)
            prs = Presentation(buf)
            buf = io.BytesIO()
            prs.save(buf)

        buf.seek(0)
        with zipfile.ZipFile(buf) as z:
            names = z.namelist()

        # A saved .pptx must contain no duplicate entries; stdlib zipfile
        # only warns (doesn't raise) on duplicate writes, so we check the
        # name list directly.
        assert len(names) == len(set(names)), sorted(
            name for name in names if names.count(name) > 1
        )

    def it_emits_exactly_one_slideLayout7_entry_on_default_deck(self):
        # The FU-3 symptom was specifically about slideLayout7.xml. Pin
        # that the default deck's slideLayout7 partname appears exactly
        # once after a save-reopen-save cycle.
        buf = io.BytesIO()
        Presentation().save(buf)
        buf.seek(0)
        prs = Presentation(buf)
        buf2 = io.BytesIO()
        prs.save(buf2)

        buf2.seek(0)
        with zipfile.ZipFile(buf2) as z:
            names = z.namelist()

        assert names.count("ppt/slideLayouts/slideLayout7.xml") == 1
