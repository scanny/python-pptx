"""Regression tests for FU-4 — `_rel_ref_count` must count `@r:embed` references.

When two shapes on the same slide share an image (same image-part `r:embed`),
deleting one of them must NOT drop the shared relationship. Before the fix
:meth:`pptx.opc.package.XmlPart._rel_ref_count` only scanned ``@r:id``
attributes, so two `a:blip/@r:embed` references were seen as zero, the rel was
prematurely removed on ``drop_rel``, and the second picture was left with a
dangling reference (or the image part was garbage-collected on save).
"""

# pyright: reportPrivateUsage=false

from __future__ import annotations

import io
from typing import TYPE_CHECKING, cast

from pptx import Presentation
from pptx.opc.package import XmlPart
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.oxml.xmlchemy import BaseOxmlElement
from pptx.shapes.picture import Picture
from pptx.util import Inches

from .unitutil.file import absjoin, test_file_dir

if TYPE_CHECKING:
    from pptx.presentation import Presentation as PresentationCls

IMAGE_FILE = absjoin(test_file_dir, "python-powered.png")


def _xml_part_with_element(xml: str) -> XmlPart:
    """Build a free-standing :class:`XmlPart` whose tree is the given XML."""
    part = XmlPart.__new__(XmlPart)
    part._element = cast(BaseOxmlElement, parse_xml(xml))
    return part


def _add_two_pictures_sharing_image(prs: "PresentationCls") -> tuple[Picture, Picture]:
    """Create a deck slide carrying two pictures that share an image part."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic1 = slide.shapes.add_picture(IMAGE_FILE, Inches(1), Inches(1), Inches(1), Inches(1))
    pic2 = slide.shapes.add_picture(IMAGE_FILE, Inches(3), Inches(1), Inches(1), Inches(1))
    return pic1, pic2


class DescribeRelRefCountEmbedCounting:
    """Unit-level checks on ``XmlPart._rel_ref_count``."""

    def it_counts_r_embed_references(self):
        # -- A `p:sld` tree with two `a:blip/@r:embed` attributes pointing at the same rId.
        xml = (
            "<p:sld %s>"
            "  <p:cSld><p:spTree>"
            '    <p:pic><p:blipFill><a:blip r:embed="rId9"/></p:blipFill></p:pic>'
            '    <p:pic><p:blipFill><a:blip r:embed="rId9"/></p:blipFill></p:pic>'
            "  </p:spTree></p:cSld>"
            "</p:sld>"
        ) % nsdecls("p", "a", "r")
        part = _xml_part_with_element(xml)

        # -- both references are counted ------------------------------------------------
        assert part._rel_ref_count("rId9") == 2

    def it_counts_r_link_references(self):
        xml = (
            "<p:sld %s>"
            "  <p:cSld><p:spTree>"
            "    <p:pic><p:nvPicPr><p:nvPr>"
            '      <a:videoFile r:link="rId7"/>'
            "    </p:nvPr></p:nvPicPr></p:pic>"
            "  </p:spTree></p:cSld>"
            "</p:sld>"
        ) % nsdecls("p", "a", "r")
        part = _xml_part_with_element(xml)

        assert part._rel_ref_count("rId7") == 1

    def it_counts_mixed_r_id_embed_link_together(self):
        xml = (
            "<p:sld %s>"
            "  <p:cSld><p:spTree>"
            '    <p:pic><p:blipFill><a:blip r:embed="rIdX"/></p:blipFill></p:pic>'
            "    <p:pic><p:nvPicPr><p:nvPr>"
            '      <a:videoFile r:link="rIdX"/>'
            "    </p:nvPr></p:nvPicPr></p:pic>"
            '    <a:hlinkClick r:id="rIdX"/>'
            "  </p:spTree></p:cSld>"
            "</p:sld>"
        ) % nsdecls("p", "a", "r")
        part = _xml_part_with_element(xml)

        assert part._rel_ref_count("rIdX") == 3

    def it_returns_zero_when_rId_absent(self):
        xml = (
            "<p:sld %s>"
            "  <p:cSld><p:spTree>"
            '    <p:pic><p:blipFill><a:blip r:embed="rId1"/></p:blipFill></p:pic>'
            "  </p:spTree></p:cSld>"
            "</p:sld>"
        ) % nsdecls("p", "a", "r")
        part = _xml_part_with_element(xml)

        assert part._rel_ref_count("rId999") == 0

    def it_only_counts_the_rId_being_queried(self):
        xml = (
            "<p:sld %s>"
            "  <p:cSld><p:spTree>"
            '    <p:pic><p:blipFill><a:blip r:embed="rIdA"/></p:blipFill></p:pic>'
            '    <p:pic><p:blipFill><a:blip r:embed="rIdB"/></p:blipFill></p:pic>'
            "  </p:spTree></p:cSld>"
            "</p:sld>"
        ) % nsdecls("p", "a", "r")
        part = _xml_part_with_element(xml)

        assert part._rel_ref_count("rIdA") == 1
        assert part._rel_ref_count("rIdB") == 1


class DescribeSharedImageRoundTrip:
    """End-to-end regression: shared image survives partial deletion + save."""

    def it_keeps_the_image_part_when_only_one_of_two_pictures_is_deleted(self):
        # -- build a deck with two pictures sharing the same image file ----------------
        prs = Presentation()
        pic1, pic2 = _add_two_pictures_sharing_image(prs)

        # -- both pictures share the same rId (get_or_add_image_part dedupes) ----------
        rId_a = pic1._pic.blip_rId
        rId_b = pic2._pic.blip_rId
        assert rId_a is not None, "test precondition: pic1 must have an rId"
        assert rId_a == rId_b, "test precondition: two pictures should share one rId"
        shared_rId = rId_a
        slide_part = pic1.part

        # -- delete one picture; this invokes drop_rel on the shared rId ---------------
        pic1.delete()

        # -- the rel must still be present because pic2 still refs it ------------------
        assert (
            shared_rId in slide_part.rels
        ), "shared rel was dropped prematurely — pic2 would be left dangling"

        # -- save and reopen; the image part must survive because pic2 still refs it --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        slide2 = prs2.slides[0]
        pics = [s for s in slide2.shapes if isinstance(s, Picture)]
        assert len(pics) == 1, f"expected 1 picture after delete+save, got {len(pics)}"

        # -- its image resolves (i.e. image part not GC'd) and has non-empty bytes ---
        surviving_image = pics[0].image
        assert surviving_image is not None, "image part was garbage-collected on save"
        assert len(surviving_image.blob) > 0

    def it_drops_the_image_part_when_all_referencing_pictures_are_deleted(self):
        # -- complementary case: when every reference is removed, the rel IS dropped --
        prs = Presentation()
        pic1, pic2 = _add_two_pictures_sharing_image(prs)
        shared_rId = pic1._pic.blip_rId
        assert shared_rId is not None
        slide_part = pic1.part

        pic1.delete()
        pic2.delete()

        assert (
            shared_rId not in slide_part.rels
        ), "rel should be dropped once every reference is removed"


class DescribeSharedImageXmlSurvivesSave:
    """Assert post-save XML still carries the ``r:embed`` for the surviving picture."""

    def it_preserves_r_embed_on_the_surviving_blip(self):
        prs = Presentation()
        pic1, _ = _add_two_pictures_sharing_image(prs)
        pic1.delete()

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        slide2 = prs2.slides[0]

        # -- the slide XML still contains a blip with an r:embed attribute ------------
        slide_xml = slide2.part.blob.decode()
        assert "r:embed=" in slide_xml

        # -- and that rId resolves to an actual image part in the package ------------
        pics = [s for s in slide2.shapes if isinstance(s, Picture)]
        assert len(pics) == 1
        img = pics[0].image
        assert img is not None
        assert img.blob  # non-empty bytes
