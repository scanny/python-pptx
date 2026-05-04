# pyright: reportPrivateUsage=false

"""FU-7 — ``GraphicFrame.delete()`` drops content-specific slide rels.

Before FU-7, :class:`~pptx.shapes.graphfrm.GraphicFrame` relied on the
base :meth:`~pptx.shapes.base.BaseShape.delete` which only removes the
``p:graphicFrame`` XML from the ``p:spTree`` — it did not drop any of
the slide-part relationships carried by the frame. Deleting a chart-
bearing, OLE-bearing, SmartArt-bearing, or 3D-model-bearing graphic
frame therefore left the slide's ``.rels`` file pointing at the
now-orphaned target parts, and the save-time
:meth:`~pptx.opc.package.OpcPackage.iter_parts` reachability walk
consequently preserved them in the output zip.

FU-7 adds a :meth:`GraphicFrame.delete` override that drops every rel
carried from inside the ``p:graphicFrame``:

* ``c:chart/@r:id`` — classic ChartPart (plus any embedded xlsx it
  transitively references)
* ``cx:chart/@r:id`` — Office 2016+ extended chartex part
* ``p:oleObj/@r:id`` — the embedded or linked OLE payload part
* ``a:blip/@r:embed`` inside ``p:oleObj/p:pic`` — the icon image for
  an OLE object shown as an icon
* the four ``dgm:relIds`` attributes (``r:dm`` / ``r:lo`` / ``r:qs`` /
  ``r:cs``) — the four parts of a SmartArt graphic
* ``am3d:model3D/@r:embed`` — the embedded 3D-model media part

Tables are self-contained XML with no external rels, so a table graphic
frame passes through ``super().delete()`` untouched.

These tests cover the behaviour end-to-end where the library already
authors the content kind (chart, OLE, table) and cover the enumeration
helper synthetically for kinds the library doesn't yet author
(chartex, SmartArt, 3D model).
"""

from __future__ import annotations

import io
import zipfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.shapes.graphfrm import GraphicFrame
from pptx.spec import (
    GRAPHIC_DATA_URI_CHART,
    GRAPHIC_DATA_URI_CHARTEX,
    GRAPHIC_DATA_URI_MODEL_3D,
    GRAPHIC_DATA_URI_SMART_ART,
    GRAPHIC_DATA_URI_TABLE,
)
from pptx.util import Inches

from .unitutil.cxml import element


def _zip_entries(buf: io.BytesIO) -> list[str]:
    """Return the list of zip-member names in the saved package ``buf``."""
    buf.seek(0)
    with zipfile.ZipFile(buf) as zf:
        return zf.namelist()


def _slide_rels_xml(buf: io.BytesIO, slide_name: str = "slide1") -> str:
    """Return the decoded ``ppt/slides/_rels/<slide_name>.xml.rels`` text."""
    buf.seek(0)
    with zipfile.ZipFile(buf) as zf:
        return zf.read(f"ppt/slides/_rels/{slide_name}.xml.rels").decode()


class DescribeFU7GraphicFrameDeleteRels:
    """FU-7: ``GraphicFrame.delete()`` releases its content-specific rels."""

    # -- classic chart (end-to-end via add_chart) ------------------------

    def it_drops_the_chart_rel_on_delete(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ("a", "b", "c")
        data.add_series("series", (1, 2, 3))
        shape = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(4),
            data,
        )
        before_rels = set(slide.part.rels)

        shape.delete()

        assert len(slide.part.rels) < len(before_rels)

        buf = io.BytesIO()
        prs.save(buf)
        entries = _zip_entries(buf)
        assert not any("ppt/charts/chart" in n for n in entries)
        assert not any("ppt/embeddings/" in n for n in entries)
        assert "charts/chart" not in _slide_rels_xml(buf)

    def it_drops_the_chart_rel_via_clear_shapes(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ("a", "b")
        data.add_series("series", (1, 2))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(4),
            data,
        )

        slide.shapes.clear()

        buf = io.BytesIO()
        prs.save(buf)
        entries = _zip_entries(buf)
        assert not any("ppt/charts/chart" in n for n in entries)
        assert not any("ppt/embeddings/" in n for n in entries)

    # -- OLE object (end-to-end via add_ole_object) ----------------------

    def it_drops_both_OLE_rels_on_delete(self):
        # -- add_ole_object produces a graphicFrame with two rels: the
        # -- oleObj @r:id (embedded object part) and the a:blip @r:embed
        # -- (icon image). Both must be dropped on delete.
        from pathlib import Path

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        test_files = Path(__file__).parent / "test_files"
        icon_path = test_files / "python-icon.jpeg"

        shape = slide.shapes.add_ole_object(
            io.BytesIO(b"dummy OLE payload"),
            prog_id="Excel.Sheet.12",
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(2),
            icon_file=str(icon_path),
            extension="bin",
        )
        before_rels = set(slide.part.rels)

        shape.delete()

        assert len(slide.part.rels) <= len(before_rels) - 2

        buf = io.BytesIO()
        prs.save(buf)
        entries = _zip_entries(buf)
        assert not any(n.startswith("ppt/embeddings/") for n in entries)
        assert not any(n.startswith("ppt/media/") for n in entries)

    # -- chartex, SmartArt, 3D (enumeration via cxml-constructed element) -

    def it_enumerates_the_chartex_rId_for_drop(self):
        graphicFrame = element(
            "p:graphicFrame/a:graphic/a:graphicData{uri=%s}/cx:chart{r:id=rId7}"
            % GRAPHIC_DATA_URI_CHARTEX
        )
        rIds = GraphicFrame(graphicFrame, None)._rIds  # type: ignore[arg-type]
        assert "rId7" in rIds

    def it_enumerates_the_four_SmartArt_rIds_for_drop(self):
        graphicFrame = element(
            "p:graphicFrame/a:graphic/a:graphicData{uri=%s}/"
            "dgm:relIds{r:dm=rId_dm,r:lo=rId_lo,r:qs=rId_qs,r:cs=rId_cs}"
            % GRAPHIC_DATA_URI_SMART_ART
        )
        rIds = set(GraphicFrame(graphicFrame, None)._rIds)  # type: ignore[arg-type]
        assert {"rId_dm", "rId_lo", "rId_qs", "rId_cs"}.issubset(rIds)

    def it_enumerates_the_model_3d_embed_rId_for_drop(self):
        graphicFrame = element(
            "p:graphicFrame/a:graphic/a:graphicData{uri=%s}/"
            "am3d:model3D{r:embed=rId9,ext=glb}" % GRAPHIC_DATA_URI_MODEL_3D
        )
        rIds = GraphicFrame(graphicFrame, None)._rIds  # type: ignore[arg-type]
        assert "rId9" in rIds

    def it_enumerates_both_OLE_rels_for_drop(self):
        graphicFrame = element(
            "p:graphicFrame/a:graphic/a:graphicData{uri=http://schemas.openxmlformats.org/"
            "presentationml/2006/ole}/p:oleObj{r:id=rId5}/p:pic/p:blipFill/a:blip{r:embed=rId6}"
        )
        rIds = set(GraphicFrame(graphicFrame, None)._rIds)  # type: ignore[arg-type]
        assert {"rId5", "rId6"}.issubset(rIds)

    # -- tables pass through untouched -----------------------------------

    def it_passes_through_for_a_table_graphic_frame(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table_shape = slide.shapes.add_table(
            rows=2, cols=2, left=Inches(1), top=Inches(1), width=Inches(4), height=Inches(2)
        )
        before_rels = set(slide.part.rels)

        table_shape.delete()

        assert set(slide.part.rels) == before_rels
        assert not any(s is table_shape for s in slide.shapes)

    def it_enumerates_no_rIds_for_a_table_graphic_frame(self):
        graphicFrame = element(
            "p:graphicFrame/a:graphic/a:graphicData{uri=%s}/a:tbl" % GRAPHIC_DATA_URI_TABLE
        )
        rIds = GraphicFrame(graphicFrame, None)._rIds  # type: ignore[arg-type]
        assert rIds == []

    def it_references_no_rIds_for_an_empty_classic_chart_frame(self):
        # -- graphicData with chart uri but no c:chart child (malformed but
        # -- shouldn't crash _rIds)
        graphicFrame = element(
            "p:graphicFrame/a:graphic/a:graphicData{uri=%s}" % GRAPHIC_DATA_URI_CHART
        )
        rIds = GraphicFrame(graphicFrame, None)._rIds  # type: ignore[arg-type]
        assert rIds == []
