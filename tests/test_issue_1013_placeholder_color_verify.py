# pyright: reportPrivateUsage=false

"""Regression test for issue #1013 — "color is empty when set via placeholder".

Issue #1013 (https://github.com/scanny/python-pptx/issues/1013) reports that
when a caller sets text on a placeholder and then queries the run's font
color, the color comes back empty. The symptom is a direct consequence of
how ``Font.color`` has always worked: it only inspects the run's own
``a:rPr/a:solidFill`` element, so a run whose color is inherited from the
enclosing paragraph, the shape's list-style, or the slide-master's
``p:txStyles`` reports ``color.type is None`` (and raises on ``color.rgb``)
even though PowerPoint renders that run in a perfectly concrete color.

Two common authoring patterns produce #1013 symptoms:

  1. Reading an untouched placeholder — e.g. ``slide.placeholders[1].text =
     "Hello"`` then ``run.font.color.rgb``. The run has no explicit color;
     the effective color comes from the master's ``p:bodyStyle`` (typically
     a ``schemeClr val="tx1"`` → ``(0,0,0)``). ``run.font.color.type`` is
     ``None``.
  2. Setting color at paragraph-level — e.g.
     ``paragraph.font.color.rgb = RGBColor(0xFF, 0, 0)``. The color lands
     on ``a:pPr/a:defRPr``, not on the run's ``a:rPr``, so each run inside
     that paragraph still reports ``color.type is None`` even though the
     paragraph (and therefore the run) obviously renders red.

Wave 6 #938 (``feat/issue-938-effective-font-color``, commit ``cec3edc1``)
shipped the public API that resolves #1013: :attr:`.Font.effective_color`
returns the |RGBColor| PowerPoint would actually render, walking the
inheritance chain the spec defines:

  1. the run's own ``a:rPr/a:solidFill``
  2. the enclosing paragraph's ``a:pPr/a:defRPr/a:solidFill``
  3. the enclosing text body's
     ``a:lstStyle/a:lvl{N}pPr/a:defRPr/a:solidFill`` (level-matched to the
     paragraph's ``@lvl``)
  4. the slide-master's ``p:txStyles`` (``bodyStyle`` / ``otherStyle`` /
     ``titleStyle``) for the matching level

Scheme colors (``a:schemeClr``) encountered anywhere in the walk are
resolved against the slide-master's theme, and any ``a:lumMod`` /
``a:lumOff`` siblings are applied. So a caller who was previously forced
to inspect ``a:rPr`` XML by hand can now write::

    rgb = run.font.effective_color

and get the concrete rendered color back, regardless of where in the style
hierarchy the color was actually set — closing the #1013 gap.

#1013 is therefore verified-and-closed by #938. This suite pins that
resolution against the two placeholder-flavoured scenarios the #1013
reporter's framing emphasizes:

  * an untouched placeholder whose effective color is the master's
    inherited default (explicit RGB and theme-scheme variants);
  * a placeholder where the caller set the color at the
    ``paragraph.font.color`` level (the most common "I set a color, why
    is color.rgb still empty?" case);

plus the adjacent sanity-checks that keep the API usable:

  * a run with an explicit ``run.font.color.rgb`` still reports that RGB
    on ``effective_color`` — the run's own setting beats any inherited
    color, matching how PowerPoint resolves the chain;
  * the same ``effective_color`` survives ``Presentation.save`` + reopen
    (the reporter's deck was an already-saved ``.pptx``);
  * the legacy ``font.color`` contract is **unchanged** — ``color.type is
    None`` and ``color.rgb`` still raises on an inherited-color run, so
    callers who need to distinguish "explicit" from "inherited" can still
    do so. #1013 adds a new read path; it does not mutate the old one.

Any future regression that drops one of the inheritance steps (paragraph
``defRPr``, ``lstStyle``, or the master's ``p:txStyles`` fallback), or
that fails to resolve scheme colors when walking the chain from a
placeholder, will reproduce the #1013 symptom and be caught here.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Pt


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of ``PartFactory.part_type_for``.

    Mirrors the guard used by ``tests/test_issue_720_ea_font_verify.py`` and
    friends — ``Presentation(stream)`` reopen dispatches through
    ``PartFactory.part_type_for``, which earlier test modules overwrite with
    mocks and do not restore.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


def _retint_body_style_lvl1(prs, scheme_val):
    """Rewrite the master's ``p:bodyStyle/a:lvl1pPr/a:defRPr/a:solidFill``
    to point at ``scheme_val`` (e.g. ``"accent1"``).

    The default template's ``p:bodyStyle`` carries an
    ``a:schemeClr val="tx1"`` that resolves to black — useful for the
    inherited-black scenario but not distinctive enough on its own. This
    helper lets a test pin the full "scheme color walked from master,
    resolved against theme" path with an accent color that's unambiguously
    not black.
    """
    master_elt = prs.slide_masters[0]._element
    txStyles = master_elt.find(qn("p:txStyles"))
    assert txStyles is not None
    bodyStyle = txStyles.find(qn("p:bodyStyle"))
    assert bodyStyle is not None
    lvl1 = bodyStyle.find(qn("a:lvl1pPr"))
    assert lvl1 is not None
    defRPr = lvl1.find(qn("a:defRPr"))
    assert defRPr is not None
    solidFill = defRPr.find(qn("a:solidFill"))
    assert solidFill is not None
    schemeClr = solidFill.find(qn("a:schemeClr"))
    assert schemeClr is not None
    schemeClr.set("val", scheme_val)


class DescribeIssue1013PlaceholderColorVerify:
    """#1013 placeholder color-read verify-and-close via #938.

    Pins that :attr:`.Font.effective_color` resolves the placeholder
    inheritance chain — the read path the #1013 reporter was looking for
    when ``font.color.rgb`` came back empty.
    """

    # -- API surface ------------------------------------------------------

    def it_exposes_effective_color_on_Font(self):
        """Guard: the public API that resolves #1013 is present.

        A caller reaching for a colour-read on an inherited-colour run
        needs :attr:`.Font.effective_color` — a refactor that removed or
        renamed it would silently reintroduce #1013.
        """
        from pptx.text.text import Font

        assert hasattr(Font, "effective_color")

    # -- untouched placeholder (inherited from master) -------------------

    def it_reads_the_inherited_color_of_an_untouched_title_placeholder(self):
        """The simplest #1013 scenario — a fresh title placeholder.

        No color was ever assigned; the rendered color is the master's
        inherited text colour (a ``schemeClr val="tx1"`` resolving to
        black). :attr:`.Font.color` reports ``type is None`` (the legacy
        contract), but :attr:`.Font.effective_color` now returns the
        concrete ``RGBColor(0, 0, 0)`` the #1013 reporter expected.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.placeholders[0]
        title.text_frame.text = "Hello"
        run = title.text_frame.paragraphs[0].runs[0]

        # -- legacy contract stays "empty" on inherited colour --
        assert run.font.color.type is None

        # -- #938 resolves the inheritance chain all the way to the master --
        assert run.font.effective_color == RGBColor(0, 0, 0)

    def it_reads_the_inherited_color_of_an_untouched_body_placeholder(self):
        """The content-placeholder variant of the same scenario.

        A content/body placeholder's inherited color comes from the
        master's ``p:bodyStyle`` rather than ``p:titleStyle``. The walk
        must consult that fall-back too.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        body.text_frame.text = "Hello body"
        run = body.text_frame.paragraphs[0].runs[0]

        assert run.font.color.type is None
        assert run.font.effective_color == RGBColor(0, 0, 0)

    def it_resolves_a_theme_scheme_color_through_the_master_body_style(self):
        """Scheme-color resolution on the master → theme walk.

        Retargeting the master's ``bodyStyle/lvl1pPr/defRPr/solidFill``
        from ``tx1`` to ``accent1`` must flow through to
        :attr:`.Font.effective_color` on a body placeholder — any failure
        in the scheme-colour lookup (or in the master→theme hop) would
        reproduce the #1013 empty-color symptom.
        """
        prs = Presentation()
        # -- default theme's accent1 == (79, 129, 189) --
        _retint_body_style_lvl1(prs, "accent1")

        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        body.text_frame.text = "Accent1 text"
        run = body.text_frame.paragraphs[0].runs[0]

        assert run.font.color.type is None
        assert run.font.effective_color == RGBColor(79, 129, 189)

    # -- paragraph-level color (pPr/defRPr) -------------------------------

    def it_resolves_color_set_at_paragraph_font_color_level(self):
        """The canonical #1013 "I set a color, why is it empty?" scenario.

        Assigning :attr:`.Font.color` on the paragraph — i.e.
        ``paragraph.font.color.rgb = …`` — writes the solidFill onto the
        paragraph's ``a:pPr/a:defRPr``, not onto any run's ``a:rPr``. So
        the runs inside that paragraph still report ``color.type is None``
        (the legacy contract) while obviously rendering the paragraph's
        color. :attr:`.Font.effective_color` walks one step up and returns
        the concrete RGB.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        body.text_frame.text = "Hello red"
        paragraph = body.text_frame.paragraphs[0]
        paragraph.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

        run = paragraph.runs[0]
        # -- run's own rPr is untouched; legacy .color is empty --
        assert run.font.color.type is None
        # -- paragraph-level color is the source of truth --
        assert paragraph.font.color.rgb == RGBColor(0xFF, 0x00, 0x00)
        # -- and now readable from the run's effective_color --
        assert run.font.effective_color == RGBColor(0xFF, 0x00, 0x00)

    # -- explicit run color beats paragraph/master ------------------------

    def it_returns_the_run_rPr_color_when_set_directly_on_the_run(self):
        """Sanity: an explicit run color beats inherited sources.

        The #1013 fix does not change the win-order — a run that carries
        its own ``a:rPr/a:solidFill`` renders in that color regardless of
        what's inherited from the paragraph or master, and
        :attr:`.Font.effective_color` returns the run's own colour.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        body.text_frame.text = "Hello blue"
        paragraph = body.text_frame.paragraphs[0]

        # -- inherited (master default) would be black; paragraph says red… --
        paragraph.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
        # -- …but the run's own explicit color is blue and wins. --
        run = paragraph.runs[0]
        run.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)

        assert run.font.color.type is not None
        assert run.font.color.rgb == RGBColor(0x00, 0x00, 0xFF)
        assert run.font.effective_color == RGBColor(0x00, 0x00, 0xFF)

    # -- round-trip through save + reopen ---------------------------------

    def it_round_trips_effective_color_through_save_and_reload(self, _restore_part_factory):
        """The #1013 reporter's deck was already saved; the fix must
        survive save + reopen.

        Authors a placeholder whose color is set at paragraph-level (the
        classic #1013 trap), saves to a ``.pptx``, reopens, and confirms
        the reloaded run still reports the right ``effective_color``
        with no explicit run rPr.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        body.text_frame.text = "Paragraph red"
        p0 = body.text_frame.paragraphs[0]
        p0.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
        # -- set size too so we have proof the placeholder text survived --
        p0.font.size = Pt(24)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        body2 = prs2.slides[0].placeholders[1]
        run2 = body2.text_frame.paragraphs[0].runs[0]

        # -- the text is there --
        assert run2.text == "Paragraph red"
        # -- legacy contract still says "inherited" on the run --
        assert run2.font.color.type is None
        # -- but effective_color walks up to the paragraph and returns red --
        assert run2.font.effective_color == RGBColor(0xFF, 0x00, 0x00)

    def it_round_trips_inherited_master_color_through_save_and_reload(self, _restore_part_factory):
        """An untouched placeholder's inherited colour survives save+reopen.

        Complements the paragraph-level round-trip by pinning the
        master-walk branch: the ``p:bodyStyle`` colour must still resolve
        correctly on a reopened deck with no explicit colour anywhere on
        the run or paragraph.
        """
        prs = Presentation()
        _retint_body_style_lvl1(prs, "accent1")
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        body.text_frame.text = "Inherited accent"

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        body2 = prs2.slides[0].placeholders[1]
        run2 = body2.text_frame.paragraphs[0].runs[0]

        assert run2.text == "Inherited accent"
        assert run2.font.color.type is None
        assert run2.font.effective_color == RGBColor(79, 129, 189)

    # -- XML-shape invariants ---------------------------------------------

    def it_does_not_write_an_rPr_solidFill_when_color_is_set_on_paragraph(self):
        """Pin the XML shape: paragraph-level colour lives on ``a:pPr/a:defRPr``.

        A regression that silently shadowed the paragraph setting onto
        every run's ``a:rPr`` would make ``color.type`` non-None on the
        run and change the ``font.color`` contract. That would hide #1013
        by accident rather than resolving it through a read-side API —
        not the fix the project committed to.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        body.text_frame.text = "Hello"
        p0 = body.text_frame.paragraphs[0]
        p0.font.color.rgb = RGBColor(0x12, 0x34, 0x56)

        run = p0.runs[0]
        rPr = run._r.get_or_add_rPr()
        # -- run's own a:rPr carries no solidFill --
        assert rPr.find(qn("a:solidFill")) is None

        # -- but the paragraph's a:pPr/a:defRPr does, with the right RGB --
        pPr = p0._pPr
        assert pPr is not None
        defRPr = pPr.find(qn("a:defRPr"))
        assert defRPr is not None
        solidFill = defRPr.find(qn("a:solidFill"))
        assert solidFill is not None
        srgbClr = solidFill.find(qn("a:srgbClr"))
        assert srgbClr is not None
        assert srgbClr.get("val").lower() == "123456"

    # -- legacy contract unchanged ----------------------------------------

    def it_preserves_the_legacy_color_rgb_raises_contract_on_inherited_runs(self):
        """Legacy ``font.color`` behaviour is intentionally **unchanged**.

        Callers who rely on ``color.type is None`` / ``color.rgb`` raising
        to distinguish explicit from inherited color still see that.
        :attr:`.Font.effective_color` is a new, additive read path —
        not a silent mutation of the old one.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        body.text_frame.text = "Hello"
        run = body.text_frame.paragraphs[0].runs[0]

        assert run.font.color.type is None
        with pytest.raises(AttributeError):
            _ = run.font.color.rgb

        # -- effective_color, by contrast, always resolves to a concrete RGB --
        assert isinstance(run.font.effective_color, RGBColor)

    # -- diagnostic: XML confirms the inheritance chain -------------------

    def it_resolves_a_tx1_tinted_body_style_via_lumMod_lumOff(self):
        """Diagnostic: the inheritance walk applies ``a:lumMod``/``a:lumOff``.

        #1013's scope is "color I can see but can't read" — one common
        authored shape that previously tripped the old ``font.color``
        path is a tinted/shaded master style: ``<a:schemeClr val="tx1">
        <a:lumMod val="50000"/><a:lumOff val="50000"/></a:schemeClr>``
        (PowerPoint's "Lighter 50%" variant on the tx1 text colour).

        :attr:`.Font.effective_color` delegates tint/shade resolution to
        :meth:`.ColorFormat.to_rgb`, so this scenario must round-trip —
        ``tx1 = (0,0,0)`` luminance-adjusted to 50% yields
        ``(0x7F, 0x7F, 0x7F)``. A regression that stripped the lumMod/
        lumOff siblings during the walk would fall back to raw ``tx1``
        (black) and reintroduce the #1013 symptom for tinted styles.
        """
        prs = Presentation()
        # -- rewrite body-style defRPr's solidFill to tx1 + lumMod/lumOff --
        master_elt = prs.slide_masters[0]._element
        txStyles = master_elt.find(qn("p:txStyles"))
        assert txStyles is not None
        bodyStyle = txStyles.find(qn("p:bodyStyle"))
        assert bodyStyle is not None
        lvl1 = bodyStyle.find(qn("a:lvl1pPr"))
        assert lvl1 is not None
        defRPr = lvl1.find(qn("a:defRPr"))
        assert defRPr is not None
        solidFill = defRPr.find(qn("a:solidFill"))
        assert solidFill is not None
        schemeClr = solidFill.find(qn("a:schemeClr"))
        assert schemeClr is not None
        # -- inject <a:lumMod val="50000"/><a:lumOff val="50000"/> --
        from pptx.oxml import parse_xml

        a_ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        schemeClr.append(parse_xml(f'<a:lumMod {a_ns} val="50000"/>'))
        schemeClr.append(parse_xml(f'<a:lumOff {a_ns} val="50000"/>'))

        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        body.text_frame.text = "Tinted tx1"
        run = body.text_frame.paragraphs[0].runs[0]

        assert run.font.color.type is None
        # -- tx1 is (0,0,0); 50% luminance lift -> mid-grey (~ 0x80) --
        rgb = run.font.effective_color
        assert rgb is not None
        assert rgb != RGBColor(0, 0, 0)  # -- lumMod/lumOff must have taken effect --
        # -- allow small rounding wiggle: each channel sits near 0x80 --
        for channel in rgb:
            assert 0x70 <= channel <= 0x90
