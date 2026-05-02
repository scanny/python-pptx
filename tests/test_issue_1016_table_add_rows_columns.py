# pyright: reportPrivateUsage=false

"""Regression test for issue #1016 — ``table.columns.add()`` / ``table.rows.add()``.

Issue #1016 (https://github.com/scanny/python-pptx/issues/1016) reported that
``table.columns.add()`` and ``table.rows.add()`` — APIs the reporter had
relied on in python-pptx 0.6.23 for an Excel-to-PowerPoint table generator —
produced ``AttributeError: '_ColumnCollection' object has no attribute 'add'``
(and the same for ``_RowCollection``) in v1.0.0. The reporter cited
PowerPoint's own ``Rows.Add`` / ``Columns.Add`` methods
(https://learn.microsoft.com/en-us/previous-versions/office/developer/office-2010/ff745415(v=office.14))
as the expected shape of the API and asked whether removal was intentional.

That API had never actually existed in any released python-pptx version; the
0.6.x ``table.rows.add()`` support the reporter remembered was a local patch
or a misremembering. The fork shipped the table-mutation API across three
independent PRs:

* #832 (Wave 2, ``feat/issue-832-table-row-add``) added
  :meth:`._RowCollection.add` — append a row at the bottom of the table, one
  cell per column, inheriting the last-row height (or defaulting to
  370,840 EMU ≈ 0.4 inch when the table has no rows). The containing
  graphic-frame height is recomputed.
* #837 (Wave 3, ``feat/issue-837-table-row-delete``) added the delete
  counterpart — :meth:`._Row.delete` and :meth:`._RowCollection.remove`.
* #895 (Wave 6, ``feat/issue-895-table-column-mutation``) added the column
  variant — :meth:`._ColumnCollection.add` (append a new ``a:gridCol`` and
  an empty ``a:tc`` in every row, inheriting the last-column width or
  defaulting to 914,400 EMU = 1 inch), :meth:`._Column.delete`, and
  :meth:`._ColumnCollection.remove`.

Together #832 + #895 satisfy the #1016 reporter's exact ask — `table.columns.add()`
and `table.rows.add()` on an existing table — and #837 closes the loop with a
symmetric delete. #1016 is therefore a duplicate of the #832 / #895 feature
PRs (with #837 bracketing the delete path) and is verified-and-closed by this
suite.

The scenarios below pin the exact Excel-to-PowerPoint workflow the reporter
described: author a table, then call ``rows.add()`` / ``columns.add()``
procedurally while reading the spreadsheet, and save. Both directions are
exercised, including save + reopen round-trip, to guard against silent
regression that would reopen the ticket.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.table import _Column, _ColumnCollection, _Row, _RowCollection
from pptx.util import Emu, Inches


@pytest.fixture
def _restore_part_factory():
    """Guard against pollution from tests that mutate `PartFactory.part_type_for`.

    See the identical fixture in ``tests/test_issue_791_table_row_delete.py``
    and ``tests/test_issue_937_table_cell_layout.py`` for the rationale —
    other regression suites mock out the part registry and, if their
    teardown leaks, this suite's ``Presentation.save`` + reopen round-trips
    encounter a ``Mock`` in place of ``SlidePart`` and explode with
    ``AttributeError: Mock object has no attribute 'slide'``.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


def _new_2x2_table_shape(prs):
    """Return the graphic-frame shape for a fresh 2x2 table on a blank slide.

    The 2x2 starting size mirrors the #1016 reporter's "author a small table,
    then grow it as we read the spreadsheet" workflow more closely than the
    7x3 shape used by the #791 suite.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_table(
        rows=2,
        cols=2,
        left=Inches(1),
        top=Inches(1),
        width=Inches(4),
        height=Inches(2),
    )
    table = shape.table
    for r in range(2):
        for c in range(2):
            table.cell(r, c).text = f"r{r}c{c}"
    return shape


class DescribeIssue1016TableAddRowsColumns(object):
    """#1016 ``table.columns.add()`` / ``table.rows.add()`` verify-and-close.

    The three wave PRs that together resolve #1016:

    * Wave 2 #832 — :meth:`._RowCollection.add`
    * Wave 3 #837 — :meth:`._Row.delete` / :meth:`._RowCollection.remove`
    * Wave 6 #895 — :meth:`._ColumnCollection.add` /
      :meth:`._Column.delete` / :meth:`._ColumnCollection.remove`
    """

    # -- public-surface smoke: the attributes the reporter got
    # -- ``AttributeError`` on must exist on the class -----------------------

    def it_exposes_add_on_RowCollection(self):
        assert hasattr(_RowCollection, "add")
        assert callable(_RowCollection.add)

    def it_exposes_add_on_ColumnCollection(self):
        assert hasattr(_ColumnCollection, "add")
        assert callable(_ColumnCollection.add)

    def it_exposes_remove_on_RowCollection(self):
        # -- the delete counterpart landed by #837
        assert hasattr(_RowCollection, "remove")
        assert callable(_RowCollection.remove)

    def it_exposes_delete_on_Row(self):
        assert hasattr(_Row, "delete")
        assert callable(_Row.delete)

    def it_exposes_remove_on_ColumnCollection(self):
        # -- the delete counterpart landed by #895
        assert hasattr(_ColumnCollection, "remove")
        assert callable(_ColumnCollection.remove)

    def it_exposes_delete_on_Column(self):
        assert hasattr(_Column, "delete")
        assert callable(_Column.delete)

    # -- reporter's rows.add() workflow ------------------------------------

    def it_adds_a_row_to_an_existing_table(self, _restore_part_factory):
        # -- the reporter's call: `table.rows.add()` on an existing table.
        prs = Presentation()
        shape = _new_2x2_table_shape(prs)
        table = shape.table

        new_row = table.rows.add()

        assert isinstance(new_row, _Row)
        assert len(table.rows) == 3
        # -- the new row has one cell per column
        assert len(new_row.cells) == 2
        # -- existing row content is undisturbed
        assert [table.cell(r, c).text for r in range(2) for c in range(2)] == [
            "r0c0",
            "r0c1",
            "r1c0",
            "r1c1",
        ]
        # -- new row cells are empty
        assert new_row.cells[0].text == ""
        assert new_row.cells[1].text == ""

    def it_inherits_row_height_from_the_prior_last_row_by_default(self, _restore_part_factory):
        prs = Presentation()
        shape = _new_2x2_table_shape(prs)
        table = shape.table
        prior_last_height = Emu(table.rows[len(table.rows) - 1].height)

        new_row = table.rows.add()

        assert Emu(new_row.height) == prior_last_height

    def it_accepts_an_explicit_row_height(self, _restore_part_factory):
        prs = Presentation()
        shape = _new_2x2_table_shape(prs)
        table = shape.table

        new_row = table.rows.add(height=Inches(1))

        assert Emu(new_row.height) == Inches(1)

    def it_grows_the_graphic_frame_height_by_the_new_row_height(self, _restore_part_factory):
        prs = Presentation()
        shape = _new_2x2_table_shape(prs)
        table = shape.table
        height_before = Emu(shape.height)

        new_row = table.rows.add()

        assert Emu(shape.height) == height_before + Emu(new_row.height)
        assert Emu(shape.height) == Emu(sum(row.height for row in table.rows))

    # -- reporter's columns.add() workflow ---------------------------------

    def it_adds_a_column_to_an_existing_table(self, _restore_part_factory):
        # -- the reporter's call: `table.columns.add()` on an existing table.
        prs = Presentation()
        shape = _new_2x2_table_shape(prs)
        table = shape.table

        new_col = table.columns.add()

        assert isinstance(new_col, _Column)
        assert len(table.columns) == 3
        # -- the new column is represented by an empty `a:tc` in every row
        assert all(len(row.cells) == 3 for row in table.rows)
        # -- existing cell content is undisturbed
        assert [table.cell(r, c).text for r in range(2) for c in range(2)] == [
            "r0c0",
            "r0c1",
            "r1c0",
            "r1c1",
        ]
        # -- new column cells are empty
        for r in range(2):
            assert table.cell(r, 2).text == ""

    def it_inherits_column_width_from_the_prior_last_column_by_default(self, _restore_part_factory):
        prs = Presentation()
        shape = _new_2x2_table_shape(prs)
        table = shape.table
        prior_last_width = Emu(table.columns[len(table.columns) - 1].width)

        new_col = table.columns.add()

        assert Emu(new_col.width) == prior_last_width

    def it_accepts_an_explicit_column_width(self, _restore_part_factory):
        prs = Presentation()
        shape = _new_2x2_table_shape(prs)
        table = shape.table

        new_col = table.columns.add(width=Inches(2))

        assert Emu(new_col.width) == Inches(2)

    def it_grows_the_graphic_frame_width_by_the_new_column_width(self, _restore_part_factory):
        prs = Presentation()
        shape = _new_2x2_table_shape(prs)
        table = shape.table
        width_before = Emu(shape.width)

        new_col = table.columns.add()

        assert Emu(shape.width) == width_before + Emu(new_col.width)
        assert Emu(shape.width) == Emu(sum(col.width for col in table.columns))

    # -- the reporter's Excel-driven "grow procedurally" scenario ----------

    def it_supports_the_reporter_excel_driven_grow_workflow(self, _restore_part_factory):
        # -- simulate the #1016 reporter's "read the spreadsheet, call
        # -- `table.rows.add()` / `table.columns.add()` as we go" loop.
        # -- Start with a minimal 1x1 table (the reporter creates then
        # -- grows), append 3 extra rows and 2 extra columns procedurally,
        # -- and fill in the body. The final shape should be a 4x3 table
        # -- with every cell labelled, as a proxy for "the spreadsheet
        # -- we just transferred".
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_table(
            rows=1,
            cols=1,
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(0.5),
        )
        table = shape.table

        # -- grow horizontally first (procedural columns-from-spreadsheet)
        for _ in range(2):
            table.columns.add()
        # -- then grow vertically (procedural rows-from-spreadsheet)
        for _ in range(3):
            table.rows.add()

        assert len(table.rows) == 4
        assert len(table.columns) == 3
        # -- fill in each cell, mirroring the reporter's write loop
        for r in range(4):
            for c in range(3):
                table.cell(r, c).text = f"R{r}C{c}"
        # -- body survives
        assert [[table.cell(r, c).text for c in range(3)] for r in range(4)] == [
            [f"R{r}C{c}" for c in range(3)] for r in range(4)
        ]

    # -- round-trip: save + reopen -----------------------------------------

    def it_round_trips_rows_add_through_save_reopen(self, _restore_part_factory):
        prs = Presentation()
        shape = _new_2x2_table_shape(prs)
        table = shape.table
        new_row = table.rows.add(height=Inches(1))
        new_row.cells[0].text = "new-row-cell-0"
        new_row.cells[1].text = "new-row-cell-1"
        expected_height = shape.height

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        shape2 = next(s for s in prs2.slides[0].shapes if s.has_table)
        table2 = shape2.table

        assert len(table2.rows) == 3
        assert len(table2.columns) == 2
        assert table2.cell(2, 0).text == "new-row-cell-0"
        assert table2.cell(2, 1).text == "new-row-cell-1"
        # -- graphic-frame height is the authored sum-of-row-heights
        assert shape2.height == expected_height

    def it_round_trips_columns_add_through_save_reopen(self, _restore_part_factory):
        prs = Presentation()
        shape = _new_2x2_table_shape(prs)
        table = shape.table
        table.columns.add(width=Inches(2))
        # -- fill the new column so we can spot it after reopen
        table.cell(0, 2).text = "new-col-cell-0"
        table.cell(1, 2).text = "new-col-cell-1"
        expected_width = shape.width

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        shape2 = next(s for s in prs2.slides[0].shapes if s.has_table)
        table2 = shape2.table

        assert len(table2.rows) == 2
        assert len(table2.columns) == 3
        assert table2.cell(0, 2).text == "new-col-cell-0"
        assert table2.cell(1, 2).text == "new-col-cell-1"
        # -- graphic-frame width is the authored sum-of-column-widths
        assert shape2.width == expected_width

    def it_round_trips_a_grow_then_shrink_sequence(self, _restore_part_factory):
        # -- the #837 / #895 delete counterparts must round-trip too, since
        # -- a real Excel-to-PowerPoint script will sometimes pre-grow and
        # -- trim. This guards the "add a column, then delete the first
        # -- column, then add a row" kind of mixed sequence against regression.
        prs = Presentation()
        shape = _new_2x2_table_shape(prs)
        table = shape.table

        # -- grow to 2x3, then delete the leftmost column (#895 delete)
        table.columns.add()
        table.columns[0].delete()
        # -- grow to 3x2, then delete the first row (#837 delete)
        table.rows.add()
        table.rows[0].delete()
        # -- label the surviving cells so we can verify after reopen
        for r in range(len(table.rows)):
            for c in range(len(table.columns)):
                table.cell(r, c).text = f"S{r}{c}"

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        table2 = next(s for s in prs2.slides[0].shapes if s.has_table).table

        assert len(table2.rows) == 2
        assert len(table2.columns) == 2
        assert [[table2.cell(r, c).text for c in range(2)] for r in range(2)] == [
            ["S00", "S01"],
            ["S10", "S11"],
        ]
