# pyright: reportPrivateUsage=false

"""Regression test for issue #1017 — "How do I adjust an elbow connector?".

Issue #1017 (https://github.com/scanny/python-pptx/issues/1017) asks how to
bend an elbow / bent connector through python-pptx. The reporter wanted the
equivalent of PowerPoint's yellow adjustment handle on a ``bentConnector3``
— shift the "elbow" of the connector along its long axis from the default
midpoint (50%) to some other normalized position.

This behaviour is delivered by the ``Connector.adjustments`` collection
added on branch ``feat/issue-946-connector-adjustments`` (merged into
integration/waves-1-2 as commit ``ac37f3bc``). Each ``bentConnectorN`` /
``curvedConnectorN`` preset exposes its ``a:avLst`` adjustments as a
mutable collection — ``connector.adjustments[0] = 0.25`` rewrites the
``a:gd`` child of ``a:avLst`` and is round-tripped across save / reopen.

This regression test exercises the flow #1017's reporter specifically
asked for: take a real ``.pptx`` file that contains a bent-connector
shape, programmatically bend it by assigning to ``adjustments[0]``, save
the package, reopen it, and confirm the new bend position survived.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.shapes.connector import Connector, ConnectorAdjustmentCollection
from pptx.util import Emu


# -- A minimal "real" bent-connector PPTX built by python-pptx itself, then --
# -- re-loaded so the saved ZIP is what we're reading. This mirrors the  --
# -- reporter's setup (they load a .pptx authored in PowerPoint).        --
def _make_bent_connector_pptx() -> io.BytesIO:
    """Return an in-memory .pptx with a single ``bentConnector3`` shape."""
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    # -- Emu(914400) == 1 inch; begin at (1", 1"), end at (3", 3") --
    slide.shapes.add_connector(
        MSO_CONNECTOR.ELBOW, Emu(914400), Emu(914400), Emu(2743200), Emu(2743200)
    )
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


class DescribeIssue1017ElbowConnectorAdjust(object):
    """End-to-end coverage of the #1017 flow against a real saved .pptx."""

    def it_bends_an_elbow_connector_via_the_adjustments_api(self):
        """Reporter's core ask: bend the elbow and have it stick after save.

        Loads a package that contains a ``bentConnector3``, asserts the
        default adjustment reads as 0.5, bends it to 0.25, saves, reopens,
        and confirms the new value survived — both as the collection value
        and as a literal ``a:gd`` child of ``a:avLst`` in the saved XML.
        """
        # -- load the baseline package (mimics loading a real authored .pptx) --
        prs = Presentation(_make_bent_connector_pptx())
        slide = prs.slides[0]
        connector = slide.shapes[0]

        # -- sanity: a:prstGeom[@prst=bentConnector3] has exactly one adjustment --
        assert isinstance(connector, Connector)
        assert connector.shape_type == MSO_SHAPE_TYPE.LINE
        adjustments = connector.adjustments
        assert isinstance(adjustments, ConnectorAdjustmentCollection)
        assert len(adjustments) == 1
        # -- default adj1 for bentConnector3 is 50000 (50%) --
        assert adjustments[0] == 0.5

        # -- bend it 25% along the long axis (the "#1017 one-liner") --
        connector.adjustments[0] = 0.25
        assert connector.adjustments[0] == 0.25

        # -- round-trip: save, reopen, confirm the bend persisted --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        reopened = Presentation(buf)
        reopened_connector = reopened.slides[0].shapes[0]
        assert isinstance(reopened_connector, Connector)
        assert len(reopened_connector.adjustments) == 1
        assert reopened_connector.adjustments[0] == 0.25

        # -- and the XML actually carries an a:gd with the right fmla --
        spPr = reopened_connector._element.spPr
        prstGeom = spPr.prstGeom
        assert prstGeom is not None
        gds = prstGeom.findall(qn("a:avLst") + "/" + qn("a:gd"))
        assert len(gds) == 1
        assert gds[0].get("name") == "adj1"
        assert gds[0].get("fmla") == "val 25000"

    def it_reads_a_prebent_connector_from_an_authored_avLst(self):
        """Also covers the "I already bent it in PowerPoint" direction:

        if the incoming ``a:avLst`` already carries a non-default
        ``a:gd[@name=adj1]``, ``Connector.adjustments[0]`` reports that
        authored value (not the preset default).
        """
        # -- build a fresh deck, then inject a pre-authored avLst directly --
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.ELBOW, Emu(0), Emu(0), Emu(914400), Emu(914400)
        )
        # -- author adj1=33333 (≈33.3%) directly on a:avLst, the way     --
        # -- PowerPoint would if the user dragged the yellow handle      --
        prstGeom = connector._element.spPr.prstGeom
        assert prstGeom is not None
        prstGeom.rewrite_guides((("adj1", 33333),))

        # -- round-trip before reading, so we're reading the saved XML --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        reopened = Presentation(buf)
        reopened_connector = reopened.slides[0].shapes[0]
        assert isinstance(reopened_connector, Connector)
        assert reopened_connector.adjustments[0] == pytest.approx(0.33333)

    def it_round_trips_multiple_adjustments_on_a_five_segment_connector(
        self
    ):
        """A ``bentConnector5`` has three adjustments (adj1/2/3). Setting
        just one of them must not clobber the other two — all three
        values must be written to XML because the python model manages
        ``a:avLst`` as a whole.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        # -- MSO_CONNECTOR only exposes STRAIGHT / ELBOW / CURVE so we --
        # -- start from an elbow and manually switch prst so we don't  --
        # -- depend on bentConnector5 being an authoring preset.       --
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.ELBOW, Emu(0), Emu(0), Emu(914400), Emu(914400)
        )
        connector._element.spPr.prstGeom.set("prst", "bentConnector5")

        # -- defaults for bentConnector5: adj1=adj2=adj3=50000 (0.5) --
        adjustments = connector.adjustments
        assert len(adjustments) == 3
        assert [adjustments[i] for i in range(3)] == [0.5, 0.5, 0.5]

        # -- set only adj2; adj1 and adj3 must land in XML at default --
        adjustments[1] = 0.2

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        reopened = Presentation(buf)
        reopened_connector = reopened.slides[0].shapes[0]
        assert isinstance(reopened_connector, Connector)
        r_adj = reopened_connector.adjustments
        assert [r_adj[i] for i in range(3)] == [0.5, 0.2, 0.5]
