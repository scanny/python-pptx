"""Regression tests for issue #1022 — `ActionSetting.screen_tip` round-trip.

The reporter observed that setting ``click_action.screen_tip`` stores the
tooltip in XML but PowerPoint does not show it on hover.

Root cause: PowerPoint only *displays* a ScreenTip when the hyperlink
element carries an actionable target (URL, slide jump, embedded sound, or
``ppaction://`` action verb). A lone ``<a:hlinkClick tooltip=".."/>``
element with no ``r:id`` / ``action`` is spec-valid but behaviorally inert
— PowerPoint ignores the tooltip at display time. This matches
PowerPoint's own Insert Hyperlink UI, which will not accept a ScreenTip
without a target.

These tests pin the XML-level behavior so the library side of the
contract (what actually gets written) stays stable, and verify the
documented workaround (pair the tooltip with a URL or slide jump)
produces a PowerPoint-renderable tooltip.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches


class DescribeIssue1022ScreenTipRoundTrip(object):
    """Regression suite for issue #1022 — screen_tip visibility in PowerPoint."""

    def it_writes_tooltip_only_hlinkClick_when_no_URL_is_set(self):
        # -- Setting `screen_tip` with no URL writes a spec-valid
        # -- <a:hlinkClick tooltip="..."/> with no r:id. This is what
        # -- PowerPoint ignores at display time.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )

        shape.click_action.screen_tip = "Tooltip only"

        hlinkClick = shape._element.nvSpPr.cNvPr.hlinkClick
        assert hlinkClick is not None
        assert hlinkClick.tooltip == "Tooltip only"
        assert hlinkClick.rId is None or hlinkClick.rId == ""
        assert hlinkClick.action is None

    def it_pairs_tooltip_with_URL_when_both_are_set(self):
        # -- The documented workaround: assign a URL in addition to the
        # -- tooltip. PowerPoint then renders the ScreenTip on hover.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )

        shape.click_action.hyperlink.address = "https://example.com/"
        shape.click_action.screen_tip = "Click to open example.com"

        hlinkClick = shape._element.nvSpPr.cNvPr.hlinkClick
        assert hlinkClick is not None
        assert hlinkClick.tooltip == "Click to open example.com"
        assert hlinkClick.rId  # truthy — a real relationship is present

    def it_pairs_tooltip_with_slide_jump(self):
        # -- Alternative workaround: pair with a slide jump. PowerPoint
        # -- renders the ScreenTip on hover when the click action is
        # -- `hlinksldjump`.
        prs = Presentation()
        slide1 = prs.slides.add_slide(prs.slide_layouts[6])
        slide2 = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide1.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )

        shape.click_action.target_slide = slide2
        shape.click_action.screen_tip = "Go to slide 2"

        hlinkClick = shape._element.nvSpPr.cNvPr.hlinkClick
        assert hlinkClick is not None
        assert hlinkClick.tooltip == "Go to slide 2"
        assert hlinkClick.action == "ppaction://hlinksldjump"
        assert hlinkClick.rId  # a real rel targeting slide2

    def it_round_trips_a_tooltip_paired_with_a_URL(self):
        # -- Save and reload. The tooltip + URL survive intact and the
        # -- relationship target is preserved.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )
        shape.click_action.hyperlink.address = "https://example.com/"
        shape.click_action.screen_tip = "Visit site"

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        rshape = reloaded.slides[0].shapes[0]
        assert rshape.click_action.screen_tip == "Visit site"
        assert rshape.click_action.hyperlink.address == "https://example.com/"

    def it_round_trips_a_tooltip_only_hlinkClick_without_error(self):
        # -- Even in the PowerPoint-invisible case (tooltip only, no
        # -- target), the library must not crash and must preserve the
        # -- authored tooltip. That way callers who authored the file
        # -- with an older python-pptx can still read their
        # -- screen_tip value back, and a follow-on save doesn't
        # -- silently drop it.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )
        shape.click_action.screen_tip = "Tooltip only"

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        rshape = reloaded.slides[0].shapes[0]
        assert rshape.click_action.screen_tip == "Tooltip only"

    def it_clears_the_tooltip_attribute_when_assigned_None(self):
        # -- Clearing the tooltip must not remove the hyperlink element
        # -- itself, because it may still carry a URL or other action
        # -- that the caller wants to keep.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )
        shape.click_action.hyperlink.address = "https://example.com/"
        shape.click_action.screen_tip = "Remove me"

        shape.click_action.screen_tip = None

        hlinkClick = shape._element.nvSpPr.cNvPr.hlinkClick
        assert hlinkClick is not None  # preserved
        assert hlinkClick.tooltip is None
        assert hlinkClick.rId  # URL relationship preserved

    def it_writes_hlinkHover_on_the_hover_action(self):
        # -- The same caveat applies to `hover_action.screen_tip`: the
        # -- tooltip lands on `a:hlinkHover` instead of `a:hlinkClick`,
        # -- but PowerPoint's display rule is the same — a target
        # -- (URL, slide jump, or sound) is required for the tooltip
        # -- to render.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )

        shape.hover_action.hyperlink.address = "https://example.com/"
        shape.hover_action.screen_tip = "Hover to open"

        cNvPr = shape._element.nvSpPr.cNvPr
        hlinkHover = cNvPr.find(qn("a:hlinkHover"))
        assert hlinkHover is not None
        assert hlinkHover.get("tooltip") == "Hover to open"
        # -- and a r:id attribute is present (truthy rId)
        rId = hlinkHover.get(qn("r:id"))
        assert rId  # non-empty relationship id
