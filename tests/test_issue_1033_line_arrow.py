# pyright: reportPrivateUsage=false

"""Regression test for issue #1033 — set arrow type of LINE object.

Issue #1033 (https://github.com/scanny/python-pptx/issues/1033) asks for
a way to set the arrow type on a LINE shape (an ``MSO_SHAPE.LINE`` /
``a:prstGeom prst="line"`` auto-shape). Two prior waves deliver the
pieces needed to satisfy the ask:

  * ``feat/issue-375-line-arrows`` (Wave 1, commit ``f95be72e``) added
    ``LineFormat.begin_arrow`` / ``LineFormat.end_arrow`` sub-objects
    with read/write ``type`` / ``width`` / ``length`` properties mapped
    to the ``a:headEnd`` / ``a:tailEnd`` DrawingML elements, plus the
    ``MSO_LINE_END_TYPE`` / ``MSO_LINE_END_WIDTH`` / ``MSO_LINE_END_LENGTH``
    enumerations.
  * ``fix/issue-749-auto-shape-type-line`` (Wave 2, commit ``7a743394``)
    added ``MSO_SHAPE.LINE`` so a line auto-shape can actually be added
    via ``shapes.add_shape(MSO_SHAPE.LINE, ...)``.

This regression test wires those together: it constructs the exact
flow the #1033 reporter asked for — "add a line, decorate its end with
a triangle arrow" — round-trips it through ``Presentation.save`` and
``Presentation()`` reopen, and confirms every arrow attribute survives.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.enum.dml import MSO_LINE_END_LENGTH, MSO_LINE_END_TYPE, MSO_LINE_END_WIDTH
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches


class DescribeIssue1033LineArrow(object):
    """The #1033 flow: add a LINE shape, set its end arrow to a triangle,
    save, reopen, confirm the arrow (type + width + length) survived.
    """

    def it_round_trips_an_end_arrow_on_a_LINE_shape(self):
        """The reporter's core ask: decorate a line with an arrowhead.

        Exercises ``MSO_SHAPE.LINE`` (from #749) together with
        ``LineFormat.end_arrow`` (from #375) — the two pieces #1033
        depended on — and rounds the combination through save + reopen.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # -- blank layout
        line = slide.shapes.add_shape(
            MSO_SHAPE.LINE, Inches(1), Inches(1), Inches(4), Inches(0)
        )

        # -- the shape really is a `prst="line"` auto-shape --
        assert line.auto_shape_type == MSO_SHAPE.LINE

        # -- baseline: no arrow attributes set --
        assert line.line.end_arrow.type is None
        assert line.line.end_arrow.width is None
        assert line.line.end_arrow.length is None

        # -- author the end arrow (the #1033 ask) --
        line.line.end_arrow.type = MSO_LINE_END_TYPE.TRIANGLE
        line.line.end_arrow.width = MSO_LINE_END_WIDTH.MEDIUM
        line.line.end_arrow.length = MSO_LINE_END_LENGTH.MEDIUM

        # -- visible immediately --
        assert line.line.end_arrow.type is MSO_LINE_END_TYPE.TRIANGLE
        assert line.line.end_arrow.width is MSO_LINE_END_WIDTH.MEDIUM
        assert line.line.end_arrow.length is MSO_LINE_END_LENGTH.MEDIUM

        # -- save + reopen round-trip --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        prs2 = Presentation(buf)
        slide2 = prs2.slides[0]
        line2 = slide2.shapes[0]

        # -- the LINE shape survived as a `prst="line"` auto-shape --
        assert line2.auto_shape_type == MSO_SHAPE.LINE

        # -- every arrow attribute preserved --
        assert line2.line.end_arrow.type is MSO_LINE_END_TYPE.TRIANGLE
        assert line2.line.end_arrow.width is MSO_LINE_END_WIDTH.MEDIUM
        assert line2.line.end_arrow.length is MSO_LINE_END_LENGTH.MEDIUM

        # -- and the opposite end was not implicitly decorated --
        assert line2.line.begin_arrow.type is None

    def it_round_trips_arrows_on_both_ends_of_a_LINE(self):
        """The reporter also asked implicitly about the opposite end: make
        sure ``begin_arrow`` and ``end_arrow`` are independent and both
        survive the save/reopen trip on a LINE shape.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        line = slide.shapes.add_shape(
            MSO_SHAPE.LINE, Inches(1), Inches(2), Inches(4), Inches(0)
        )

        line.line.begin_arrow.type = MSO_LINE_END_TYPE.OVAL
        line.line.end_arrow.type = MSO_LINE_END_TYPE.STEALTH

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        line2 = Presentation(buf).slides[0].shapes[0]
        assert line2.auto_shape_type == MSO_SHAPE.LINE
        assert line2.line.begin_arrow.type is MSO_LINE_END_TYPE.OVAL
        assert line2.line.end_arrow.type is MSO_LINE_END_TYPE.STEALTH
