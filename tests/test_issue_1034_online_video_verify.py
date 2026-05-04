# pyright: reportPrivateUsage=false

"""Regression test for issue #1034 — "Online Video (insert URL-linked video)".

Issue #1034 (https://github.com/scanny/python-pptx/issues/1034) asked for a
way to insert an online video — e.g. a YouTube or Vimeo URL — into a slide
without embedding the media bytes, matching PowerPoint's *Insert > Video >
Online Video* workflow. The existing :meth:`.SlideShapes.add_movie` method
unconditionally copies the media bytes into the package as a ``MediaPart``,
so there was no public surface for an external-URL video.

#1034 is **verified and closed** by Wave 17 #839
(``feat/issue-839-url-video``, commit ``377b9412``), which shipped
:meth:`.SlideShapes.add_movie_link`::

    slide.shapes.add_movie_link(
        url,                # e.g. "https://www.youtube.com/embed/<ID>"
        poster_frame_image, # required — path or binary file-like
        left, top, width, height,
    )

The emitted ``p:pic`` matches PowerPoint's online-video output: both the
``a:videoFile`` and the ``p14:media`` descriptors carry ``r:link`` pointing
at the same external-mode ``VIDEO`` relationship (``TargetMode="External"``)
on the slide part. Only the poster-frame image is embedded in the ``.pptx``;
the video bytes stay at their URL, so decks stay small.

This suite pins the #1034-reporter-facing contract: a basic YouTube embed,
a Vimeo embed, a custom poster image, save + reopen preserving both the URL
and the poster, and a cross-reference that the implementation the reporter
needs is the one #839 shipped. Any regression that drops
``add_movie_link``, stops emitting the external-mode ``VIDEO`` relationship,
or loses the poster-frame round-trip will reproduce the #1034 symptom and
be caught here.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.shapes.picture import Movie
from pptx.shapes.shapetree import SlideShapes
from pptx.util import Inches

# --- Minimal 1x1 transparent PNG used as a poster frame. ------------------
_PNG_BYTES = (
    b"\x89PNG\r\n\x1a\n"
    b"\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
    b"\x00\x00\x00\rIDATx\x9cc\xf8\xcf\xc0\x00\x00\x00\x03\x00\x01\xde\xfb\xc4f"
    b"\x00\x00\x00\x00IEND\xaeB\x60\x82"
)

# --- Second poster frame — a minimal 1x1 red JPEG — to pin a "custom" image -
# that's different-on-the-bytes from the _PNG_BYTES fallback.
_JPEG_BYTES = (
    b"\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00"
    b"\xff\xdb\x00C\x00\x08\x06\x06\x07\x06\x05\x08\x07\x07\x07\t\t\x08\n\x0c"
    b"\x14\r\x0c\x0b\x0b\x0c\x19\x12\x13\x0f\x14\x1d\x1a\x1f\x1e\x1d\x1a\x1c\x1c "
    b"$.' \",#\x1c\x1c(7),01444\x1f'9=82<.342\xff\xc0\x00\x0b\x08\x00\x01\x00\x01"
    b"\x01\x01\x11\x00\xff\xc4\x00\x1f\x00\x00\x01\x05\x01\x01\x01\x01\x01\x01\x00"
    b"\x00\x00\x00\x00\x00\x00\x00\x01\x02\x03\x04\x05\x06\x07\x08\t\n\x0b\xff\xc4"
    b"\x00\xb5\x10\x00\x02\x01\x03\x03\x02\x04\x03\x05\x05\x04\x04\x00\x00\x01}\x01"
    b"\x02\x03\x00\x04\x11\x05\x12!1A\x06\x13Qa\x07\"q\x142\x81\x91\xa1\x08#B\xb1\xc1"
    b"\x15R\xd1\xf0$3br\x82\t\n\x16\x17\x18\x19\x1a%&'()*456789:CDEFGHIJSTUVWXYZ"
    b"cdefghijstuvwxyz\x83\x84\x85\x86\x87\x88\x89\x8a\x92\x93\x94\x95\x96\x97\x98"
    b"\x99\x9a\xa2\xa3\xa4\xa5\xa6\xa7\xa8\xa9\xaa\xb2\xb3\xb4\xb5\xb6\xb7\xb8\xb9"
    b"\xba\xc2\xc3\xc4\xc5\xc6\xc7\xc8\xc9\xca\xd2\xd3\xd4\xd5\xd6\xd7\xd8\xd9\xda"
    b"\xe1\xe2\xe3\xe4\xe5\xe6\xe7\xe8\xe9\xea\xf1\xf2\xf3\xf4\xf5\xf6\xf7\xf8\xf9"
    b"\xfa\xff\xda\x00\x08\x01\x01\x00\x00?\x00\xfb\xd0\xff\xd9"
)


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of ``PartFactory.part_type_for``.

    Mirrors the guard used by ``tests/test_issue_720_ea_font_verify.py`` —
    ``Presentation(stream)`` reopen dispatches through
    ``PartFactory.part_type_for``, which earlier test modules overwrite with
    mocks and do not restore.
    """
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue1034OnlineVideoVerify:
    """#1034 online-video insert verify-and-close via #839.

    Pins that :meth:`.SlideShapes.add_movie_link` — the API shipped by
    #839 — covers the reporter-facing scenarios in #1034: YouTube embed,
    Vimeo embed, custom poster-frame image, and a save-and-reopen
    round-trip preserving both URL and poster.
    """

    # -- basic YouTube URL embed -----------------------------------------

    def it_embeds_a_youtube_video_via_add_movie_link(self):
        """The canonical #1034 scenario: add a YouTube URL to a slide.

        The emitted ``p:pic`` must reference the URL via an external-mode
        ``VIDEO`` relationship, and no media bytes should be copied into
        the package (``Movie.blob`` / ``Movie.content_type`` come back
        ``None`` on a URL-linked shape).
        """
        url = "https://www.youtube.com/embed/dQw4w9WgXcQ"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        movie = slide.shapes.add_movie_link(
            url, io.BytesIO(_PNG_BYTES), Inches(1), Inches(1), Inches(4), Inches(3)
        )

        assert isinstance(movie, Movie)
        # -- external URL relationship, VIDEO reltype --
        rId = movie._element.xpath("./p:nvPicPr/p:nvPr/a:videoFile/@r:link")[0]
        rel = slide.part.rels[rId]
        assert rel.reltype == RT.VIDEO
        assert rel.is_external
        assert rel.target_ref == url
        # -- no media part materialised for URL-linked shape --
        assert movie.blob is None
        assert movie.content_type is None

    # -- Vimeo URL embed --------------------------------------------------

    def it_embeds_a_vimeo_video_via_add_movie_link(self):
        """Vimeo URLs (``player.vimeo.com/video/<ID>``) work the same way.

        :meth:`.SlideShapes.add_movie_link` is URL-agnostic — it does no
        host-specific validation, so any ``https://`` URL PowerPoint's
        online-video player can load round-trips through the same code
        path. Pinning Vimeo alongside YouTube confirms the API is not
        silently YouTube-only.
        """
        url = "https://player.vimeo.com/video/123456789"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        movie = slide.shapes.add_movie_link(
            url, io.BytesIO(_PNG_BYTES), Inches(1), Inches(1), Inches(4), Inches(3)
        )

        assert isinstance(movie, Movie)
        rId = movie._element.xpath("./p:nvPicPr/p:nvPr/a:videoFile/@r:link")[0]
        rel = slide.part.rels[rId]
        assert rel.reltype == RT.VIDEO
        assert rel.is_external
        assert rel.target_ref == url

    # -- custom poster-frame image ---------------------------------------

    def it_accepts_a_custom_poster_frame_image(self):
        """A caller-supplied poster image is embedded and linked to the shape.

        ``add_movie_link`` requires an explicit poster image (there is no
        media part from which to derive a default). The caller can supply
        a PNG, JPEG, or any format :meth:`Image.from_file` understands —
        the bytes land in an ``ImagePart`` inside the package and the
        ``p:pic``'s ``p:blipFill/a:blip @r:embed`` points at it.
        """
        url = "https://www.youtube.com/embed/dQw4w9WgXcQ"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        movie = slide.shapes.add_movie_link(
            url, io.BytesIO(_JPEG_BYTES), Inches(1), Inches(1), Inches(4), Inches(3)
        )

        poster = movie.poster_frame
        assert poster is not None
        # -- the custom JPEG bytes round-tripped through the ImagePart --
        assert poster.blob == _JPEG_BYTES
        # -- poster content-type reflects the supplied JPEG --
        assert poster.content_type == CT.JPEG

    # -- save + reopen round-trip ----------------------------------------

    def it_round_trips_url_and_poster_through_save_and_reload(
        self, _restore_part_factory
    ):
        """Save + reopen preserves both the URL link and the poster image.

        The reporter's deck must survive ``Presentation.save`` and the
        subsequent reload: the external-mode ``VIDEO`` relationship keeps
        its ``TargetMode="External"`` and URL target, the poster-frame
        ``ImagePart`` is still present in the package, and the ``Movie``
        proxy still short-circuits ``blob`` / ``content_type`` to ``None``
        because no ``MediaPart`` was ever created.
        """
        url = "https://www.youtube.com/embed/dQw4w9WgXcQ"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.add_movie_link(
            url, io.BytesIO(_JPEG_BYTES), Inches(2), Inches(3), Inches(5), Inches(4)
        )

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        slide2 = prs2.slides[0]
        movies = [sh for sh in slide2.shapes if isinstance(sh, Movie)]
        assert len(movies) == 1
        movie = movies[0]

        # -- external URL relationship survives the save + reopen --
        rId = movie._element.xpath("./p:nvPicPr/p:nvPr/a:videoFile/@r:link")[0]
        rel = slide2.part.rels[rId]
        assert rel.reltype == RT.VIDEO
        assert rel.is_external
        assert rel.target_ref == url
        # -- poster image survives --
        poster = movie.poster_frame
        assert poster is not None
        assert poster.blob == _JPEG_BYTES
        # -- shape geometry preserved --
        assert movie.left == Inches(2)
        assert movie.top == Inches(3)
        assert movie.width == Inches(5)
        assert movie.height == Inches(4)
        # -- no media part was ever embedded --
        assert movie.blob is None
        assert movie.content_type is None

    # -- cross-reference #839 --------------------------------------------

    def it_cross_references_the_839_resolution_api_surface(self):
        """Guard: the public API that resolves #1034 is the one #839 shipped.

        #1034 is closed by Wave 17 #839
        (``feat/issue-839-url-video``), which shipped
        :meth:`.SlideShapes.add_movie_link`. A refactor that removed or
        renamed that method would silently reintroduce #1034 — this guard
        pins the entry point by name and signature.
        """
        import inspect

        assert hasattr(SlideShapes, "add_movie_link")
        sig = inspect.signature(SlideShapes.add_movie_link)
        # -- (self, url, poster_frame_image, left, top, width, height) --
        assert list(sig.parameters) == [
            "self",
            "url",
            "poster_frame_image",
            "left",
            "top",
            "width",
            "height",
        ]
