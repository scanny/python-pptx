# pyright: reportPrivateUsage=false

"""Regression test for issue #1035 — ``add_chart`` return type annotation.

Issue #1035 (https://github.com/scanny/python-pptx/issues/1035) reported
that ``SlideShapes.add_chart`` declared a return type of |Chart| while
the runtime behaviour — documented in the method's own docstring — is
that a |GraphicFrame| is returned and the caller reaches the chart
object through ``graphic_frame.chart``.

The fix was a pure annotation/docs correction: the return type is now
|GraphicFrame|. This test pins both the runtime contract (the returned
object is a |GraphicFrame|, its ``chart`` attribute is a |Chart|) and
the static annotation (read off ``__annotations__`` and resolved against
the defining module's namespace) so a future edit that regresses either
surface will fail loudly.
"""

from __future__ import annotations

from pptx import Presentation
from pptx.chart.chart import Chart
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.shapetree import _BaseGroupShapes
from pptx.util import Inches


def _chart_data() -> CategoryChartData:
    cd = CategoryChartData()
    cd.categories = ["A", "B", "C"]
    cd.add_series("S0", (1.0, 2.0, 3.0))
    return cd


class DescribeIssue1035AddChartReturnType:
    """`SlideShapes.add_chart` returns a `GraphicFrame`, not a `Chart`."""

    def it_returns_a_GraphicFrame_at_runtime(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        shape = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(6),
            Inches(4),
            _chart_data(),
        )

        assert isinstance(shape, GraphicFrame), (
            f"add_chart must return a GraphicFrame; got {type(shape).__name__}"
        )
        # -- the Chart is reachable via the `.chart` property, NOT directly --
        assert isinstance(shape.chart, Chart)

    def it_declares_GraphicFrame_as_its_return_annotation(self):
        # -- With `from __future__ import annotations` every annotation is a --
        # -- string; evaluating the return annotation in the module's       --
        # -- namespace confirms it resolves to the `GraphicFrame` class.    --
        from pptx.shapes import shapetree as _shapetree

        annotations = _BaseGroupShapes.add_chart.__annotations__
        return_ann = annotations["return"]
        assert return_ann == "GraphicFrame", (
            f"raw return annotation must be 'GraphicFrame' (got {return_ann!r})"
        )

        # -- resolve the string against the module's symbol table (plus any --
        # -- TYPE_CHECKING-only names) and confirm it is the real class.    --
        ns = dict(vars(_shapetree))
        ns["GraphicFrame"] = GraphicFrame
        resolved = eval(return_ann, ns)  # noqa: S307 -- trusted string
        assert resolved is GraphicFrame
