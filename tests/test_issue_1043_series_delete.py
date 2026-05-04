# pyright: reportPrivateUsage=false

"""Regression test for issue #1043 — exclude columns from chart series.

Issue #1043 (https://github.com/scanny/python-pptx/issues/1043) asks for a
way to exclude a column of the embedded workbook from the chart series —
the equivalent of PowerPoint's right-click "Select Data" dialog unchecking
a series. OOXML models series selection at the ``c:ser`` level: a column
that should not be plotted simply has no corresponding ``c:ser`` child of
the enclosing ``c:{x}Chart``. The resolution is therefore a series-level
delete — implemented as ``_BaseSeries.delete()``.

This test pins the contract from the reporter's perspective:

* Chart with 5 series — delete two of them — remaining series keep
  their original names and values.
* The deletion round-trips through ``Presentation.save`` + reopen.
* Deleting from the high index first is the caller's stable pattern,
  matching the equivalent ``del lst[i]`` idiom.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def _add_chart_with_five_series(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    cd = CategoryChartData()
    cd.categories = ["A", "B", "C"]
    cd.add_series("S0", (1.0, 2.0, 3.0))
    cd.add_series("S1", (4.0, 5.0, 6.0))
    cd.add_series("S2", (7.0, 8.0, 9.0))
    cd.add_series("S3", (10.0, 11.0, 12.0))
    cd.add_series("S4", (13.0, 14.0, 15.0))
    gf = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        cd,
    )
    return gf.chart


class DescribeIssue1043SeriesDelete:
    """``_BaseSeries.delete()`` removes the ``c:ser`` from its plot."""

    def it_deletes_a_single_series_from_a_five_series_chart(self):
        prs = Presentation()
        chart = _add_chart_with_five_series(prs)
        assert [s.name for s in chart.series] == ["S0", "S1", "S2", "S3", "S4"]

        chart.series[2].delete()

        assert [s.name for s in chart.series] == ["S0", "S1", "S3", "S4"]

    def it_deletes_non_contiguous_series_high_to_low(self):
        # -- deleting [3] then [1] from a 5-series chart leaves 0, 2, 4 --
        prs = Presentation()
        chart = _add_chart_with_five_series(prs)

        chart.series[3].delete()
        chart.series[1].delete()

        assert [s.name for s in chart.series] == ["S0", "S2", "S4"]

    def it_preserves_values_on_the_remaining_series(self):
        prs = Presentation()
        chart = _add_chart_with_five_series(prs)

        chart.series[1].delete()
        chart.series[2].delete()  # this was originally S3 before the first delete

        remaining = list(chart.series)
        assert [s.name for s in remaining] == ["S0", "S2", "S4"]
        assert remaining[0].values == (1.0, 2.0, 3.0)
        assert remaining[1].values == (7.0, 8.0, 9.0)
        assert remaining[2].values == (13.0, 14.0, 15.0)

    def it_survives_save_and_reopen_round_trip(self):
        prs = Presentation()
        chart = _add_chart_with_five_series(prs)

        chart.series[3].delete()
        chart.series[1].delete()

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        prs2 = Presentation(buf)
        # -- slide layout 5 has title placeholders; find the graphic frame --
        slide = prs2.slides[-1]
        chart2 = next(s for s in slide.shapes if getattr(s, "has_chart", False)).chart
        assert [s.name for s in chart2.series] == ["S0", "S2", "S4"]

    def it_renders_one_fewer_c_ser_element_after_delete(self):
        prs = Presentation()
        chart = _add_chart_with_five_series(prs)
        chartSpace = chart._chartSpace
        assert len(chartSpace.xpath(".//c:ser")) == 5

        chart.series[0].delete()

        assert len(chartSpace.xpath(".//c:ser")) == 4
