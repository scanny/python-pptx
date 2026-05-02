"""Regression tests for issue #1058.

Passing float coordinates to connector-creation helpers or to the
``Connector.begin_x`` / ``begin_y`` / ``end_x`` / ``end_y`` setters must not
produce a corrupted ``.pptx``.

The underlying XML attributes (``a:off/@x``, ``a:off/@y``, ``a:ext/@cx``,
``a:ext/@cy``) are typed ``xsd:long`` (ST_Coordinate) and PowerPoint refuses
to open a file where they serialize as ``"1000.5"``. The fix coerces
floats with ``int(round(...))`` at both entry points — the connector
constructor (``CT_Connector.new_cxnSp``) and the four position setters on
``Connector``.
"""

from __future__ import annotations

import re

import pytest

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches

_FRACTIONAL_ATTR_RE = re.compile(r'\b(?:x|y|cx|cy)="-?\d+\.')


def _assert_no_fractional_coord(xml: str) -> None:
    """Fail if any int-typed position/extent attribute contains a decimal point."""
    match = _FRACTIONAL_ATTR_RE.search(xml)
    assert match is None, f"found non-integer coord attribute in XML: {xml}"


class DescribeIssue1058ConnectorFloatCoords:
    """Float coordinates must not produce corrupt connector XML."""

    def it_coerces_float_coords_passed_to_add_connector(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, 1000.5, 2000.0, 5000.75, 6000.5
        )

        _assert_no_fractional_coord(connector._element.xml)

    def it_coerces_a_float_begin_x_setter(self):
        connector = self._fresh_connector()

        connector.begin_x = 1500.75

        _assert_no_fractional_coord(connector._element.xml)

    def it_coerces_a_float_begin_y_setter(self):
        connector = self._fresh_connector()

        connector.begin_y = 1500.25

        _assert_no_fractional_coord(connector._element.xml)

    def it_coerces_a_float_end_x_setter(self):
        connector = self._fresh_connector()

        connector.end_x = 3500.5

        _assert_no_fractional_coord(connector._element.xml)

    def it_coerces_a_float_end_y_setter(self):
        connector = self._fresh_connector()

        connector.end_y = 3500.5

        _assert_no_fractional_coord(connector._element.xml)

    def it_rounds_float_coords_to_the_nearest_integer(self):
        connector = self._fresh_connector()

        connector.begin_x = 1500.75  # -> 1501
        connector.end_x = 3500.4  # -> 3500

        # -- begin_x of a non-flipped connector is stored as `a:off/@x` --
        xml = connector._element.xml
        assert 'x="1501"' in xml, xml
        # -- end_x derives from x + cx; with end_x=3500 and x=1501, cx=1999 --
        assert 'cx="1999"' in xml, xml

    @pytest.mark.parametrize("value", [1500.25, 1500.75, 0.5, -0.5])
    def it_accepts_any_float_without_raising(self, value: float):
        connector = self._fresh_connector()

        # -- None of these should raise `TypeError` from the ST_Coordinate validator. --
        connector.begin_x = value
        connector.begin_y = value
        connector.end_x = value
        connector.end_y = value

        _assert_no_fractional_coord(connector._element.xml)

    def it_saves_a_clean_pptx_when_floats_are_used(self, tmp_path):
        """End-to-end: the resulting file must round-trip through Presentation()."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, 1000.5, 2000.0, 5000.75, 6000.5
        )
        connector.begin_x = 1500.5
        connector.end_y = 3500.25

        out_path = tmp_path / "issue_1058.pptx"
        prs.save(str(out_path))

        # -- re-open and read back; corrupt XML would blow up on parse --
        reloaded = Presentation(str(out_path))
        reloaded_slide = reloaded.slides[0]
        # -- last shape is the connector we added --
        reloaded_connector = list(reloaded_slide.shapes)[-1]
        _assert_no_fractional_coord(reloaded_connector._element.xml)

    # -- fixture helpers --

    @staticmethod
    def _fresh_connector():
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        return slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(1), Inches(3), Inches(3)
        )
