# pyright: reportPrivateUsage=false

"""Verify-close regression test for issue #1101 — "Chart features not yet supportable".

Issue #1101 (https://github.com/scanny/python-pptx/issues/1101) was a
meta-issue cataloguing chart features the reporter could not reach
through python-pptx. Most of those asks have since shipped on this
fork:

* Combo / secondary axis — shipped via #338 (``Chart.add_plot``) and
  #141 / #470 (:attr:`Chart.secondary_value_axis`).
* Trendlines — shipped via #299 / #617 (``_BaseSeries.add_trendline``
  / :class:`Trendline`, available to every category-chart series type).
* Error bars — shipped via #544 / #607 (``_BaseSeries.set_error_bars``
  / :class:`ErrorBars`, category and XY).
* Per-point formatting — shipped via #450 / #638
  (:attr:`Point.format` for fill/line, :attr:`DataLabel.number_format`
  for per-point number formatting).
* Axis title formatting — shipped (partially) via #764
  (``Chart.set_axis_title`` / ``Chart.set_title`` convenience wrappers
  plus :attr:`AxisTitle.text_frame` for formatting).

This suite is the umbrella regression pin that exercises every one of
those features in a single end-to-end flow, with a fresh
``Presentation()`` → author → save → reopen → assert round-trip. A
future refactor that silently broke any one of the features listed in
#1101 would be caught here alongside the feature-specific verify
suites (``tests/test_issue_470_combo_charts_verify.py``,
``tests/test_issue_617_column_trendline_verify.py``,
``tests/test_issue_607_xy_error_bars_verify.py``,
``tests/test_issue_1068_individual_data_labels.py``).

Scope
-----
The end-to-end flow in :meth:`it_round_trips_the_full_1101_feature_set`
builds *one* chart that exercises:

1. A column + line combo via :meth:`Chart.add_plot`.
2. A linear trendline on the column series
   (:meth:`_BaseSeries.add_trendline`).
3. A fixed-value error-bars block on the line series
   (:meth:`_BaseSeries.set_error_bars`).
4. A per-point fill colour override on the column series
   (:attr:`Point.format.fill`) and a per-point number format on the
   same point's data label (:attr:`DataLabel.number_format`).
5. A chart title and a category/value axis title set through the
   single-call convenience wrappers
   (:meth:`Chart.set_title` / :meth:`Chart.set_axis_title`).
6. A full save-and-reopen pass that re-reads every one of the above
   through the public API.

Secondary-value-axis round-trip is pinned separately in
:meth:`it_round_trips_a_pre_authored_secondary_axis_chart` using the
documented pre-authored-XML approach (per #470 scope note, authoring
a ``c:valAx`` secondary axis from scratch on a library-created chart
is still a follow-up — the supported workflow is to start from a
chartSpace that already has the secondary axis).
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.chart import Chart
from pptx.chart.data import CategoryChartData
from pptx.chart.plot import BarPlot, LinePlot
from pptx.chart.series import ErrorBars, Trendline
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_ERROR_BAR_INCLUDE,
    XL_ERROR_BAR_TYPE,
    XL_TRENDLINE_TYPE,
)
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches

# -- fixtures ------------------------------------------------------------


@pytest.fixture
def _restore_part_factory():
    """Guard against pollution from tests that mutate ``PartFactory.part_type_for``.

    Mirrors the identical fixture used in ``test_issue_470_combo_charts_verify``
    and ``test_issue_303_multi_axes_verify``. Some unit tests in the wider
    suite replace ``PartFactory.part_type_for[CT.PML_SLIDE]`` with a Mock and
    do not restore it, which breaks round-trip tests that open a pptx
    produced by the same process.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


# -- helpers -------------------------------------------------------------


def _roundtrip(prs):
    """Save `prs` into an in-memory stream and reopen it."""
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)


def _chart_on_first_slide(prs):
    """Return the single chart on the first slide of `prs`."""
    for shp in prs.slides[0].shapes:
        if shp.has_chart:
            return shp.chart
    raise AssertionError("presentation's first slide has no chart")


def _combo_chartSpace_with_secondary_axis_xml():
    """Return chartSpace XML for a column-bar + line combo w/ secondary valAx.

    Mirrors the identical fixture in
    ``test_issue_470_combo_charts_verify`` — a minimal but
    PowerPoint-shaped combo chart that carries a secondary ``c:valAx``
    so the secondary-axis branch of the #1101 checklist can be
    exercised without depending on authoring-from-scratch support that
    is still a follow-up item on this fork (see
    ``docs/dev/analysis/combo-chart.rst``).
    """
    return (
        '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml'
        '/2006/chart">'
        "<c:chart><c:plotArea>"
        "<c:barChart>"
        '<c:barDir val="col"/>'
        '<c:grouping val="clustered"/>'
        '<c:ser><c:idx val="0"/><c:order val="0"/>'
        "<c:tx><c:v>Revenue</c:v></c:tx>"
        "</c:ser>"
        '<c:axId val="111"/>'
        '<c:axId val="222"/>'
        "</c:barChart>"
        "<c:lineChart>"
        '<c:grouping val="standard"/>'
        '<c:ser><c:idx val="1"/><c:order val="1"/>'
        "<c:tx><c:v>Trend</c:v></c:tx>"
        "</c:ser>"
        '<c:axId val="111"/>'
        '<c:axId val="333"/>'
        "</c:lineChart>"
        "<c:catAx>"
        '<c:axId val="111"/>'
        '<c:scaling><c:orientation val="minMax"/></c:scaling>'
        '<c:delete val="0"/>'
        '<c:axPos val="b"/>'
        '<c:crossAx val="222"/>'
        "</c:catAx>"
        "<c:valAx>"
        '<c:axId val="222"/>'
        '<c:scaling><c:orientation val="minMax"/></c:scaling>'
        '<c:delete val="0"/>'
        '<c:axPos val="l"/>'
        '<c:crossAx val="111"/>'
        "</c:valAx>"
        "<c:valAx>"
        '<c:axId val="333"/>'
        "<c:scaling>"
        '<c:orientation val="minMax"/>'
        '<c:max val="100"/>'
        "</c:scaling>"
        '<c:delete val="0"/>'
        '<c:axPos val="r"/>'
        '<c:crossAx val="111"/>'
        "</c:valAx>"
        "</c:plotArea></c:chart></c:chartSpace>"
    )


# -- verification suite --------------------------------------------------


class DescribeIssue1101ChartFeaturesVerify:
    """#1101 umbrella verify — every listed feature ships and round-trips."""

    def it_round_trips_the_full_1101_feature_set(self, _restore_part_factory):
        # ------------------------------------------------------------------
        # -- author: combo (column + line), trendline, error bars, per-point
        # -- override, chart + axis titles, on one chart.
        # ------------------------------------------------------------------
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        # --- 1) start with a column chart ---
        bar_data = CategoryChartData()
        bar_data.categories = ["Q1", "Q2", "Q3", "Q4"]
        bar_data.add_series("Revenue", (10.0, 20.0, 30.0, 40.0))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(8),
            Inches(5),
            bar_data,
        ).chart

        # --- 2) overlay a line plot (combo chart) ---
        line_data = CategoryChartData()
        line_data.categories = ["Q1", "Q2", "Q3", "Q4"]
        line_data.add_series("Trend", (5.0, 15.0, 25.0, 35.0))
        chart.add_plot(XL_CHART_TYPE.LINE, line_data)

        # --- 3) add a linear trendline to the column series ---
        column_series = chart.plots[0].series[0]
        column_series.add_trendline(
            trendline_type=XL_TRENDLINE_TYPE.LINEAR,
            display_equation=True,
            display_r_squared=True,
        )

        # --- 4) add a fixed-value error bars block to the line series ---
        line_series = chart.plots[1].series[0]
        line_series.set_error_bars(
            type_=XL_ERROR_BAR_TYPE.FIXED_VALUE,
            value=2.5,
            include=XL_ERROR_BAR_INCLUDE.BOTH,
        )

        # --- 5) per-point fill on the column series + per-point number ---
        # --- format on the same point's data label --------------------
        chart.plots[0].has_data_labels = True
        point = column_series.points[2]  # -- Q3's column, value 30
        fill = point.format.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0xFF, 0x00, 0x00)
        column_series.points[2].data_label.number_format = "#,##0.00"

        # --- 6) chart + axis titles via the #764 convenience wrappers ---
        chart.set_title("Quarterly Revenue")
        chart.set_axis_title("category", "Quarter")
        chart.set_axis_title("value", "USD (millions)")

        # ------------------------------------------------------------------
        # -- save + reopen --------------------------------------------------
        # ------------------------------------------------------------------
        reopened = _roundtrip(prs)
        chart2 = _chart_on_first_slide(reopened)

        # --- combo shape survived ---
        assert len(chart2.plots) == 2
        plot_kinds = {type(p).__name__ for p in chart2.plots}
        assert plot_kinds == {"BarPlot", "LinePlot"}
        assert isinstance(chart2.plots[0], BarPlot)
        assert isinstance(chart2.plots[1], LinePlot)
        names = [s.name for s in chart2.series]
        assert "Revenue" in names
        assert "Trend" in names

        # --- trendline survived on the column series ---
        column2 = chart2.plots[0].series[0]
        assert len(column2.trendlines) == 1
        tl = column2.trendlines[0]
        assert isinstance(tl, Trendline)
        assert tl.trendline_type == XL_TRENDLINE_TYPE.LINEAR
        assert tl.display_equation is True
        assert tl.display_r_squared is True

        # --- error bars survived on the line series ---
        line2 = chart2.plots[1].series[0]
        assert line2.has_error_bars is True
        bars = line2.error_bars
        assert isinstance(bars, ErrorBars)
        assert bars.type == XL_ERROR_BAR_TYPE.FIXED_VALUE
        assert bars.include == XL_ERROR_BAR_INCLUDE.BOTH
        assert bars.value == 2.5

        # --- per-point fill on column[2] survived ---
        p2 = chart2.plots[0].series[0].points[2]
        assert p2.format.fill.fore_color.rgb == RGBColor(0xFF, 0x00, 0x00)
        # --- per-point number format on that point's data label ---
        p2_label = chart2.plots[0].series[0].points[2].data_label
        assert p2_label.number_format == "#,##0.00"
        assert p2_label.number_format_is_linked is False

        # --- chart + axis titles survived ---
        assert chart2.has_title is True
        assert chart2.chart_title.text_frame.text == "Quarterly Revenue"
        assert chart2.category_axis.has_title is True
        assert chart2.category_axis.axis_title.text_frame.text == "Quarter"
        assert chart2.value_axis.has_title is True
        assert chart2.value_axis.axis_title.text_frame.text == "USD (millions)"

    def it_round_trips_a_pre_authored_secondary_axis_chart(self):
        """Secondary-axis combo chart round-trips its distinct axes and scale.

        #470 closed with the scope note that authoring a fresh secondary
        ``c:valAx`` on a library-created chart is still a follow-up. The
        supported path is pre-authored chartSpace XML and mutation via
        :attr:`Chart.secondary_value_axis`. This case pins that path end
        to end: two plots, two value axes, independent scales, mutable.
        """
        chartSpace = parse_xml(_combo_chartSpace_with_secondary_axis_xml())
        chart = Chart(chartSpace, None)

        # -- two plots under one plot area --
        assert len(chart.plots) == 2
        plot_kinds = {type(p).__name__ for p in chart.plots}
        assert plot_kinds == {"BarPlot", "LinePlot"}
        # -- secondary value axis is readable through the public API --
        assert chart.has_secondary_value_axis is True
        assert chart.secondary_value_axis.maximum_scale == 100.0
        # -- mutation lands on the secondary axis only --
        chart.secondary_value_axis.maximum_scale = 250.0
        primary, secondary = chartSpace.xpath(".//c:valAx")
        max_elms = secondary.xpath("./c:scaling/c:max")
        assert len(max_elms) == 1
        assert float(max_elms[0].get("val")) == 250.0
        assert primary.xpath("./c:scaling/c:max") == []
        # -- line plot's axId pair binds to the secondary valAx --
        lineChart = chart._chartSpace.plotArea.xCharts[1]
        line_axIds = [e.get("val") for e in lineChart.findall(qn("c:axId"))]
        assert line_axIds[1] == "333"
        assert secondary.find(qn("c:axId")).get("val") == "333"

    def it_round_trips_a_secondary_axis_title_via_set_axis_title(self):
        """``Chart.set_axis_title("secondary_value", ...)`` writes and reads back.

        Covers the #764 / #1101 overlap: the secondary-axis *title*
        branch of the convenience wrapper is the final unit of "axis
        title formatting" on the reporter's checklist.
        """
        chartSpace = parse_xml(_combo_chartSpace_with_secondary_axis_xml())
        chart = Chart(chartSpace, None)

        chart.set_axis_title("secondary_value", "Secondary scale")

        secondary = chart.secondary_value_axis
        assert secondary.has_title is True
        assert secondary.axis_title.text_frame.text == "Secondary scale"
        # -- clearing via None works too (documented #764 behaviour) --
        chart.set_axis_title("secondary_value", None)
        assert chart.secondary_value_axis.has_title is False
