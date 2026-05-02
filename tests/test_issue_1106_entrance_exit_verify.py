# pyright: reportPrivateUsage=false

"""Regression test for issue #1106 — entrance/exit animations on shapes.

Issue #1106 (https://github.com/scanny/python-pptx/issues/1106) asked
for the ability to author *entrance* and *exit* animation effects on a
shape (for example, a shape that fades in when the slide advances and
fades out before the next click).

The issue originally tracked a parallel, property-based API on
``BaseShape`` (``shape.animation.entrance = MSO_ENTRANCE_EFFECT.FADE``,
``shape.animation.exit = MSO_EXIT_EFFECT.FADE``) that never landed.
During the Wave 5 merge that feature was *superseded* by #102's
``BaseShape.set_animation(effect_type, trigger, delay)`` MVP, which
covers the same ground via the ``MSO_ANIMATION_TYPE`` enum (in
particular, ``FADE_IN`` for entrance and ``FADE_OUT`` for exit).

This suite verifies that #102's ``set_animation`` satisfies #1106's
use cases so the issue can be closed as *resolved by #102*. The tests
are adapted from the 45-test suite on the (unmerged)
``feat/issue-1106-entrance-exit-animations`` branch so that the valuable
behavioral coverage from that work lands as a regression net pinning
entrance + exit authoring through the #102 API.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.animation import AnimationEffect
from pptx.enum.animation import MSO_ANIMATION_TRIGGER, MSO_ANIMATION_TYPE
from pptx.util import Inches


@pytest.fixture
def slide_with_shape():
    """Return (prs, slide, shape) for a fresh blank slide with one auto-shape."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(1))
    return prs, slide, shape


@pytest.fixture
def slide_with_two_shapes():
    """Return (prs, slide, shape_a, shape_b) for two shapes on the same slide."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    shape_a = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(1))
    shape_b = slide.shapes.add_shape(1, Inches(4), Inches(1), Inches(2), Inches(1))
    return prs, slide, shape_a, shape_b


class DescribeIssue1106EntranceExitViaSetAnimation(object):
    """#1106 entrance / exit animations verify-and-close via #102.

    Pins that ``BaseShape.set_animation(MSO_ANIMATION_TYPE.FADE_IN)``
    and ``set_animation(MSO_ANIMATION_TYPE.FADE_OUT)`` cover the
    entrance-and-exit authoring the #1106 reporter asked for.
    """

    # -- entrance authoring --------------------------------------------------

    def it_has_no_animation_on_a_fresh_shape(self, slide_with_shape):
        _, slide, shape = slide_with_shape
        assert shape.animation is None
        assert slide.has_animations is False
        assert slide.timing_xml is None

    def it_accessing_animation_does_not_create_timing(self, slide_with_shape):
        _, slide, shape = slide_with_shape
        _ = shape.animation
        assert slide.has_animations is False
        assert slide.timing_xml is None

    def it_writes_an_entrance_effect_via_FADE_IN(self, slide_with_shape):
        _, slide, shape = slide_with_shape

        effect = shape.set_animation(MSO_ANIMATION_TYPE.FADE_IN)

        assert isinstance(effect, AnimationEffect)
        assert effect.type is MSO_ANIMATION_TYPE.FADE_IN
        assert shape.animation.type is MSO_ANIMATION_TYPE.FADE_IN
        assert slide.has_animations is True

    def it_authors_an_animEffect_entrance_child_in_timing(self, slide_with_shape):
        _, slide, shape = slide_with_shape

        shape.set_animation(MSO_ANIMATION_TYPE.FADE_IN)

        # -- the slide now has a p:animEffect descendant under p:timing --
        assert slide._element.timing is not None
        animEffects = slide._element.timing.xpath(".//p:animEffect")
        assert len(animEffects) == 1
        # -- entrance is coded presetClass="entr", transition="in" --
        assert animEffects[0].get("transition") == "in"
        assert animEffects[0].get("filter") == "fade"
        entr_pars = slide._element.timing.xpath(
            './/p:cTn[@presetClass="entr" and @nodeType="clickEffect"]'
        )
        assert len(entr_pars) == 1

    def it_targets_the_right_shape_via_spTgt(self, slide_with_shape):
        _, slide, shape = slide_with_shape

        shape.set_animation(MSO_ANIMATION_TYPE.FADE_IN)

        spids = slide._element.timing.xpath(".//p:spTgt/@spid")
        assert str(shape.shape_id) in spids

    # -- exit authoring ------------------------------------------------------

    def it_writes_an_exit_effect_via_FADE_OUT(self, slide_with_shape):
        _, slide, shape = slide_with_shape

        effect = shape.set_animation(MSO_ANIMATION_TYPE.FADE_OUT)

        assert isinstance(effect, AnimationEffect)
        assert effect.type is MSO_ANIMATION_TYPE.FADE_OUT
        assert shape.animation.type is MSO_ANIMATION_TYPE.FADE_OUT
        assert slide.has_animations is True

    def it_authors_an_animEffect_exit_child_in_timing(self, slide_with_shape):
        _, slide, shape = slide_with_shape

        shape.set_animation(MSO_ANIMATION_TYPE.FADE_OUT)

        animEffects = slide._element.timing.xpath(".//p:animEffect")
        assert len(animEffects) == 1
        assert animEffects[0].get("transition") == "out"
        assert animEffects[0].get("filter") == "fade"
        exit_pars = slide._element.timing.xpath(
            './/p:cTn[@presetClass="exit" and @nodeType="clickEffect"]'
        )
        assert len(exit_pars) == 1

    # -- coexistence (#1106's use case) -------------------------------------

    def it_supports_entrance_and_exit_on_the_same_slide(self, slide_with_two_shapes):
        """#1106's use case: one deck containing both entrance and exit effects.

        The MVP's per-shape ``set_animation`` *replaces* any prior
        effect on that shape, so covering the "both effects live
        together" scenario means authoring them on two shapes on the
        same slide — which is the multi-shape choreography the #1106
        reporter asked to author.
        """
        _, slide, shape_a, shape_b = slide_with_two_shapes

        shape_a.set_animation(MSO_ANIMATION_TYPE.FADE_IN)
        shape_b.set_animation(MSO_ANIMATION_TYPE.FADE_OUT)

        # -- both effects landed in the same p:timing tree --
        animEffects = slide._element.timing.xpath(".//p:animEffect")
        assert len(animEffects) == 2
        transitions = sorted(e.get("transition") for e in animEffects)
        assert transitions == ["in", "out"]
        # -- each shape sees its own effect --
        assert shape_a.animation.type is MSO_ANIMATION_TYPE.FADE_IN
        assert shape_b.animation.type is MSO_ANIMATION_TYPE.FADE_OUT

    def it_keeps_entrance_and_exit_independent_across_shapes(
        self, slide_with_two_shapes
    ):
        _, _, shape_a, shape_b = slide_with_two_shapes

        shape_a.set_animation(MSO_ANIMATION_TYPE.FADE_IN)
        shape_b.set_animation(MSO_ANIMATION_TYPE.FADE_OUT)

        # -- mutating one shape does not disturb the other --
        shape_a.set_animation(MSO_ANIMATION_TYPE.FLY_IN)
        assert shape_a.animation.type is MSO_ANIMATION_TYPE.FLY_IN
        assert shape_b.animation.type is MSO_ANIMATION_TYPE.FADE_OUT

    # -- removal / NONE semantics -------------------------------------------

    def it_removes_the_effect_when_set_animation_called_with_None(
        self, slide_with_shape
    ):
        _, slide, shape = slide_with_shape
        shape.set_animation(MSO_ANIMATION_TYPE.FADE_IN)
        assert shape.animation is not None

        result = shape.set_animation(None)

        assert result is None
        assert shape.animation is None
        assert slide._element.timing.xpath(".//p:animEffect") == []

    def it_rejects_MSO_ANIMATION_TYPE_NONE_as_an_author_value(
        self, slide_with_shape
    ):
        _, _, shape = slide_with_shape
        with pytest.raises(ValueError, match="NONE"):
            shape.set_animation(MSO_ANIMATION_TYPE.NONE)

    def it_replaces_prior_effect_without_duplicating_on_same_shape(
        self, slide_with_shape
    ):
        _, slide, shape = slide_with_shape

        shape.set_animation(MSO_ANIMATION_TYPE.FADE_IN)
        shape.set_animation(MSO_ANIMATION_TYPE.FADE_OUT)

        # -- MVP: exactly one effect targets this shape id --
        spid_refs = slide._element.timing.xpath(
            './/p:spTgt[@spid="%d"]' % shape.shape_id
        )
        assert len(spid_refs) == 1
        assert shape.animation.type is MSO_ANIMATION_TYPE.FADE_OUT

    # -- round-trip ----------------------------------------------------------

    def it_round_trips_entrance_and_exit_through_save_reopen(
        self, slide_with_two_shapes
    ):
        prs, _, shape_a, shape_b = slide_with_two_shapes
        shape_a.set_animation(MSO_ANIMATION_TYPE.FADE_IN)
        shape_b.set_animation(MSO_ANIMATION_TYPE.FADE_OUT)
        spid_a, spid_b = shape_a.shape_id, shape_b.shape_id

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        slide2 = prs2.slides[0]

        # -- find both reopened shapes by id --
        shape_a2 = next(s for s in slide2.shapes if s.shape_id == spid_a)
        shape_b2 = next(s for s in slide2.shapes if s.shape_id == spid_b)

        assert shape_a2.animation is not None
        assert shape_a2.animation.type is MSO_ANIMATION_TYPE.FADE_IN
        assert shape_b2.animation is not None
        assert shape_b2.animation.type is MSO_ANIMATION_TYPE.FADE_OUT
        # -- both animEffects survived the round-trip --
        animEffects = slide2._element.timing.xpath(".//p:animEffect")
        assert len(animEffects) == 2

    def it_round_trips_removal_through_save_reopen(self, slide_with_shape):
        prs, _, shape = slide_with_shape
        shape.set_animation(MSO_ANIMATION_TYPE.FADE_IN)
        shape.set_animation(None)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        shape2 = prs2.slides[0].shapes[-1]

        assert shape2.animation is None

    # -- trigger / delay plumbing (from 1106's "duration" coverage) ---------

    def it_writes_entrance_with_AFTER_PREVIOUS_trigger(self, slide_with_shape):
        _, _, shape = slide_with_shape

        effect = shape.set_animation(
            MSO_ANIMATION_TYPE.FADE_IN,
            trigger=MSO_ANIMATION_TRIGGER.AFTER_PREVIOUS,
        )

        assert effect.trigger is MSO_ANIMATION_TRIGGER.AFTER_PREVIOUS

    def it_writes_exit_with_a_nonzero_delay(self, slide_with_shape):
        _, _, shape = slide_with_shape

        effect = shape.set_animation(
            MSO_ANIMATION_TYPE.FADE_OUT, delay=1250
        )

        assert effect.delay == 1250

    def it_rejects_negative_delay(self, slide_with_shape):
        _, _, shape = slide_with_shape
        with pytest.raises(ValueError, match="delay"):
            shape.set_animation(MSO_ANIMATION_TYPE.FADE_IN, delay=-1)
