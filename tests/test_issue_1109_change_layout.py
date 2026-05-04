"""Regression tests for issue #1109 — ``Slide.slide_layout`` setter.

The reporter asked for a way to re-point an existing slide at a
different slide layout, including a layout belonging to a *different*
slide master. python-pptx 1.0 exposed ``Slide.slide_layout`` as read-only.

These tests pin the minimum viable behavior: assigning a new layout
rewrites the underlying ``RT.SLIDE_LAYOUT`` relationship, the new layout
reads back through :attr:`Slide.slide_layout`, and the change survives a
save-and-reopen round-trip, including the cross-master case.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT


class DescribeIssue1109ChangeLayout(object):
    """Regression suite for issue #1109 — `Slide.slide_layout` setter."""

    def it_re_points_a_slide_at_a_same_master_layout(self):
        """Assigning a new layout from the same master swaps the rel losslessly."""
        prs = Presentation()
        original_layout = prs.slide_layouts[0]  # Title Slide
        target_layout = prs.slide_layouts[5]  # Title Only
        slide = prs.slides.add_slide(original_layout)
        assert slide.slide_layout is original_layout

        slide.slide_layout = target_layout

        # -- getter round-trips the new layout --
        assert slide.slide_layout is target_layout
        # -- the underlying rel now points at the new layout-part --
        assert slide.part.part_related_by(RT.SLIDE_LAYOUT) is target_layout.part

    def it_re_points_a_slide_at_a_cross_master_layout(self):
        """Assigning a layout from a different slide master is supported."""
        prs = Presentation("features/steps/test_files/prs-slide-masters.pptx")
        assert len(prs.slide_masters) == 2
        layout_a = prs.slide_masters[0].slide_layouts[0]
        layout_b = prs.slide_masters[1].slide_layouts[0]
        slide = prs.slides.add_slide(layout_a)
        assert slide.slide_layout.slide_master is prs.slide_masters[0]

        slide.slide_layout = layout_b

        assert slide.slide_layout is layout_b
        assert slide.slide_layout.slide_master is prs.slide_masters[1]

    def it_round_trips_a_same_master_layout_swap(self):
        prs = Presentation()
        original_layout = prs.slide_layouts[0]
        target_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(original_layout)

        slide.slide_layout = target_layout

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reopened = Presentation(buf)
        # -- target_layout is the 6th layout in the first master --
        assert reopened.slides[0].slide_layout is reopened.slide_layouts[5]

    def it_round_trips_a_cross_master_layout_swap(self):
        prs = Presentation("features/steps/test_files/prs-slide-masters.pptx")
        layout_a = prs.slide_masters[0].slide_layouts[0]
        layout_b = prs.slide_masters[1].slide_layouts[0]
        slide = prs.slides.add_slide(layout_a)
        assert slide.slide_layout.slide_master is prs.slide_masters[0]

        slide.slide_layout = layout_b

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reopened = Presentation(buf)
        reopened_slide = reopened.slides[0]
        # -- the reopened slide's layout is backed by the 2nd master --
        assert reopened_slide.slide_layout.slide_master is reopened.slide_masters[1]

    def it_drops_the_prior_layout_rel_after_the_swap(self):
        """The replaced slide-to-layout rel is no longer on the slide part."""
        prs = Presentation()
        original_layout = prs.slide_layouts[0]
        target_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(original_layout)

        slide.slide_layout = target_layout

        # -- exactly one SLIDE_LAYOUT rel remains --
        layout_rels = [r for r in slide.part.rels.values() if r.reltype == RT.SLIDE_LAYOUT]
        assert len(layout_rels) == 1
        assert layout_rels[0].target_part is target_layout.part

    def it_is_a_noop_when_assigned_the_current_layout(self):
        """Assigning the existing layout re-uses the existing rel."""
        prs = Presentation()
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        before_rIds = set(slide.part.rels)

        slide.slide_layout = layout

        after_rIds = set(slide.part.rels)
        assert before_rIds == after_rIds
        assert slide.slide_layout is layout

    def it_raises_on_a_foreign_presentation_layout(self):
        """Cross-presentation layout assignment raises ValueError."""
        prs_a = Presentation()
        prs_b = Presentation()
        slide = prs_a.slides.add_slide(prs_a.slide_layouts[0])
        foreign_layout = prs_b.slide_layouts[5]

        with pytest.raises(ValueError, match="same presentation"):
            slide.slide_layout = foreign_layout
