# pyright: reportPrivateUsage=false

"""Regression tests for issue #1111 — `Font.color` getter mutates XML.

Reading `run.font.color` (and any of its attributes like `.type`, `.rgb`,
`.theme_color`, `str()`, etc.) must not add `<a:solidFill>` (or any other
child) to the run's `<a:rPr>`. Doing so breaks the theme-color inheritance
chain: PowerPoint interprets an explicit empty `<a:solidFill/>` as "override
inherited color with nothing", after which the run can never again inherit
from the paragraph / layout / master.
"""

from __future__ import annotations

import pytest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.oxml.ns import qn


def _rPr(run):
    """Return the `<a:rPr>` element on `run`, or |None| if absent."""
    return run._r.find(qn("a:rPr"))


def _solidFill(run):
    """Return `<a:solidFill>` child of this run's `<a:rPr>`, or |None|."""
    rPr = _rPr(run)
    if rPr is None:
        return None
    return rPr.find(qn("a:solidFill"))


@pytest.fixture
def run_with_text():
    """Yield a freshly-authored run inside a default slide-layout title."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tf = slide.shapes.title.text_frame
    tf.text = "regression"
    return tf.paragraphs[0].runs[0]


class DescribeFontColorGetterNoMutation:
    """Reading `font.color` must not write `<a:solidFill>` into `<a:rPr>`."""

    def it_does_not_add_solidFill_on_plain_color_access(self, run_with_text):
        run = run_with_text

        _ = run.font.color  # read

        assert _solidFill(run) is None

    def it_does_not_add_solidFill_on_color_type_access(self, run_with_text):
        run = run_with_text

        color_type = run.font.color.type

        assert color_type is None
        assert _solidFill(run) is None

    def it_does_not_add_solidFill_on_rgb_access(self, run_with_text):
        run = run_with_text

        color = run.font.color
        with pytest.raises(AttributeError):
            _ = color.rgb

        assert _solidFill(run) is None

    def it_does_not_add_solidFill_on_theme_color_access(self, run_with_text):
        run = run_with_text

        color = run.font.color
        with pytest.raises(AttributeError):
            _ = color.theme_color

        assert _solidFill(run) is None

    def it_does_not_add_solidFill_on_alpha_read(self, run_with_text):
        run = run_with_text

        # -- reading `alpha` on a color with no fill returns 1.0 (the OOXML
        # -- default); it must not mutate.
        alpha = run.font.color.alpha

        assert alpha == 1.0
        assert _solidFill(run) is None

    def it_does_not_add_solidFill_after_multiple_reads(self, run_with_text):
        run = run_with_text

        c1 = run.font.color
        c2 = run.font.color
        _ = c1.type
        _ = c2.type

        assert _solidFill(run) is None

    def it_preserves_theme_color_inheritance_after_read(self, run_with_text):
        """The full user scenario from the issue.

        Reading the run's color, then writing the pptx to a buffer, must
        leave the run's `<a:rPr>` without a `<a:solidFill>` so the run
        continues to inherit its color from the enclosing placeholder /
        layout / master.
        """
        run = run_with_text

        _ = run.font.color.type  # this used to mutate

        # -- examine XML directly; no need to serialize a full pptx to
        # -- verify the XML-level invariant.
        assert _solidFill(run) is None, (
            "reading run.font.color must not add <a:solidFill> to <a:rPr>; "
            "doing so breaks theme-color inheritance"
        )

    # -- writes still work as before -----------------------------------

    def it_still_creates_solidFill_on_rgb_write(self, run_with_text):
        run = run_with_text

        run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

        assert _solidFill(run) is not None
        assert run.font.color.rgb == RGBColor(0xFF, 0x00, 0x00)

    def it_still_creates_solidFill_on_theme_color_write(self, run_with_text):
        run = run_with_text

        run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1

        assert _solidFill(run) is not None
        assert run.font.color.theme_color == MSO_THEME_COLOR.ACCENT_1

    def it_promotes_to_solid_then_accepts_brightness(self, run_with_text):
        run = run_with_text

        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        run.font.color.brightness = 0.5

        assert run.font.color.brightness == 0.5
