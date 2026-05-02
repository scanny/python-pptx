# pyright: reportPrivateUsage=false

"""Regression test for issue #1112 — ``SlideShapes.add_picture`` + SVG input.

Issue #1112 (https://github.com/scanny/python-pptx/issues/1112) reported
that ``slide.shapes.add_picture(<some.svg>, ...)`` surfaced an opaque
``PIL.UnidentifiedImageError`` — the caller had no actionable signal
that the library simply cannot embed an SVG, and no hint that a
pre-rasterize workaround exists.

Wave 2 #652 (``feat/issue-652-svg-placeholder-handling``) addressed the
root cause by sniffing the image bytes at the package-level image-part
factory and raising the new
:class:`pptx.exc.UnsupportedImageTypeError` (with a message pointing
callers at the pre-rasterize workaround documented under the
"Inserting SVG images" section of
``docs/user/placeholders-using.rst``). Because
:meth:`SlideShapes.add_picture` funnels through
``slide.part.get_or_add_image_part`` →
``package.get_or_add_image_part`` →
``_ImageParts.get_or_add_image_part``, the #652 fix automatically
covers the :meth:`~.SlideShapes.add_picture` entry-point the #1112
reporter used — on top of the :meth:`~.PicturePlaceholder.insert_picture`
path #652 was explicitly scoped to.

#1112 is therefore a duplicate of #652 and is verified-and-closed by
this suite. The scenarios below pin the reporter's exact workflow
(``slide.shapes.add_picture(<svg>)``) end-to-end from the public API,
confirm the new :class:`UnsupportedImageTypeError` is raised (not the
pre-fix opaque ``PIL.UnidentifiedImageError``), and verify that the
error message steers the caller toward the documented pre-rasterize
workaround. The pre-rasterized-PNG happy path is exercised to prove the
sniff does not false-positive on ordinary raster input.
"""

from __future__ import annotations

import io
import pathlib

import pytest

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import UnsupportedImageTypeError
from pptx.util import Inches

# -- fixtures -----------------------------------------------------------

_MINIMAL_SVG = (
    b'<?xml version="1.0" encoding="UTF-8"?>\n'
    b'<svg xmlns="http://www.w3.org/2000/svg" '
    b'viewBox="0 0 100 100" width="100" height="100">\n'
    b'  <circle cx="50" cy="50" r="40" fill="#4B8BBE"/>\n'
    b"</svg>\n"
)


@pytest.fixture
def svg_file(tmp_path: pathlib.Path) -> str:
    """Return a filesystem path to a minimal SVG image."""
    svg_path = tmp_path / "logo.svg"
    svg_path.write_bytes(_MINIMAL_SVG)
    return str(svg_path)


# -- the verification suite ---------------------------------------------


class DescribeIssue1112SvgAddPictureVerify(object):
    """#1112 verify-and-close: ``SlideShapes.add_picture(<svg>)`` raises clearly.

    Pins that the Wave 2 #652 SVG sniff at the image-parts boundary
    reaches the :meth:`SlideShapes.add_picture` entry-point the #1112
    reporter used (not just :meth:`PicturePlaceholder.insert_picture`,
    which #652's own scenario covers), and that the raised
    :class:`UnsupportedImageTypeError` carries the pre-rasterize
    workaround hint documented under the "Inserting SVG images" section
    of ``docs/user/placeholders-using.rst``.
    """

    # -- reporter's exact call ---------------------------------------------

    def it_raises_UnsupportedImageTypeError_from_add_picture_on_svg_path(self, svg_file: str):
        """The #1112 symptom path: ``add_picture(<some.svg>)`` raises clearly.

        Pre-#652 this surfaced as ``PIL.UnidentifiedImageError`` with no
        indication of the underlying cause. The fix raises the
        library-level :class:`UnsupportedImageTypeError` so callers can
        except it by type, not by probing an opaque Pillow message.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        with pytest.raises(UnsupportedImageTypeError) as exc_info:
            slide.shapes.add_picture(svg_file, Inches(1), Inches(1))

        # -- the filename is surfaced so the caller can log/report which
        # -- specific input was rejected --
        assert "logo.svg" in str(exc_info.value)

    def it_raises_UnsupportedImageTypeError_from_add_picture_on_svg_stream(self):
        """File-like SVG input is rejected identically — reporters commonly
        pass ``BytesIO`` from HTTP fetches or in-memory generation."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        svg_stream = io.BytesIO(_MINIMAL_SVG)

        with pytest.raises(UnsupportedImageTypeError):
            slide.shapes.add_picture(svg_stream, Inches(1), Inches(1))

    def it_raises_UnsupportedImageTypeError_with_explicit_dimensions(self, svg_file: str):
        """The error path fires before any dimension/scaling math runs.

        Pins that the SVG sniff short-circuits the add-picture flow at
        the image-part boundary, well ahead of ``Image.from_file`` and
        the auto-scaling logic. No geometry-dependent failure modes can
        mask the clear type error.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        with pytest.raises(UnsupportedImageTypeError):
            slide.shapes.add_picture(svg_file, Inches(1), Inches(1), Inches(4), Inches(3))

    # -- workaround-hint contract ------------------------------------------

    def it_steers_the_caller_to_the_pre_rasterize_workaround(self, svg_file: str):
        """The raised message cites the documented workaround.

        The user guide's "Inserting SVG images" section at
        ``docs/user/placeholders-using.rst`` documents pre-rasterizing
        the SVG to PNG (``cairosvg``, ``svglib`` + ``reportlab``, etc.)
        as the supported workaround. The exception message mirrors that
        guidance so the first-time reader sees the path forward without
        having to consult the docs.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        with pytest.raises(UnsupportedImageTypeError) as exc_info:
            slide.shapes.add_picture(svg_file, Inches(1), Inches(1))

        message = str(exc_info.value)
        # -- cites the SVG format as the cause --
        assert "SVG" in message or "svg" in message
        # -- names the recommended workaround --
        assert "PNG" in message or "png" in message
        assert "rasterize" in message.lower()
        # -- cross-refs the user-guide section --
        assert "user guide" in message.lower() or "Inserting SVG" in message

    def it_raises_the_pptx_error_not_the_underlying_pillow_error(self, svg_file: str):
        """The public exception is a |PythonPptxError| subclass.

        Important contract: the reporter's `except PIL.UnidentifiedImageError`
        sites migrate to ``except UnsupportedImageTypeError`` (or its
        parent ``PythonPptxError``) — not to a third-party exception.
        """
        from pptx.exc import PythonPptxError

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        with pytest.raises(UnsupportedImageTypeError) as exc_info:
            slide.shapes.add_picture(svg_file, Inches(1), Inches(1))

        # -- the type is a library-level error, not a third-party one --
        assert isinstance(exc_info.value, PythonPptxError)

    # -- no-regression on the happy path -----------------------------------

    def it_still_accepts_a_pre_rasterized_png(self):
        """A PNG still embeds cleanly via ``add_picture``.

        The SVG sniff must not false-positive on ordinary raster input.
        This is the user-facing "workaround" in action — the reporter
        pre-rasterizes to PNG (using the rasterizer of their choice) and
        feeds the resulting file through the same ``add_picture`` call.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        pic = slide.shapes.add_picture(
            "tests/test_files/python-powered.png",
            Inches(1),
            Inches(1),
            Inches(2),
            Inches(1),
        )

        assert pic.shape_type == MSO_SHAPE_TYPE.PICTURE
        # -- ends up embedded in the slide's spTree --
        assert pic._element in list(slide.shapes._spTree)

    def it_round_trips_the_pre_rasterized_png_through_save_reopen(self):
        """The workaround survives ``Presentation.save`` + reopen.

        End-to-end pin: author → save → reopen and confirm the picture
        is still there. This closes the loop on the #1112 reporter's
        workflow once they migrate from the SVG file to a PNG.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(
            "tests/test_files/python-powered.png",
            Inches(1),
            Inches(1),
            Inches(2),
            Inches(1),
        )

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        pics = [s for s in prs2.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert len(pics) == 1

    # -- package is not mutated by the failed call -------------------------

    def it_leaves_the_slide_untouched_when_the_svg_is_rejected(self, svg_file: str):
        """The failed ``add_picture`` is atomic — no shape is appended.

        Pins that the SVG sniff fires before any ``p:pic`` is grafted
        onto the shape tree, so the caller can recover (e.g. swap the
        SVG for a PNG fallback and retry) without having to reset the
        slide.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        before = list(slide.shapes._spTree)

        with pytest.raises(UnsupportedImageTypeError):
            slide.shapes.add_picture(svg_file, Inches(1), Inches(1))

        after = list(slide.shapes._spTree)
        assert before == after

    # -- public API surface -----------------------------------------------

    def it_exposes_UnsupportedImageTypeError_on_the_pptx_exc_module(self):
        """The exception class the caller needs to import is reachable.

        Belt-and-braces: the reporter's migration is
        ``from pptx.exc import UnsupportedImageTypeError``; pin that the
        public import path works.
        """
        from pptx import exc

        assert hasattr(exc, "UnsupportedImageTypeError")
        assert issubclass(exc.UnsupportedImageTypeError, exc.PythonPptxError)
