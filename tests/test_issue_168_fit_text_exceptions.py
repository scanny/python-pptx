# pyright: reportPrivateUsage=false

"""Regression test for issue #168 — ``TextFrame.fit_text()`` exceptions.

Issue #168 (https://github.com/scanny/python-pptx/issues/168) reports that
:meth:`TextFrame.fit_text` raises opaque exceptions in several common edge
cases:

* The platform's installed-fonts directory can't be auto-discovered
  (e.g. Linux) — ``FontFiles._font_directories()`` raised a bare
  ``OSError: unsupported operating system`` with no guidance for the
  caller.
* The requested ``font_family`` / ``bold`` / ``italic`` tuple has no
  matching font in the system's font directory — ``FontFiles.find()``
  raised a bare ``KeyError`` tuple.
* The supplied ``font_file`` path doesn't exist, isn't readable, or
  isn't a TrueType/OpenType file — Pillow's
  ``ImageFont.truetype()`` raised ``OSError: cannot open resource``
  / ``OSError: unknown file format`` from deep inside the layout loop.
* ``max_size`` is ``0`` or negative — the binary-search tree over the
  range ``1..max_size`` was empty, producing ``IndexError: pop from
  empty list``.
* The text-frame's effective rendering area is non-positive because
  the left/right (or top/bottom) margins exceed the shape's width
  (height) — the layout loop could not find any fitting size and
  surfaced an unhelpful "nothing fits" message.

All of these now raise :class:`pptx.exc.TextLayoutError` with an
actionable message pointing the caller at the recommended workaround
(supply an explicit ``font_file`` or turn autofit off).

The layout-loop "no point size fits" path (#773) and the
"wrap-at-smaller-size" path (#936) continue to be covered by their own
regression tests; this file adds coverage for the remaining #168
failure modes.
"""

from __future__ import annotations

from os.path import abspath, dirname, join
from typing import cast

import pytest

from pptx import Presentation
from pptx.exc import TextLayoutError
from pptx.text.fonts import FontFiles
from pptx.text.layout import TextFitter
from pptx.util import Inches

FONT_FILE = abspath(
    join(
        dirname(__file__),
        "..",
        "features",
        "steps",
        "test_files",
        "calibriz.ttf",
    )
)


def _text_frame_with_text(text: str = "Hello"):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    tb.text_frame.text = text
    return tb.text_frame


class DescribeIssue168FitTextExceptions(object):
    """Regression scenarios for #168 ``TextFrame.fit_text`` exceptions."""

    # ------------------------------------------------------------------
    # TextFrame.fit_text() wrapping
    # ------------------------------------------------------------------

    def it_raises_TextLayoutError_when_platform_fonts_cannot_be_discovered(
        self, monkeypatch: pytest.MonkeyPatch
    ):
        # -- stub `_font_directories` to raise the same OSError that the
        # -- Linux code path raises, so the behavior is exercised
        # -- regardless of which OS the tests run on.
        def raise_unsupported(cls):
            raise OSError("unsupported operating system")

        # -- also clear the memoized cache so the stub actually runs --
        monkeypatch.setattr(FontFiles, "_font_files", None)
        monkeypatch.setattr(
            FontFiles,
            "_font_directories",
            classmethod(raise_unsupported),
        )

        text_frame = _text_frame_with_text("Hello world!")

        with pytest.raises(TextLayoutError) as exc_info:
            text_frame.fit_text()

        msg = str(exc_info.value)
        assert "cannot auto-discover installed fonts" in msg
        assert "font_file" in msg
        assert "auto_size" in msg
        # -- original OSError is chained as __cause__ for debuggability --
        assert isinstance(exc_info.value.__cause__, OSError)

    def it_raises_TextLayoutError_when_font_family_not_installed(
        self, monkeypatch: pytest.MonkeyPatch
    ):
        # -- simulate "fonts are discoverable but none match the request":
        # -- prime `_font_files` with an empty dict so `FontFiles.find()`
        # -- raises KeyError when it tries to look up the requested family.
        monkeypatch.setattr(FontFiles, "_font_files", {})

        text_frame = _text_frame_with_text("Hello")

        with pytest.raises(TextLayoutError) as exc_info:
            text_frame.fit_text(font_family="NonExistentFontXYZ")

        msg = str(exc_info.value)
        assert "NonExistentFontXYZ" in msg
        assert "font_file" in msg
        assert "auto_size" in msg
        # -- original KeyError chained as __cause__ --
        assert isinstance(exc_info.value.__cause__, KeyError)

    def it_raises_TextLayoutError_when_font_file_path_is_missing(self):
        text_frame = _text_frame_with_text("Hello")

        with pytest.raises(TextLayoutError) as exc_info:
            text_frame.fit_text(font_file="/nonexistent/path/to/font.ttf")

        msg = str(exc_info.value)
        assert "cannot read font file" in msg
        assert "font.ttf" in msg
        assert "font_file" in msg
        assert "auto_size" in msg
        assert isinstance(exc_info.value.__cause__, OSError)

    def it_raises_TextLayoutError_when_font_file_is_not_a_font(self, tmp_path):
        # -- a real file that Pillow cannot decode as a TrueType font --
        bogus = tmp_path / "not-a-font.ttf"
        bogus.write_bytes(b"this is plainly not a TrueType file")

        text_frame = _text_frame_with_text("Hello")

        with pytest.raises(TextLayoutError) as exc_info:
            text_frame.fit_text(font_file=str(bogus))

        msg = str(exc_info.value)
        assert "cannot read font file" in msg
        assert str(bogus) in msg

    def it_raises_TextLayoutError_when_margins_exceed_shape_width(self):
        # -- default margins total ~0.2" L+R and ~0.1" T+B; a 0.1" x 0.1"
        # -- textbox will have a negative effective rendering area.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(0.1), Inches(0.1))
        text_frame = tb.text_frame
        text_frame.text = "Hello"

        with pytest.raises(TextLayoutError) as exc_info:
            text_frame.fit_text(font_file=FONT_FILE)

        msg = str(exc_info.value)
        assert "no positive rendering area" in msg
        assert "margins" in msg
        assert "auto_size" in msg

    def it_raises_TextLayoutError_when_set_margins_exceed_shape_height(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        text_frame = tb.text_frame
        text_frame.text = "Hello"
        # -- deliberately set top+bottom margins larger than the shape's height --
        text_frame.margin_top = Inches(2)
        text_frame.margin_bottom = Inches(2)

        with pytest.raises(TextLayoutError) as exc_info:
            text_frame.fit_text(font_file=FONT_FILE)

        assert "no positive rendering area" in str(exc_info.value)

    # ------------------------------------------------------------------
    # TextFitter.best_fit_font_size() guard on `max_size`
    # ------------------------------------------------------------------

    @pytest.mark.parametrize("max_size", [0, -1, -17])
    def it_raises_TextLayoutError_for_non_positive_max_size(self, max_size: int):
        with pytest.raises(TextLayoutError) as exc_info:
            TextFitter.best_fit_font_size("Hello", (500000, 500000), max_size, FONT_FILE)

        msg = str(exc_info.value)
        assert "max_size must be >= 1" in msg
        assert str(max_size) in msg

    def it_still_succeeds_for_min_valid_max_size_of_1(self):
        # -- guard must not fire when max_size is exactly 1 (the smallest
        # -- valid candidate); the fitter returns 1 on "trivial fit" and
        # -- TextLayoutError on "single-EMU width".
        size = TextFitter.best_fit_font_size("hi", (500000, 500000), 1, FONT_FILE)
        assert size == 1

    # ------------------------------------------------------------------
    # verify-and-pin: previously-fixed scenarios still behave
    # ------------------------------------------------------------------

    def it_still_raises_TextLayoutError_when_no_point_size_fits_773(self):
        # -- single EMU wide: no glyph can render; TextFitter should still
        # -- raise TextLayoutError with the pre-existing "cannot be fit"
        # -- message (issue #773).
        with pytest.raises(TextLayoutError, match="cannot be fit"):
            TextFitter.best_fit_font_size(
                "Supercalifragilisticexpialidocious",
                (1, 1),
                18,
                FONT_FILE,
            )

    def it_still_wraps_at_smaller_size_when_max_overflows_936(self):
        # -- the pre-existing #936 path: text wraps at a smaller size when
        # -- the longest word overflows at `max_size` — must not be broken
        # -- by the new OSError wrapping inside `_best_fit_font_size`.
        size = TextFitter.best_fit_font_size(
            "Supercalifragilisticexpialidocious is a long word",
            (500000, 2000000),
            36,
            FONT_FILE,
        )
        assert isinstance(size, int)
        assert 1 <= size < 36

    def it_fit_text_still_succeeds_on_happy_path(self):
        # -- supply a valid font_file; the method must succeed without
        # -- touching the platform-font discovery path.
        text_frame = _text_frame_with_text("Hello world, fit me!")
        text_frame.fit_text(font_file=FONT_FILE)

        # -- after fit_text, word_wrap is True and every run carries a
        # -- concrete font size.
        assert text_frame.word_wrap is True
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                size = run.font.size
                assert size is not None
                size_pt = cast(float, size.pt)
                assert 1 <= size_pt <= 18
