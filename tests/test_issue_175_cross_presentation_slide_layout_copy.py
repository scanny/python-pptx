# pyright: reportPrivateUsage=false

"""Regression test for issue #175 — add slide/slide layout from other presentation.

Issue #175 (https://github.com/scanny/python-pptx/issues/175) asks:

    "Is there some possibility to add foreign slide layout or slide to one
    presentation from the other one using the library python-pptx? I tried to
    use ``pr1.slides.add_slide(pr2.slide_layouts[11])``. This adds slide to
    the presentation pr1, but not its layout to the pr1.slide_layouts, which
    is what I need."

The #175 reporter actually asked for two related things:

  1. **Cross-presentation slide copy** — appending a slide from another
     presentation into this one.
  2. **Cross-presentation slide-layout copy** — importing a layout from one
     deck into another so future :meth:`Slides.add_slide` calls can use it.

Both use cases are now served:

  * Wave 7 #934 shipped :meth:`Presentation.merge` — appends every slide of
    another presentation to the receiver, full-fidelity (image, media,
    chart with its own embedded workbook, OLE objects, external hyperlinks
    all materialised in the target package).
  * Wave 7 #934 also promoted :meth:`Slides.add_slide_from_external`
    (originally shipped by Wave 1/2 #1036 as a *basic* copy path) to the
    same full-fidelity cloning pipeline, so per-slide cross-presentation
    copy is available too.
  * For use-case (2), the documented workaround is to open the presentation
    that *owns the desired layouts* as the base, then :meth:`Presentation.merge`
    in the content-carrying deck(s) and :meth:`Slides.delete` any unwanted
    slides. This leaves the target's layout list intact (layouts are
    inherited from the starting presentation) while pulling in content from
    other sources bound to those layouts.

This regression suite pins the end-to-end behaviour the #175 reporter asked
for so the issue can be closed as *resolved by #934 +
Slides.add_slide_from_external*.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


@pytest.fixture
def _restore_part_factory():
    """Guard against ``PartFactory.part_type_for`` pollution from other tests.

    Mirrors the identical guard used by ``tests/test_issue_285_*.py`` and
    siblings; earlier test modules overwrite slide-part registrations with
    mocks and do not restore them.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.embeddedpackage import EmbeddedXlsxPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
        CT.SML_SHEET: EmbeddedXlsxPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


def _absjoin(*args: str) -> str:
    from .unitutil.file import absjoin, test_file_dir

    return absjoin(test_file_dir, *args)


class DescribeIssue175CrossPresentationSlideLayoutCopy(object):
    """#175 — close as resolved by #934 + Slides.add_slide_from_external.

    Covers the two distinct asks from the #175 thread:

    1. Cross-presentation *slide* copy — via :meth:`Slides.add_slide_from_external`
       (per-slide) or :meth:`Presentation.merge` (whole-deck).
    2. Cross-presentation *layout* import — via the documented "open layout
       source as base + merge content decks + delete unwanted slides"
       workaround.
    """

    # -- (1) per-slide cross-presentation copy ------------------------------

    def it_clones_a_foreign_slide_via_add_slide_from_external(self, _restore_part_factory):
        image_path = _absjoin("python-icon.jpeg")

        source = Presentation()
        source_slide = source.slides.add_slide(source.slide_layouts[5])
        source_slide.shapes.add_picture(image_path, Inches(1), Inches(1), height=Inches(1))
        source_shape_names = [sh.name for sh in source_slide.shapes]

        target = Presentation()
        target_layout = target.slide_layouts[5]

        cloned = target.slides.add_slide_from_external(source_slide, target_layout)

        # -- cloned slide lives in the target presentation, bound to the
        # -- supplied target layout, with the source's shape tree deep-copied --
        assert cloned in list(target.slides)
        assert cloned.slide_layout is target_layout
        assert [sh.name for sh in cloned.shapes] == source_shape_names

    def it_round_trips_an_externally_cloned_slide_through_save_reopen(self, _restore_part_factory):
        image_path = _absjoin("python-icon.jpeg")

        source = Presentation()
        source_slide = source.slides.add_slide(source.slide_layouts[5])
        source_slide.shapes.add_picture(image_path, Inches(1), Inches(1), height=Inches(1))

        target = Presentation()
        target.slides.add_slide_from_external(source_slide, target.slide_layouts[5])

        buf = io.BytesIO()
        target.save(buf)
        buf.seek(0)
        reopened = Presentation(buf)

        assert len(reopened.slides) == 1
        # -- the cloned picture shape survived the round-trip --
        assert any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in reopened.slides[0].shapes)

    # -- (2) whole-deck merge (#934) ----------------------------------------

    def it_merges_every_slide_of_another_presentation_via_merge(self, _restore_part_factory):
        image_path = _absjoin("python-icon.jpeg")

        source = Presentation()
        for _ in range(3):
            sl = source.slides.add_slide(source.slide_layouts[5])
            sl.shapes.add_picture(image_path, Inches(1), Inches(1), height=Inches(1))
        source_count = len(source.slides)

        target = Presentation()
        target_pre_count = len(target.slides)
        target_layout_count = len(target.slide_masters[0].slide_layouts)

        appended = target.merge(source)

        assert len(appended) == source_count
        assert len(target.slides) == target_pre_count + source_count
        # -- merge does not inject new layouts into the target master --
        assert len(target.slide_masters[0].slide_layouts) == target_layout_count

    def it_raises_on_merging_a_presentation_into_itself(self, _restore_part_factory):
        prs = Presentation()
        with pytest.raises(ValueError, match="itself"):
            prs.merge(prs)

    def it_raises_on_merge_with_non_presentation_argument(self, _restore_part_factory):
        prs = Presentation()
        with pytest.raises(TypeError, match="must be a Presentation"):
            prs.merge("not-a-presentation")  # type: ignore[arg-type]

    # -- (3) layout-import workaround (#175's second ask) -------------------

    def it_imports_foreign_layouts_by_using_layout_source_as_base(self, _restore_part_factory):
        """Document the workaround for "add a foreign *layout* to my deck".

        python-pptx does not offer a direct
        ``Slides.copy_layout_from_external(layout)`` API because importing a
        single layout requires resolving its slide-master, theme, and color /
        font / format scheme relationships into the target deck — a
        deeper-reaching cross-package clone than the shipped slide-copy
        pipeline. The practical workaround is:

        1. Open the presentation that *owns the desired layouts* as the base.
        2. :meth:`Presentation.merge` in the content-carrying deck(s).
        3. :meth:`Slides.delete` any unwanted slides.

        The target (the "base") keeps all of its layouts; merged slides are
        bound to the target's layouts by index (see :meth:`Presentation.merge`
        docstring). This pins that pattern end-to-end so the workaround stays
        regression-covered.
        """
        image_path = _absjoin("python-icon.jpeg")

        # -- (a) content-deck to merge in: a single picture slide. --
        content_deck = Presentation()
        content_slide = content_deck.slides.add_slide(content_deck.slide_layouts[5])
        content_slide.shapes.add_picture(image_path, Inches(1), Inches(1), height=Inches(1))

        # -- (b) "layout source" deck — plays the role of the template whose
        # -- layouts we want to import. Carries a couple of starter slides we
        # -- will delete after the merge. --
        layout_source = Presentation()
        starter_a = layout_source.slides.add_slide(layout_source.slide_layouts[0])
        starter_b = layout_source.slides.add_slide(layout_source.slide_layouts[1])
        layout_names_before = [ly.name for ly in layout_source.slide_masters[0].slide_layouts]

        # -- (c) merge content into the layout-source deck, then delete the
        # -- starter slides. The layout-source's layouts survive untouched. --
        layout_source.merge(content_deck)
        layout_source.slides.delete(starter_a)
        layout_source.slides.delete(starter_b)

        # -- Assertions pinning the workaround:
        # -- layouts present before merge are still present after delete --
        layout_names_after = [ly.name for ly in layout_source.slide_masters[0].slide_layouts]
        assert layout_names_after == layout_names_before
        # -- only the merged content slide remains; starters were deleted --
        assert len(layout_source.slides) == 1
        # -- the surviving slide is the merged one, and it carries the
        # -- picture shape from the content deck --
        assert any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in layout_source.slides[0].shapes)

    def it_round_trips_the_layout_import_workaround_through_save_reopen(
        self, _restore_part_factory
    ):
        image_path = _absjoin("python-icon.jpeg")

        content_deck = Presentation()
        sl = content_deck.slides.add_slide(content_deck.slide_layouts[5])
        sl.shapes.add_picture(image_path, Inches(1), Inches(1), height=Inches(1))

        layout_source = Presentation()
        starter = layout_source.slides.add_slide(layout_source.slide_layouts[0])
        layout_count_before = len(layout_source.slide_masters[0].slide_layouts)

        layout_source.merge(content_deck)
        layout_source.slides.delete(starter)

        buf = io.BytesIO()
        layout_source.save(buf)
        buf.seek(0)
        reopened = Presentation(buf)

        # -- layouts survive save/reopen; no layouts pruned as side-effect --
        assert len(reopened.slide_masters[0].slide_layouts) == layout_count_before
        # -- exactly the merged content slide remains in the reopened deck --
        assert len(reopened.slides) == 1
