# pyright: reportPrivateUsage=false

"""Regression test for issue #234 — ``FillFormat.blip_fill``.

Issue #234 (https://github.com/scanny/python-pptx/issues/234) asks for a
way to fill an auto-shape (or table cell, slide background, etc.) with a
picture — the "Picture or texture fill" option in PowerPoint's "Format
Shape" pane. The feature is implemented by
``FillFormat.blip_fill(image_file)``, which adds (or reuses) an
``ImagePart`` on the containing part and rewrites the underlying
``EG_FillProperties`` as ``<a:blipFill><a:blip r:embed="..."/>
<a:stretch><a:fillRect/></a:stretch></a:blipFill>``.

This test exercises the end-to-end flow:

  1. Add a plain rectangular auto-shape to a fresh slide.
  2. Call ``shape.fill.blip_fill("tests/test_files/python-powered.png")``.
  3. Round-trip the presentation through save / reopen.
  4. Verify the reloaded shape's fill type is ``MSO_FILL.PICTURE`` and
     that the image relationship survived round-trip.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches


class DescribeIssue234BlipFill:
    """Round-trip regression for ``FillFormat.blip_fill`` (issue #234)."""

    def it_fills_an_auto_shape_with_a_picture(self):
        image_path = "tests/test_files/python-powered.png"
        prs = Presentation()
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2)
        )

        # -- act: apply a picture fill --
        shape.fill.blip_fill(image_path)

        # -- assert: fill.type is PICTURE --
        assert shape.fill.type == MSO_FILL.PICTURE
        # -- assert: an a:blipFill/a:blip element exists with r:embed --
        blips = shape._element.xpath(".//a:blipFill/a:blip")
        assert len(blips) == 1
        rId = blips[0].rEmbed
        assert rId is not None
        # -- assert: image part was registered on the slide part --
        slide_part = slide.part
        assert slide_part.related_part(rId) is not None

        # -- round-trip the presentation through save / reopen --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)
        reloaded_shape = reloaded.slides[0].shapes[0]

        # -- assert: reloaded shape still has a picture fill --
        assert reloaded_shape.fill.type == MSO_FILL.PICTURE

        # -- assert: the embedded image bytes match the source file --
        with open(image_path, "rb") as f:
            expected_bytes = f.read()
        reloaded_blips = reloaded_shape._element.xpath(".//a:blipFill/a:blip")
        assert len(reloaded_blips) == 1
        reloaded_rId = reloaded_blips[0].rEmbed
        reloaded_image_part = reloaded.slides[0].part.related_part(reloaded_rId)
        assert reloaded_image_part.blob == expected_bytes

    def it_reuses_an_existing_image_part_when_bytes_match(
        self
    ):
        image_path = "tests/test_files/python-powered.png"
        prs = Presentation()
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        shape_a = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(2)
        )
        shape_b = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(4), Inches(1), Inches(2), Inches(2)
        )

        shape_a.fill.blip_fill(image_path)
        shape_b.fill.blip_fill(image_path)

        # -- same image bytes -> the same underlying ImagePart is reused --
        blip_a = shape_a._element.xpath(".//a:blipFill/a:blip")[0]
        blip_b = shape_b._element.xpath(".//a:blipFill/a:blip")[0]
        slide_part = slide.part
        # -- both rIds (may differ) resolve to the same ImagePart instance --
        assert slide_part.related_part(blip_a.rEmbed) is slide_part.related_part(
            blip_b.rEmbed
        )

    def it_raises_on_fills_created_without_a_part_reference(self):
        from pptx.dml.fill import FillFormat

        from .unitutil.cxml import element

        fill = FillFormat.from_fill_parent(element("p:spPr"))

        with pytest.raises(ValueError, match="blip_fill requires a part reference"):
            fill.blip_fill("tests/test_files/python-powered.png")
