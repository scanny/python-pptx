# pyright: reportPrivateUsage=false

"""Regression test for issue #235 — add support for 3-D pie chart type.

Issue #235 (https://github.com/scanny/python-pptx/issues/235) asked for
``Shapes.add_chart`` to accept :attr:`XL_CHART_TYPE.THREE_D_PIE` and emit
a PowerPoint-readable chart. On pre-fork python-pptx the factory raised
``NotImplementedError: XML writer for chart type THREE_D_PIE (-4102) not
yet implemented`` for every ``THREE_D_*`` enum member.

Wave 7 #266 (``feat/issue-266-3d-chart-types``) landed the missing
writers — including ``_Pie3DChartXmlWriter`` — so ``ChartXmlWriter`` now
dispatches both :attr:`XL_CHART_TYPE.THREE_D_PIE` and
:attr:`XL_CHART_TYPE.THREE_D_PIE_EXPLODED` to a builder that emits a
schema-valid ``c:pie3DChart`` wrapper plus a ``c:view3D`` sibling on
``c:chart`` using PowerPoint's default angles (``rotX=15``, ``rotY=20``,
``rAngAx=1``, ``depthPercent=100``).

This suite pins the #235 resolution so the issue can be closed. Each
test exercises the real API against a fresh ``Presentation()`` via
``Shapes.add_chart`` — not a mock, and not the raw ``ChartXmlWriter``.

Known limitation (out of scope for #235): reading a ``c:pie3DChart``
back and iterating ``chart.plots`` still raises ``ValueError`` because
``PlotFactory`` has no ``Pie3DPlot``. That read-access gap is a generic
3D-chart concern (applies to every 3D chart type except
``THREE_D_AREA``); the #266 MVP contract is write-only, and #235's
original ask — "add_chart accepts THREE_D_PIE" — is satisfied.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

# -- fixtures -----------------------------------------------------------


@pytest.fixture
def _restore_part_factory():
    """Guard against pollution from tests that mutate `PartFactory.part_type_for`.

    ``tests/opc/test_package.py::DescribePartFactory`` overwrites
    ``PartFactory.part_type_for[CT.PML_SLIDE]`` with a Mock and does
    not restore it. Round-trip tests that call ``Presentation.save``
    + ``Presentation()`` reopen depend on the real registrations being
    in place. Mirrors the identical guard elsewhere in the
    regression-test suite.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


def _pie_chart_data():
    """Return a three-slice ``CategoryChartData`` for pie-chart scenarios."""
    data = CategoryChartData()
    data.categories = ["Apples", "Bananas", "Cherries"]
    data.add_series("Fruit", (25, 35, 40))
    return data


# -- the verification suite ---------------------------------------------


class DescribeIssue235ThreeDPieVerify(object):
    """#235 verify-and-close: ``add_chart(THREE_D_PIE, ...)`` works end-to-end.

    Pins that both ``XL_CHART_TYPE.THREE_D_PIE`` and
    ``XL_CHART_TYPE.THREE_D_PIE_EXPLODED`` are accepted by
    :meth:`Shapes.add_chart`, emit a ``c:pie3DChart`` wrapper with a
    ``c:view3D`` sibling on ``c:chart``, and round-trip through
    ``Presentation.save`` + reopen intact. #266 shipped the writer;
    this suite is the #235-specific regression pin.
    """

    # -- authoring path -----------------------------------------------------

    def it_accepts_THREE_D_PIE_as_an_add_chart_type(self):
        """``add_chart`` no longer raises ``NotImplementedError`` for THREE_D_PIE."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        # -- act: this call raised NotImplementedError on pre-fork python-pptx --
        graphic_frame = slide.shapes.add_chart(
            XL_CHART_TYPE.THREE_D_PIE,
            Inches(1),
            Inches(1),
            Inches(6),
            Inches(4),
            _pie_chart_data(),
        )

        # -- assert: a real chart is returned with the expected XML shape --
        chart = graphic_frame.chart
        chartSpace = chart._chartSpace
        assert chartSpace.find(f".//{qn('c:pie3DChart')}") is not None
        # -- c:view3D is a sibling of c:plotArea under c:chart --
        c_chart = chartSpace.find(qn("c:chart"))
        assert c_chart is not None
        view3D = c_chart.find(qn("c:view3D"))
        assert view3D is not None
        # -- PowerPoint's default 3D angles --
        assert view3D.find(qn("c:rotX")).get("val") == "15"
        assert view3D.find(qn("c:rotY")).get("val") == "20"
        assert view3D.find(qn("c:rAngAx")).get("val") == "1"
        assert view3D.find(qn("c:depthPercent")).get("val") == "100"

    def it_accepts_THREE_D_PIE_EXPLODED_as_an_add_chart_type(self):
        """``add_chart`` accepts the exploded variant and emits c:explosion."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        graphic_frame = slide.shapes.add_chart(
            XL_CHART_TYPE.THREE_D_PIE_EXPLODED,
            Inches(1),
            Inches(1),
            Inches(6),
            Inches(4),
            _pie_chart_data(),
        )

        chart = graphic_frame.chart
        chartSpace = chart._chartSpace
        # -- pie3DChart wrapper still present --
        assert chartSpace.find(f".//{qn('c:pie3DChart')}") is not None
        # -- the exploded variant injects a c:explosion child on the series --
        explosion_elems = chartSpace.xpath(".//c:pie3DChart/c:ser/c:explosion")
        assert len(explosion_elems) == 1
        assert explosion_elems[0].get("val") == "25"

    def it_does_not_emit_explosion_for_plain_THREE_D_PIE(self):
        """The non-exploded variant has no ``c:explosion`` child."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        graphic_frame = slide.shapes.add_chart(
            XL_CHART_TYPE.THREE_D_PIE,
            Inches(1),
            Inches(1),
            Inches(6),
            Inches(4),
            _pie_chart_data(),
        )

        chart = graphic_frame.chart
        assert chart._chartSpace.xpath(".//c:pie3DChart/c:ser/c:explosion") == []

    def it_carries_the_series_values_into_the_pie3DChart(self):
        """The series categories and values survive the XML emission."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        graphic_frame = slide.shapes.add_chart(
            XL_CHART_TYPE.THREE_D_PIE,
            Inches(1),
            Inches(1),
            Inches(6),
            Inches(4),
            _pie_chart_data(),
        )

        chartSpace = graphic_frame.chart._chartSpace
        assert chartSpace.find(f".//{qn('c:pie3DChart')}") is not None
        # -- three categories authored, three c:pt elements in c:cat --
        cat_pt_count = len(chartSpace.xpath(".//c:pie3DChart/c:ser/c:cat//c:pt"))
        assert cat_pt_count == 3
        # -- three numeric values --
        val_pt_count = len(chartSpace.xpath(".//c:pie3DChart/c:ser/c:val//c:pt"))
        assert val_pt_count == 3

    # -- round-trip through save + reopen -----------------------------------

    def it_round_trips_a_THREE_D_PIE_chart_through_save_reopen(self, _restore_part_factory):
        """A ``THREE_D_PIE`` chart survives ``Presentation.save`` + reopen."""
        # -- arrange: author a deck with a 3-D pie chart --
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.add_chart(
            XL_CHART_TYPE.THREE_D_PIE,
            Inches(1),
            Inches(1),
            Inches(6),
            Inches(4),
            _pie_chart_data(),
        )

        # -- act: save to a buffer and reopen --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        # -- assert: the reopened chart still has the pie3DChart wrapper --
        graphic_frame = next(s for s in prs2.slides[0].shapes if s.has_chart)
        reloaded_chart = graphic_frame.chart
        chartSpace = reloaded_chart._chartSpace
        assert chartSpace.find(f".//{qn('c:pie3DChart')}") is not None
        view3D = chartSpace.find(f".//{qn('c:view3D')}")
        assert view3D is not None
        assert view3D.find(qn("c:rotX")).get("val") == "15"

    def it_round_trips_a_THREE_D_PIE_EXPLODED_chart_through_save_reopen(
        self, _restore_part_factory
    ):
        """The exploded variant also round-trips through save + reopen."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.add_chart(
            XL_CHART_TYPE.THREE_D_PIE_EXPLODED,
            Inches(1),
            Inches(1),
            Inches(6),
            Inches(4),
            _pie_chart_data(),
        )

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        graphic_frame = next(s for s in prs2.slides[0].shapes if s.has_chart)
        chartSpace = graphic_frame.chart._chartSpace
        # -- c:explosion is preserved through round-trip --
        explosion_elems = chartSpace.xpath(".//c:pie3DChart/c:ser/c:explosion")
        assert len(explosion_elems) == 1
        assert explosion_elems[0].get("val") == "25"

    # -- regression: pre-fork error is gone ---------------------------------

    def it_does_not_raise_NotImplementedError_for_THREE_D_PIE(self):
        """The exact pre-fork error path no longer fires.

        Before #266, ``ChartData.xml_bytes(XL_CHART_TYPE.THREE_D_PIE)``
        raised ``NotImplementedError`` with the message
        ``"XML writer for chart type THREE_D_PIE (-4102) not yet
        implemented"``. This regression pins that the factory now
        dispatches successfully.
        """
        data = _pie_chart_data()

        # -- act: go straight through the ChartData -> ChartXmlWriter path --
        xml_bytes = data.xml_bytes(XL_CHART_TYPE.THREE_D_PIE)

        # -- assert: a real XML document came back, not an exception --
        assert xml_bytes.startswith(b"<?xml")
        assert b"pie3DChart" in xml_bytes
