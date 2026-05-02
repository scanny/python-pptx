# pyright: reportPrivateUsage=false

"""Regression test for issue #246 — "ability to delete or replace shape".

Issue #246 (https://github.com/scanny/python-pptx/issues/246) asked for a way
to (a) delete an existing shape on a slide and (b) replace an existing picture
with a different image while preserving its position in the slide.

Part (a) was shipped by #41 as ``BaseShape.delete()``. This branch adds
``BaseShape.replace_with(other_shape)`` — a small convenience wrapper that
copies the existing shape's position and size onto a newly-added replacement
shape, moves the replacement into the existing shape's z-order slot, then
deletes the existing shape. This is the minimum ergonomic gap between the
#41 primitive and the "swap picture while preserving layout" flow the #246
reporter described.

The scenario below round-trips a real package and uses only the public API.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.util import Inches


class DescribeIssue246RegressionShapeReplace(object):
    """Regression scenarios for #246 "ability to delete or replace shape"."""

    def it_deletes_a_shape_via_BaseShape_delete(self):
        # -- #41 primitive: add three shapes, delete the middle one --
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shapes = slide.shapes
        sp1 = shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        sp2 = shapes.add_textbox(Inches(3), Inches(1), Inches(2), Inches(1))
        sp3 = shapes.add_textbox(Inches(5), Inches(1), Inches(2), Inches(1))
        sp1_name, sp3_name = sp1.name, sp3.name

        sp2.delete()

        # -- roundtrip through save/reopen --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        names = [s.name for s in prs2.slides[0].shapes]
        assert sp1_name in names
        assert sp3_name in names
        assert len(names) == 2

    def it_replaces_an_image_preserving_position_and_zorder(
        self
    ):
        # -- #246 primary ask: replace picture A with picture B at same spot --
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shapes = slide.shapes

        # -- a background shape (behind) so we can verify z-order --
        bg = shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(1))
        bg_name = bg.name

        old_pic = shapes.add_picture(
            "tests/test_files/python-powered.png",
            Inches(2),
            Inches(3),
            Inches(4),
            Inches(2),
        )
        original_left = old_pic.left
        original_top = old_pic.top
        original_width = old_pic.width
        original_height = old_pic.height

        # -- a foreground textbox after the picture --
        fg = shapes.add_textbox(Inches(0), Inches(6), Inches(1), Inches(1))
        fg_name = fg.name

        # -- caller adds the new picture (appended at end / topmost) --
        new_pic = shapes.add_picture(
            "tests/test_files/monty-truth.png",
            Inches(0),  # placeholder position, will be overwritten
            Inches(0),
            Inches(1),
            Inches(1),
        )

        # -- one-call replacement --
        old_pic.replace_with(new_pic)

        # -- new picture inherits original picture's geometry --
        assert new_pic.left == original_left
        assert new_pic.top == original_top
        assert new_pic.width == original_width
        assert new_pic.height == original_height

        # -- new picture sits in the old picture's z-order slot (between bg and fg) --
        z_names = [s.name for s in slide.shapes]
        assert z_names.index(bg_name) < z_names.index(new_pic.name)
        assert z_names.index(new_pic.name) < z_names.index(fg_name)

        # -- round-trip through save/reopen --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        shapes2 = list(prs2.slides[0].shapes)
        # -- exactly one picture (the new one), at the original geometry --
        pics = [s for s in shapes2 if s.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
        assert len(pics) == 1
        assert pics[0].left == original_left
        assert pics[0].top == original_top
        assert pics[0].width == original_width
        assert pics[0].height == original_height

    def it_rejects_replacing_a_shape_with_itself(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        sp = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))

        with pytest.raises(ValueError, match="itself"):
            sp.replace_with(sp)
