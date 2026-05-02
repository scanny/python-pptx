# pyright: reportPrivateUsage=false

"""Regression test for issue #285 — change text without editing formatting.

Issue #285 (https://github.com/scanny/python-pptx/issues/285) asks:

    "To edit some text i currently use pharagraphs and runs and apply all
    the styles i'm interested back. It would be interesting to have a way
    to just change the text of an existing text, leaving the formatting
    untouched."

In other words: the caller wants to rewrite the *text* of an existing
text frame (typically a template placeholder like ``{NAME}``) without
losing the bold / italic / font-size / color that was authored on the
run holding the token.

Wave 6 #836 shipped :meth:`TextFrame.replace_text` and
:meth:`_Paragraph.replace_text` which do exactly this:

  * The run where the match *starts* absorbs the replacement text,
    keeping **all** of its run-level formatting (``a:rPr`` — bold,
    italic, font size, color, highlight, underline, font family, …).
  * Runs fully contained within the match are dropped.
  * If a match ends partway through a run, that run's surviving suffix
    keeps its own formatting.
  * Matches do not cross paragraph / ``a:br`` / ``a:fld`` boundaries.

This regression test pins the end-to-end behaviour #285 asked for: a
caller can take a deck, rewrite a templated token sitting in a
hand-formatted run, round-trip through ``Presentation.save`` + reopen,
and every last formatting attribute on the replacement text matches
what the author originally set — including the cross-run case
PowerPoint creates whenever it re-flows a token after an edit.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


class DescribeIssue285ReplaceTextPreservesFormatting:
    """End-to-end regression for "change text without editing formatting"."""

    def it_preserves_run_formatting_when_replacing_a_token_inside_one_run(self):
        """The simple #285 case: the whole token lives in a single run.

        Author a textbox whose single run carries bold + italic + 24-pt
        font size + a named font family + an explicit RGB color, then
        swap the token out via ``replace_text``. After round-trip every
        attribute on the rewritten run must match the original.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(5), Inches(1)
        )
        tf = textbox.text_frame

        # -- author a single run with rich formatting holding a token --
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = "Hello {NAME}!"
        run.font.bold = True
        run.font.italic = True
        run.font.size = Pt(24)
        run.font.name = "Courier New"
        run.font.color.rgb = RGBColor(0x11, 0x22, 0x33)

        # -- act: rewrite the token --
        count = tf.replace_text("{NAME}", "Alice")

        assert count == 1
        assert tf.text == "Hello Alice!"

        # -- formatting is still there (pre-save) --
        rewritten = tf.paragraphs[0].runs[0]
        assert rewritten.text == "Hello Alice!"
        assert rewritten.font.bold is True
        assert rewritten.font.italic is True
        assert rewritten.font.size == Pt(24)
        assert rewritten.font.name == "Courier New"
        assert rewritten.font.color.rgb == RGBColor(0x11, 0x22, 0x33)

        # -- round-trip through save + reopen --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        reloaded = prs2.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]

        assert reloaded.text == "Hello Alice!"
        assert reloaded.font.bold is True
        assert reloaded.font.italic is True
        assert reloaded.font.size == Pt(24)
        assert reloaded.font.name == "Courier New"
        assert reloaded.font.color.rgb == RGBColor(0x11, 0x22, 0x33)

    def it_preserves_origin_run_formatting_when_token_spans_two_runs(self):
        """The PowerPoint-splits-a-token case #285 really cares about.

        PowerPoint frequently splits a token like ``{NAME}`` across two
        ``a:r`` elements after in-editor edits, each run potentially
        carrying different formatting. ``replace_text`` must keep the
        formatting of the run where the match **starts** for the
        replacement, per the #836 spec.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(5), Inches(1)
        )
        tf = textbox.text_frame
        p = tf.paragraphs[0]

        # -- run #1 (origin): bold + 32 pt + red --
        r1 = p.add_run()
        r1.text = "{NA"
        r1.font.bold = True
        r1.font.size = Pt(32)
        r1.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

        # -- run #2 (fully inside the match): deliberately different
        # -- formatting so we can tell which run's rPr survives --
        r2 = p.add_run()
        r2.text = "ME}"
        r2.font.italic = True
        r2.font.size = Pt(12)
        r2.font.color.rgb = RGBColor(0x00, 0xFF, 0x00)

        # -- act: replace the cross-run token --
        count = tf.replace_text("{NAME}", "Bob")

        assert count == 1
        assert tf.text == "Bob"

        # -- exactly one run remains; it carries the *origin* run's rPr --
        runs = tf.paragraphs[0].runs
        assert len(runs) == 1
        surviving = runs[0]
        assert surviving.text == "Bob"
        assert surviving.font.bold is True
        assert surviving.font.italic is None  # -- r2's italic was dropped
        assert surviving.font.size == Pt(32)
        assert surviving.font.color.rgb == RGBColor(0xFF, 0x00, 0x00)

        # -- round-trip and re-assert --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        reloaded_runs = prs2.slides[0].shapes[0].text_frame.paragraphs[0].runs
        assert len(reloaded_runs) == 1
        reloaded = reloaded_runs[0]
        assert reloaded.text == "Bob"
        assert reloaded.font.bold is True
        assert reloaded.font.italic is None
        assert reloaded.font.size == Pt(32)
        assert reloaded.font.color.rgb == RGBColor(0xFF, 0x00, 0x00)

    def it_preserves_trailing_run_formatting_when_match_ends_mid_run(self):
        """A match that ends partway through a run: the surviving suffix
        keeps that run's own formatting, independent of the origin run.

        This is the third #285 sub-case — the most common after an edit
        that inserted plain text after a formatted token (e.g., an
        italic ``{NAME}`` followed by plain `"'s report"`).
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(5), Inches(1)
        )
        tf = textbox.text_frame
        p = tf.paragraphs[0]

        # -- origin run: bold --
        r1 = p.add_run()
        r1.text = "{NA"
        r1.font.bold = True

        # -- trailing run: italic + underline; match ends partway
        # -- through it, so its surviving suffix (" tail") keeps these --
        r2 = p.add_run()
        r2.text = "ME} tail"
        r2.font.italic = True
        r2.font.underline = True

        count = tf.replace_text("{NAME}", "X")
        assert count == 1
        assert tf.text == "X tail"

        runs = tf.paragraphs[0].runs
        assert len(runs) == 2

        # -- replacement run inherits the origin rPr --
        assert runs[0].text == "X"
        assert runs[0].font.bold is True
        assert runs[0].font.italic is None
        assert runs[0].font.underline is None

        # -- surviving suffix keeps the *trailing* run's rPr --
        assert runs[1].text == " tail"
        assert runs[1].font.bold is None
        assert runs[1].font.italic is True
        assert runs[1].font.underline is True

        # -- round-trip assertions --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        reloaded = prs2.slides[0].shapes[0].text_frame.paragraphs[0].runs
        assert len(reloaded) == 2
        assert reloaded[0].text == "X"
        assert reloaded[0].font.bold is True
        assert reloaded[1].text == " tail"
        assert reloaded[1].font.italic is True
        assert reloaded[1].font.underline is True
