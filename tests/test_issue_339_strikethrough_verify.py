# pyright: reportPrivateUsage=false

"""Regression test for issue #339 — Font.strikethrough support.

Issue #339 (https://github.com/scanny/python-pptx/issues/339) asked for a
public way to read and write the ``a:rPr/@strike`` attribute on a run — the
OOXML element PowerPoint uses to render single-line (``sngStrike``) or
double-line (``dblStrike``) strikethrough on text. Before the fix
``Font.underline`` existed but its sibling ``Font.strikethrough`` was
missing, so a caller producing a styled run had no supported way to author
strikethrough from python-pptx.

The feature is shipped (see ``FEATURES.md``):

  * :attr:`pptx.text.text.Font.strikethrough` — read/write tri-state
    (``True`` / ``False`` / ``None``) that also accepts a
    :class:`pptx.enum.text.MSO_TEXT_STRIKE_TYPE` member for the full
    single-line / double-line / no-strike vocabulary.
  * :class:`pptx.enum.text.MSO_TEXT_STRIKE_TYPE` (alias ``MSO_STRIKE``)
    with members ``NONE`` (``noStrike``), ``SINGLE_LINE`` (``sngStrike``),
    and ``DOUBLE_LINE`` (``dblStrike``).

This regression suite pins the reporter's workflow end-to-end:

  * The untouched default is |None| (inherit from the style hierarchy).
  * ``font.strikethrough = True`` writes ``@strike="sngStrike"`` and reads
    back as |True|.
  * ``font.strikethrough = False`` writes ``@strike="noStrike"`` (an
    explicit opt-out, *not* the same as removing the attribute) and reads
    back as |False|.
  * Each :class:`MSO_TEXT_STRIKE_TYPE` value round-trips through the
    setter/getter, including the ``DOUBLE_LINE`` case that the bool-only
    shortcut can't express.
  * ``font.strikethrough = None`` drops the ``@strike`` attribute so the
    style-hierarchy default reasserts itself.
  * The value survives :meth:`Presentation.save` + reopen.
  * ``TextFrame.replace_text`` preserves strikethrough formatting on the
    origin run — a mailmerge-style edit does not silently drop the flag.

Any future regression that drops the setter, swaps the attribute spelling,
or forgets to preserve formatting across ``replace_text`` will surface
here.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.enum.text import MSO_TEXT_STRIKE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of ``PartFactory.part_type_for``.

    Earlier test modules in the suite patch
    :class:`pptx.opc.package.PartFactory` with mocks and do not always
    restore the production mapping; ``Presentation(stream)`` reopens
    dispatch through ``part_type_for`` and need the concrete part classes
    to be registered.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


def _fresh_run(text: str = "Hello"):
    """Return a ``(prs, run)`` pair with one textbox run carrying `text`."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    run = textbox.text_frame.paragraphs[0].add_run()
    run.text = text
    return prs, run


class DescribeIssue339Strikethrough:
    """Pins the ``Font.strikethrough`` authoring path resolved by the fork."""

    def it_reads_None_on_an_untouched_run(self):
        """A brand-new run inherits its strike setting — getter returns |None|.

        This is the authoring default: no ``@strike`` attribute is written
        until the caller assigns a value, so the run participates in the
        normal style-hierarchy cascade (paragraph → body → master → theme).
        """
        _prs, run = _fresh_run()

        assert run.font.strikethrough is None
        rPr = run._r.get_or_add_rPr()
        assert rPr.get("strike") is None

    def it_writes_single_line_when_assigned_True(self):
        """``font.strikethrough = True`` → ``a:rPr/@strike="sngStrike"``.

        The bool-shortcut mirrors :attr:`Font.underline`'s True-means-single
        contract and is the most common spelling callers reach for.
        """
        _prs, run = _fresh_run()

        run.font.strikethrough = True

        assert run.font.strikethrough is True
        rPr = run._r.get_or_add_rPr()
        assert rPr.get("strike") == "sngStrike"

    def it_writes_no_strike_when_assigned_False(self):
        """``font.strikethrough = False`` → explicit ``@strike="noStrike"``.

        |False| is an explicit opt-out, *not* the same as clearing the
        attribute: it writes ``noStrike`` so the run overrides an
        inherited strike. Use |None| to revert to inheritance.
        """
        _prs, run = _fresh_run()

        run.font.strikethrough = False

        assert run.font.strikethrough is False
        rPr = run._r.get_or_add_rPr()
        assert rPr.get("strike") == "noStrike"

    @pytest.mark.parametrize(
        ("value", "xml_value", "getter_value"),
        [
            (MSO_TEXT_STRIKE_TYPE.NONE, "noStrike", False),
            (MSO_TEXT_STRIKE_TYPE.SINGLE_LINE, "sngStrike", True),
            (MSO_TEXT_STRIKE_TYPE.DOUBLE_LINE, "dblStrike", MSO_TEXT_STRIKE_TYPE.DOUBLE_LINE),
        ],
    )
    def it_accepts_every_MSO_TEXT_STRIKE_TYPE_member(self, value, xml_value, getter_value):
        """Each enum member round-trips through the setter/getter.

        ``DOUBLE_LINE`` is the case the bool-only shortcut can't express —
        callers who want double-strike must go through the enum. NONE and
        SINGLE_LINE normalise to |False| / |True| on read (matching
        :attr:`Font.underline`'s convention), while DOUBLE_LINE surfaces
        as the enum member directly.
        """
        _prs, run = _fresh_run()

        run.font.strikethrough = value

        assert run.font.strikethrough == getter_value
        rPr = run._r.get_or_add_rPr()
        assert rPr.get("strike") == xml_value

    def it_clears_the_strike_attribute_when_assigned_None(self):
        """Assigning |None| drops ``@strike`` so inheritance reasserts.

        Complements the explicit-False case: |None| removes the attribute
        entirely rather than writing ``noStrike``, restoring the run's
        participation in the style cascade.
        """
        _prs, run = _fresh_run()
        run.font.strikethrough = MSO_TEXT_STRIKE_TYPE.DOUBLE_LINE
        rPr = run._r.get_or_add_rPr()
        assert rPr.get("strike") == "dblStrike"

        run.font.strikethrough = None

        assert run.font.strikethrough is None
        assert rPr.get("strike") is None

    def it_round_trips_strikethrough_through_save_and_reload(self, _restore_part_factory):
        """End-to-end: every strike flavour survives save + reopen.

        Writes one run per flavour — ``True`` (single), ``False`` (explicit
        none), ``MSO_TEXT_STRIKE_TYPE.DOUBLE_LINE`` (double), and an
        untouched run (inherit) — saves to a ``BytesIO`` stream, reopens,
        and asserts each run reports the expected value. Any serializer /
        deserializer regression that renames the attribute, drops it, or
        loses the enum mapping will fail here.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(3))
        tf = textbox.text_frame

        run_single = tf.paragraphs[0].add_run()
        run_single.text = "single"
        run_single.font.size = Pt(18)
        run_single.font.strikethrough = True

        p2 = tf.add_paragraph()
        run_none = p2.add_run()
        run_none.text = "none"
        run_none.font.strikethrough = False

        p3 = tf.add_paragraph()
        run_double = p3.add_run()
        run_double.text = "double"
        run_double.font.strikethrough = MSO_TEXT_STRIKE_TYPE.DOUBLE_LINE

        p4 = tf.add_paragraph()
        run_inherit = p4.add_run()
        run_inherit.text = "inherit"

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        reloaded_tf = prs2.slides[0].shapes[0].text_frame
        assert reloaded_tf.paragraphs[0].runs[0].text == "single"
        assert reloaded_tf.paragraphs[0].runs[0].font.strikethrough is True
        assert reloaded_tf.paragraphs[1].runs[0].text == "none"
        assert reloaded_tf.paragraphs[1].runs[0].font.strikethrough is False
        assert reloaded_tf.paragraphs[2].runs[0].text == "double"
        assert (
            reloaded_tf.paragraphs[2].runs[0].font.strikethrough == MSO_TEXT_STRIKE_TYPE.DOUBLE_LINE
        )
        assert reloaded_tf.paragraphs[3].runs[0].text == "inherit"
        assert reloaded_tf.paragraphs[3].runs[0].font.strikethrough is None

    def it_preserves_strikethrough_across_replace_text(self):
        """``TextFrame.replace_text`` keeps the origin run's strike flag.

        The mailmerge-style edit path (:meth:`TextFrame.replace_text`,
        issue #836) rewrites the text of the match's origin run in place;
        the surrounding formatting — including ``@strike`` — must survive
        untouched. A regression that rebuilt the run from scratch would
        silently drop strikethrough and reintroduce the original #339
        symptom.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        tf = textbox.text_frame
        run = tf.paragraphs[0].add_run()
        run.text = "Hello {NAME}"
        run.font.strikethrough = MSO_TEXT_STRIKE_TYPE.DOUBLE_LINE

        count = tf.replace_text("{NAME}", "World")

        assert count == 1
        replaced = tf.paragraphs[0].runs[0]
        assert replaced.text == "Hello World"
        assert replaced.font.strikethrough == MSO_TEXT_STRIKE_TYPE.DOUBLE_LINE
        rPr = replaced._r.get_or_add_rPr()
        assert rPr.get("strike") == "dblStrike"

    def it_writes_the_strike_attribute_directly_on_a_rPr(self):
        """Pin the XML shape: ``@strike`` lives on ``a:rPr`` itself.

        PowerPoint reads the strike setting from the run's own
        ``a:rPr/@strike`` attribute; the value is an attribute, not a
        child element. This assertion catches a regression that ever
        tried to write ``<a:strike .../>`` or similar.
        """
        _prs, run = _fresh_run()

        run.font.strikethrough = MSO_TEXT_STRIKE_TYPE.SINGLE_LINE

        rPr = run._r.get_or_add_rPr()
        assert rPr.tag == qn("a:rPr")
        assert rPr.get("strike") == "sngStrike"
        # -- belt-and-braces: no rogue `a:strike` child element --
        assert rPr.find(qn("a:strike")) is None
