# pyright: reportPrivateUsage=false

"""Regression test for issue #343 — support Chinese fonts.

Issue #343 (https://github.com/scanny/python-pptx/issues/343) asks for a way
to author PowerPoint runs whose Chinese (Simplified / Traditional) characters
render in a caller-chosen Chinese typeface such as SimSun (宋体) or
Microsoft YaHei (微软雅黑). Before the Wave 1 work this could not be done
reliably through the public API: ``Font.name`` only writes the ``a:latin``
font-slot on ``a:rPr``, and PowerPoint's per-script dispatcher routes CJK
code points to the ``a:ea`` (East-Asian) slot instead, so a value assigned
to ``Font.name`` is silently ignored for Chinese glyphs. The reporter's
deck therefore always fell back to the theme EA font (typically Microsoft's
default, not the caller's choice), regardless of what ``Font.name`` was
set to.

The underlying gap is identical to the one reported in issue #720 (which
specifically compared WPS vs. PowerPoint behaviour for a SimSun run), and
both issues share the same fix. Wave 1 #337
(``feat/issue-337-font-ea-cs-slots``) added the two missing accessors —
:attr:`.Font.name_ea` (East-Asian) and :attr:`.Font.name_cs`
(complex-script) — alongside the existing :attr:`.Font.name` (Latin),
exposing the full three-slot font-hint contract that PowerPoint actually
reads. Wave 9 #720's verify suite then pinned that contract end-to-end.
Between them, #337 and #720 fully resolve #343: a caller who wants Chinese
text rendered in SimSun or 微软雅黑 can now write ::

    run.font.name_ea = "SimSun"       # or "微软雅黑"

and the value survives through the XML, through
:meth:`.Presentation.save`, and through reopen — the concrete ask in #343.

This regression suite is a thin, Chinese-typeface-focused pin on that
resolution. It overlaps intentionally with the generic three-slot
contract from ``test_issue_720_ea_font_verify.py`` but fixes the
Chinese-specific angles #343 explicitly raises:

  * Chinese typeface names written in **Latin transliteration**
    (``"SimSun"``) are accepted on :attr:`.Font.name_ea` and round-trip
    unchanged.
  * Chinese typeface names written in **native Chinese characters**
    (``"微软雅黑"`` — Microsoft YaHei) are accepted on
    :attr:`.Font.name_ea` and round-trip unchanged, including through
    ``Presentation.save`` + reopen where a naïve string-handling bug
    would surface as mojibake.
  * The emitted ``a:ea/@typeface`` attribute carries the exact Chinese
    characters the caller supplied (verified against the live XML), so
    PowerPoint's font-name lookup sees the same bytes the caller meant.
  * Mixed Latin + Chinese runs use the three-slot pattern (``a:latin``
    for the Latin prefix, ``a:ea`` for the Chinese suffix) and both
    slots land in the saved ``.pptx``.

Any future regression that truncates, re-encodes, or drops the Chinese
``@typeface`` — for example a UTF-8 vs. UTF-16 mismatch in the
serializer, or a refactor that routes ``name_ea`` back through
``a:latin`` — will reproduce the #343 symptom and be caught here.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# -- Chinese typefaces the #343 reporter cares about --
SIMSUN = "SimSun"  # 宋体 — Simplified Chinese serif, bundled with Windows
MS_YAHEI = "微软雅黑"  # Microsoft YaHei — Windows default CJK UI face


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of ``PartFactory.part_type_for``.

    Mirrors the guard used by
    ``tests/test_issue_720_ea_font_verify.py`` and friends —
    ``Presentation(stream)`` reopen dispatches through
    ``PartFactory.part_type_for``, which earlier test modules overwrite
    with mocks and do not restore.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue343ChineseFontsVerify:
    """Pins the Chinese-typeface authoring path resolved by #337 + #720."""

    def it_accepts_a_latin_transliterated_chinese_typeface_on_name_ea(self):
        """The simplest #343 ask: ``run.font.name_ea = "SimSun"``.

        SimSun (宋体) is the canonical example from the #343 thread — a
        Chinese typeface the user expects their CJK text to render in.
        With #337 in place the assignment writes ``a:ea/@typeface`` and
        the getter reads it back, so the deck author can explicitly
        target a Chinese face rather than accepting whatever EA face
        the theme happens to supply.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = textbox.text_frame.paragraphs[0].add_run()
        run.text = "你好世界"  # -- "Hello world" in Simplified Chinese --

        run.font.name_ea = SIMSUN

        assert run.font.name_ea == SIMSUN

    def it_accepts_a_native_chinese_typeface_name_on_name_ea(self):
        """A Chinese typeface name written in Chinese characters.

        Microsoft YaHei ships under its native name 微软雅黑 on Chinese
        Windows installations. A #343 caller who types the Chinese name
        directly (rather than the pinyin / English alias) expects the
        assignment to survive the round-trip intact — a naïve
        ASCII-only code path would mangle the value.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = textbox.text_frame.paragraphs[0].add_run()
        run.text = "微软雅黑示例"

        run.font.name_ea = MS_YAHEI

        assert run.font.name_ea == MS_YAHEI
        assert run.font.name_ea == "微软雅黑"  # -- belt-and-braces: exact characters --

    def it_writes_the_chinese_typeface_into_a_ea_typeface_attribute(self):
        """Pin the XML shape: ``a:ea/@typeface`` carries the exact string.

        PowerPoint looks up the face by the ``@typeface`` attribute on
        ``a:ea``. If the serializer ever started HTML-escaping,
        ASCII-folding, or base-encoding that attribute the Chinese
        characters would no longer match the installed font name and
        PowerPoint would silently fall back to the theme face — the
        original #343 symptom.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = textbox.text_frame.paragraphs[0].add_run()
        run.text = "微软雅黑"
        run.font.name_ea = MS_YAHEI

        rPr = run._r.get_or_add_rPr()
        ea_children = rPr.findall(qn("a:ea"))

        assert len(ea_children) == 1
        assert ea_children[0].get("typeface") == MS_YAHEI
        assert ea_children[0].get("typeface") == "微软雅黑"

    def it_pairs_latin_and_chinese_faces_for_a_mixed_script_run(self):
        """Mixed Latin + Chinese runs use the three-slot pattern.

        The #343 reporter's real decks mix English and Chinese in a
        single text frame. The supported authoring is: set
        ``font.name`` for the Latin face and ``font.name_ea`` for the
        Chinese face — PowerPoint then dispatches each character to
        the matching slot.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = textbox.text_frame.paragraphs[0].add_run()
        run.text = "Hello 世界"

        run.font.name = "Arial"
        run.font.name_ea = SIMSUN

        assert run.font.name == "Arial"
        assert run.font.name_ea == SIMSUN
        # -- the mixed-script author doesn't touch the CS slot --
        assert run.font.name_cs is None

    def it_round_trips_a_chinese_typeface_through_save_and_reload(self, _restore_part_factory):
        """End-to-end: a SimSun + 微软雅黑 deck survives save + reopen.

        This is the acid test for #343. Both a Latin-transliterated
        Chinese name (``SimSun`` on one run) and a native-Chinese
        typeface name (``微软雅黑`` on another) must come back from
        :func:`pptx.Presentation` unchanged, with the Chinese bytes
        intact.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
        tf = textbox.text_frame

        run_a = tf.paragraphs[0].add_run()
        run_a.text = "你好"
        run_a.font.name = "Arial"
        run_a.font.name_ea = SIMSUN
        run_a.font.size = Pt(24)

        p2 = tf.add_paragraph()
        run_b = p2.add_run()
        run_b.text = "微软雅黑"
        run_b.font.name_ea = MS_YAHEI

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        reloaded_tf = prs2.slides[0].shapes[0].text_frame
        reloaded_a = reloaded_tf.paragraphs[0].runs[0]
        reloaded_b = reloaded_tf.paragraphs[1].runs[0]

        assert reloaded_a.text == "你好"
        assert reloaded_a.font.name == "Arial"
        assert reloaded_a.font.name_ea == SIMSUN
        assert reloaded_a.font.size == Pt(24)

        assert reloaded_b.text == "微软雅黑"
        assert reloaded_b.font.name_ea == MS_YAHEI
        assert reloaded_b.font.name_ea == "微软雅黑"

    def it_lets_a_chinese_typeface_be_cleared_with_None(self):
        """Assigning ``None`` to ``name_ea`` drops the override cleanly.

        A #343 caller who experimentally sets a Chinese face and then
        wants to revert to the theme default uses ``name_ea = None``.
        The ``a:ea`` element is removed; other slots (and their
        typefaces) are untouched.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = textbox.text_frame.paragraphs[0].add_run()
        run.text = "Hello 世界"
        run.font.name = "Arial"
        run.font.name_ea = SIMSUN

        run.font.name_ea = None

        assert run.font.name == "Arial"
        assert run.font.name_ea is None
        rPr = run._r.get_or_add_rPr()
        assert rPr.find(qn("a:ea")) is None
        assert rPr.find(qn("a:latin")) is not None
