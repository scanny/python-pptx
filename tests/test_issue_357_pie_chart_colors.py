# pyright: reportPrivateUsage=false

"""Regression test for issue #357 — customize colors of pie chart slices.

Issue #357 (https://github.com/scanny/python-pptx/issues/357) asks for
a way to customize the fill color of individual pie-chart slices — i.e.
per-slice (per-data-point) color control on a ``c:pieChart/c:ser``.

Coverage map:

* Per-slice **fill color** via
  ``Chart.plots[0].series[0].points[n].format.fill``. The ``Point.format``
  property was added in commit ``de58d605`` ("cht: add Point.format") and
  ``CategoryPoints`` has covered pie series since ``PieSeries`` began
  inheriting from ``_BaseCategorySeries`` — together those shipped the
  exact API #357 asks for. This regression test pins the end-to-end
  round trip through save / reload so a future refactor of the
  ``c:ser / c:dPt / c:spPr / a:solidFill / a:srgbClr`` subtree keeps
  per-slice color control working for pie charts.

* Per-slice **line color** via ``points[n].format.line`` — the same
  ``ChartFormat`` wrapper exposes line color, so a caller who wants
  a colored border on a slice gets it in the same call pattern.

* ``c:dPt`` **element shape** — the per-slice element must have
  ``c:idx`` matching the slice index and a ``c:spPr`` child carrying
  the solid-fill sub-tree, so PowerPoint attributes the color to the
  right slice.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.chart.series import PieSeries
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

SLICE_COLORS = (
    RGBColor(0xFF, 0x00, 0x00),
    RGBColor(0x00, 0xFF, 0x00),
    RGBColor(0x00, 0x00, 0xFF),
    RGBColor(0xFF, 0xFF, 0x00),
)


def _build_pie_chart():
    """Build a 4-slice pie chart, return (prs, chart_shape)."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_data = CategoryChartData()
    chart_data.categories = ["A", "B", "C", "D"]
    chart_data.add_series("s1", (10.0, 20.0, 30.0, 40.0))
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.PIE,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    )
    return prs, chart_shape


def _set_per_slice_colors(series):
    for i, color in enumerate(SLICE_COLORS):
        fill = series.points[i].format.fill
        fill.solid()
        fill.fore_color.rgb = color


def _reloaded_chart(prs):
    for shp in prs.slides[0].shapes:
        if shp.has_chart:
            return shp.chart
    raise AssertionError("reloaded presentation contains no chart")


def _roundtrip(prs):
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)


class DescribeIssue357PieChartColors(object):
    """Per-slice color control on a pie chart, end-to-end."""

    def it_accesses_points_on_a_pie_series(self):
        prs, chart_shape = _build_pie_chart()
        series = chart_shape.chart.plots[0].series[0]

        assert isinstance(series, PieSeries)
        assert len(series.points) == 4

    def it_sets_fill_on_each_slice(self):
        """The core #357 request: paint each slice a different color."""
        prs, chart_shape = _build_pie_chart()
        series = chart_shape.chart.plots[0].series[0]

        _set_per_slice_colors(series)

        for i, expected in enumerate(SLICE_COLORS):
            assert series.points[i].format.fill.fore_color.rgb == expected

    def it_round_trips_per_slice_fill_through_save_and_reload(self):
        prs, chart_shape = _build_pie_chart()
        series = chart_shape.chart.plots[0].series[0]
        _set_per_slice_colors(series)

        reloaded = _reloaded_chart(_roundtrip(prs))
        reloaded_series = reloaded.plots[0].series[0]

        assert isinstance(reloaded_series, PieSeries)
        for i, expected in enumerate(SLICE_COLORS):
            assert reloaded_series.points[i].format.fill.fore_color.rgb == expected

    def it_emits_one_dPt_per_slice_with_matching_idx_and_spPr(self):
        """Per-slice color must land on a ``c:dPt[c:idx/@val=i]/c:spPr``."""
        prs, chart_shape = _build_pie_chart()
        series = chart_shape.chart.plots[0].series[0]

        _set_per_slice_colors(series)

        dPts = series._ser.findall(qn("c:dPt"))
        assert len(dPts) == 4
        for i, dPt in enumerate(dPts):
            idx = dPt.find(qn("c:idx"))
            assert idx is not None and int(idx.get("val")) == i
            spPr = dPt.find(qn("c:spPr"))
            assert spPr is not None
            srgbClr = spPr.find(qn("a:solidFill") + "/" + qn("a:srgbClr"))
            assert srgbClr is not None
            assert srgbClr.get("val") == str(SLICE_COLORS[i])

    def it_also_supports_per_slice_line_color(self):
        """``points[i].format.line`` is the same ChartFormat wrapper, pins border."""
        prs, chart_shape = _build_pie_chart()
        series = chart_shape.chart.plots[0].series[0]

        series.points[0].format.line.color.rgb = RGBColor(0x12, 0x34, 0x56)

        reloaded = _reloaded_chart(_roundtrip(prs))
        reloaded_point = reloaded.plots[0].series[0].points[0]
        assert reloaded_point.format.line.color.rgb == RGBColor(0x12, 0x34, 0x56)
