# pyright: reportPrivateUsage=false

"""Regression test for issue #410 — embedded 3D-model passthrough.

Issue #410 (https://github.com/scanny/python-pptx/issues/410) asks for
read-only support of PowerPoint 365's "Insert > 3D Models" feature.
Authoring 3D models requires camera / lighting / scene-graph machinery
that is out of scope for this wave; what *is* shipped is detection and
round-trip preservation of an already-authored 3D model.

The fixture below constructs a minimal ``.pptx`` that contains an
``am3d:model3D`` graphic-frame pointing at a stub ``.glb`` part. Real
PowerPoint 3D-model files are tens of kilobytes of binary glTF; for the
round-trip test a short opaque byte sequence is sufficient since the
library does not interpret the payload. The test then:

1. Loads the fixture with :class:`~pptx.Presentation`.
2. Verifies :attr:`GraphicFrame.has_model_3d` is True.
3. Verifies :attr:`GraphicFrame.model_3d_xml` returns a well-formed XML
   string that still names the ``am3d:model3D`` tag.
4. Verifies :attr:`GraphicFrame.model_3d.embedded_rel_id` / ``.ext`` /
   ``.media_blob`` all resolve.
5. Saves to a ``BytesIO``, reopens it, re-runs the same assertions, and
   asserts the embedded ``.glb`` bytes survive verbatim.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from pptx import Presentation
from pptx.shapes.graphfrm import GraphicFrame
from pptx.util import Emu

# -- opaque byte payload standing in for a binary glTF (.glb) blob. Real
# -- PowerPoint-authored models run 10-200 KB; the library's passthrough
# -- does not care about the content, only that the bytes survive.
_GLB_STUB = b"glTF-stub\x00\x01\x02\x03binary-payload"

# -- content type strings (Microsoft extension — not in ECMA-376) --
_MODEL_3D_URI = "http://schemas.microsoft.com/office/drawing/2016/12/model3D"
_MODEL_3D_RELTYPE = "http://schemas.microsoft.com/office/2017/06/relationships/model3D"
_MODEL_3D_CT = "model/gltf-binary"

# -- XML fragment injected into the slide to declare an `am3d:model3D` graphic --
# -- frame. `rId1` is always the slide-layout rel; we append `rId2` for model3D --
_SLIDE_XML_TEMPLATE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    ' xmlns:am3d="http://schemas.microsoft.com/office/drawing/2017/model3d"'
    ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    "<p:cSld><p:spTree>"
    '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
    "<p:grpSpPr/>"
    "<p:graphicFrame>"
    "<p:nvGraphicFramePr>"
    '<p:cNvPr id="2" name="3D Model 1"/>'
    '<p:cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr>'
    "<p:nvPr/>"
    "</p:nvGraphicFramePr>"
    '<p:xfrm><a:off x="914400" y="914400"/><a:ext cx="3657600" cy="3657600"/></p:xfrm>'
    "<a:graphic>"
    f'<a:graphicData uri="{_MODEL_3D_URI}">'
    '<am3d:model3D r:embed="rId2" ext="glb">'
    "<am3d:camera/>"
    "</am3d:model3D>"
    "</a:graphicData>"
    "</a:graphic>"
    "</p:graphicFrame>"
    "</p:spTree></p:cSld>"
    "<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>"
    "</p:sld>"
)

# -- slide1.xml.rels replacement adding the `rId2` model3D rel --
_SLIDE_RELS_TEMPLATE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1"'
    ' Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"'
    ' Target="../slideLayouts/slideLayout7.xml"/>'
    f'<Relationship Id="rId2" Type="{_MODEL_3D_RELTYPE}" Target="../media/model1.glb"/>'
    "</Relationships>"
)


def _build_3d_model_pptx() -> io.BytesIO:
    """Construct an in-memory `.pptx` containing one slide with a 3D model.

    Starts from a fresh :class:`~pptx.Presentation`, adds a blank slide, writes it
    to a BytesIO, then unzips that zip and surgically rewrites `ppt/slides/slide1.xml`
    and its rels to add an ``am3d:model3D`` graphicFrame. Also adds the media
    (`.glb`) part and a content-type override.
    """
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    base = io.BytesIO()
    prs.save(base)
    base.seek(0)

    # -- rewrite into a new zip with the 3D-model bits added --
    out = io.BytesIO()
    with zipfile.ZipFile(base, "r") as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            name = item.filename
            data = src.read(name)
            if name == "ppt/slides/slide1.xml":
                data = _SLIDE_XML_TEMPLATE.encode("utf-8")
            elif name == "ppt/slides/_rels/slide1.xml.rels":
                data = _SLIDE_RELS_TEMPLATE.encode("utf-8")
            elif name == "[Content_Types].xml":
                # -- append override for the new .glb part before </Types> --
                extra = (
                    f'<Override PartName="/ppt/media/model1.glb"' f' ContentType="{_MODEL_3D_CT}"/>'
                ).encode("utf-8")
                data = data.replace(b"</Types>", extra + b"</Types>")
            dst.writestr(item, data)
        # -- add the stub .glb part itself --
        dst.writestr("ppt/media/model1.glb", _GLB_STUB)

    out.seek(0)
    return out


class DescribeIssue410Model3DPassthrough(object):
    """Exercises the read-only 3D-model detection and round-trip surface."""

    def it_detects_and_surfaces_an_embedded_3D_model(self):
        pkg = _build_3d_model_pptx()
        prs = Presentation(pkg)
        slide = prs.slides[0]

        graphic_frames = [s for s in slide.shapes if isinstance(s, GraphicFrame) and s.has_model_3d]
        assert len(graphic_frames) == 1, "expected exactly one 3D-model frame"

        frame = graphic_frames[0]
        assert frame.has_model_3d is True

        xml_str = frame.model_3d_xml
        assert xml_str is not None
        assert "model3D" in xml_str
        assert 'ext="glb"' in xml_str

        model = frame.model_3d
        assert model.embedded_rel_id == "rId2"
        assert model.ext == "glb"
        assert model.media_blob == _GLB_STUB

    def it_round_trips_the_3D_model_through_save_and_reopen(self):
        pkg = _build_3d_model_pptx()
        prs1 = Presentation(pkg)
        buf = io.BytesIO()
        prs1.save(buf)
        buf.seek(0)

        prs2 = Presentation(buf)
        slide = prs2.slides[0]

        frames = [s for s in slide.shapes if isinstance(s, GraphicFrame) and s.has_model_3d]
        assert len(frames) == 1
        frame = frames[0]
        assert frame.has_model_3d is True

        # -- the model3D XML survives, including its child elements --
        xml_str = frame.model_3d_xml
        assert xml_str is not None
        assert "model3D" in xml_str
        assert "camera" in xml_str  # -- nested child preserved verbatim

        # -- the embedded .glb bytes survive verbatim --
        assert frame.model_3d.media_blob == _GLB_STUB
        assert frame.model_3d.ext == "glb"

    def but_has_model_3d_is_False_on_a_non_model_graphic_frame(self):
        """A table-carrying graphic-frame reports `has_model_3d == False`."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_table(
            rows=2,
            cols=2,
            left=Emu(0),
            top=Emu(0),
            width=Emu(1_000_000),
            height=Emu(500_000),
        )

        frame = slide.shapes[0]
        assert isinstance(frame, GraphicFrame)
        assert frame.has_model_3d is False
        assert frame.model_3d_xml is None
        with pytest.raises(ValueError):
            frame.model_3d
