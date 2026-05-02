# pyright: reportPrivateUsage=false

"""Regression test for issue #427 — ``SlideShapes.add_movie(autoplay=True)``.

Issue #427 (https://github.com/scanny/python-pptx/issues/427) asks for a
one-line way to insert a movie (or audio clip) that begins playing
automatically when the slide is shown, rather than PowerPoint's default
click-to-play behavior.

Wave 5 (#811) added the underlying ``Movie.start_condition`` / ``Movie.start_time``
read/write properties that manipulate the ``p:cond`` node inside the
``p:video`` timing subtree. This issue is the user-facing convenience
layer on top of that: ``add_movie(..., autoplay=True)`` sets
``start_condition="withPrevious"`` on the newly created shape, which is
what PowerPoint emits when a user ticks the "Start: Automatically"
checkbox in the Video / Audio Tools ribbon.

This test end-to-end covers:

  1. ``add_movie(..., autoplay=True)`` produces a shape whose
     ``Movie.start_condition == "withPrevious"`` and round-trips the
     condition across save + reopen.
  2. ``add_movie(...)`` without the ``autoplay`` kwarg preserves the
     pre-existing "onClick" default — no regression for callers that
     rely on click-to-play.
  3. The kwarg works for audio clips too (mime_type="audio/...") since
     the underlying ``p:video`` timing node is identical.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.picture import Movie
from pptx.util import Inches

if TYPE_CHECKING:
    from pptx.presentation import Presentation as _PresentationT
    from pptx.slide import Slide


def _blank_slide() -> tuple[_PresentationT, Slide]:
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[6])  # blank layout


def _movie_shape(slide: Slide) -> Movie:
    media = next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.MEDIA)
    assert isinstance(media, Movie)
    return media


class DescribeIssue427MovieAutoplay:
    """End-to-end coverage for the ``autoplay`` kwarg on ``SlideShapes.add_movie``."""

    def it_sets_start_condition_withPrevious_when_autoplay_True(self):
        _prs, slide = _blank_slide()
        audio = io.BytesIO(b"\x52\x49\x46\x46" + b"\x00" * 36)  # fake wav
        slide.shapes.add_movie(
            audio,
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(2),
            mime_type="audio/wav",
            autoplay=True,
        )
        movie = _movie_shape(slide)
        assert movie.start_condition == "withPrevious"

    def it_leaves_start_condition_onClick_by_default(self):
        """Omitting `autoplay` preserves PowerPoint's default click-to-play."""
        _prs, slide = _blank_slide()
        audio = io.BytesIO(b"\x52\x49\x46\x46" + b"\x00" * 36)
        slide.shapes.add_movie(
            audio,
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(2),
            mime_type="audio/wav",
        )
        movie = _movie_shape(slide)
        assert movie.start_condition == "onClick"

    def it_respects_autoplay_False_explicitly(self):
        """Explicit `autoplay=False` is identical to omitting the kwarg."""
        _prs, slide = _blank_slide()
        audio = io.BytesIO(b"\x52\x49\x46\x46" + b"\x00" * 36)
        slide.shapes.add_movie(
            audio,
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(2),
            mime_type="audio/wav",
            autoplay=False,
        )
        movie = _movie_shape(slide)
        assert movie.start_condition == "onClick"

    def it_round_trips_autoplay_through_save_and_reload(self):
        """Start-condition survives a save + reopen cycle."""
        prs, slide = _blank_slide()
        audio = io.BytesIO(b"\x52\x49\x46\x46" + b"\x00" * 36)
        slide.shapes.add_movie(
            audio,
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(2),
            mime_type="audio/wav",
            autoplay=True,
        )

        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        prs2 = Presentation(buffer)
        movie = _movie_shape(prs2.slides[0])

        assert movie.start_condition == "withPrevious"

    def it_supports_autoplay_for_video_mime_types(self):
        """The kwarg applies equally to video clips (the default mime_type)."""
        _prs, slide = _blank_slide()
        video = io.BytesIO(b"\x00" * 128)
        slide.shapes.add_movie(
            video,
            left=Inches(1),
            top=Inches(1),
            width=Inches(4),
            height=Inches(3),
            autoplay=True,
        )
        movie = _movie_shape(slide)
        assert movie.start_condition == "withPrevious"

    def it_is_overridable_via_direct_property_assignment_afterwards(self):
        """`autoplay=True` is a shorthand — callers can still reassign later.

        This protects callers that mix ``autoplay`` with finer-grained
        control (e.g. switching to ``"afterPrevious"`` after the shape is
        created).
        """
        _prs, slide = _blank_slide()
        audio = io.BytesIO(b"\x52\x49\x46\x46" + b"\x00" * 36)
        slide.shapes.add_movie(
            audio,
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(2),
            mime_type="audio/wav",
            autoplay=True,
        )
        movie = _movie_shape(slide)
        assert movie.start_condition == "withPrevious"

        movie.start_condition = "afterPrevious"
        assert movie.start_condition == "afterPrevious"
