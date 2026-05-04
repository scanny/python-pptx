# pyright: reportPrivateUsage=false

"""Regression / verify test for issue #430 — ``SlideShapes.add_movie`` accepts MP3 audio.

Issue #430 (https://github.com/scanny/python-pptx/issues/430) reported that
passing an MP3 file to :meth:`SlideShapes.add_movie` raised. The method's
name is a bit misleading — PowerPoint's ``p:pic`` media element is used
for both video *and* audio clips, distinguished only by ``<a:videoFile>``
vs ``<a:audioFile>`` children.

Resolved in the fork-era audio-MIME work (see ``FEATURES.md`` under
"Movies and audio"): the ``mime_type`` kwarg is honored and any
``audio/*`` prefix causes ``add_movie`` to emit the ``<a:audioFile>``
child and register an ``audio/*`` content-type override, so the
resulting ``.pptx`` opens cleanly in PowerPoint as an audio clip.

This test pins the behavior end-to-end so a future refactor of the
movie / audio path cannot silently regress it:

  1. Passing an MP3-shaped payload by path with ``mime_type="audio/mpeg"``
     succeeds (no exception) and produces a ``p:pic`` with an
     ``<a:audioFile>`` child (not ``<a:videoFile>``).
  2. The same works when passing a file-like object instead of a path
     (the original reporter's call signature).
  3. The resulting ``.pptx`` round-trips through
     ``Presentation.save`` / ``Presentation(...)`` without corruption
     and the audio clip shape survives with its ``<a:audioFile>`` tag
     intact.
  4. Variant spellings used in the wild (``"audio/mp3"``, ``"audio/mpeg"``)
     both yield audio shapes (``<a:audioFile>``), confirming the
     ``audio/*`` prefix sniff — not an exact whitelist match — drives
     the audio-vs-video choice.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

import pytest

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.picture import Movie
from pptx.util import Inches

if TYPE_CHECKING:
    from pathlib import Path

    from pptx.presentation import Presentation as _PresentationT
    from pptx.slide import Slide


# -- bytes that look plausibly like an MP3 header; the library does not
# -- interrogate the file for its type, so any non-empty payload suffices.
_MP3_BYTES = b"ID3\x03\x00\x00\x00" + b"\x00" * 1024


def _blank_slide() -> tuple[_PresentationT, Slide]:
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[6])  # blank layout


def _media_shape(slide: Slide) -> Movie:
    media = next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.MEDIA)
    assert isinstance(media, Movie)
    return media


class DescribeIssue430AddMovieMP3:
    """End-to-end coverage for ``add_movie`` with MP3 / audio MIME types."""

    def it_accepts_an_mp3_path_with_audio_mpeg_mime_type(self, tmp_path: Path):
        mp3_path = tmp_path / "narration.mp3"
        mp3_path.write_bytes(_MP3_BYTES)
        _prs, slide = _blank_slide()

        slide.shapes.add_movie(
            str(mp3_path),
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(1),
            mime_type="audio/mpeg",
        )

        movie = _media_shape(slide)
        pic = movie._element
        audio_nodes = pic.xpath("./p:nvPicPr/p:nvPr/a:audioFile")
        video_nodes = pic.xpath("./p:nvPicPr/p:nvPr/a:videoFile")
        assert len(audio_nodes) == 1
        assert len(video_nodes) == 0

    def it_accepts_an_mp3_file_like_with_audio_mpeg_mime_type(self):
        """Reporter's original call signature: an open file, not a path."""
        _prs, slide = _blank_slide()

        slide.shapes.add_movie(
            io.BytesIO(_MP3_BYTES),
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(1),
            mime_type="audio/mpeg",
        )

        movie = _media_shape(slide)
        pic = movie._element
        assert pic.xpath("./p:nvPicPr/p:nvPr/a:audioFile")
        assert not pic.xpath("./p:nvPicPr/p:nvPr/a:videoFile")

    def it_round_trips_an_mp3_clip_through_save_and_reload(self, tmp_path: Path):
        prs, slide = _blank_slide()
        slide.shapes.add_movie(
            io.BytesIO(_MP3_BYTES),
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(1),
            mime_type="audio/mpeg",
        )

        out = tmp_path / "with-mp3.pptx"
        prs.save(str(out))

        prs2 = Presentation(str(out))
        movie = _media_shape(prs2.slides[0])
        pic = movie._element
        # -- audio tag survives round-trip; the shape is still recognized --
        assert pic.xpath("./p:nvPicPr/p:nvPr/a:audioFile")
        assert not pic.xpath("./p:nvPicPr/p:nvPr/a:videoFile")

    @pytest.mark.parametrize("mime", ["audio/mpeg", "audio/mp3", "audio/wav", "audio/x-wav"])
    def it_treats_any_audio_prefixed_mime_type_as_audio(self, mime: str):
        """The audio / video choice is driven by the ``audio/*`` prefix.

        Historically there has been no consensus on a single MP3 MIME
        spelling (``audio/mpeg`` vs ``audio/mp3``); likewise WAV
        (``audio/wav`` vs ``audio/x-wav``). All should produce audio
        shapes so callers don't have to guess PowerPoint's preferred
        spelling.
        """
        _prs, slide = _blank_slide()

        slide.shapes.add_movie(
            io.BytesIO(_MP3_BYTES),
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(1),
            mime_type=mime,
        )

        movie = _media_shape(slide)
        pic = movie._element
        assert pic.xpath(
            "./p:nvPicPr/p:nvPr/a:audioFile"
        ), f"mime_type={mime!r} should produce an <a:audioFile>, got a <a:videoFile>"
