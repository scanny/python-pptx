# pyright: reportPrivateUsage=false

"""Regression test for issue #440 — Chart data alignment (row vs column series).

Issue #440 (https://github.com/scanny/python-pptx/issues/440) asked for a
way to read whether a chart's source data is laid out with series in
**columns** (PowerPoint's default) or **rows** (the post-"Switch
Row/Column" layout). In OOXML there is no persisted ``@switchRowCol``
attribute — the orientation is implicit in the shape of the ``c:f``
cell-range references under each ``c:ser`` (``c:cat`` / ``c:tx`` /
``c:val``). Without a supported surface area, a caller inspecting the
XML has to replicate this reasoning themselves.

#440 is **essentially a duplicate of #828**. The fix shipped in Wave 11
as :attr:`.Chart.series_in_rows` (``feat(chart): #828 add
Chart.series_in_rows read-only accessor``, commit 575de804), which
parses the first series's category and series-name references and
reports:

  * ``False`` — categories run vertically (a single column, multiple
    rows), i.e. series-in-columns. This is the layout python-pptx
    writes when authoring from :class:`.CategoryChartData` and the
    layout PowerPoint defaults to.
  * ``True`` — categories run horizontally (a single row, multiple
    columns), i.e. series-in-rows. This is the state produced by
    *Chart Design > Switch Row/Column* in PowerPoint.
  * ``None`` — the orientation cannot be determined (no series,
    XY/scatter or bubble chart with no ``c:cat``, inline literals
    only). Callers should treat this as "unknown", not "default".

The property is read-only: PowerPoint's UI toggle rewrites every
sub-reference on every ``c:ser`` (names, categories, values) and
emulating that without a live workbook would silently diverge from
what PowerPoint re-authors on next save. Callers who need to flip
orientation programmatically should re-author via
:meth:`.Chart.replace_data` with the data transposed.

This suite pins the reporter's concrete asks from the #440 thread,
plus the fall-through and None-returning edge cases from the #828
implementation, plus a save/reopen round-trip to confirm the
orientation signal survives the ``.pptx`` zip. ``tests/chart/test_chart.py``
holds the fine-grained unit tests (under ``DescribeChart``) that pin
the XML-shape → value mapping on minimal hand-built ``c:chartSpace``
fragments; this module re-verifies #440 from the user's perspective
via the public API and the real chart-authoring path.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches

# -- fixtures -----------------------------------------------------------


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of ``PartFactory.part_type_for``.

    ``tests/opc/test_package.py::DescribePartFactory`` overwrites
    ``PartFactory.part_type_for[CT.PML_SLIDE]`` with a ``Mock`` and does not
    restore it, so a later ``Presentation(stream)`` reopen dispatches to the
    mock and fails. Mirrors the fixture used by other verify suites
    (e.g. ``tests/test_issue_470_combo_charts_verify.py``).
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


# -- helpers -----------------------------------------------------------


def _category_chart(chart_type=XL_CHART_TYPE.BAR_CLUSTERED):
    """Return ``(prs, chart)`` for a freshly authored multi-series bar chart.

    Uses the real ``add_chart`` path — the series, category, and value
    references are written by :class:`._BaseSeriesXmlWriter` the same way
    python-pptx writes every :class:`.CategoryChartData` chart (series-
    in-columns, i.e. default orientation).
    """
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_data = CategoryChartData()
    chart_data.categories = ["East", "West", "Mid"]
    chart_data.add_series("Q1", (1.0, 2.0, 3.0))
    chart_data.add_series("Q2", (4.0, 5.0, 6.0))
    graphic_frame = slide.shapes.add_chart(
        chart_type, Inches(1), Inches(1), Inches(6), Inches(4), chart_data
    )
    return prs, graphic_frame.chart


def _transpose_to_series_in_rows(chart):
    """Rewrite *chart*'s ``c:f`` refs in-place to the series-in-rows layout.

    PowerPoint rewrites every series's name / category / value reference
    when the user toggles *Chart Design > Switch Row/Column*. We mirror
    that layout:

      * series names move from row 1 of columns B..N to column A of rows
        2..N (one cell per series).
      * the category range becomes a single horizontal run across row 1
        of columns B..N (was vertical in column A rows 2..M).
      * each series's value range becomes a horizontal run in its own
        row (row = ``ser_idx + 2``, columns B..N).

    This matches the real XML PowerPoint emits and is what the #828
    implementation detects as ``series_in_rows is True``.
    """
    sers = chart._chartSpace.plotArea.findall(".//" + qn("c:ser"))
    # -- number of categories = N-columns across row 1 after the switch. --
    # -- For our helper, each series shares the same category span — we   --
    # -- pick it up from the original c:cat on the first series to keep   --
    # -- the test-helper decoupled from the caller's category count.      --
    first_cat = sers[0].find(qn("c:cat"))
    first_cat_f = first_cat.find(".//" + qn("c:f"))
    # -- original was Sheet1!$A$2:$A$N where N = len(categories)+1. --
    # -- parse the trailing row number to recover N-1. --
    # -- e.g. "Sheet1!$A$2:$A$4" -> 4 -> 3 categories. --
    tail = first_cat_f.text.rsplit("$", 1)[-1]
    n_categories = int(tail) - 1
    # -- column letters B.. for n_categories columns (A..Z; tests stay small) --
    last_col = chr(ord("B") + n_categories - 1)
    for ser_idx, ser in enumerate(sers):
        row = ser_idx + 2  # -- first series in row 2, next in row 3, ... --
        tx_f = ser.find(qn("c:tx") + "/" + qn("c:strRef") + "/" + qn("c:f"))
        if tx_f is not None:
            tx_f.text = "Sheet1!$A$%d" % row
        cat_f = ser.find(qn("c:cat") + "/" + qn("c:strRef") + "/" + qn("c:f"))
        if cat_f is not None:
            cat_f.text = "Sheet1!$B$1:$%s$1" % last_col
        val_f = ser.find(qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:f"))
        if val_f is not None:
            val_f.text = "Sheet1!$B$%d:$%s$%d" % (row, last_col, row)


def _xy_chart():
    """Return ``(prs, chart)`` for a fresh XY / scatter chart (no c:cat)."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_data = XyChartData()
    series = chart_data.add_series("XY")
    series.add_data_point(1.0, 2.0)
    series.add_data_point(2.0, 3.0)
    series.add_data_point(3.0, 4.0)
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER, Inches(1), Inches(1), Inches(6), Inches(4), chart_data
    )
    return prs, graphic_frame.chart


def _bubble_chart():
    """Return ``(prs, chart)`` for a fresh bubble chart (no c:cat)."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_data = BubbleChartData()
    series = chart_data.add_series("Bubble")
    series.add_data_point(1.0, 2.0, 10.0)
    series.add_data_point(2.0, 3.0, 20.0)
    series.add_data_point(3.0, 4.0, 30.0)
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.BUBBLE, Inches(1), Inches(1), Inches(6), Inches(4), chart_data
    )
    return prs, graphic_frame.chart


def _single_category_chart_with_tx_row_1():
    """Return a ``Chart`` whose single-cell c:cat falls through to a row-1 c:tx.

    The #828 implementation applies a dedicated fall-through when the
    first series's ``c:cat`` range resolves to a single cell — the one
    category doesn't establish an orientation on its own. The fall-through
    decides on the series-name reference: row 1 ⇒ default, elsewhere ⇒
    switched. This helper pins the row-1 branch of that fall-through by
    building the minimal ``c:chartSpace`` directly (authoring a real
    single-category ``CategoryChartData`` still emits a vertical
    ``c:cat`` range when multiple series are added, which defeats the
    fall-through).
    """
    from pptx.chart.chart import Chart

    chartSpace_xml = (
        '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml'
        '/2006/chart">'
        "<c:chart><c:plotArea><c:barChart><c:ser>"
        '<c:idx val="0"/><c:order val="0"/>'
        "<c:tx><c:strRef><c:f>Sheet1!$B$1</c:f></c:strRef></c:tx>"
        "<c:cat><c:strRef><c:f>Sheet1!$A$2</c:f></c:strRef></c:cat>"
        "<c:val><c:numRef><c:f>Sheet1!$B$2</c:f></c:numRef></c:val>"
        "</c:ser></c:barChart></c:plotArea></c:chart>"
        "</c:chartSpace>"
    )
    return Chart(parse_xml(chartSpace_xml), None)


def _single_category_chart_with_tx_row_2():
    """Companion of :func:`_single_category_chart_with_tx_row_1`.

    Pins the "elsewhere" branch of the single-cell c:cat fall-through —
    series name on row 2 ⇒ series_in_rows is True.
    """
    from pptx.chart.chart import Chart

    chartSpace_xml = (
        '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml'
        '/2006/chart">'
        "<c:chart><c:plotArea><c:barChart><c:ser>"
        '<c:idx val="0"/><c:order val="0"/>'
        "<c:tx><c:strRef><c:f>Sheet1!$A$2</c:f></c:strRef></c:tx>"
        "<c:cat><c:strRef><c:f>Sheet1!$B$1</c:f></c:strRef></c:cat>"
        "<c:val><c:numRef><c:f>Sheet1!$B$2</c:f></c:numRef></c:val>"
        "</c:ser></c:barChart></c:plotArea></c:chart>"
        "</c:chartSpace>"
    )
    return Chart(parse_xml(chartSpace_xml), None)


def _empty_chart():
    """Return a ``Chart`` with zero series — the "no ``c:ser``" fall-through."""
    from pptx.chart.chart import Chart

    chartSpace_xml = (
        '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml'
        '/2006/chart">'
        "<c:chart><c:plotArea><c:barChart/></c:plotArea></c:chart>"
        "</c:chartSpace>"
    )
    return Chart(parse_xml(chartSpace_xml), None)


# -- the verification suite ---------------------------------------------


class DescribeIssue440DataAlignment:
    """#440 chart data-alignment verify-and-close via #828.

    The reporter asked for a supported way to tell whether a chart's
    data is laid out "with series in rows" versus "with series in
    columns". :attr:`.Chart.series_in_rows` — shipped by #828
    (commit 575de804) — is that surface area. The scenarios below
    pin the reporter's canonical read cases against the real
    chart-authoring pipeline plus a save/reopen round-trip.
    """

    # -- (1) default orientation: series-in-columns is False ----------

    def it_reports_False_for_a_chart_authored_by_add_chart_440(self):
        """``Chart.series_in_rows`` is ``False`` for a default-layout chart.

        ``add_chart`` (the public authoring entry point) and the
        :class:`.CategoryChartData` pipeline write series-in-columns —
        each series name in row 1 of its own column, categories
        running vertically down a single column. That is the layout
        PowerPoint defaults to when you create a new chart, and the
        layout #440 wanted to be able to distinguish from the
        post-switch layout.
        """
        _prs, chart = _category_chart()

        assert chart.series_in_rows is False

    # -- (2) switched orientation: series-in-rows is True -------------

    def it_reports_True_when_the_chart_xml_is_in_row_series_layout_440(self):
        """``Chart.series_in_rows`` is ``True`` for the post-switch layout.

        Simulates the XML PowerPoint writes after *Chart Design >
        Switch Row/Column* — each series name is a single cell in
        column A rows 2..N, the category range runs horizontally
        across row 1 columns B..N, and each series's value range
        shares its series-name row. This is the #440 reporter's
        canonical "series in rows" case.
        """
        _prs, chart = _category_chart()
        _transpose_to_series_in_rows(chart)

        assert chart.series_in_rows is True

    # -- (3) single-category fall-through to the series-name ref ------

    def it_falls_through_to_the_tx_ref_on_a_single_cat_chart_440(self):
        """Single-cell ``c:cat`` ⇒ orientation decided by series-name row.

        PowerPoint still writes ``c:cat`` even when there is exactly
        one category; the single-cell range gives no orientation
        signal on its own. The #828 implementation falls through to
        the series-name reference: row 1 ⇒ ``False`` (default),
        elsewhere ⇒ ``True`` (switched). Pinning both branches here
        so a future refactor that drops the fall-through fails loudly.
        """
        # -- row-1 branch: single-cell c:cat + c:tx at row 1 ⇒ default --
        chart_row1 = _single_category_chart_with_tx_row_1()
        assert chart_row1.series_in_rows is False

        # -- elsewhere branch: single-cell c:cat + c:tx at row 2 ⇒ switched --
        chart_row2 = _single_category_chart_with_tx_row_2()
        assert chart_row2.series_in_rows is True

    # -- (4) XY / scatter: no c:cat, no tx ref ⇒ None -----------------

    def it_reports_None_for_an_xy_scatter_chart_with_no_tx_ref_440(self):
        """XY chart with no ``c:cat`` AND no parseable ``c:tx`` ⇒ ``None``.

        XY / scatter charts store x-values in ``c:xVal`` and y-values
        in ``c:yVal`` — there is no category axis and no ``c:cat``
        reference. When there is also no ``c:tx`` or the ``c:tx`` has
        no parseable cell reference (e.g. an inline ``c:v`` literal
        instead of a ``c:strRef/c:f``), the orientation concept does
        not apply and :attr:`.Chart.series_in_rows` returns ``None``.

        Pinning this keeps a well-meaning refactor from silently
        defaulting XY charts to ``False``, which would mislead
        #440-style consumers. NB: an XY chart authored via
        :class:`.XyChartData` writes a ``c:tx/c:strRef/c:f`` pointing
        at row 1 (the series name in cell B1), which falls through
        the tx-ref branch to ``False`` — that behaviour is pinned by
        :meth:`it_reports_False_for_an_xy_chart_authored_via_add_chart_440`
        below. This test exercises the stricter "no refs at all"
        shape to guarantee the ``None`` branch is reachable.
        """
        from pptx.chart.chart import Chart

        chartSpace_xml = (
            '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml'
            '/2006/chart">'
            "<c:chart><c:plotArea><c:scatterChart><c:ser>"
            '<c:idx val="0"/><c:order val="0"/>'
            "</c:ser></c:scatterChart></c:plotArea></c:chart>"
            "</c:chartSpace>"
        )
        chart = Chart(parse_xml(chartSpace_xml), None)

        assert chart.series_in_rows is None

    def it_reports_False_for_an_xy_chart_authored_via_add_chart_440(self):
        """An ``add_chart``-built XY chart falls through to the tx ref ⇒ ``False``.

        The live authoring path writes ``c:tx/c:strRef/c:f = Sheet1!$B$1``
        for the series name. Even without a ``c:cat``, the fall-through
        detects row 1 as the default orientation and returns ``False``.
        This companion to the preceding test pins the observable
        behaviour of the real authoring path — so a #440 caller
        inspecting a live XY chart knows what to expect and a future
        refactor that suppresses the XY tx-ref fall-through is caught.
        """
        _prs, chart = _xy_chart()

        assert chart.series_in_rows is False

    # -- (5) bubble: no c:cat, no tx ref ⇒ None -----------------------

    def it_reports_None_for_a_bubble_chart_with_no_tx_ref_440(self):
        """Bubble chart with no ``c:cat`` / no ``c:tx`` ref ⇒ ``None``.

        Mirrors the XY-without-tx case for the bubble-chart family —
        same contract: no ``c:cat`` and no parseable ``c:tx`` means
        orientation is not meaningful, so the property is ``None``.
        """
        from pptx.chart.chart import Chart

        chartSpace_xml = (
            '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml'
            '/2006/chart">'
            "<c:chart><c:plotArea><c:bubbleChart><c:ser>"
            '<c:idx val="0"/><c:order val="0"/>'
            "</c:ser></c:bubbleChart></c:plotArea></c:chart>"
            "</c:chartSpace>"
        )
        chart = Chart(parse_xml(chartSpace_xml), None)

        assert chart.series_in_rows is None

    def it_reports_False_for_a_bubble_chart_authored_via_add_chart_440(self):
        """An ``add_chart``-built bubble chart falls through to tx ⇒ ``False``.

        Companion to the XY version — the live bubble authoring path
        also writes ``c:tx/c:strRef/c:f = Sheet1!$B$1``, so the
        fall-through returns ``False``. Pinning this so a #440 caller
        inspecting a freshly-authored bubble chart knows what to
        expect.
        """
        _prs, chart = _bubble_chart()

        assert chart.series_in_rows is False

    # -- (6) empty chart: no c:ser ⇒ None -----------------------------

    def it_reports_None_for_a_chart_with_no_series_440(self):
        """No ``c:ser`` ⇒ ``None`` (the "nothing to infer from" fall-through).

        A chart with an empty plot area cannot advertise an
        orientation. Pinning ``None`` here so a future refactor
        that defaults to ``False`` in the "no series" branch is
        caught — silently returning ``False`` would misreport an
        empty chart as "default orientation" to a #440 caller.
        """
        chart = _empty_chart()

        assert chart.series_in_rows is None

    # -- (7) round-trip through save + reopen preserves orientation ---

    def it_round_trips_the_orientation_signal_through_save_and_reopen_440(
        self, _restore_part_factory
    ):
        """Orientation detection survives a ``Presentation.save`` + reopen.

        This is the test #440 really needs — a chart authored in the
        default layout reads as ``False`` after save/reopen, and a
        chart whose XML has been flipped to the series-in-rows layout
        reads as ``True`` after save/reopen. Round-trip fidelity of
        the orientation signal is what makes :attr:`.series_in_rows`
        useful for callers inspecting a ``.pptx`` file written by
        somebody else (which is the #440 scenario: read an existing
        file, figure out whether it's in default or switched layout,
        and decide what to do next).
        """
        # -- default orientation round-trip --
        prs_default, chart_default = _category_chart()
        assert chart_default.series_in_rows is False
        buf = io.BytesIO()
        prs_default.save(buf)
        buf.seek(0)
        reopened_default = Presentation(buf)
        reloaded_default = next(s for s in reopened_default.slides[0].shapes if s.has_chart).chart
        assert reloaded_default.series_in_rows is False

        # -- switched orientation round-trip --
        prs_switched, chart_switched = _category_chart()
        _transpose_to_series_in_rows(chart_switched)
        assert chart_switched.series_in_rows is True
        buf2 = io.BytesIO()
        prs_switched.save(buf2)
        buf2.seek(0)
        reopened_switched = Presentation(buf2)
        reloaded_switched = next(s for s in reopened_switched.slides[0].shapes if s.has_chart).chart
        assert reloaded_switched.series_in_rows is True

    # -- (8) public-surface cross-reference guard ---------------------

    def it_is_the_series_in_rows_surface_that_resolves_440_828(self):
        """Cross-reference: #440 resolves via :attr:`.Chart.series_in_rows`.

        Pins the public-surface contract: ``Chart.series_in_rows`` is
        a read-only property (no setter) because PowerPoint rewrites
        every series sub-reference on the toggle — emulating that
        without a live workbook would silently diverge from what
        PowerPoint re-authors on next save. A future refactor that
        adds a setter (and lies about what it does) regresses the
        #440/#828 contract and is caught here.
        """
        from pptx.chart.chart import Chart

        prop = Chart.__dict__.get("series_in_rows")
        assert isinstance(prop, property)
        assert prop.fget is not None
        assert prop.fset is None, (
            "Chart.series_in_rows is deliberately read-only per #828 — "
            "PowerPoint rewrites every c:ser sub-reference on the "
            "Switch Row/Column toggle and emulating that without a "
            "live workbook would diverge. Callers should use "
            "Chart.replace_data with transposed data to flip orientation."
        )

    # -- (9) xCharts cover the property too — line + pie --------------

    def it_reports_False_for_non_bar_category_charts_too_440(self):
        """Line / pie also use ``c:cat`` ⇒ property works on any category chart.

        The #440 ask was phrased in terms of "a chart" — not "a bar
        chart". :attr:`.Chart.series_in_rows` looks up the first
        ``c:ser`` anywhere under the plot area, so it works
        identically on line, pie, column, bar, area, radar, and
        doughnut. Pinning two non-bar category chart types here so a
        plot-type-specific regression is caught.
        """
        _prs, line_chart = _category_chart(XL_CHART_TYPE.LINE)
        assert line_chart.series_in_rows is False

        # -- pie has exactly one series; the single c:ser still carries --
        # -- a vertical c:cat range (categories are pie slices). --
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        pie_data = CategoryChartData()
        pie_data.categories = ["A", "B", "C"]
        pie_data.add_series("Share", (30.0, 45.0, 25.0))
        pie_chart = slide.shapes.add_chart(
            XL_CHART_TYPE.PIE, Inches(1), Inches(1), Inches(4), Inches(4), pie_data
        ).chart
        assert pie_chart.series_in_rows is False
