# pyright: reportPrivateUsage=false

"""Regression test for issue #449 — ``insert_chart`` on Content Placeholder.

Issue #449 (https://github.com/scanny/python-pptx/issues/449) asks for the
ability to insert a chart into a "Content" placeholder (the generic
multiple-purpose placeholder that accepts text, chart, table, picture, media,
or SmartArt — the one PowerPoint shows with six icons in the middle).

A Content placeholder carries no explicit ``type`` attribute on its ``p:ph``
element (or carries ``type="obj"``), so its :attr:`ph_type` resolves to
``PP_PLACEHOLDER.OBJECT``. ``_SlidePlaceholderFactory`` does not map ``OBJECT``
to a specialized subclass, so such a placeholder is loaded as a generic
:class:`pptx.shapes.placeholder.SlidePlaceholder`.

The fix ships in two steps:

  * Issue #199 (Wave 2) lifted ``insert_chart`` from
    ``ChartPlaceholder`` onto ``_BaseSlidePlaceholder``, so every slide
    placeholder — including ``SlidePlaceholder`` — inherits it.
  * Issue #333 (Wave 3, commit ``ea36e6b5``) extended the same treatment to
    ``insert_picture`` and ``insert_table`` and updated the
    ``SlidePlaceholder`` docstring to advertise the full content-insertion
    API.

This test locks the specific scenario #449 describes: open a presentation,
reach for a Content placeholder on a "Title and Content" layout slide, call
``insert_chart`` on it, save, reopen, and confirm the chart survives.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.chart import Chart
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.shapes.placeholder import (
    PlaceholderGraphicFrame,
    SlidePlaceholder,
    _BaseSlidePlaceholder,
)


@pytest.fixture
def _restore_part_factory():
    """Guard against cross-test mutation of ``PartFactory.part_type_for``.

    Same pattern as ``tests/test_issue_784_replace_audio.py`` — other
    suites replace registry entries with ``Mock`` objects and don't
    restore them, which makes real ``Presentation()`` construction
    return Mocks when test-module ordering places this module after
    those suites.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.comments import CommentAuthorsPart, CommentsPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_COMMENTS: CommentsPart,
        CT.PML_COMMENT_AUTHORS: CommentAuthorsPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
        CT.WAV: MediaPart,
        CT.AUDIO_MPEG: MediaPart,
        CT.AUDIO_WAV: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue449ContentPlaceholderInsertChart:
    """Verify #449 — `insert_chart` works on a generic Content placeholder."""

    def it_exposes_insert_chart_on_SlidePlaceholder(self):
        """The method is inherited, so `SlidePlaceholder` must expose it."""
        # ---verify the method lives on `_BaseSlidePlaceholder` (the lift
        # done by #199 that #333 normalized) and is visible on the
        # generic `SlidePlaceholder` subclass returned for Content
        # placeholders.
        assert hasattr(SlidePlaceholder, "insert_chart")
        assert SlidePlaceholder.insert_chart is _BaseSlidePlaceholder.insert_chart

    def it_loads_a_Content_placeholder_as_SlidePlaceholder(self, _restore_part_factory):
        """Confirm the factory path #449 depends on.

        A Content placeholder has ``ph_type == OBJECT`` and must not be
        mapped to a specialized subclass — it has to land on
        ``SlidePlaceholder`` so the inherited `insert_chart` applies.
        """
        prs = Presentation()
        # ---layout 1 is "Title and Content" which has an OBJECT
        # (Content) placeholder at idx=1.
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        content_ph = slide.placeholders[1]

        assert content_ph.placeholder_format.type == PP_PLACEHOLDER.OBJECT
        assert type(content_ph) is SlidePlaceholder

    def it_inserts_a_chart_into_a_Content_placeholder(self, _restore_part_factory):
        """The #449 happy path."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        content_ph = slide.placeholders[1]
        assert type(content_ph) is SlidePlaceholder

        chart_data = CategoryChartData()
        chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
        chart_data.add_series("Sales", (1.2, 2.3, 3.4, 4.5))

        result = content_ph.insert_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, chart_data)

        # ---return type is the placeholder-wrapping graphic-frame, not a
        # raw chart; the chart is reached via `.chart`---
        assert isinstance(result, PlaceholderGraphicFrame)
        assert isinstance(result.chart, Chart)
        assert result.chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
        # ---the placeholder-ness is preserved on the replacement shape;
        # shape_type reflects the underlying graphic-frame content
        # (CHART) even though the shape is still a placeholder.
        assert result.is_placeholder is True
        assert result.shape_type == MSO_SHAPE_TYPE.CHART

    def it_preserves_the_chart_across_save_and_reload(self, _restore_part_factory):
        """A round-trip regression guard for #449."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        content_ph = slide.placeholders[1]

        chart_data = CategoryChartData()
        chart_data.categories = ["A", "B", "C"]
        chart_data.add_series("s1", (10.0, 20.0, 30.0))
        content_ph.insert_chart(XL_CHART_TYPE.BAR_CLUSTERED, chart_data)

        # ---round-trip---
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        reloaded_slide = reloaded.slides[0]
        # ---placeholder list now has a graphic-frame at idx 1 in place
        # of the original Content placeholder---
        ph_by_idx = {ph.placeholder_format.idx: ph for ph in reloaded_slide.placeholders}
        assert 1 in ph_by_idx
        frame = ph_by_idx[1]
        assert frame.has_chart is True
        assert frame.chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED
        # ---category and value data survive the round-trip---
        plot = frame.chart.plots[0]
        assert list(plot.categories) == ["A", "B", "C"]
        assert list(plot.series[0].values) == [10.0, 20.0, 30.0]

    @pytest.mark.parametrize("layout_idx", [1, 3, 4, 7])
    def it_works_on_every_layout_with_a_Content_placeholder(
        self, layout_idx, _restore_part_factory
    ):
        """All Content-placeholder-bearing layouts in the default template.

        Layouts 1 (Title and Content), 3 (Two Content), 4 (Comparison),
        and 7 (Content with Caption) all carry at least one OBJECT /
        Content placeholder. Each must accept ``insert_chart``.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])

        # ---grab the first Content (OBJECT) placeholder on the slide---
        content_phs = [
            ph for ph in slide.placeholders
            if ph.placeholder_format.type == PP_PLACEHOLDER.OBJECT
        ]
        assert content_phs, f"layout {layout_idx} has no OBJECT placeholder"
        content_ph = content_phs[0]
        assert type(content_ph) is SlidePlaceholder

        chart_data = CategoryChartData()
        chart_data.categories = ["x", "y"]
        chart_data.add_series("s", (1.0, 2.0))

        result = content_ph.insert_chart(XL_CHART_TYPE.PIE, chart_data)

        assert isinstance(result, PlaceholderGraphicFrame)
        assert result.chart.chart_type == XL_CHART_TYPE.PIE
