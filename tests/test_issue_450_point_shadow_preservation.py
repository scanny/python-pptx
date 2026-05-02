# pyright: reportPrivateUsage=false

"""Regression test for issue #450 -- shadow lost when coloring a data point.

Issue #450 (https://github.com/scanny/python-pptx/issues/450) reports that
assigning a custom color to a single ``CategoryPoint`` on a chart series
wipes out the shadow the series had previously been given. PowerPoint
treats a data-point-level ``c:spPr`` as a *complete* replacement for the
series-level shape properties -- anything missing from the point's
``c:spPr`` is no longer rendered. The original code created the point's
``c:spPr`` fresh with only the fill, so any ``a:effectLst`` (shadow) or
``a:ln`` (line) authored on the series was silently dropped from the
point's rendering.

The fix (``CT_DPt._new_spPr``) seeds a freshly-created point ``c:spPr``
with deep copies of the series ``c:spPr``'s non-fill children
(``a:ln``, ``a:effectLst``, ``a:effectDag``, ``a:scene3d``, ``a:sp3d``),
so the user's color override is additive rather than destructive.

This module pins the reporter's exact flow -- author a series-level
shadow, override one data point's color, then round-trip through
``Presentation.save`` -- and asserts the shadow survives on the point.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Inches, Pt


def _make_chart():
    """Return a freshly authored clustered-bar chart with one three-point series."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_data = CategoryChartData()
    chart_data.categories = ["A", "B", "C"]
    chart_data.add_series("S1", (10, 20, 30))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    ).chart
    return prs, chart


class DescribeIssue450PointShadowPreservation(object):
    """End-to-end regression suite for issue #450."""

    def it_preserves_series_shadow_on_a_colored_point(self):
        """The reporter's exact flow: author a series-level shadow, then
        override a single point's color. Shadow must still be present on
        the point's ``c:spPr``.
        """
        prs, chart = _make_chart()
        series = chart.plots[0].series[0]

        # -- author a series shadow via the public API --
        series.format.shadow.blur_radius = Emu(63500)
        series.format.shadow.distance = Emu(38100)
        series.format.shadow.direction = 45.0
        series.format.shadow.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # -- override a single point's color --
        point = series.points[1]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor(0xFF, 0x00, 0x00)

        # -- the point's own spPr now carries both the fill AND the
        # -- inherited shadow, so PowerPoint will render both --
        dPt = series._ser.xpath('c:dPt[c:idx/@val="1"]')[0]
        assert dPt.xpath("c:spPr/a:solidFill/a:srgbClr/@val") == ["FF0000"]
        assert dPt.xpath("c:spPr/a:effectLst/a:outerShdw") != []
        assert dPt.xpath("c:spPr/a:effectLst/a:outerShdw/@blurRad") == ["63500"]
        assert dPt.xpath("c:spPr/a:effectLst/a:outerShdw/@dist") == ["38100"]

        # -- and the series spPr was not mutated by the point override --
        ser_spPr = series._ser.xpath("c:spPr")[0]
        assert ser_spPr.xpath("a:effectLst/a:outerShdw/@blurRad") == ["63500"]

    def it_preserves_series_outline_on_a_colored_point(self):
        """The same inheritance applies to ``a:ln``. If the series has an
        explicit outline and the user overrides one point's color, the
        outline must stay on that point.
        """
        prs, chart = _make_chart()
        series = chart.plots[0].series[0]

        # -- series fill + outline --
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor(0x22, 0x88, 0xCC)
        series.format.line.color.rgb = RGBColor(0x00, 0x00, 0x00)
        series.format.line.width = Pt(1.5)

        point = series.points[0]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor(0xFF, 0x99, 0x00)

        dPt = series._ser.xpath('c:dPt[c:idx/@val="0"]')[0]
        assert dPt.xpath("c:spPr/a:solidFill/a:srgbClr/@val") == ["FF9900"]
        # -- outline survived the override --
        assert dPt.xpath("c:spPr/a:ln/@w") == ["19050"]
        assert dPt.xpath("c:spPr/a:ln/a:solidFill/a:srgbClr/@val") == ["000000"]

    def it_round_trips_point_shadow_through_save_and_reopen(self):
        """Author a series-level shadow, override a point's color, save
        the deck to a stream, reopen, and assert the shadow is still on
        the point.
        """
        prs, chart = _make_chart()
        series = chart.plots[0].series[0]

        series.format.shadow.blur_radius = Emu(76200)
        series.format.shadow.distance = Emu(25400)
        series.format.shadow.direction = 135.0
        series.format.shadow.color.rgb = RGBColor(0x11, 0x22, 0x33)

        point = series.points[2]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor(0xAB, 0xCD, 0xEF)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        prs2 = Presentation(buf)
        graphic_frame = next(shape for shape in prs2.slides[0].shapes if shape.has_chart)
        chart2 = graphic_frame.chart
        series2 = chart2.plots[0].series[0]

        dPt = series2._ser.xpath('c:dPt[c:idx/@val="2"]')[0]
        assert dPt.xpath("c:spPr/a:solidFill/a:srgbClr/@val") == ["ABCDEF"]
        outerShdw = dPt.xpath("c:spPr/a:effectLst/a:outerShdw")
        assert outerShdw != []
        assert outerShdw[0].get("blurRad") == "76200"
        assert outerShdw[0].get("dist") == "25400"
        assert dPt.xpath("c:spPr/a:effectLst/a:outerShdw/a:srgbClr/@val") == ["112233"]

    def it_does_not_add_spurious_children_when_series_has_no_spPr(self):
        """Baseline: without any series-level overrides, the point's
        ``c:spPr`` should only contain the fill the user added -- no
        fabricated inheritance children.
        """
        prs, chart = _make_chart()
        series = chart.plots[0].series[0]
        # -- no series-level format set --

        point = series.points[0]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor(0x44, 0x55, 0x66)

        dPt = series._ser.xpath('c:dPt[c:idx/@val="0"]')[0]
        assert dPt.xpath("c:spPr/a:solidFill/a:srgbClr/@val") == ["445566"]
        assert dPt.xpath("c:spPr/a:ln") == []
        assert dPt.xpath("c:spPr/a:effectLst") == []

    def it_does_not_inherit_series_fill_onto_the_point(self):
        """Only non-fill children propagate. The point's fill MUST be the
        override; the series fill must not sneak in as a sibling.
        """
        prs, chart = _make_chart()
        series = chart.plots[0].series[0]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor(0x22, 0x88, 0xCC)

        point = series.points[1]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor(0xFF, 0x00, 0x00)

        dPt = series._ser.xpath('c:dPt[c:idx/@val="1"]')[0]
        # -- exactly one solidFill, and it's the override color --
        srgbClr_vals = dPt.xpath("c:spPr/a:solidFill/a:srgbClr/@val")
        assert srgbClr_vals == ["FF0000"]
