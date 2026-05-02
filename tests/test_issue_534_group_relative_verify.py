# pyright: reportPrivateUsage=false

"""Regression test for issue #534 — shape position/size relative to parent GroupShape.

Issue #534 (https://github.com/scanny/python-pptx/issues/534) asked for a way
to report a shape's position and size **relative to** (i.e. composed through)
its parent :class:`.GroupShape`. The OOXML reality is that PowerPoint stores
each group-child's ``a:off`` / ``a:ext`` in the enclosing group's local
coordinate system (``a:chOff`` / ``a:chExt``); when the group has been resized
on the slide the group's ``a:ext`` diverges from its ``a:chExt`` and the
child's raw numbers no longer match where it actually renders. The reporter
(and several follow-on issues such as #838) wanted the slide-relative
rectangle PowerPoint renders at, not the group-local one.

Wave 3 #925 (``fix/issue-925-group-shape-transform``) shipped exactly that
API surface — four read-only companion properties on :class:`.BaseShape`
that walk every enclosing ``p:grpSp`` ancestor and apply its
``a:chOff``/``a:chExt`` → ``a:off``/``a:ext`` linear transform:

* :attr:`.BaseShape.effective_left`
* :attr:`.BaseShape.effective_top`
* :attr:`.BaseShape.effective_width`
* :attr:`.BaseShape.effective_height`

The pre-existing :attr:`.BaseShape.left` / :attr:`~.BaseShape.top` /
:attr:`~.BaseShape.width` / :attr:`~.BaseShape.height` continue to return the
group-local raw values for backwards compatibility — giving the caller both
frames of reference: raw (parent-relative / group-local) and effective
(slide-relative).

Wave 9 #838 already verified the straight-forward resized-group scenario.
#534 is the sibling "group-local ask"; this suite pins the complementary
angles the #534 reporter's framing emphasized — scenarios #838's suite does
not cover:

* an **asymmetric 2D resize** (both x and y scale simultaneously, each by a
  distinct ratio) — the common real-world case when a user grabs a corner
  handle;
* a group authored with a **non-zero** ``a:chOff`` (child-frame origin not at
  the slide-coord origin), so the ``(x - chOff_x)`` translation term is
  non-zero and meaningfully contributes to the composited ``effective_left``;
* **picture shapes** (``p:pic``) as group children — the #534 reporter's
  wording was shape-kind agnostic and the transform must apply to every
  ``BaseShape`` subclass, not just auto-shapes;
* the dual-report contract — **both** the raw (parent/group-local) and the
  effective (slide-relative) readings are simultaneously correct and both
  are accessible on the same shape object, which is the literal "report
  relative to parent" reading the #534 title admits.

Any future regression that breaks slide-relative compositing for children of
a 2D-resized group, or for group children anchored at a non-zero ``a:chOff``,
or for non-``p:sp`` shapes, will trip here.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

import pytest

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

if TYPE_CHECKING:
    from pptx.shapes.group import GroupShape


# -- helpers ----------------------------------------------------------------


def _resize_group_on_slide(grp: GroupShape, new_cx: int, new_cy: int) -> None:
    """Scale `grp`'s slide-frame ``a:ext`` to (new_cx, new_cy) WITHOUT touching
    its ``a:chExt`` — mimicking how PowerPoint records a user-resized group.

    After :meth:`SlideShapes.add_group_shape` + ``recalculate_extents``, the
    group's ``a:ext`` equals its ``a:chExt`` and children render at their raw
    local size. Divergence is how PowerPoint persists a user-dragged corner
    resize.
    """
    xfrm = grp._element.xfrm
    assert xfrm is not None
    ext = xfrm.find(qn("a:ext"))
    assert ext is not None
    ext.set("cx", str(new_cx))
    ext.set("cy", str(new_cy))


def _offset_group_child_frame(grp: GroupShape, new_ch_x: int, new_ch_y: int) -> None:
    """Translate `grp`'s ``a:chOff`` (child-frame origin) in place.

    PowerPoint may author a group whose child-frame origin is not at the
    slide-coord origin — e.g. a group imported from another deck, or a group
    whose children sit away from the group's own top-left. This helper
    simulates that by rewriting ``a:chOff`` directly. Children's raw
    ``a:off`` are NOT touched; the effect is to make the ``(x - chOff_x)``
    translation term of the effective-geometry transform non-zero.
    """
    xfrm = grp._element.xfrm
    assert xfrm is not None
    chOff = xfrm.find(qn("a:chOff"))
    assert chOff is not None
    chOff.set("x", str(new_ch_x))
    chOff.set("y", str(new_ch_y))


# -- fixtures ---------------------------------------------------------------


@pytest.fixture
def group_with_asymmetric_2d_resize():
    """Return ``(prs, slide, grp, shape)`` for a group resized 2:1 in x and
    3:1 in y — the typical corner-handle drag that scales both axes at once.

    Children are authored at ``Inches(1,1)`` size ``Inches(2,2)`` so that
    each axis has a different scale factor and the x/y compositing cannot
    accidentally share math.
    """
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(2))
    grp = slide.shapes.add_group_shape([shape])

    raw_cx = grp.width
    raw_cy = grp.height
    # -- halve x, third y: asymmetric 2D scale (sx=0.5, sy=1/3). --
    _resize_group_on_slide(grp, raw_cx // 2, raw_cy // 3)
    return prs, slide, grp, shape


@pytest.fixture
def group_with_nonzero_chOff():
    """Return ``(prs, slide, grp, shape)`` for a group whose child-frame origin
    has been translated to ``(Inches(10), Inches(10))``. The child's raw ``a:off``
    is untouched — this only shifts the group's ``a:chOff``, making the
    effective-geometry translation term non-zero.
    """
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3), Inches(3), Inches(2), Inches(2))
    grp = slide.shapes.add_group_shape([shape])
    # -- group now has a:off == a:chOff at (Inches(3),Inches(3)) and            --
    # -- a:ext == a:chExt at (Inches(2),Inches(2)). Keep ext unchanged         --
    # -- (so sx = sy = 1.0) and translate chOff to force the                  --
    # -- ``(x - chOff_x)`` term to dominate.                                   --
    _offset_group_child_frame(grp, int(Inches(10)), int(Inches(10)))
    return prs, slide, grp, shape


# -- test suite -------------------------------------------------------------


class DescribeIssue534GroupRelativePositionSize(object):
    """#534 shape-position-relative-to-parent-GroupShape verify-and-close via #925.

    Pins the complementary scenarios the #534 framing emphasizes — an
    asymmetric 2D resize, a non-zero child-frame origin, non-auto-shape
    children (pictures), and the dual-report contract where raw and
    effective values co-exist on the same shape.
    """

    # -- asymmetric 2D resize --------------------------------------------

    def it_composites_x_and_y_scales_independently_on_effective_left_top(
        self, group_with_asymmetric_2d_resize
    ):
        # -- sx = 0.5, sy = 1/3. Child local (Inches(1), Inches(1)) with       --
        # -- chOff == (Inches(1), Inches(1)) and off == (Inches(1), Inches(1)) --
        # -- composites back to the group's own anchor on both axes.           --
        _, _, grp, shape = group_with_asymmetric_2d_resize
        assert shape.effective_left == grp.left
        assert shape.effective_top == grp.top

    def it_composites_x_and_y_scales_independently_on_effective_width_height(
        self, group_with_asymmetric_2d_resize
    ):
        # -- raw size (Inches(2), Inches(2)) scaled by (sx=0.5, sy=1/3)        --
        # --   effective_width  = Inches(2) * 0.5   = Inches(1)                --
        # --   effective_height = Inches(2) * 1/3                              --
        _, _, _, shape = group_with_asymmetric_2d_resize
        assert shape.effective_width == Inches(1)
        # -- exact EMU: floor-divided within the impl's ``int(round(...))``.  --
        # -- Inches(2) = 1,828,800 EMU; /3 rounds to 609,600.                 --
        assert shape.effective_height == Emu(round(int(Inches(2)) / 3))

    def it_preserves_raw_values_under_asymmetric_2d_resize(self, group_with_asymmetric_2d_resize):
        # -- raw values are group-local and unaffected by the group resize.    --
        # -- This is the #534 reporter's "relative to parent GroupShape" read.  --
        _, _, _, shape = group_with_asymmetric_2d_resize
        assert shape.left == Inches(1)
        assert shape.top == Inches(1)
        assert shape.width == Inches(2)
        assert shape.height == Inches(2)

    # -- non-zero a:chOff translation ------------------------------------

    def it_applies_the_chOff_translation_term_to_effective_left_top(self, group_with_nonzero_chOff):
        # -- group: off = (Inches(3), Inches(3)),  ext = chExt -> sx = sy = 1 --
        # -- chOff has been translated to (Inches(10), Inches(10)); child's   --
        # -- raw off stays at (Inches(3), Inches(3)).                         --
        # --   effective_left = off_x + (x - chOff_x) * sx                    --
        # --                 = Inches(3) + (Inches(3) - Inches(10)) * 1.0     --
        # --                 = Inches(3) - Inches(7) = -Inches(4)             --
        # -- Same arithmetic on y.                                            --
        _, _, _, shape = group_with_nonzero_chOff
        assert shape.effective_left == Inches(3) - Inches(7)
        assert shape.effective_top == Inches(3) - Inches(7)

    def it_preserves_effective_width_height_when_only_chOff_changes(self, group_with_nonzero_chOff):
        # -- chOff only shifts the translation; the scale factors are untouched --
        # -- so effective_{width,height} equal raw {width,height}.              --
        _, _, _, shape = group_with_nonzero_chOff
        assert shape.effective_width == shape.width == Inches(2)
        assert shape.effective_height == shape.height == Inches(2)

    # -- non-auto-shape children (picture) -------------------------------

    def it_composites_the_transform_for_a_picture_group_child(self, tmp_path):
        # -- the #534 ask was shape-kind agnostic: a Picture inside a resized  --
        # -- group must also composite through the group's transform.         --
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        # -- write a tiny PNG to disk for add_picture() (must be a real file).  --
        png_path = tmp_path / "px.png"
        # -- minimal 1x1 PNG (PIL round-trip keeps it portable) --
        from PIL import Image

        Image.new("RGB", (1, 1), (0, 0, 0)).save(str(png_path))

        pic = slide.shapes.add_picture(str(png_path), Inches(2), Inches(2), Inches(4), Inches(4))
        grp = slide.shapes.add_group_shape([pic])

        # -- Resize the group 2:1 in x and leave y unchanged.                  --
        raw_cx = grp.width
        raw_cy = grp.height
        _resize_group_on_slide(grp, raw_cx // 2, raw_cy)

        # -- raw still reports group-local authoring-time values --
        assert pic.left == Inches(2)
        assert pic.width == Inches(4)
        # -- effective composites through sx = 0.5 --
        # --   effective_left = off_x + (x - chOff_x) * sx                     --
        # --                 = Inches(2) + (Inches(2) - Inches(2)) * 0.5       --
        # --                 = Inches(2)                                       --
        assert pic.effective_left == grp.left == Inches(2)
        assert pic.effective_width == Inches(2)  # Inches(4) * 0.5
        # -- y axis passes through unchanged (sy = 1.0) --
        assert pic.effective_top == Inches(2)
        assert pic.effective_height == Inches(4)

    # -- dual-report contract ---------------------------------------------

    def it_reports_both_raw_and_effective_geometry_on_the_same_shape(
        self, group_with_asymmetric_2d_resize
    ):
        # -- the literal "report relative to parent GroupShape" reading of    --
        # -- #534's title: raw (group-local / parent-relative) and effective  --
        # -- (slide-relative) are both accessible on the same ``BaseShape``,  --
        # -- independently and without mutating each other.                   --
        _, _, _, shape = group_with_asymmetric_2d_resize

        raw = (shape.left, shape.top, shape.width, shape.height)
        eff = (
            shape.effective_left,
            shape.effective_top,
            shape.effective_width,
            shape.effective_height,
        )

        # -- raw and effective are distinct (the whole point of the issue).   --
        assert raw != eff
        # -- re-reading raw again after touching effective does not change it. --
        assert (shape.left, shape.top, shape.width, shape.height) == raw
        # -- assigning to raw setters leaves effective consistent with the    --
        # -- new raw value (raw writes flow through ``a:off`` / ``a:ext``).    --
        shape.left = Inches(2)
        assert shape.left == Inches(2)
        # -- with sx = 0.5 and the child-frame still (chOff.x = Inches(1))    --
        # -- the new effective_left is Inches(1) + (Inches(2) - Inches(1)) * 0.5 --
        # --                         = Inches(1) + Inches(0.5) = Inches(1.5)  --
        assert shape.effective_left == Inches(1) + Inches(1) // 2

    # -- round-trip survives save + reopen --------------------------------

    def it_round_trips_effective_geometry_through_save_reopen(
        self, group_with_asymmetric_2d_resize
    ):
        prs, _, _, shape_orig = group_with_asymmetric_2d_resize
        orig_eff = (
            shape_orig.effective_left,
            shape_orig.effective_top,
            shape_orig.effective_width,
            shape_orig.effective_height,
        )
        shape_id = shape_orig.shape_id

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        slide2 = prs2.slides[0]

        grp2 = next(s for s in slide2.shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP)
        shape2 = next(s for s in grp2.shapes if s.shape_id == shape_id)

        # -- effective geometry survives save + reopen unchanged --
        assert (
            shape2.effective_left,
            shape2.effective_top,
            shape2.effective_width,
            shape2.effective_height,
        ) == orig_eff
        # -- raw group-local values survive unchanged too --
        assert (shape2.left, shape2.top, shape2.width, shape2.height) == (
            Inches(1),
            Inches(1),
            Inches(2),
            Inches(2),
        )
