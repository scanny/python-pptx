# pyright: reportPrivateUsage=false

"""Regression test for issue #556 — replace title text, keep formatting.

Issue #556 (https://github.com/scanny/python-pptx/issues/556) asks how
to replace a slide's **title** text — in particular, a title authored
like ``My title __ placeholder __`` where the bare ``__`` tokens are
double-underscore placeholders the user wants rewritten in-place —
**without** having to rebuild the paragraph from scratch and lose the
run-level formatting (bold, italic, font family, size, colour, …) that
the author applied around the token.

This is the same "rewrite text, keep run formatting" problem surfaced
by issue #285 and solved in Wave 6 by issue #836, which shipped
:meth:`TextFrame.replace_text` / :meth:`_Paragraph.replace_text` on the
``src/pptx/text/text.py`` ``TextFrame`` / ``_Paragraph`` proxies. Those
methods:

  * Search across consecutive ``a:r`` runs in a paragraph so a token
    broken into two runs by PowerPoint's autoformatter is still matched.
  * Keep the ``a:rPr`` of the run where the match *starts* for the
    replacement text.
  * Drop runs fully contained within the match; a run that the match
    ends inside keeps its formatting on its surviving suffix.
  * Never cross ``a:br`` / ``a:fld`` / paragraph boundaries.

See ``tests/test_issue_285_replace_text_preserve_format.py`` for the
textbox-facing regression of the same capability and
``tests/text/test_text.py::Describe_Paragraph_replace_text`` for the
unit-level contract.

This file pins the #556-specific scenario: the shape under test is the
*title placeholder* obtained via ``slide.shapes.title`` (a
:class:`~pptx.shapes.placeholder.SlidePlaceholder`), the token is the
double-underscore literal ``"__"`` asked about by the reporter, and the
surrounding text carries rich run-level formatting that must survive
both the replacement and a full ``Presentation.save`` + reopen
round-trip. Two occurrences of ``__`` in the same title are both
rewritten by a single ``replace_text`` call, per the documented "every
occurrence" contract.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt


class DescribeIssue556TitleReplaceTextPreservesFormatting:
    """End-to-end regression: rewrite ``__`` placeholders in the title."""

    def it_replaces_both_double_underscore_tokens_and_keeps_formatting(self):
        """Both ``__`` placeholders are rewritten in one call.

        Author a title-bearing slide with a single run holding
        ``"My title __ placeholder __"`` and a full palette of run-level
        formatting (bold + italic + 28-pt size + named font family +
        explicit RGB colour + underline). A single
        ``title.text_frame.replace_text("__", "NAME")`` call must rewrite
        **both** occurrences of ``__`` to ``NAME``, return ``2``, and
        leave the run's ``a:rPr`` entirely intact on the rewritten run.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        assert title is not None  # -- layout[0] has a title placeholder

        tf = title.text_frame
        # -- start from a single authored run carrying rich formatting --
        tf.text = "My title __ placeholder __"
        run = tf.paragraphs[0].runs[0]
        run.font.bold = True
        run.font.italic = True
        run.font.size = Pt(28)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(0x11, 0x22, 0x33)
        run.font.underline = True

        # -- act: targeted replacement of the double-underscore token --
        count = tf.replace_text("__", "NAME")

        # -- both tokens rewritten, text reflects the substitution --
        assert count == 2
        assert tf.text == "My title NAME placeholder NAME"

        # -- a single run still carries the original formatting --
        rewritten = tf.paragraphs[0].runs[0]
        assert rewritten.text == "My title NAME placeholder NAME"
        assert rewritten.font.bold is True
        assert rewritten.font.italic is True
        assert rewritten.font.size == Pt(28)
        assert rewritten.font.name == "Calibri"
        assert rewritten.font.color.rgb == RGBColor(0x11, 0x22, 0x33)
        assert rewritten.font.underline is True

    def it_preserves_formatting_through_save_and_reopen(self):
        """Formatting and replacement survive a save + reopen round-trip.

        Same #556 scenario as the first case, re-asserted after
        ``Presentation.save`` + reload, which is the actual shipping path
        most ``python-pptx`` callers care about: populate a template's
        title, replace the placeholder tokens, write to disk, read back.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        assert title is not None

        tf = title.text_frame
        tf.text = "My title __ placeholder __"
        run = tf.paragraphs[0].runs[0]
        run.font.bold = True
        run.font.italic = True
        run.font.size = Pt(28)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(0x11, 0x22, 0x33)
        run.font.underline = True

        count = tf.replace_text("__", "NAME")
        assert count == 2

        # -- round-trip through save + reopen --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        title2 = prs2.slides[0].shapes.title
        assert title2 is not None
        tf2 = title2.text_frame

        assert tf2.text == "My title NAME placeholder NAME"
        reloaded = tf2.paragraphs[0].runs[0]
        assert reloaded.text == "My title NAME placeholder NAME"
        assert reloaded.font.bold is True
        assert reloaded.font.italic is True
        assert reloaded.font.size == Pt(28)
        assert reloaded.font.name == "Calibri"
        assert reloaded.font.color.rgb == RGBColor(0x11, 0x22, 0x33)
        assert reloaded.font.underline is True

    def it_preserves_surrounding_run_formatting_when_token_is_its_own_run(self):
        """PowerPoint-splits-the-placeholder case #556 cares about.

        After editing in PowerPoint, a templated title commonly ends up
        with the placeholder in its own ``a:r`` element, flanked by the
        static prose in two differently formatted runs. ``replace_text``
        must only rewrite the token's run and leave the prose runs'
        ``a:rPr`` untouched — so bold/italic/underline on the prose is
        preserved independent of whatever formatting the placeholder
        run carried.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        assert title is not None

        tf = title.text_frame
        # -- start with a clean paragraph and build three runs manually --
        tf.text = ""
        p = tf.paragraphs[0]

        # -- run #1: bold prose preceding the placeholder --
        r1 = p.add_run()
        r1.text = "My title "
        r1.font.bold = True

        # -- run #2: the placeholder token, italic + red --
        r2 = p.add_run()
        r2.text = "__"
        r2.font.italic = True
        r2.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

        # -- run #3: underlined prose after the placeholder --
        r3 = p.add_run()
        r3.text = " placeholder"
        r3.font.underline = True

        count = tf.replace_text("__", "NAME")
        assert count == 1
        assert tf.text == "My title NAME placeholder"

        runs = tf.paragraphs[0].runs
        assert len(runs) == 3

        # -- leading prose run untouched --
        assert runs[0].text == "My title "
        assert runs[0].font.bold is True
        assert runs[0].font.italic is None
        assert runs[0].font.underline is None

        # -- placeholder run carries its own formatting onto "NAME" --
        assert runs[1].text == "NAME"
        assert runs[1].font.italic is True
        assert runs[1].font.color.rgb == RGBColor(0xFF, 0x00, 0x00)
        assert runs[1].font.bold is None
        assert runs[1].font.underline is None

        # -- trailing prose run untouched --
        assert runs[2].text == " placeholder"
        assert runs[2].font.underline is True
        assert runs[2].font.bold is None
        assert runs[2].font.italic is None

        # -- round-trip assertions --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        title2 = prs2.slides[0].shapes.title
        assert title2 is not None
        reloaded_runs = title2.text_frame.paragraphs[0].runs
        assert len(reloaded_runs) == 3
        assert reloaded_runs[0].text == "My title "
        assert reloaded_runs[0].font.bold is True
        assert reloaded_runs[1].text == "NAME"
        assert reloaded_runs[1].font.italic is True
        assert reloaded_runs[1].font.color.rgb == RGBColor(0xFF, 0x00, 0x00)
        assert reloaded_runs[2].text == " placeholder"
        assert reloaded_runs[2].font.underline is True
