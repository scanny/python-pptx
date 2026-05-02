# pyright: reportPrivateUsage=false

"""Regression test for issue #560 — customize colors of series data labels.

Issue #560 (https://github.com/scanny/python-pptx/issues/560) asks for
a way to customize the colors of data labels on a chart series — both
the font color of the label text and the fill / border color of the
containing rectangle — at the per-series and per-point level.

Coverage map:

* Per-point **font color** via ``Series.points[i].data_label.font.color.rgb``.
  Exists since the original ``DataLabel.font`` property; this test locks
  down the end-to-end round trip through save / reload so future refactors
  of the ``c:dLbl / c:txPr / a:p / a:pPr / a:defRPr`` subtree keep the
  font-color flow working.

* Per-point **fill / line** via ``Series.points[i].data_label.format.fill``
  and ``.format.line`` — shipped by #716
  (``feat/issue-716-datalabel-border``), verified here in the issue #560
  regression harness.

* Per-series **font color** via ``Series.data_labels.font.color.rgb`` —
  exercises ``DataLabels.font`` on ``c:dLbls/c:txPr/a:p/a:pPr/a:defRPr``.

* Per-series **fill / line** via ``Series.data_labels.format.fill`` and
  ``.format.line`` — new in this branch. Before this change ``DataLabels``
  had no ``.format`` property, so there was no clean way to color all the
  labels on a series in one shot without dropping into oxml.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.chart.datalabel import DataLabels
from pptx.dml.chtfmt import ChartFormat
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


@pytest.fixture
def _restore_part_factory():
    """Guard against pollution from tests that mutate `PartFactory.part_type_for`.

    `tests/opc/test_package.py::DescribePartFactory` overwrites
    `PartFactory.part_type_for[CT.PML_SLIDE]` (and friends) with Mock objects
    and does not restore them. Without this guard, running this module after
    ``tests/opc/test_package.py`` causes ``prs.slides[0]`` to return a Mock
    instead of a real Slide. Same pattern as the fixture in
    ``tests/test_comments.py`` / ``tests/test_issue_400_animation_umbrella.py``.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.comments import CommentAuthorsPart, CommentsPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_COMMENTS: CommentsPart,
        CT.PML_COMMENT_AUTHORS: CommentAuthorsPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


def _build_chart_and_roundtrip():
    """Build a single-series column chart with data labels, save, reload."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_data = CategoryChartData()
    chart_data.categories = ["A", "B", "C"]
    chart_data.add_series("s1", (1.0, 2.0, 3.0))
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    )
    chart = chart_shape.chart
    chart.plots[0].has_data_labels = True
    series = chart.plots[0].series[0]

    # ---per-point font color (issue #560 core request)---
    series.points[1].data_label.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    # ---per-point fill + line (shipped by #716, re-exercised here)---
    series.points[0].data_label.format.fill.solid()
    series.points[0].data_label.format.fill.fore_color.rgb = RGBColor(0x00, 0xFF, 0x00)
    series.points[0].data_label.format.line.color.rgb = RGBColor(0x00, 0x00, 0xFF)

    # ---per-series font color---
    series.data_labels.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    # ---per-series fill + line (new in this branch)---
    series.data_labels.format.fill.solid()
    series.data_labels.format.fill.fore_color.rgb = RGBColor(0x11, 0x22, 0x33)
    series.data_labels.format.line.color.rgb = RGBColor(0x44, 0x55, 0x66)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)


def _reloaded_series(prs):
    for shp in prs.slides[0].shapes:
        if shp.has_chart:
            return shp.chart.plots[0].series[0]
    raise AssertionError("reloaded presentation contains no chart")


class DescribeIssue560DataLabelColors(object):
    def it_round_trips_per_point_font_color(self, _restore_part_factory):
        prs = _build_chart_and_roundtrip()
        series = _reloaded_series(prs)

        assert series.points[1].data_label.font.color.rgb == RGBColor(0xFF, 0x00, 0x00)

    def it_round_trips_per_point_fill_and_line(self, _restore_part_factory):
        prs = _build_chart_and_roundtrip()
        series = _reloaded_series(prs)

        dl = series.points[0].data_label
        assert dl.format.fill.fore_color.rgb == RGBColor(0x00, 0xFF, 0x00)
        assert dl.format.line.color.rgb == RGBColor(0x00, 0x00, 0xFF)

    def it_round_trips_per_series_font_color(self, _restore_part_factory):
        prs = _build_chart_and_roundtrip()
        series = _reloaded_series(prs)

        assert series.data_labels.font.color.rgb == RGBColor(0x80, 0x80, 0x80)

    def it_round_trips_per_series_fill_and_line(self, _restore_part_factory):
        """New with issue #560: DataLabels.format for per-series fill / line.

        Regression for the #560 gap — before this branch, ``DataLabels``
        had no ``.format`` property, so there was no series-level API for
        the containing rectangle's fill or border color.
        """
        prs = _build_chart_and_roundtrip()
        series = _reloaded_series(prs)

        fmt = series.data_labels.format
        assert isinstance(fmt, ChartFormat)
        assert fmt.fill.fore_color.rgb == RGBColor(0x11, 0x22, 0x33)
        assert fmt.line.color.rgb == RGBColor(0x44, 0x55, 0x66)

    def it_returns_a_ChartFormat_from_DataLabels_format(self, _restore_part_factory):
        """DataLabels.format is a ChartFormat wrapping c:dLbls."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        chart_data = CategoryChartData()
        chart_data.categories = ["A", "B"]
        chart_data.add_series("s1", (1.0, 2.0))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(6),
            Inches(4),
            chart_data,
        ).chart
        chart.plots[0].has_data_labels = True
        series = chart.plots[0].series[0]

        data_labels = series.data_labels

        assert isinstance(data_labels, DataLabels)
        fmt = data_labels.format
        assert isinstance(fmt, ChartFormat)
        # ---lazyproperty: same object on subsequent access---
        assert data_labels.format is fmt
