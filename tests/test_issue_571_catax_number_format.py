# pyright: reportPrivateUsage=false

"""Regression test for issue #571 — number_format for category axis tick labels.

Issue #571 (https://github.com/scanny/python-pptx/issues/571) reports that
setting ``chart.category_axis.tick_labels.number_format`` appeared to have no
effect on the rendered tick labels, even though the XML correctly reflected
the assignment (``<c:numFmt formatCode="..." sourceLinked="0"/>`` appearing
under ``c:catAx``).

The underlying support — ``CategoryAxis`` inheriting ``tick_labels`` from
``_BaseAxis`` and ``CT_CatAx`` defining a ``c:numFmt`` child — has been in
the codebase all along; the reporter's rendered-output confusion is the
expected PowerPoint behaviour that category-axis tick labels are plain
strings (the category names), so a numeric format code has no visible
effect unless the categories are themselves numbers (as on a date axis or
on a "numeric" category axis that PowerPoint treats as a value scale).

This regression test pins the public API for both the read path and the
write path on ``CategoryAxis.tick_labels.number_format`` and round-trips the
assignment through ``Presentation.save`` + reopen, so any future change that
drops ``c:numFmt`` from ``CT_CatAx``'s ``_tag_seq`` or reroutes
``TickLabels.number_format`` away from the shared ``_BaseAxis.tick_labels``
plumbing will reproduce the #571 symptom and be caught here.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.axis import CategoryAxis, TickLabels
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


@pytest.fixture
def _restore_part_factory():
    """Guard against pollution from other tests that mutate `PartFactory.part_type_for`.

    See the identical fixture in ``tests/test_issue_539_chart_color_cycle.py``
    for the rationale — ``Presentation(stream)`` reopen dispatches through
    ``PartFactory.part_type_for``, which earlier tests may have emptied.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.comments import CommentAuthorsPart, CommentsPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_COMMENTS: CommentsPart,
        CT.PML_COMMENT_AUTHORS: CommentAuthorsPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue571CategoryAxisNumberFormat(object):
    """Pins ``CategoryAxis.tick_labels.number_format`` round-trip behaviour."""

    def it_defaults_number_format_to_General_when_unset(self, category_axis):
        assert isinstance(category_axis, CategoryAxis)
        assert isinstance(category_axis.tick_labels, TickLabels)
        # -- fresh catAx has no c:numFmt child, so the getter falls back to
        # -- the "General" sentinel defined by the OOXML spec.
        assert category_axis.tick_labels.number_format == "General"

    def it_can_set_and_read_back_number_format_on_category_axis(self, category_axis):
        category_axis.tick_labels.number_format = "0.00%"

        assert category_axis.tick_labels.number_format == "0.00%"

    def it_writes_formatCode_to_catAx_numFmt_xml(self, category_axis):
        category_axis.tick_labels.number_format = "$#,##0;$-#,##0"

        numFmts = category_axis._element.xpath("c:numFmt")
        assert len(numFmts) == 1
        assert numFmts[0].get("formatCode") == "$#,##0;$-#,##0"

    def it_clears_number_format_is_linked_when_setting_number_format(self, category_axis):
        # -- assigning number_format must also flip source-linked off, otherwise
        # -- PowerPoint silently overrides the custom format with the spreadsheet's.
        category_axis.tick_labels.number_format = "0.00%"

        assert category_axis.tick_labels.number_format_is_linked is False
        numFmt = category_axis._element.xpath("c:numFmt")[0]
        assert numFmt.get("sourceLinked") == "0"

    def it_round_trips_number_format_through_save_and_reload(self, _restore_part_factory):
        """Save a .pptx with a custom catAx number_format and reopen it."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        chart_data = CategoryChartData()
        chart_data.categories = ["A", "B", "C"]
        chart_data.add_series("Series 1", (1, 2, 3))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(3),
            chart_data,
        ).chart

        chart.category_axis.tick_labels.number_format = "0.00%"

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs_reloaded = Presentation(buf)
        reloaded_chart = prs_reloaded.slides[0].shapes[1].chart

        assert reloaded_chart.category_axis.tick_labels.number_format == "0.00%"
        assert reloaded_chart.category_axis.tick_labels.number_format_is_linked is False

    # fixtures -------------------------------------------------------

    @pytest.fixture
    def category_axis(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        chart_data = CategoryChartData()
        chart_data.categories = ["A", "B", "C"]
        chart_data.add_series("Series 1", (1, 2, 3))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(3),
            chart_data,
        ).chart
        return chart.category_axis
