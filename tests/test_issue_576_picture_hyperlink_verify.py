# pyright: reportPrivateUsage=false

"""Verify-close regression test for issue #576 — hyperlink on an image shape.

Issue #576 (https://github.com/scanny/python-pptx/issues/576) asked whether
a caller can attach a hyperlink to a :class:`~pptx.shapes.picture.Picture`
— either an external URL or a jump to another slide in the same deck. The
reporter observed that the documentation centred on run-level hyperlinks
and wasn't sure whether picture shapes carried their own click behaviour.

The capability has always been present: :class:`Picture` inherits from
:class:`~pptx.shapes.base.BaseShape`, which exposes
:attr:`~pptx.shapes.base.BaseShape.click_action` (and, since Wave 14,
:attr:`~pptx.shapes.base.BaseShape.hover_action`). Both return an
:class:`~pptx.action.ActionSetting` bound to the picture's
``p:nvPicPr/p:cNvPr`` element — the same surface used by auto-shapes and
connectors — so the full hyperlink / slide-jump / screen-tip / sound API
works on pictures out of the box.

The canonical authoring pattern is therefore::

    pic = slide.shapes.add_picture("logo.png", Inches(1), Inches(1))

    # -- external URL --
    pic.click_action.hyperlink.address = "https://example.com"

    # -- or a jump to another slide in the same deck --
    pic.click_action.target_slide = prs.slides[2]

    # -- optional ScreenTip and mouse-over hyperlink --
    pic.click_action.screen_tip = "Open our home page"
    pic.hover_action.hyperlink.address = "https://example.com/hover"

#576 is verified-and-closed by the pre-existing ``BaseShape.click_action``
surface plus the #734 ``hover_action`` / ``screen_tip`` / ``set_sound``
additions. This suite pins the scenarios the #576 thread calls out:

* URL hyperlink on a picture, round-tripped through save + reopen.
* Slide-jump on a picture, round-tripped through save + reopen.
* Hover-action hyperlink on a picture (added after #734).
* ScreenTip on a picture's click action.
* Clearing a picture's hyperlink (assigning ``None``) leaves no stale
  ``a:hlinkClick`` element behind.
* The ``a:hlinkClick`` is written inside ``p:nvPicPr/p:cNvPr`` — the
  correct location PowerPoint reads from for the picture element.

Any future regression that drops picture-level click/hover actions,
writes the hyperlink to the wrong element, or fails to round-trip either
flavour of action will be caught here.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.action import ActionSetting
from pptx.enum.action import PP_ACTION
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of ``PartFactory.part_type_for``.

    Mirrors the guard used by ``tests/test_issue_1020_fill_alpha_freeform_verify.py``
    and friends — ``Presentation(stream)`` reopen dispatches through
    ``PartFactory.part_type_for``, which earlier test modules overwrite with
    mocks and do not restore.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


IMAGE_PATH = "tests/test_files/python-powered.png"


def _add_picture(prs):
    """Return a freshly-added ``Picture`` on a new blank slide.

    Uses a blank layout (``slide_layouts[6]``) so no inherited placeholder
    interferes with the picture's own ``p:nvPicPr/p:cNvPr`` surface, which
    is where the click/hover hyperlink elements are authored.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return slide.shapes.add_picture(IMAGE_PATH, Inches(1), Inches(1), Inches(2), Inches(2))


def _first_picture(slide):
    """Return the first ``Picture`` on ``slide``."""
    for shp in slide.shapes:
        if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return shp
    raise AssertionError("no picture on slide")


class DescribeIssue576PictureHyperlink:
    """#576 picture hyperlink verify-and-close via ``BaseShape.click_action``.

    Pins the per-shape click / hover / screen-tip API on
    :class:`~pptx.shapes.picture.Picture`, including round-trip through
    ``Presentation.save`` + reopen.
    """

    # -- API surface ------------------------------------------------------

    def it_exposes_click_action_and_hover_action_on_Picture(self):
        """Guard: the public API that resolves #576 is present.

        A caller reaching for a picture-level hyperlink needs both
        ``click_action`` and ``hover_action`` to be ``ActionSetting``
        instances bound to the picture's own ``p:cNvPr`` element. A
        refactor that moved either off ``BaseShape`` would silently
        reintroduce #576.
        """
        prs = Presentation()
        pic = _add_picture(prs)

        assert isinstance(pic.click_action, ActionSetting)
        assert isinstance(pic.hover_action, ActionSetting)
        # -- both sit on the same cNvPr so they share ownership of tooltip and sound --
        assert pic.click_action._element is pic.hover_action._element

    # -- URL hyperlink ---------------------------------------------------

    def it_attaches_a_URL_hyperlink_to_a_picture(self):
        """The canonical "how do I link a picture to a URL?" recipe.

        The #576 thread's lead question. Assigning an address to
        ``pic.click_action.hyperlink`` writes an ``a:hlinkClick`` into
        ``p:nvPicPr/p:cNvPr`` and reads back as the same URL.
        """
        prs = Presentation()
        pic = _add_picture(prs)

        pic.click_action.hyperlink.address = "https://example.com"

        assert pic.click_action.hyperlink.address == "https://example.com"
        assert pic.click_action.action == PP_ACTION.HYPERLINK
        # -- hlinkClick is authored inside p:nvPicPr/p:cNvPr --
        cNvPr = pic._element.find(qn("p:nvPicPr") + "/" + qn("p:cNvPr"))
        assert cNvPr is not None
        assert cNvPr.find(qn("a:hlinkClick")) is not None

    def it_round_trips_a_URL_hyperlink_through_save_and_reopen(self, _restore_part_factory):
        """Save + reopen preserves the URL and action type.

        The #576 reporter's deck is saved and redistributed; a regression
        that dropped the rel on save or failed to re-hydrate the address
        on reopen would reintroduce the visible symptom.
        """
        prs = Presentation()
        pic = _add_picture(prs)
        pic.click_action.hyperlink.address = "https://example.com"

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        pic2 = _first_picture(prs2.slides[0])
        assert pic2.click_action.hyperlink.address == "https://example.com"
        assert pic2.click_action.action == PP_ACTION.HYPERLINK

    # -- slide-jump ------------------------------------------------------

    def it_attaches_a_slide_jump_hyperlink_to_a_picture(self):
        """The other half of the #576 ask — jump to another slide.

        Assigning a ``Slide`` to ``pic.click_action.target_slide`` writes a
        ``ppaction://hlinksldjump`` action whose rId points at the target
        slide part and reads back as the same ``Slide``.
        """
        prs = Presentation()
        pic = _add_picture(prs)
        target_slide = prs.slides.add_slide(prs.slide_layouts[6])

        pic.click_action.target_slide = target_slide

        assert pic.click_action.action == PP_ACTION.NAMED_SLIDE
        assert pic.click_action.target_slide == target_slide

    def it_round_trips_a_slide_jump_through_save_and_reopen(self, _restore_part_factory):
        """Save + reopen preserves the slide-jump.

        The ``hlinksldjump`` rel and the ``p:nvPicPr/p:cNvPr/a:hlinkClick``
        carrying its rId must both survive the package-assembly round trip.
        """
        prs = Presentation()
        pic = _add_picture(prs)
        target_slide = prs.slides.add_slide(prs.slide_layouts[6])
        pic.click_action.target_slide = target_slide

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        pic2 = _first_picture(prs2.slides[0])
        assert pic2.click_action.action == PP_ACTION.NAMED_SLIDE
        assert pic2.click_action.target_slide == prs2.slides[1]

    # -- hover action ----------------------------------------------------

    def it_attaches_a_hover_hyperlink_to_a_picture(self):
        """Pictures support mouse-over actions, not just click actions.

        ``BaseShape.hover_action`` (added for #734) exposes the
        ``a:hlinkMouseOver`` / ``a:hlinkHover`` slot on the picture's
        ``p:cNvPr`` — parallel to ``click_action`` but backed by the
        hover element.
        """
        prs = Presentation()
        pic = _add_picture(prs)

        pic.hover_action.hyperlink.address = "https://example.com/hover"

        assert pic.hover_action.hyperlink.address == "https://example.com/hover"
        # -- click-side still reads as no-action --
        assert pic.click_action.action == PP_ACTION.NONE
        cNvPr = pic._element.find(qn("p:nvPicPr") + "/" + qn("p:cNvPr"))
        assert cNvPr is not None
        assert cNvPr.find(qn("a:hlinkHover")) is not None
        assert cNvPr.find(qn("a:hlinkClick")) is None

    def it_round_trips_a_hover_hyperlink_through_save_and_reopen(self, _restore_part_factory):
        """Save + reopen preserves the hover address.

        The hover rel is separately registered from the click rel; a
        regression that confused the two slots would drop the hover on save.
        """
        prs = Presentation()
        pic = _add_picture(prs)
        pic.hover_action.hyperlink.address = "https://example.com/hover"

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        pic2 = _first_picture(prs2.slides[0])
        assert pic2.hover_action.hyperlink.address == "https://example.com/hover"

    # -- screen tip ------------------------------------------------------

    def it_sets_a_screen_tip_on_a_picture_click_action(self):
        """ScreenTip (tooltip) on a picture click action.

        The ``tooltip`` attribute of ``a:hlinkClick`` is surfaced via
        :attr:`ActionSetting.screen_tip`. A caller must be able to set it
        on a picture the same way as on an auto-shape.
        """
        prs = Presentation()
        pic = _add_picture(prs)
        pic.click_action.hyperlink.address = "https://example.com"

        pic.click_action.screen_tip = "Click to open our home page"

        assert pic.click_action.screen_tip == "Click to open our home page"

    # -- clearing the hyperlink -----------------------------------------

    def it_clears_a_URL_hyperlink_by_assigning_None(self):
        """Assigning ``None`` to ``hyperlink.address`` removes the action.

        A regression that left a stale ``a:hlinkClick`` behind would be
        observable as ``action != PP_ACTION.NONE`` after the clear.
        """
        prs = Presentation()
        pic = _add_picture(prs)
        pic.click_action.hyperlink.address = "https://example.com"
        assert pic.click_action.hyperlink.address == "https://example.com"

        pic.click_action.hyperlink.address = None

        assert pic.click_action.hyperlink.address is None
        assert pic.click_action.action == PP_ACTION.NONE
        # -- and no stale element is left behind --
        cNvPr = pic._element.find(qn("p:nvPicPr") + "/" + qn("p:cNvPr"))
        assert cNvPr is not None
        assert cNvPr.find(qn("a:hlinkClick")) is None

    def it_clears_a_slide_jump_by_assigning_None_to_target_slide(self):
        """Assigning ``None`` to ``target_slide`` removes the slide jump."""
        prs = Presentation()
        pic = _add_picture(prs)
        target_slide = prs.slides.add_slide(prs.slide_layouts[6])
        pic.click_action.target_slide = target_slide
        assert pic.click_action.action == PP_ACTION.NAMED_SLIDE

        pic.click_action.target_slide = None

        assert pic.click_action.action == PP_ACTION.NONE
        assert pic.click_action.target_slide is None
