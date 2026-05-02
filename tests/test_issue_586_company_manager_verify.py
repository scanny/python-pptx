"""Verify-close regression test for issue #586 — company/manager extended properties.

Issue #586 (https://github.com/scanny/python-pptx/issues/586) asked for
read/write access to the **Company** and **Manager** fields of the
application-level document properties stored in ``/docProps/app.xml`` —
the properties PowerPoint surfaces under *File → Info → Properties →
Advanced → Properties*. These are distinct from the Dublin-Core
properties exposed by :attr:`Presentation.core_properties`
(``/docProps/core.xml``) and from the user-defined
:attr:`Presentation.custom_properties` (``/docProps/custom.xml``).

The feature was shipped per ``FEATURES.md`` under the CalVer 2026.05.0
release line as :attr:`Presentation.extended_properties` — an
:class:`ExtendedPropertiesPart` proxy exposing each curated app-metadata
field as an attribute (``company``, ``manager``, ``application``,
``app_version``, ``presentation_format``, ``template``,
``hyperlink_base``, ``slide_count``).

This suite is the verify-close pass for #586. It pins:

* Reading ``company`` and ``manager`` from a deck that has them set.
* Writing ``company`` and ``manager``, saving, reopening, reading back.
* Accessibility of the other extended-properties fields
  (``application``, ``app_version``, and the slide count) the spec
  defines under the same root element.
* Non-interference with :attr:`Presentation.core_properties` — writes
  to the extended part must not disturb the core-properties part.
"""

from __future__ import annotations

import io
import zipfile
from typing import TYPE_CHECKING

from pptx import Presentation

if TYPE_CHECKING:
    from pptx.presentation import Presentation as PresentationAPI


class DescribeIssue586ExtendedProperties:
    """Verify-close regression suite for ``Presentation.extended_properties`` (issue #586)."""

    # -- reading --------------------------------------------------------------

    def it_reads_company_from_a_deck_that_has_one_set(self):
        # -- seed a deck with a <Company> value, save, reopen, and confirm
        # -- the reader surfaces the exact string the writer wrote.
        prs = Presentation()
        prs.extended_properties.company = "Acme Industries"

        prs2 = _round_trip(prs)

        assert prs2.extended_properties.company == "Acme Industries"

    def it_reads_manager_from_a_deck_that_has_one_set(self):
        prs = Presentation()
        prs.extended_properties.manager = "Daphne Director"

        prs2 = _round_trip(prs)

        assert prs2.extended_properties.manager == "Daphne Director"

    def it_returns_empty_string_when_company_and_manager_are_absent(self):
        # -- the stock default template does not carry a <Company> or
        # -- <Manager> element; reading must return "" (not None) so the
        # -- attribute is safe to use in string-interpolation contexts.
        prs = Presentation()

        ep = prs.extended_properties

        assert ep.company == ""
        assert ep.manager == ""

    # -- writing --------------------------------------------------------------

    def it_writes_company_and_manager_and_round_trips_through_save_and_reopen(self):
        # -- the #586 reporter's ask verbatim: set both fields on a deck
        # -- from Python and have PowerPoint (and python-pptx on reopen)
        # -- find the values intact in the saved app.xml.
        prs = Presentation()
        prs.extended_properties.company = "Example Inc"
        prs.extended_properties.manager = "Alex Manager"

        prs2 = _round_trip(prs)

        assert prs2.extended_properties.company == "Example Inc"
        assert prs2.extended_properties.manager == "Alex Manager"

    def it_writes_the_values_into_docProps_app_xml_on_save(self):
        # -- pin the wire format: the strings end up inside <Company> and
        # -- <Manager> children of the <Properties> root in app.xml, which
        # -- is what downstream consumers (PowerPoint, SharePoint, etc.)
        # -- actually read.
        prs = Presentation()
        prs.extended_properties.company = "Globex Corp"
        prs.extended_properties.manager = "Morgan Manager"

        buf = io.BytesIO()
        prs.save(buf)

        buf.seek(0)
        with zipfile.ZipFile(buf) as zf:
            app_xml = zf.read("docProps/app.xml").decode("utf-8")
        assert "<Company>Globex Corp</Company>" in app_xml
        assert "<Manager>Morgan Manager</Manager>" in app_xml

    def it_supports_overwriting_existing_company_and_manager(self):
        # -- second write replaces first; round-trip pins the new value,
        # -- not a duplicated element.
        prs = Presentation()
        prs.extended_properties.company = "First Co"
        prs.extended_properties.manager = "First Mgr"
        prs.extended_properties.company = "Second Co"
        prs.extended_properties.manager = "Second Mgr"

        prs2 = _round_trip(prs)

        assert prs2.extended_properties.company == "Second Co"
        assert prs2.extended_properties.manager == "Second Mgr"

        buf = io.BytesIO()
        prs2.save(buf)
        buf.seek(0)
        with zipfile.ZipFile(buf) as zf:
            app_xml = zf.read("docProps/app.xml").decode("utf-8")
        assert app_xml.count("<Company>") == 1
        assert app_xml.count("<Manager>") == 1

    # -- sibling extended-properties fields -----------------------------------

    def it_exposes_the_other_extended_properties_fields(self):
        # -- the #586 ask is scoped to company/manager, but the part also
        # -- exposes the neighbouring app.xml fields from the same
        # -- <Properties> root. Pin that they are reachable and
        # -- round-trip cleanly.
        prs = Presentation()
        ep = prs.extended_properties

        # -- the default template carries Application / AppVersion /
        # -- PresentationFormat values seeded by the original authoring
        # -- tool; pin that they read through as plain strings.
        assert isinstance(ep.application, str)
        assert isinstance(ep.app_version, str)
        assert isinstance(ep.presentation_format, str)
        # -- the slide count is exposed as an int; it's automatically
        # -- refreshed on save to match the real <sldIdLst> (see #131),
        # -- but pre-save it reflects whatever the template carries.
        assert isinstance(ep.slide_count, int)

        # -- writes to the sibling string fields also round-trip --
        ep.application = "custom-app"
        ep.app_version = "99.0000"
        ep.presentation_format = "Custom Format"

        prs2 = _round_trip(prs)

        assert prs2.extended_properties.application == "custom-app"
        assert prs2.extended_properties.app_version == "99.0000"
        assert prs2.extended_properties.presentation_format == "Custom Format"

    def it_reflects_the_real_slide_count_in_app_xml_after_save(self):
        # -- cross-feature confirmation: the <Slides> count the #131 fix
        # -- refreshes at save-time is the same <Slides> element the #586
        # -- extended-properties API reads. Adding three slides must
        # -- surface a slide_count of 3 on reopen.
        prs = Presentation()
        for _ in range(3):
            prs.slides.add_slide(prs.slide_layouts[5])

        prs2 = _round_trip(prs)

        assert prs2.extended_properties.slide_count == 3
        assert len(prs2.slides) == 3

    # -- isolation from core_properties ---------------------------------------

    def it_does_not_disturb_core_properties_when_extended_is_written(self):
        # -- app.xml (extended) and core.xml (Dublin-Core) are separate
        # -- package parts. Writing into one must leave the other
        # -- untouched, both in-memory and through save + reopen.
        prs = Presentation()
        prs.core_properties.author = "Original Author"
        prs.core_properties.title = "Original Title"
        prs.core_properties.subject = "Original Subject"
        prs.core_properties.keywords = "alpha,beta,gamma"

        # -- now write only to the extended part --
        prs.extended_properties.company = "Test Corp"
        prs.extended_properties.manager = "Test Manager"

        # -- in-memory: core properties survive the extended write --
        assert prs.core_properties.author == "Original Author"
        assert prs.core_properties.title == "Original Title"
        assert prs.core_properties.subject == "Original Subject"
        assert prs.core_properties.keywords == "alpha,beta,gamma"

        prs2 = _round_trip(prs)

        # -- on reopen: both parts carry their respective values --
        assert prs2.core_properties.author == "Original Author"
        assert prs2.core_properties.title == "Original Title"
        assert prs2.core_properties.subject == "Original Subject"
        assert prs2.core_properties.keywords == "alpha,beta,gamma"
        assert prs2.extended_properties.company == "Test Corp"
        assert prs2.extended_properties.manager == "Test Manager"

    def it_does_not_disturb_extended_properties_when_core_is_written(self):
        # -- the symmetric guarantee: writing core.xml must not clobber
        # -- already-set company/manager in app.xml.
        prs = Presentation()
        prs.extended_properties.company = "Keeper Co"
        prs.extended_properties.manager = "Keeper Mgr"

        # -- now write into core.xml --
        prs.core_properties.author = "New Author"
        prs.core_properties.comments = "All hands"

        prs2 = _round_trip(prs)

        assert prs2.extended_properties.company == "Keeper Co"
        assert prs2.extended_properties.manager == "Keeper Mgr"
        assert prs2.core_properties.author == "New Author"
        assert prs2.core_properties.comments == "All hands"


# -- helpers ------------------------------------------------------------------


def _round_trip(prs: "PresentationAPI") -> "PresentationAPI":
    """Save `prs` to an in-memory buffer and reopen it.

    Returns the reopened :class:`Presentation`. Used to pin that a write
    survives the XML-serialization-and-parse-back cycle, not just the
    in-memory proxy state.
    """
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)
