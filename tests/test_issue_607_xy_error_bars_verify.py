# pyright: reportPrivateUsage=false

"""Regression test for issue #607 — vertical error bars on an XY (scatter) chart.

Issue #607 (https://github.com/scanny/python-pptx/issues/607) asked how to
add vertical (Y-direction) error bars to the data points of an XY scatter
chart. The reporter could see ``c:errBars`` in the schema but there was no
Python-API hook to configure it on a series: pre-fork python-pptx did not
expose ``_BaseSeries.error_bars`` / ``.set_error_bars`` / ``.has_error_bars``
nor the matching :class:`XL_ERROR_BAR_DIRECTION` enum.

Wave 2 #544 (commit ``05072c75``, "feat(chart): #544 add error-bar support
for chart series") shipped the missing API. :class:`_BaseSeries` — and
therefore every concrete series class including :class:`XySeries` and
:class:`BubbleSeries` — now provides:

- ``has_error_bars`` — boolean query
- ``error_bars`` — read/write |ErrorBars| proxy (``None`` when absent; assign
  ``None`` to remove)
- ``set_error_bars(type_, value, include, direction)`` — attach a fresh
  ``c:errBars`` block configured from high-level enum arguments

The ``direction`` keyword accepts :attr:`XL_ERROR_BAR_DIRECTION.X` or
:attr:`XL_ERROR_BAR_DIRECTION.Y` (or ``None`` to omit ``c:errDir``, which is
fine on category-axis charts since PowerPoint then infers Y). For XY scatter
and bubble charts — which have *two* value axes rather than a category axis
plus a value axis — the direction is not implicit: specifying
:attr:`XL_ERROR_BAR_DIRECTION.Y` is the canonical way to get the vertical
whiskers shown in the issue's screenshot.

This suite pins the #607 resolution so the issue can be closed. Tests
exercise the real public API against a fresh ``Presentation()`` — not mocks,
not the raw oxml layer — and include a save/reload round-trip to confirm
the ``c:errBars`` subtree survives serialization.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.data import BubbleChartData, XyChartData
from pptx.chart.series import ErrorBars, XySeries
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_ERROR_BAR_DIRECTION,
    XL_ERROR_BAR_INCLUDE,
    XL_ERROR_BAR_TYPE,
)
from pptx.oxml.ns import qn
from pptx.util import Inches

# -- helpers --------------------------------------------------------------


def _build_scatter_chart(chart_type=XL_CHART_TYPE.XY_SCATTER):
    """Return ``(prs, chart)`` for a fresh single-series XY scatter chart."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_data = XyChartData()
    series = chart_data.add_series("Series 1")
    series.add_data_point(1.0, 2.7)
    series.add_data_point(2.0, 3.2)
    series.add_data_point(3.0, 4.1)
    series.add_data_point(4.0, 5.8)
    graphic_frame = slide.shapes.add_chart(
        chart_type,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    )
    return prs, graphic_frame.chart


def _build_bubble_chart():
    """Return ``(prs, chart)`` for a fresh single-series bubble chart."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_data = BubbleChartData()
    series = chart_data.add_series("Series 1")
    series.add_data_point(1.0, 2.0, 10)
    series.add_data_point(2.0, 4.0, 15)
    series.add_data_point(3.0, 6.0, 20)
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.BUBBLE,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    )
    return prs, graphic_frame.chart


_ALL_SCATTER_TYPES = [
    XL_CHART_TYPE.XY_SCATTER,
    XL_CHART_TYPE.XY_SCATTER_LINES,
    XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
    XL_CHART_TYPE.XY_SCATTER_SMOOTH,
    XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS,
]


# -- the verification suite -----------------------------------------------


class DescribeIssue607XyErrorBarsVerify(object):
    """#607 verify-and-close: vertical error bars on an XY chart work end-to-end.

    Pins that :meth:`_BaseSeries.set_error_bars` accepts
    :attr:`XL_ERROR_BAR_DIRECTION.Y` on XY-scatter and bubble series, emits
    a conforming ``c:errBars`` subtree with ``c:errDir val="y"``, and
    survives a ``Presentation.save`` + reopen round-trip intact.
    """

    # -- starting state -----------------------------------------------------

    def it_starts_without_error_bars_on_a_fresh_xy_series(self):
        """A brand-new XY scatter series has no ``c:errBars`` child."""
        _, chart = _build_scatter_chart()
        series = chart.series[0]

        assert isinstance(series, XySeries)
        assert series.has_error_bars is False
        assert series.error_bars is None

    # -- authoring path: the #607 ask ---------------------------------------

    @pytest.mark.parametrize("chart_type", _ALL_SCATTER_TYPES)
    def it_attaches_vertical_error_bars_to_every_xy_scatter_variant(self, chart_type):
        """``set_error_bars(direction=Y)`` works on all five XY scatter variants."""
        _, chart = _build_scatter_chart(chart_type)
        series = chart.series[0]

        # -- the #607 ask: add vertical error bars --
        error_bars = series.set_error_bars(
            type_=XL_ERROR_BAR_TYPE.FIXED_VALUE,
            value=0.5,
            include=XL_ERROR_BAR_INCLUDE.BOTH,
            direction=XL_ERROR_BAR_DIRECTION.Y,
        )

        assert isinstance(error_bars, ErrorBars)
        assert series.has_error_bars is True
        assert series.error_bars.direction == XL_ERROR_BAR_DIRECTION.Y
        assert series.error_bars.type == XL_ERROR_BAR_TYPE.FIXED_VALUE
        assert series.error_bars.value == 0.5
        assert series.error_bars.include == XL_ERROR_BAR_INCLUDE.BOTH

    def it_emits_a_conforming_errBars_subtree_with_errDir_y(self):
        """The emitted ``c:errBars`` subtree matches the OOXML schema shape."""
        _, chart = _build_scatter_chart()
        series = chart.series[0]

        series.set_error_bars(
            type_=XL_ERROR_BAR_TYPE.FIXED_VALUE,
            value=0.5,
            direction=XL_ERROR_BAR_DIRECTION.Y,
        )

        errBars = series._element.errBars
        assert errBars is not None
        # -- the crucial bit for #607: c:errDir carries val="y" --
        errDir = errBars.find(qn("c:errDir"))
        assert errDir is not None
        assert errDir.get("val") == "y"
        # -- other expected children --
        assert errBars.find(qn("c:errBarType")).get("val") == "both"
        assert errBars.find(qn("c:errValType")).get("val") == "fixedVal"
        assert errBars.find(qn("c:val")).get("val") == "0.5"

    def it_supports_horizontal_error_bars_with_direction_X(self):
        """XY charts also allow ``direction=X`` for horizontal whiskers."""
        _, chart = _build_scatter_chart()
        series = chart.series[0]

        series.set_error_bars(
            type_=XL_ERROR_BAR_TYPE.PERCENT,
            value=5.0,
            direction=XL_ERROR_BAR_DIRECTION.X,
        )

        assert series.error_bars.direction == XL_ERROR_BAR_DIRECTION.X
        errDir = series._element.errBars.find(qn("c:errDir"))
        assert errDir.get("val") == "x"

    def it_supports_percent_and_stdev_magnitude_types(self):
        """Magnitude computation enums other than FIXED_VALUE also round-trip."""
        _, chart = _build_scatter_chart()
        series = chart.series[0]

        series.set_error_bars(
            type_=XL_ERROR_BAR_TYPE.STDEV,
            value=2.0,
            direction=XL_ERROR_BAR_DIRECTION.Y,
        )

        assert series.error_bars.type == XL_ERROR_BAR_TYPE.STDEV
        assert series.error_bars.value == 2.0
        assert series._element.errBars.find(qn("c:errValType")).get("val") == "stdDev"

    def it_supports_plus_values_and_minus_values_include_modes(self):
        """``include=PLUS_VALUES`` / ``MINUS_VALUES`` draw only one side."""
        _, chart = _build_scatter_chart()
        series = chart.series[0]

        series.set_error_bars(
            type_=XL_ERROR_BAR_TYPE.FIXED_VALUE,
            value=1.0,
            include=XL_ERROR_BAR_INCLUDE.PLUS_VALUES,
            direction=XL_ERROR_BAR_DIRECTION.Y,
        )

        assert series.error_bars.include == XL_ERROR_BAR_INCLUDE.PLUS_VALUES
        assert series._element.errBars.find(qn("c:errBarType")).get("val") == "plus"

    # -- round-trip --------------------------------------------------------

    def it_round_trips_vertical_error_bars_through_save_reload(self):
        """The ``c:errBars`` subtree survives ``Presentation.save`` + reopen."""
        prs, chart = _build_scatter_chart()
        series = chart.series[0]
        series.set_error_bars(
            type_=XL_ERROR_BAR_TYPE.PERCENT,
            value=7.5,
            include=XL_ERROR_BAR_INCLUDE.BOTH,
            direction=XL_ERROR_BAR_DIRECTION.Y,
        )

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        reloaded_chart = None
        for shp in reloaded.slides[0].shapes:
            if shp.has_chart:
                reloaded_chart = shp.chart
                break
        assert reloaded_chart is not None
        reloaded_series = reloaded_chart.series[0]

        assert reloaded_series.has_error_bars is True
        error_bars = reloaded_series.error_bars
        assert error_bars.direction == XL_ERROR_BAR_DIRECTION.Y
        assert error_bars.type == XL_ERROR_BAR_TYPE.PERCENT
        assert error_bars.value == 7.5
        assert error_bars.include == XL_ERROR_BAR_INCLUDE.BOTH

    # -- replace / remove --------------------------------------------------

    def it_replaces_existing_error_bars_when_set_error_bars_called_twice(self):
        """Calling ``set_error_bars`` twice overwrites rather than stacking."""
        _, chart = _build_scatter_chart()
        series = chart.series[0]
        series.set_error_bars(
            type_=XL_ERROR_BAR_TYPE.FIXED_VALUE,
            value=1.0,
            direction=XL_ERROR_BAR_DIRECTION.Y,
        )

        series.set_error_bars(
            type_=XL_ERROR_BAR_TYPE.PERCENT,
            value=10.0,
            direction=XL_ERROR_BAR_DIRECTION.Y,
        )

        assert len(series._element.xpath("c:errBars")) == 1
        assert series.error_bars.type == XL_ERROR_BAR_TYPE.PERCENT
        assert series.error_bars.value == 10.0

    def it_removes_error_bars_when_assigned_None(self):
        """Assigning ``None`` to ``series.error_bars`` drops the ``c:errBars`` child."""
        _, chart = _build_scatter_chart()
        series = chart.series[0]
        series.set_error_bars(
            type_=XL_ERROR_BAR_TYPE.FIXED_VALUE,
            value=0.5,
            direction=XL_ERROR_BAR_DIRECTION.Y,
        )
        assert series.has_error_bars is True

        series.error_bars = None

        assert series.has_error_bars is False
        assert series.error_bars is None

    # -- bubble charts -----------------------------------------------------

    def it_also_works_on_bubble_charts(self):
        """Bubble charts — which also have two value axes — support Y error bars."""
        prs, chart = _build_bubble_chart()
        series = chart.series[0]

        series.set_error_bars(
            type_=XL_ERROR_BAR_TYPE.STDEV,
            value=1.0,
            direction=XL_ERROR_BAR_DIRECTION.Y,
        )

        assert series.has_error_bars is True
        assert series.error_bars.direction == XL_ERROR_BAR_DIRECTION.Y

        # -- and it round-trips --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)
        reloaded_chart = next(shp.chart for shp in reloaded.slides[0].shapes if shp.has_chart)
        reloaded_series = reloaded_chart.series[0]
        assert reloaded_series.has_error_bars is True
        assert reloaded_series.error_bars.direction == XL_ERROR_BAR_DIRECTION.Y
        assert reloaded_series.error_bars.type == XL_ERROR_BAR_TYPE.STDEV
