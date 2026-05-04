# pyright: reportPrivateUsage=false

"""Verify-close regression test for issue #613 — ``add_table(height=)`` semantics.

Issue #613 (https://github.com/scanny/python-pptx/issues/613) reported that
``SlideShapes.add_table(rows, cols, left, top, width, height)`` looked as if
``height`` was the *per-row* height rather than the *total* table height.
The misconception is understandable — PowerPoint's rendered table may later
grow taller than the authored extent — but the argument has always been the
total (authored) table height, and the row heights are derived by dividing
that total by ``rows`` (with the last row absorbing any integer-division
remainder).

This suite pins that contract end-to-end against the public API so a future
refactor of the ``CT_Table.new_tbl`` / ``SlideShapes.add_table`` pipeline
can't silently flip the meaning of ``height`` under callers. It also
documents the sum-equality invariant for widths and covers the 5-row / 3-col
shape the reporter described.
"""

from __future__ import annotations

import io
from typing import cast

from pptx import Presentation
from pptx.shapes.graphfrm import GraphicFrame
from pptx.util import Emu, Inches


class DescribeIssue613AddTableHeightVerify:
    """Verify-close regression suite for ``SlideShapes.add_table`` sizing."""

    def it_interprets_height_as_the_total_table_height(self):
        # -- the reporter's canonical case: 5 rows, 3 cols, 5-inch total
        # -- height. Total-height semantics mean each row is ~1 inch, and
        # -- shape.height equals the height passed in.
        _, slide = _fresh_slide()
        left, top = Inches(1), Inches(1)
        width, height = Inches(6), Inches(5)
        rows, cols = 5, 3

        gf = slide.shapes.add_table(rows, cols, left, top, width, height)

        assert gf.height == int(height)
        assert gf.width == int(width)

    def it_distributes_height_evenly_across_rows(self):
        # -- each row height is approximately height / rows. The last row
        # -- absorbs any remainder so sum(row_heights) == height.
        _, slide = _fresh_slide()
        gf = slide.shapes.add_table(
            rows=5, cols=3, left=Inches(1), top=Inches(1), width=Inches(6), height=Inches(5)
        )
        table = gf.table

        expected_row_height = int(Inches(5)) // 5
        # -- first four rows are exactly height // rows; the fifth absorbs the
        # -- integer-division remainder (here zero because 5 inches divides
        # -- evenly into 5 rows).
        assert [r.height for r in table.rows][:4] == [expected_row_height] * 4
        # -- sum of row heights equals the authored table height exactly --
        assert sum(r.height for r in table.rows) == int(Inches(5))

    def it_distributes_width_evenly_across_columns(self):
        # -- symmetric invariant for columns: sum of column widths equals
        # -- the authored table width, and all but the last are
        # -- width // cols.
        _, slide = _fresh_slide()
        gf = slide.shapes.add_table(
            rows=2, cols=4, left=Inches(1), top=Inches(1), width=Inches(8), height=Inches(2)
        )
        table = gf.table

        col_widths = [c.width for c in table.columns]
        assert col_widths[:3] == [int(Inches(8)) // 4] * 3
        assert sum(col_widths) == int(Inches(8))

    def it_preserves_the_integer_division_remainder_in_the_last_row(self):
        # -- pick a size that does NOT divide evenly so the remainder shows
        # -- up. Seven rows into 914 400 EMU (1 inch): 914 400 // 7 = 130 628
        # -- with remainder 4. The last row carries the remainder so the
        # -- sum still matches.
        _, slide = _fresh_slide()
        gf = slide.shapes.add_table(
            rows=7, cols=2, left=Inches(1), top=Inches(1), width=Inches(4), height=Inches(1)
        )
        table = gf.table

        heights = [r.height for r in table.rows]
        # -- all but the last row equal the floor-division value --
        assert heights[:-1] == [Inches(1) // 7] * 6
        # -- the last row is slightly taller so the total is preserved --
        assert heights[-1] == Inches(1) - 6 * (Inches(1) // 7)
        assert sum(heights) == int(Inches(1))

    def it_survives_a_save_reopen_roundtrip(self):
        # -- the total-height contract is a property of the authored XML,
        # -- so it must survive a save + reopen with no drift.
        prs, slide = _fresh_slide()
        slide.shapes.add_table(
            rows=5, cols=3, left=Inches(1), top=Inches(1), width=Inches(6), height=Inches(5)
        )

        reloaded = _roundtrip(prs)
        gf = cast(GraphicFrame, next(s for s in reloaded.slides[0].shapes if s.has_table))

        assert gf.height == int(Inches(5))
        assert gf.width == int(Inches(6))
        assert sum(r.height for r in gf.table.rows) == int(Inches(5))

    def it_matches_the_graphic_frame_ext_in_the_XML(self):
        # -- pin the underlying XML surface too: p:graphicFrame/p:xfrm/a:ext
        # -- must carry cy == height, not height * rows. This is the exact
        # -- attribute PowerPoint reads for the shape's on-screen extent.
        _, slide = _fresh_slide()
        gf = slide.shapes.add_table(
            rows=5, cols=3, left=Inches(1), top=Inches(1), width=Inches(6), height=Inches(5)
        )

        ext = gf._element.xpath("./p:xfrm/a:ext")[0]
        assert int(ext.get("cx")) == int(Inches(6))
        assert int(ext.get("cy")) == int(Inches(5))

    def it_accepts_emu_valued_heights(self):
        # -- Emu inputs go through unchanged; confirm the contract holds
        # -- for raw EMU (not just Inches).
        _, slide = _fresh_slide()
        height_emu = Emu(3_000_000)
        gf = slide.shapes.add_table(
            rows=3, cols=2, left=Emu(0), top=Emu(0), width=Emu(4_000_000), height=height_emu
        )

        assert gf.height == int(height_emu)
        assert sum(r.height for r in gf.table.rows) == int(height_emu)


# -- helpers --------------------------------------------------------------


def _fresh_slide():
    """Return ``(prs, slide)`` — a new ``Presentation`` with one blank slide."""
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    return prs, slide


def _roundtrip(prs):
    """Serialize `prs` to a ``BytesIO`` and reopen it."""
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)
