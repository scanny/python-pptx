# pyright: reportPrivateUsage=false

"""Regression test for issue #620 — dynamically replicating a slide.

Issue #620 (https://github.com/scanny/python-pptx/issues/620) asks for a
way to duplicate an existing slide (with placeholders, a picture, a chart,
and a table) and then mutate content on the copy without affecting the
source. The reporter's workflow is a staple of template-driven reporting
pipelines: load a deck, find a "golden" slide, copy it N times, fill in
the per-row specifics on each copy.

This verify suite pins the resolution: Wave 1 ``#132`` shipped
:meth:`Slides.duplicate`, which returns a newly added slide that is a
deep-copy of the source's shape tree while sharing image/chart/media
parts by package reference (see the ``Slides.duplicate`` docstring in
``src/pptx/slide.py`` and the ``feat: #132`` entry in ``HISTORY.rst``).
The duplicate can be mutated independently of the source; edits to the
copy's title (or any other shape) do *not* leak back to the source. A
``save`` + reopen round-trip preserves the independent shape trees and
the relationship graph (image, chart, table) on both the source and the
duplicate.

Cross-references: Wave 1 ``#132`` ships the foundational
``Slides.duplicate`` method verified here. See also:

  * #696 (move a slide from one .pptx to another) — companion
    cross-presentation flavour via
    :meth:`Slides.add_slide_from_external`.
  * #403 (merge the slides of one presentation into another) — whole-deck
    flavour via :meth:`Presentation.merge`.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

import pytest

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

if TYPE_CHECKING:
    from pptx.presentation import Presentation as PresentationClass
    from pptx.slide import Slide


@pytest.fixture
def _restore_part_factory():
    """Guard against ``PartFactory.part_type_for`` pollution from other tests.

    Mirrors the identical guard used by ``tests/test_issue_175_*.py``,
    ``tests/test_issue_403_*.py``, and ``tests/test_issue_696_*.py``;
    earlier test modules overwrite slide-part registrations with mocks and
    do not restore them, which breaks ``Presentation(stream)`` reopen in
    save/reopen round-trip assertions.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.embeddedpackage import EmbeddedXlsxPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
        CT.SML_SHEET: EmbeddedXlsxPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


# -- markers used to distinguish source vs duplicate content after edits --
SOURCE_TITLE = "Golden Source Title"
DUPLICATE_TITLE = "Copy-of Title (edited)"


def _absjoin(*args: str) -> str:
    from .unitutil.file import absjoin, test_file_dir

    return absjoin(test_file_dir, *args)


def _build_rich_source_slide(prs: PresentationClass) -> Slide:
    """Return a new slide populated with all four #620 content kinds.

    The reporter's scenario calls out "text placeholders, a picture, a
    chart, and a table" — we build a single slide that has all of them,
    plus a title placeholder carrying ``SOURCE_TITLE``.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[5])

    # -- (a) title placeholder carrying the source marker --
    title = slide.shapes.title
    assert title is not None  # layout 5 is "Title Only"
    title.text_frame.text = SOURCE_TITLE

    # -- (b) an extra text placeholder (free-standing textbox) --
    tb = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(6), Inches(1))
    tb.text_frame.paragraphs[0].text = "Body text on source"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)

    # -- (c) a picture --
    slide.shapes.add_picture(_absjoin("python-icon.jpeg"), Inches(1), Inches(1.5), height=Inches(1))

    # -- (d) a chart (column-clustered with one series) --
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3"]
    cd.add_series("Revenue", (1.0, 2.0, 3.0))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(3),
        Inches(1.5),
        Inches(4),
        Inches(3),
        cd,
    )

    # -- (e) a 2x2 table --
    slide.shapes.add_table(
        rows=2, cols=2, left=Inches(1), top=Inches(5), width=Inches(4), height=Inches(1)
    )

    return slide


def _slide_title_text(slide: Slide) -> str:
    """Return the title-placeholder text for `slide`, or ``""`` if none."""
    title = slide.shapes.title
    if title is None:
        return ""
    return title.text_frame.text


def _shape_kinds(slide: Slide) -> dict[str, int]:
    """Histogram the shape kinds on `slide` keyed by diagnostic name."""
    counts = {"picture": 0, "chart": 0, "table": 0, "placeholder": 0, "textbox": 0}
    for sh in slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            counts["picture"] += 1
        elif getattr(sh, "has_chart", False):
            counts["chart"] += 1
        elif getattr(sh, "has_table", False):
            counts["table"] += 1
        elif sh.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER:
            counts["placeholder"] += 1
        elif sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            counts["textbox"] += 1
    return counts


class DescribeIssue620DuplicateAndEdit(object):
    """#620 — close as resolved by Wave 1 #132.

    Pins the reporter's "replicate a slide dynamically, then edit the
    copy" workflow end-to-end on a slide that carries all four content
    kinds the reporter called out: text placeholders, a picture, a chart,
    and a table.
    """

    # -- (1) the headline scenario: duplicate carries all four content kinds --

    def it_duplicates_a_slide_with_placeholder_picture_chart_and_table(self, _restore_part_factory):
        prs = Presentation()
        source = _build_rich_source_slide(prs)
        pre_count = len(prs.slides)

        duplicate = prs.slides.duplicate(source)

        # -- the duplicate is a distinct, newly appended slide --
        assert duplicate is not source
        assert duplicate in list(prs.slides)
        assert len(prs.slides) == pre_count + 1

        # -- the duplicate carries the same shape kinds as the source --
        src_kinds = _shape_kinds(source)
        dup_kinds = _shape_kinds(duplicate)
        assert (
            dup_kinds == src_kinds
        ), f"duplicate shape histogram {dup_kinds} != source {src_kinds}"
        # -- each of the four reporter-named kinds is present on the duplicate --
        assert dup_kinds["picture"] >= 1
        assert dup_kinds["chart"] >= 1
        assert dup_kinds["table"] >= 1
        # -- source title text also appears on the duplicate (pre-edit) --
        assert _slide_title_text(duplicate) == SOURCE_TITLE

    # -- (2) editing the duplicate does not mutate the source ---------------

    def it_does_not_leak_edits_on_the_duplicate_back_to_the_source(self, _restore_part_factory):
        """The #620 reporter's precise ask: edit content on the copy
        without disturbing the original. Mutate the duplicate's title and
        assert the source still shows ``SOURCE_TITLE``.
        """
        prs = Presentation()
        source = _build_rich_source_slide(prs)
        duplicate = prs.slides.duplicate(source)

        # -- edit the duplicate's title --
        dup_title = duplicate.shapes.title
        assert dup_title is not None
        dup_title.text_frame.text = DUPLICATE_TITLE

        assert _slide_title_text(duplicate) == DUPLICATE_TITLE
        # -- CRITICAL: source is untouched --
        assert _slide_title_text(source) == SOURCE_TITLE

    # -- (3) save + reopen round-trip preserves both slides independently ----

    def it_round_trips_the_edited_duplicate_through_save_and_reopen(self, _restore_part_factory):
        """The duplicate with an edited title survives save + reopen, and
        the source slide remains distinct in the reopened deck."""
        prs = Presentation()
        source = _build_rich_source_slide(prs)
        duplicate = prs.slides.duplicate(source)
        dup_title = duplicate.shapes.title
        assert dup_title is not None
        dup_title.text_frame.text = DUPLICATE_TITLE

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reopened = Presentation(buf)

        # -- both slides survive --
        assert len(reopened.slides) == 2

        # -- identify source vs duplicate by title text in the reopened deck --
        reopened_titles = [_slide_title_text(sl) for sl in reopened.slides]
        assert SOURCE_TITLE in reopened_titles
        assert DUPLICATE_TITLE in reopened_titles

        # -- each slide in the reopened deck still carries picture, chart,
        # -- and table shapes --
        for sl in reopened.slides:
            kinds = _shape_kinds(sl)
            assert kinds["picture"] >= 1, f"no picture on reopened slide {kinds}"
            assert kinds["chart"] >= 1, f"no chart on reopened slide {kinds}"
            assert kinds["table"] >= 1, f"no table on reopened slide {kinds}"

    # -- (4) source's own content is unchanged after the duplicate ----------

    def it_does_not_mutate_the_source_slide_shape_tree(self, _restore_part_factory):
        """Pin that ``Slides.duplicate`` is non-destructive on the source.

        The reporter implicitly relies on the source remaining a reusable
        "template" — duplicating it multiple times in a row should yield
        identical copies from an identical source.
        """
        prs = Presentation()
        source = _build_rich_source_slide(prs)
        src_kinds_before = _shape_kinds(source)
        src_title_before = _slide_title_text(source)
        src_shape_count_before = len(source.shapes)

        # -- duplicate twice in a row --
        dup1 = prs.slides.duplicate(source)
        dup2 = prs.slides.duplicate(source)

        assert dup1 is not dup2
        assert _shape_kinds(source) == src_kinds_before
        assert _slide_title_text(source) == src_title_before
        assert len(source.shapes) == src_shape_count_before
        # -- both duplicates have the same shape histogram as the source --
        assert _shape_kinds(dup1) == src_kinds_before
        assert _shape_kinds(dup2) == src_kinds_before
