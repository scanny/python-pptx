"""Regression test for issue #621 — shapes wrapped in ``mc:AlternateContent``
are ignored by ``slide.shapes``.

Issue #621 (https://github.com/scanny/python-pptx/issues/621) reports that
PowerPoint wraps shapes which require an Office 2010+ extension namespace
(modern charts — ``cx:chart`` sunburst/treemap/funnel/etc., modern comments,
math-equation-bearing shapes, and so on) inside an
``<mc:AlternateContent>``/``<mc:Choice>`` block with an ``<mc:Fallback>``
sibling holding a pre-2010 representation. Before the fix, ``slide.shapes``
iterated only direct ``p:spTree`` children whose tag matched a known shape
tag, so any wrapped shape silently vanished from the public iteration — the
shape count was wrong, ``shape.name`` lookups failed, and every shape-level
walker (``descendants()``, ``iter_leaf_shapes()``, ``shapes.clear()``,
``ShapeTree`` indexing) skipped the wrapped content.

The library now descends into ``mc:AlternateContent``:

* The first ``mc:Choice`` whose subtree contains any recognized shape tag
  wins — its shapes are yielded as if they were direct ``p:spTree`` children.
* If no ``mc:Choice`` yields a shape (e.g. because the choice requires a
  namespace extension whose shapes this library hasn't modelled, or the
  block has no ``mc:Choice`` at all), the ``mc:Fallback`` subtree is walked
  instead so that *something* is surfaced.
* Branches that are not yielded from are preserved as-is on the element
  tree so the wrapper round-trips on save.

This module pins that behavior at the public-API level (``slide.shapes``)
with fixtures modeled on what PowerPoint actually writes.
"""

from __future__ import annotations

from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.shapes.graphfrm import GraphicFrame
from pptx.slide import Slide
from pptx.spec import GRAPHIC_DATA_URI_CHARTEX


def _slide_with_spTree_xml(sp_tree_children_xml: str) -> Slide:
    """Build a minimal ``Slide`` proxy whose ``p:spTree`` contains the given children.

    Only what is needed to exercise ``slide.shapes`` iteration. The slide proxy is
    instantiated with ``part=None`` — the shape factory never needs the part for
    tag-only dispatch in this suite.
    """
    xml = (
        "<p:sld %s>\n"
        "  <p:cSld>\n"
        "    <p:spTree>\n"
        "      <p:nvGrpSpPr>\n"
        '        <p:cNvPr id="1" name=""/>\n'
        "        <p:cNvGrpSpPr/>\n"
        "        <p:nvPr/>\n"
        "      </p:nvGrpSpPr>\n"
        "      <p:grpSpPr/>\n"
        "      %s\n"
        "    </p:spTree>\n"
        "  </p:cSld>\n"
        "</p:sld>"
    ) % (nsdecls("p", "a", "r", "mc"), sp_tree_children_xml)
    sld = parse_xml(xml)
    return Slide(sld, None)  # type: ignore[arg-type]


class DescribeIssue621:
    """Regression suite for ``mc:AlternateContent`` visibility in ``slide.shapes``."""

    def it_surfaces_a_chartex_graphicFrame_wrapped_in_Choice(self):
        """Sunburst/treemap/etc. charts are written wrapped so older PPT falls back."""
        wrapper = (
            "<mc:AlternateContent>\n"
            '  <mc:Choice xmlns:cx="http://schemas.microsoft.com/office/drawing/2014'
            '/chartex" Requires="cx1">\n'
            "    <p:graphicFrame>\n"
            "      <p:nvGraphicFramePr>\n"
            '        <p:cNvPr id="3" name="Sunburst 1"/>\n'
            "        <p:cNvGraphicFramePr/>\n"
            "        <p:nvPr/>\n"
            "      </p:nvGraphicFramePr>\n"
            "      <p:xfrm>\n"
            '        <a:off x="0" y="0"/><a:ext cx="1000" cy="1000"/>\n'
            "      </p:xfrm>\n"
            "      <a:graphic>\n"
            '        <a:graphicData uri="%s"/>\n'
            "      </a:graphic>\n"
            "    </p:graphicFrame>\n"
            "  </mc:Choice>\n"
            "  <mc:Fallback>\n"
            "    <p:sp>\n"
            "      <p:nvSpPr>\n"
            '        <p:cNvPr id="4" name="Fallback Rect"/>\n'
            "        <p:cNvSpPr/>\n"
            "        <p:nvPr/>\n"
            "      </p:nvSpPr>\n"
            "      <p:spPr/>\n"
            "    </p:sp>\n"
            "  </mc:Fallback>\n"
            "</mc:AlternateContent>"
        ) % GRAPHIC_DATA_URI_CHARTEX
        slide = _slide_with_spTree_xml(wrapper)

        shapes = list(slide.shapes)

        assert len(shapes) == 1
        shape = shapes[0]
        assert isinstance(shape, GraphicFrame)
        assert shape.name == "Sunburst 1"
        assert shape.has_chart is True
        assert shape.has_chartex is True

    def it_surfaces_Choice_shapes_interleaved_with_plain_shapes(self):
        """A slide mixing plain shapes and an ``mc:AlternateContent`` yields them all."""
        mixed = (
            "<p:sp>\n"
            "  <p:nvSpPr>\n"
            '    <p:cNvPr id="2" name="Rect 1"/>\n'
            "    <p:cNvSpPr/>\n"
            "    <p:nvPr/>\n"
            "  </p:nvSpPr>\n"
            "  <p:spPr/>\n"
            "</p:sp>\n"
            "<mc:AlternateContent>\n"
            '  <mc:Choice Requires="a14">\n'
            "    <p:sp>\n"
            "      <p:nvSpPr>\n"
            '        <p:cNvPr id="3" name="EquationShape"/>\n'
            "        <p:cNvSpPr/>\n"
            "        <p:nvPr/>\n"
            "      </p:nvSpPr>\n"
            "      <p:spPr/>\n"
            "    </p:sp>\n"
            "  </mc:Choice>\n"
            "  <mc:Fallback>\n"
            "    <p:sp>\n"
            "      <p:nvSpPr>\n"
            '        <p:cNvPr id="4" name="FallbackShape"/>\n'
            "        <p:cNvSpPr/>\n"
            "        <p:nvPr/>\n"
            "      </p:nvSpPr>\n"
            "      <p:spPr/>\n"
            "    </p:sp>\n"
            "  </mc:Fallback>\n"
            "</mc:AlternateContent>\n"
            "<p:sp>\n"
            "  <p:nvSpPr>\n"
            '    <p:cNvPr id="5" name="Rect 2"/>\n'
            "    <p:cNvSpPr/>\n"
            "    <p:nvPr/>\n"
            "  </p:nvSpPr>\n"
            "  <p:spPr/>\n"
            "</p:sp>"
        )
        slide = _slide_with_spTree_xml(mixed)

        shapes = list(slide.shapes)
        names = [s.name for s in shapes]

        # -- all three surface in document order; the wrapped Choice shape is in the middle --
        assert len(shapes) == 3
        assert names == ["Rect 1", "EquationShape", "Rect 2"]

    def it_falls_back_to_Fallback_when_Choice_has_no_recognizable_shape(self):
        """When ``mc:Choice`` has only extension-namespace content, ``mc:Fallback`` wins."""
        wrapper = (
            "<mc:AlternateContent>\n"
            '  <mc:Choice xmlns:foo="http://foo" Requires="foo">\n'
            "    <foo:unknown/>\n"
            "  </mc:Choice>\n"
            "  <mc:Fallback>\n"
            "    <p:sp>\n"
            "      <p:nvSpPr>\n"
            '        <p:cNvPr id="7" name="FallbackOnly"/>\n'
            "        <p:cNvSpPr/>\n"
            "        <p:nvPr/>\n"
            "      </p:nvSpPr>\n"
            "      <p:spPr/>\n"
            "    </p:sp>\n"
            "  </mc:Fallback>\n"
            "</mc:AlternateContent>"
        )
        slide = _slide_with_spTree_xml(wrapper)

        shapes = list(slide.shapes)

        assert len(shapes) == 1
        assert shapes[0].name == "FallbackOnly"

    def it_surfaces_Fallback_when_AlternateContent_has_no_Choice(self):
        """``mc:AlternateContent`` may contain only an ``mc:Fallback`` — still surface it."""
        wrapper = (
            "<mc:AlternateContent>\n"
            "  <mc:Fallback>\n"
            "    <p:sp>\n"
            "      <p:nvSpPr>\n"
            '        <p:cNvPr id="8" name="NoChoiceShape"/>\n'
            "        <p:cNvSpPr/>\n"
            "        <p:nvPr/>\n"
            "      </p:nvSpPr>\n"
            "      <p:spPr/>\n"
            "    </p:sp>\n"
            "  </mc:Fallback>\n"
            "</mc:AlternateContent>"
        )
        slide = _slide_with_spTree_xml(wrapper)

        shapes = list(slide.shapes)

        assert len(shapes) == 1
        assert shapes[0].name == "NoChoiceShape"

    def it_preserves_the_unused_branches_on_iteration(self):
        """Iterating ``slide.shapes`` is side-effect-free — unused branches stay on the tree."""
        wrapper = (
            "<mc:AlternateContent>\n"
            '  <mc:Choice Requires="a14">\n'
            "    <p:sp>\n"
            "      <p:nvSpPr>\n"
            '        <p:cNvPr id="3" name="ChoiceShape"/>\n'
            "        <p:cNvSpPr/>\n"
            "        <p:nvPr/>\n"
            "      </p:nvSpPr>\n"
            "      <p:spPr/>\n"
            "    </p:sp>\n"
            "  </mc:Choice>\n"
            "  <mc:Fallback>\n"
            "    <p:sp>\n"
            "      <p:nvSpPr>\n"
            '        <p:cNvPr id="4" name="FallbackShape"/>\n'
            "        <p:cNvSpPr/>\n"
            "        <p:nvPr/>\n"
            "      </p:nvSpPr>\n"
            "      <p:spPr/>\n"
            "    </p:sp>\n"
            "  </mc:Fallback>\n"
            "</mc:AlternateContent>"
        )
        slide = _slide_with_spTree_xml(wrapper)
        sld = slide._element  # type: ignore[attr-defined]

        # -- walk the shapes a couple of times, then check the Fallback is still there --
        list(slide.shapes)
        list(slide.shapes)

        fallbacks = sld.xpath(".//mc:Fallback")
        assert len(fallbacks) == 1
        fallback_sps = list(fallbacks[0].iterchildren(qn("p:sp")))
        assert len(fallback_sps) == 1

    def it_counts_wrapped_shapes_correctly_via_len(self):
        """``len(slide.shapes)`` matches the visible-to-PowerPoint shape count."""
        wrapper = (
            "<p:sp>\n"
            "  <p:nvSpPr>\n"
            '    <p:cNvPr id="2" name="Rect 1"/>\n'
            "    <p:cNvSpPr/>\n"
            "    <p:nvPr/>\n"
            "  </p:nvSpPr>\n"
            "  <p:spPr/>\n"
            "</p:sp>\n"
            "<mc:AlternateContent>\n"
            '  <mc:Choice Requires="a14">\n'
            "    <p:sp>\n"
            "      <p:nvSpPr>\n"
            '        <p:cNvPr id="3" name="Equation"/>\n'
            "        <p:cNvSpPr/>\n"
            "        <p:nvPr/>\n"
            "      </p:nvSpPr>\n"
            "      <p:spPr/>\n"
            "    </p:sp>\n"
            "  </mc:Choice>\n"
            "</mc:AlternateContent>"
        )
        slide = _slide_with_spTree_xml(wrapper)

        assert len(slide.shapes) == 2
