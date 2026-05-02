# pyright: reportPrivateUsage=false

"""Regression test for issue #62 — read/write transparency on fills.

Issue #62 (https://github.com/scanny/python-pptx/issues/62) asks for the
ability to read and write the alpha (transparency / opacity) channel of
a color. In OOXML the value is expressed as an ``<a:alpha val="N"/>``
child element that may appear under any color-choice element
(``a:srgbClr``, ``a:schemeClr``, ``a:sysClr``, ``a:prstClr``,
``a:hslClr``, ``a:scrgbClr``); ``N`` is in thousandths of a percent
(``0``–``100000`` → ``0%``–``100%``) and represents opacity —
``100000`` is fully opaque, ``0`` is fully transparent.

This test wires the new ``ColorFormat.alpha`` property end-to-end: it
sets alpha on a solid fill's foreground color, round-trips through
``Presentation.save`` + reopen, and confirms the value survives.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of `PartFactory.part_type_for`.

    Mirrors the guard used by ``tests/test_issue_234_blip_fill.py``; needed
    because earlier test modules (notably ``tests/opc/test_package.py``)
    overwrite slide-part registrations with mocks and do not restore them.
    Round-trip tests that call ``Presentation.save`` + ``Presentation()``
    reopen depend on the real registrations being in place.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.comments import CommentAuthorsPart, CommentsPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_COMMENTS: CommentsPart,
        CT.PML_COMMENT_AUTHORS: CommentAuthorsPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue62FillAlpha(object):
    """Round-trip test for issue #62 transparency on a solid fill."""

    @pytest.mark.parametrize("alpha", [0.0, 0.25, 0.5, 0.75, 1.0])
    def it_round_trips_alpha_through_save_and_reopen(self, alpha, _restore_part_factory):
        # -- arrange: build a presentation with a red rectangle that has
        # -- the requested opacity applied to its fill color.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xFF, 0x00, 0x00)

        # -- act: set alpha, save, reload --
        shape.fill.fore_color.alpha = alpha
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        # -- assert: alpha survives the round trip --
        reloaded_shape = prs2.slides[0].shapes[1]
        assert reloaded_shape.fill.fore_color.alpha == alpha
        assert reloaded_shape.fill.fore_color.rgb == RGBColor(0xFF, 0x00, 0x00)

    def it_can_clear_alpha_by_assigning_None(self, _restore_part_factory):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x00, 0x80, 0x00)
        shape.fill.fore_color.alpha = 0.3
        assert shape.fill.fore_color.alpha == 0.3

        shape.fill.fore_color.alpha = None

        assert shape.fill.fore_color.alpha == 1.0

    def it_reads_and_writes_alpha_on_a_theme_color(self, _restore_part_factory):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(2))
        shape.fill.solid()
        # -- set a theme color (which writes an a:schemeClr element) --
        from pptx.enum.dml import MSO_THEME_COLOR

        shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_2
        shape.fill.fore_color.alpha = 0.6

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        reloaded = prs2.slides[0].shapes[1].fill.fore_color
        assert reloaded.theme_color == MSO_THEME_COLOR.ACCENT_2
        assert reloaded.alpha == 0.6

    def it_coexists_with_brightness_on_a_theme_color(self, _restore_part_factory):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(2))
        shape.fill.solid()
        from pptx.enum.dml import MSO_THEME_COLOR

        shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        shape.fill.fore_color.brightness = -0.25
        shape.fill.fore_color.alpha = 0.4

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        reloaded = prs2.slides[0].shapes[1].fill.fore_color
        assert reloaded.brightness == -0.25
        assert reloaded.alpha == 0.4
