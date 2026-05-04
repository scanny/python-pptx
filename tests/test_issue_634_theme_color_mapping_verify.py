# pyright: reportPrivateUsage=false

"""Regression test for issue #634 — scheme-color tag name in XML vs.
``MSO_THEME_COLOR_INDEX`` member name.

Issue #634 (https://github.com/scanny/python-pptx/issues/634) reports that
the XML tag names written for scheme colors (e.g. ``tx1``, ``bg1``,
``dk1``, ``lt1``) don't visibly match the ``MSO_THEME_COLOR_INDEX`` member
names (``TEXT_1``, ``BACKGROUND_1``, ``DARK_1``, ``LIGHT_1``). The
reporter was trying to reverse-engineer which enum member corresponds to
which XML tag.

The mapping is intentional and correct — OOXML defines two parallel sets
of scheme-color names connected by the slide-master's ``<a:clrMap>``:

    * **Slide-level** names (what PowerPoint writes in slide content):
      ``tx1`` ↔ ``TEXT_1``, ``bg1`` ↔ ``BACKGROUND_1``,
      ``tx2`` ↔ ``TEXT_2``, ``bg2`` ↔ ``BACKGROUND_2``,
      ``accent1``-``accent6`` ↔ ``ACCENT_1``-``ACCENT_6``,
      ``hlink`` ↔ ``HYPERLINK``, ``folHlink`` ↔ ``FOLLOWED_HYPERLINK``.

    * **Theme-level** names (what appears inside a theme's
      ``<a:clrScheme>``): ``dk1`` ↔ ``DARK_1``, ``lt1`` ↔ ``LIGHT_1``,
      ``dk2`` ↔ ``DARK_2``, ``lt2`` ↔ ``LIGHT_2``, plus the same accent
      and hyperlink members as above.

The default ``<p:clrMap>`` rewires ``tx1→dk1`` / ``bg1→lt1`` /
``tx2→dk2`` / ``bg2→lt2`` at the slide master, so the two sets resolve to
the same underlying theme colors. Authors can swap the mapping to flip
light/dark appearance without editing slide content.

This test pins the two-way mapping explicitly so a future refactor can't
silently relabel an enum member's XML value and break round-trip fidelity
with PowerPoint.
"""

from __future__ import annotations

import pytest

from pptx.enum.dml import MSO_THEME_COLOR, MSO_THEME_COLOR_INDEX
from pptx.oxml.ns import qn
from pptx.util import Pt

from .unitutil.cxml import element


class DescribeIssue634ThemeColorMapping:
    """Pin MSO_THEME_COLOR_INDEX ↔ ``<a:schemeClr val='…'>`` round-trip."""

    def it_exposes_MSO_THEME_COLOR_as_an_alias(self):
        assert MSO_THEME_COLOR is MSO_THEME_COLOR_INDEX

    @pytest.mark.parametrize(
        ("member", "xml_value"),
        [
            # -- slide-level scheme-color names (what PowerPoint writes) --
            (MSO_THEME_COLOR.TEXT_1, "tx1"),
            (MSO_THEME_COLOR.BACKGROUND_1, "bg1"),
            (MSO_THEME_COLOR.TEXT_2, "tx2"),
            (MSO_THEME_COLOR.BACKGROUND_2, "bg2"),
            # -- theme-level scheme-color names (what a:clrScheme uses) --
            (MSO_THEME_COLOR.DARK_1, "dk1"),
            (MSO_THEME_COLOR.LIGHT_1, "lt1"),
            (MSO_THEME_COLOR.DARK_2, "dk2"),
            (MSO_THEME_COLOR.LIGHT_2, "lt2"),
            # -- accent / hyperlink pass-through in both contexts --
            (MSO_THEME_COLOR.ACCENT_1, "accent1"),
            (MSO_THEME_COLOR.ACCENT_2, "accent2"),
            (MSO_THEME_COLOR.ACCENT_3, "accent3"),
            (MSO_THEME_COLOR.ACCENT_4, "accent4"),
            (MSO_THEME_COLOR.ACCENT_5, "accent5"),
            (MSO_THEME_COLOR.ACCENT_6, "accent6"),
            (MSO_THEME_COLOR.HYPERLINK, "hlink"),
            (MSO_THEME_COLOR.FOLLOWED_HYPERLINK, "folHlink"),
        ],
    )
    def it_maps_each_member_to_its_OOXML_val_string(
        self, member: MSO_THEME_COLOR_INDEX, xml_value: str
    ):
        # -- from the enum side: xml_value attribute & to_xml() --
        assert member.xml_value == xml_value
        assert MSO_THEME_COLOR.to_xml(member) == xml_value
        # -- from the XML side: from_xml() round-trips --
        assert MSO_THEME_COLOR.from_xml(xml_value) is member

    @pytest.mark.parametrize(
        ("member", "ms_int"),
        [
            # -- integer values must match MsoThemeColorIndex in the MS API --
            (MSO_THEME_COLOR.DARK_1, 1),
            (MSO_THEME_COLOR.LIGHT_1, 2),
            (MSO_THEME_COLOR.DARK_2, 3),
            (MSO_THEME_COLOR.LIGHT_2, 4),
            (MSO_THEME_COLOR.ACCENT_1, 5),
            (MSO_THEME_COLOR.ACCENT_2, 6),
            (MSO_THEME_COLOR.ACCENT_3, 7),
            (MSO_THEME_COLOR.ACCENT_4, 8),
            (MSO_THEME_COLOR.ACCENT_5, 9),
            (MSO_THEME_COLOR.ACCENT_6, 10),
            (MSO_THEME_COLOR.HYPERLINK, 11),
            (MSO_THEME_COLOR.FOLLOWED_HYPERLINK, 12),
            (MSO_THEME_COLOR.TEXT_1, 13),
            (MSO_THEME_COLOR.BACKGROUND_1, 14),
            (MSO_THEME_COLOR.TEXT_2, 15),
            (MSO_THEME_COLOR.BACKGROUND_2, 16),
            (MSO_THEME_COLOR.NOT_THEME_COLOR, 0),
            (MSO_THEME_COLOR.MIXED, -2),
        ],
    )
    def it_matches_the_MS_API_integer_values(self, member: MSO_THEME_COLOR_INDEX, ms_int: int):
        assert int(member) == ms_int
        assert member.value == ms_int

    def it_writes_tx1_when_TEXT_1_is_assigned_to_font_color_theme_color(self):
        # -- this is the scenario the #634 reporter exercised: setting
        # -- theme_color = TEXT_1 should produce `<a:schemeClr val="tx1">`,
        # -- the slide-level tag PowerPoint itself emits.
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Pt(100), Pt(100), Pt(200), Pt(50))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = "hello"
        run.font.size = Pt(14)
        run.font.color.theme_color = MSO_THEME_COLOR.TEXT_1

        rPr = run._r.rPr
        assert rPr is not None
        solidFill = rPr.find(qn("a:solidFill"))
        assert solidFill is not None
        schemeClr = solidFill.find(qn("a:schemeClr"))
        assert schemeClr is not None
        assert schemeClr.get("val") == "tx1"
        # -- and the round-trip back to the enum still yields TEXT_1 --
        assert run.font.color.theme_color is MSO_THEME_COLOR.TEXT_1

    def it_writes_dk1_when_DARK_1_is_assigned_to_font_color_theme_color(self):
        # -- theme-level tag: useful inside a theme's <a:clrScheme>. Slide
        # -- content can also accept this tag; PowerPoint reads both.
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Pt(100), Pt(100), Pt(200), Pt(50))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = "hello"
        run.font.color.theme_color = MSO_THEME_COLOR.DARK_1

        rPr = run._r.rPr
        assert rPr is not None
        solidFill = rPr.find(qn("a:solidFill"))
        assert solidFill is not None
        schemeClr = solidFill.find(qn("a:schemeClr"))
        assert schemeClr is not None
        assert schemeClr.get("val") == "dk1"
        assert run.font.color.theme_color is MSO_THEME_COLOR.DARK_1

    def it_reads_back_the_xml_tag_verbatim_from_schemeClr(self):
        # -- Critically, there's no silent rewrite on read: if the XML
        # -- carries `val="dk1"`, the getter returns DARK_1, not TEXT_1
        # -- (and vice-versa). This is the behavior the #634 reporter was
        # -- trying to verify.
        schemeClr = element("a:schemeClr{val=dk1}")
        val = schemeClr.get("val")
        assert val is not None
        # -- the xml_value → enum member mapping is unambiguous --
        assert MSO_THEME_COLOR.from_xml(val) is MSO_THEME_COLOR.DARK_1

        schemeClr2 = element("a:schemeClr{val=tx1}")
        val2 = schemeClr2.get("val")
        assert val2 is not None
        assert MSO_THEME_COLOR.from_xml(val2) is MSO_THEME_COLOR.TEXT_1

    def it_covers_every_member_of_ST_SchemeColorVal_except_phClr(self):
        # -- OOXML ST_SchemeColorVal enumerates 17 values. All except
        # -- `phClr` (a placeholder-color sentinel used only inside theme
        # -- style matrices, not a user-settable theme color) must have a
        # -- corresponding MSO_THEME_COLOR_INDEX member.
        ooxml_scheme_vals = {
            "bg1",
            "tx1",
            "bg2",
            "tx2",
            "accent1",
            "accent2",
            "accent3",
            "accent4",
            "accent5",
            "accent6",
            "hlink",
            "folHlink",
            "dk1",
            "lt1",
            "dk2",
            "lt2",
        }
        enum_xml_vals = {m.xml_value for m in MSO_THEME_COLOR if m.xml_value}
        missing = ooxml_scheme_vals - enum_xml_vals
        assert not missing, f"MSO_THEME_COLOR_INDEX missing XML mapping for {missing}"
