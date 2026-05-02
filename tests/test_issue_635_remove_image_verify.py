# pyright: reportPrivateUsage=false

"""Regression test for issue #635 — "Removing an image from power point".

Issue #635 (https://github.com/scanny/python-pptx/issues/635) asked for a
supported way to delete a specific image (picture shape) from an opened
``.pptx``. The reporter had an iter-over-shapes visitor that extracted the
image bytes for each ``MSO_SHAPE_TYPE.PICTURE`` shape and then — having
identified the images they wanted to delete — asked "Is there any library
through i can remove the image from pptx?"

Wave 1 #41 (``feat/issue-41-shape-delete``) shipped the supported API
this ask needs:

  * :meth:`.BaseShape.delete` — the base primitive that detaches the
    shape's XML element from its parent ``p:spTree``.
  * :meth:`.Picture.delete` — a picture-specific override that, before
    super-deleting the ``p:pic``, drops the ``r:embed`` relationship
    pointing at the image part (via the reference-counted
    :meth:`.XmlPart.drop_rel` so shared images survive). Once the last
    reference to an :class:`.ImagePart` is dropped, it is unreachable
    from the package rel graph and is garbage-collected at save time.

#635 is therefore a duplicate of #41 and is verified-and-closed by this
suite. The scenarios below exercise the reporter's exact workflow (an
image-extractor visitor, then delete a specific subset of pictures by
identity, save, reopen, confirm the rest survive) plus the
image-part-orphan semantics that make the delete actually shrink the
saved ``.pptx`` rather than leave the image bytes dangling inside the
zip.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.parts.image import ImagePart
from pptx.util import Inches


def _picture_shapes(prs):
    """Yield every ``Picture`` shape across every slide of ``prs``.

    Mirrors the reporter's ``iter_picture_shapes`` visitor: descends into
    group shapes and collects each picture-kind shape. Only the top-level
    slide pictures are produced in the scenarios below (no groups), but we
    keep the recursive walk so the helper matches the shape of the
    reporter's snippet.
    """

    def visit(shape):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for s in shape.shapes:
                yield from visit(s)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield shape

    for slide in prs.slides:
        for shape in slide.shapes:
            yield from visit(shape)


def _image_parts(prs):
    """Return the set of |ImagePart| instances reachable from ``prs``."""
    return {p for p in prs.part.package.iter_parts() if isinstance(p, ImagePart)}


class DescribeIssue635RemoveImage(object):
    """#635 delete-image verify-and-close via #41's ``Picture.delete`` API."""

    # -- reporter's workflow ------------------------------------------------

    def it_deletes_the_picture_the_reporter_targeted(self):
        # -- the #635 reporter's exact flow: iterate pictures, pick the
        # -- ones to delete by identity, call .delete(). We pin that after
        # -- the delete the slide no longer holds that picture, and that
        # -- the surviving pictures are unchanged.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        pic1 = slide.shapes.add_picture(
            "tests/test_files/python-powered.png",
            Inches(1),
            Inches(1),
            Inches(2),
            Inches(1),
        )
        pic2 = slide.shapes.add_picture(
            "tests/test_files/monty-truth.png",
            Inches(4),
            Inches(1),
            Inches(2),
            Inches(1),
        )
        pic3 = slide.shapes.add_picture(
            "tests/test_files/python-icon.jpeg",
            Inches(1),
            Inches(4),
            Inches(2),
            Inches(1),
        )
        pic1_name, pic3_name = pic1.name, pic3.name

        # -- the reporter identifies the middle picture as "image2" and
        # -- deletes it; the other two stay
        pic2.delete()

        survivors = list(_picture_shapes(prs))
        assert len(survivors) == 2
        assert {s.name for s in survivors} == {pic1_name, pic3_name}

    def it_deletes_pictures_via_the_reporter_visitor_style(self):
        # -- pin that the reporter's own iter-pictures idiom — collecting
        # -- Picture references into a list and then acting on them —
        # -- composes cleanly with Picture.delete. A naive implementation
        # -- could break here if ``delete`` mutated shapes mid-iteration.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        kept = slide.shapes.add_picture(
            "tests/test_files/python-powered.png",
            Inches(1),
            Inches(1),
            Inches(2),
            Inches(1),
        )
        gone_a = slide.shapes.add_picture(
            "tests/test_files/monty-truth.png",
            Inches(4),
            Inches(1),
            Inches(2),
            Inches(1),
        )
        gone_b = slide.shapes.add_picture(
            "tests/test_files/python-icon.jpeg",
            Inches(1),
            Inches(4),
            Inches(2),
            Inches(1),
        )
        kept_name = kept.name

        # -- snapshot the pictures first (per the reporter's pattern),
        # -- then iterate and delete a subset. Matching by identity
        # -- (shape_id) mirrors picking by filename in the reporter's code.
        pictures = list(_picture_shapes(prs))
        targets = {gone_a.shape_id, gone_b.shape_id}
        for pic in pictures:
            if pic.shape_id in targets:
                pic.delete()

        survivors = list(_picture_shapes(prs))
        assert [s.name for s in survivors] == [kept_name]

    # -- image-part-orphan semantics ---------------------------------------

    def it_orphans_the_image_part_so_it_is_gc_on_save(self):
        # -- the visible benefit of Picture.delete over a raw element
        # -- removal: the image bytes leave the package. Pin that after
        # -- save + reopen the package no longer contains the deleted
        # -- picture's image part.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        pic_to_delete = slide.shapes.add_picture(
            "tests/test_files/monty-truth.png",
            Inches(1),
            Inches(1),
            Inches(2),
            Inches(1),
        )
        pic_to_keep = slide.shapes.add_picture(
            "tests/test_files/python-powered.png",
            Inches(4),
            Inches(1),
            Inches(2),
            Inches(1),
        )
        kept_sha = pic_to_keep.image.sha1
        deleted_sha = pic_to_delete.image.sha1
        assert kept_sha != deleted_sha

        pic_to_delete.delete()

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        image_shas = {p.sha1 for p in _image_parts(prs2)}
        assert kept_sha in image_shas
        assert deleted_sha not in image_shas

    # -- XML-shape invariant -----------------------------------------------

    def it_removes_the_p_pic_from_the_p_spTree(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        pic = slide.shapes.add_picture(
            "tests/test_files/python-powered.png",
            Inches(1),
            Inches(1),
            Inches(2),
            Inches(1),
        )
        spTree = slide.shapes._spTree
        pic_elm = pic._element
        assert pic_elm in list(spTree)

        pic.delete()

        assert pic_elm not in list(spTree)

    # -- round-trip --------------------------------------------------------

    def it_round_trips_the_deleted_picture_through_save_reopen(self):
        # -- the reporter's end goal: save the file with the image gone
        # -- and have PowerPoint (modelled here by Presentation(buf)) see
        # -- the same shape tree.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        kept = slide.shapes.add_picture(
            "tests/test_files/python-powered.png",
            Inches(1),
            Inches(1),
            Inches(2),
            Inches(1),
        )
        gone = slide.shapes.add_picture(
            "tests/test_files/monty-truth.png",
            Inches(4),
            Inches(1),
            Inches(2),
            Inches(1),
        )
        kept_name = kept.name
        gone_name = gone.name

        gone.delete()

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        names = [s.name for s in prs2.slides[0].shapes]
        assert kept_name in names
        assert gone_name not in names

    # -- public API surface ------------------------------------------------

    def it_exposes_delete_on_Picture(self):
        # -- belt-and-braces: pin that the attribute the reporter needed
        # -- actually exists on the class.
        from pptx.shapes.picture import Picture

        assert hasattr(Picture, "delete")
        assert callable(Picture.delete)
