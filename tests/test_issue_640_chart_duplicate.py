# pyright: reportPrivateUsage=false

"""Regression test for issue #640 — duplicate a chart.

Issue #640 (https://github.com/scanny/python-pptx/issues/640) asks for a
way to duplicate an existing chart — specifically so the duplicate keeps
a working embedded workbook (so PowerPoint's "Edit Data" dialog still
opens on the copy) and so authoring a second variant of a chart does not
require re-building the chart XML from scratch.

Wave 5 shipped ``Chart.clone_to(shapes, x, y, cx, cy)`` for issue #877
(cross-slide chart copy). Same-slide duplication is structurally a
subset of that primitive: the target shape-tree happens to be the same
``Slide.shapes`` instance as the source chart's. This test verifies
that the #877 implementation handles the same-slide case correctly —
in particular that

* the duplicate lives in a distinct ``ChartPart`` (so editing its data
  does not rewrite the source chart's XML),
* the duplicate's embedded workbook is a distinct ``EmbeddedXlsxPart``
  (so PowerPoint's "Edit Data" dialog on the copy edits a separate
  ``.xlsx``),
* the two graphic-frames on the slide get distinct shape ids and
  distinct shape names (no same-slide collision on
  ``p:nvGraphicFramePr/p:cNvPr/@id`` or ``/@name``),
* the slide round-trips through ``save`` + ``Presentation(...)`` with
  both charts intact and still pointing at distinct parts on reload.

If #877 ever regresses in a way that breaks same-slide duplication
(e.g. by reusing the source rId / shape-id / name), this test will
catch it.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


@pytest.fixture
def _restore_part_factory():
    """Guard against pollution from other tests that mutate `PartFactory.part_type_for`.

    In particular, ``tests/opc/test_package.py::DescribePartFactory`` overwrites
    ``PartFactory.part_type_for[CT.PML_SLIDE]`` with a Mock and does not restore
    it. Without this guard, running this module after ``tests/opc/test_package.py``
    causes ``Presentation(bio).slides[0]`` to return a Mock instead of a real
    Slide on reload. See the identical fixture in
    ``tests/test_issue_400_animation_umbrella.py``.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.comments import CommentAuthorsPart, CommentsPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_COMMENTS: CommentsPart,
        CT.PML_COMMENT_AUTHORS: CommentAuthorsPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue640RegressionChartDuplicate:
    """Same-slide chart duplication via Chart.clone_to (issue #640)."""

    def it_duplicates_a_chart_onto_its_own_slide_via_clone_to(self):
        """Chart.clone_to(slide.shapes, ...) — same-slide target — produces a
        working duplicate with distinct chart + xlsx parts and no shape-id
        / shape-name collision.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        cd = CategoryChartData()
        cd.categories = ["A", "B", "C"]
        cd.add_series("S1", (10.0, 20.0, 30.0))
        source_gf = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(3),
            cd,
        )
        source_chart = source_gf.chart
        source_chart_part = source_chart.part
        source_xlsx_part = source_chart_part.chart_workbook.xlsx_part

        # --- act: clone onto the *same* slide's shape tree ---
        new_gf = source_chart.clone_to(
            slide.shapes, Inches(1), Inches(5), Inches(4), Inches(3)
        )
        cloned_chart = new_gf.chart

        # --- the duplicate is at the requested position/extents ---
        assert new_gf.left == Inches(1)
        assert new_gf.top == Inches(5)
        assert new_gf.width == Inches(4)
        assert new_gf.height == Inches(3)

        # --- distinct chart parts so editing one does not rewrite the other ---
        assert cloned_chart.part is not source_chart_part

        # --- distinct embedded xlsx parts (so "Edit Data" works per-copy) ---
        cloned_xlsx_part = cloned_chart.part.chart_workbook.xlsx_part
        assert source_xlsx_part is not None
        assert cloned_xlsx_part is not None
        assert cloned_xlsx_part is not source_xlsx_part

        # --- identical workbook bytes at copy time ---
        assert cloned_chart.workbook == source_chart.workbook

        # --- both graphic-frames are on the same slide ---
        chart_gfs = [shp for shp in slide.shapes if shp.has_chart]
        assert len(chart_gfs) == 2

        # --- no shape-id collision on the same slide (the big risk) ---
        shape_ids = [shp.shape_id for shp in slide.shapes]
        assert len(shape_ids) == len(set(shape_ids)), (
            "same-slide chart duplicate reused an existing shape id: %r" % shape_ids
        )

        # --- no shape-name collision either ---
        shape_names = [shp.name for shp in slide.shapes]
        assert len(shape_names) == len(set(shape_names)), (
            "same-slide chart duplicate reused an existing shape name: %r"
            % shape_names
        )

    def it_round_trips_a_same_slide_chart_duplicate(self, _restore_part_factory):
        """The duplicated chart survives save + reload on the same slide."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        cd = CategoryChartData()
        cd.categories = ["Q1", "Q2", "Q3", "Q4"]
        cd.add_series("Revenue", (1.0, 2.0, 3.0, 4.0))
        source_gf = slide.shapes.add_chart(
            XL_CHART_TYPE.LINE,
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(3),
            cd,
        )

        # --- duplicate onto the same slide ---
        source_gf.chart.clone_to(
            slide.shapes, Inches(1), Inches(5), Inches(5), Inches(3)
        )

        # --- round-trip through save + reload ---
        bio = io.BytesIO()
        prs.save(bio)
        bio.seek(0)
        prs2 = Presentation(bio)

        slide2 = list(prs2.slides)[0]
        chart_gfs = [shp for shp in slide2.shapes if shp.has_chart]
        assert len(chart_gfs) == 2

        # --- both reloaded charts have their own workbook ---
        assert chart_gfs[0].chart.workbook is not None
        assert chart_gfs[1].chart.workbook is not None

        # --- and the two charts live in distinct chart parts on reload ---
        assert chart_gfs[0].chart.part is not chart_gfs[1].chart.part
