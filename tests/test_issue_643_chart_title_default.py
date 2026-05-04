# pyright: reportPrivateUsage=false

"""Regression test for GitHub issue #643.

Setting ``chart.has_title = True`` must not leave a stale
``c:autoTitleDeleted val="1"`` behind — that combination tells PowerPoint
the user explicitly deleted the auto-title, which suppresses the
default auto-title behaviour (series-name for single-series charts,
"Chart Title" placeholder otherwise). python-pptx must emit a
minimal ``c:title`` (no hard-coded ``a:t`` text) and ensure
``c:autoTitleDeleted`` is absent or ``val="0"`` so PowerPoint
auto-fills the title as it would for a chart authored in the UI.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches


@pytest.fixture
def single_series_chart():
    """Return a newly-created single-series bar chart."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    data = CategoryChartData()
    data.categories = ["A", "B", "C"]
    data.add_series("Series 1", (1.0, 2.0, 3.0))
    gf = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        data,
    )
    return prs, gf.chart


class DescribeIssue643ChartTitleDefault:
    """Setting `has_title = True` must let PowerPoint auto-fill the title."""

    def it_emits_no_hard_coded_Chart_Title_text(self, single_series_chart):
        _prs, chart = single_series_chart

        chart.has_title = True

        # -- the title must not carry a hard-coded "Chart Title" run; --
        # -- PowerPoint fills in the display text from the series name --
        # -- when there's one series and the c:title has no c:tx. --
        title = chart._chartSpace.chart.title
        assert title is not None
        for a_t in title.iter(qn("a:t")):
            assert a_t.text != "Chart Title"

    def it_leaves_c_tx_absent_so_PowerPoint_auto_titles(self, single_series_chart):
        _prs, chart = single_series_chart

        chart.has_title = True

        # -- a minimal c:title (no c:tx child) is the signal to PowerPoint --
        # -- to render the auto-title (series name for single-series). --
        title = chart._chartSpace.chart.title
        assert title.find(qn("c:tx")) is None

    def it_clears_autoTitleDeleted_flag_when_re_enabling_the_title(self, single_series_chart):
        """Toggling `has_title` off then on must clear `autoTitleDeleted`.

        Otherwise PowerPoint sees ``autoTitleDeleted="1"`` and suppresses
        the auto-title — the user gets no title at all instead of the
        series name.
        """
        _prs, chart = single_series_chart

        chart.has_title = False
        # -- setter above set autoTitleDeleted val="1" --
        autoTitleDeleted = chart._chartSpace.chart.autoTitleDeleted
        assert autoTitleDeleted is not None
        assert autoTitleDeleted.val is True

        chart.has_title = True

        # -- re-enabling must clear the flag so the auto-title engages --
        autoTitleDeleted = chart._chartSpace.chart.autoTitleDeleted
        assert autoTitleDeleted is None or autoTitleDeleted.val is False

    def it_survives_a_save_reopen_round_trip(self, single_series_chart):
        prs, chart = single_series_chart

        chart.has_title = False
        chart.has_title = True

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reopened = Presentation(buf)
        # -- walk to the sole chart on the sole slide --
        reopened_chart = None
        for shape in reopened.slides[0].shapes:
            if shape.has_chart:
                reopened_chart = shape.chart
                break
        assert reopened_chart is not None
        assert reopened_chart.has_title is True
        title = reopened_chart._chartSpace.chart.title
        assert title.find(qn("c:tx")) is None
        autoTitleDeleted = reopened_chart._chartSpace.chart.autoTitleDeleted
        assert autoTitleDeleted is None or autoTitleDeleted.val is False

    def it_leaves_author_set_title_text_alone(self, single_series_chart):
        """After `has_title = True`, writing text via text_frame stays."""
        _prs, chart = single_series_chart

        chart.has_title = True
        chart.chart_title.text_frame.text = "Revenue"

        title = chart._chartSpace.chart.title
        a_t_texts = [e.text for e in title.iter(qn("a:t"))]
        assert "Revenue" in a_t_texts
        assert "Chart Title" not in a_t_texts

    def it_is_idempotent_when_autoTitleDeleted_is_already_clear(self, single_series_chart):
        _prs, chart = single_series_chart

        chart.has_title = True
        chart.has_title = True  # second call must be a no-op shape-wise

        title = chart._chartSpace.chart.title
        assert title is not None
        assert title.find(qn("c:tx")) is None
