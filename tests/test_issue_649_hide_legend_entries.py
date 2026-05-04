# pyright: reportPrivateUsage=false

"""Regression test for issue #649 — hide individual legend entries.

Issue #649 (https://github.com/scanny/python-pptx/issues/649) asks for a
way to hide specific legend entries while keeping the corresponding
series visible. PowerPoint models this at the ``c:legend`` level with
a ``c:legendEntry/c:idx`` + ``c:delete`` pair — not at the series level,
because the series must keep plotting.

This test pins the contract from the reporter's perspective:

* A chart with N series starts with no hidden legend entries.
* ``Legend.exclude_entry(i)`` hides only that one entry, leaving all
  series intact in the plot.
* ``Legend.hidden_entries`` reports the hidden indices.
* ``Legend.include_entry(i)`` re-shows a previously hidden entry.
* Excluding then re-including is a clean round-trip — the underlying
  XML has no lingering ``c:legendEntry`` elements.
* The change round-trips through :meth:`Presentation.save` + reopen.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches


def _add_chart_with_four_series(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    cd = CategoryChartData()
    cd.categories = ["A", "B", "C"]
    cd.add_series("S0", (1.0, 2.0, 3.0))
    cd.add_series("S1", (4.0, 5.0, 6.0))
    cd.add_series("S2", (7.0, 8.0, 9.0))
    cd.add_series("S3", (10.0, 11.0, 12.0))
    gf = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        cd,
    )
    chart = gf.chart
    chart.has_legend = True
    return chart


class DescribeIssue649HideLegendEntries:
    """``Legend.exclude_entry`` / ``include_entry`` / ``hidden_entries``."""

    def it_starts_with_no_hidden_entries(self):
        prs = Presentation()
        chart = _add_chart_with_four_series(prs)
        assert chart.legend.hidden_entries == ()

    def it_hides_a_single_entry_without_dropping_its_series(self):
        prs = Presentation()
        chart = _add_chart_with_four_series(prs)

        chart.legend.exclude_entry(1)

        assert chart.legend.hidden_entries == (1,)
        # -- series still plotted --
        assert [s.name for s in chart.series] == ["S0", "S1", "S2", "S3"]

    def it_hides_non_contiguous_entries(self):
        prs = Presentation()
        chart = _add_chart_with_four_series(prs)

        chart.legend.exclude_entry(0)
        chart.legend.exclude_entry(2)

        assert chart.legend.hidden_entries == (0, 2)

    def it_is_idempotent_on_repeat_exclude(self):
        prs = Presentation()
        chart = _add_chart_with_four_series(prs)

        chart.legend.exclude_entry(1)
        chart.legend.exclude_entry(1)

        assert chart.legend.hidden_entries == (1,)
        # -- only one c:legendEntry child --
        assert len(chart.legend._element.findall(qn("c:legendEntry"))) == 1

    def it_reincludes_a_previously_hidden_entry(self):
        prs = Presentation()
        chart = _add_chart_with_four_series(prs)
        chart.legend.exclude_entry(1)
        chart.legend.exclude_entry(3)

        chart.legend.include_entry(1)

        assert chart.legend.hidden_entries == (3,)

    def it_leaves_no_orphan_legendEntry_after_full_roundtrip(self):
        prs = Presentation()
        chart = _add_chart_with_four_series(prs)

        chart.legend.exclude_entry(2)
        chart.legend.include_entry(2)

        assert chart.legend.hidden_entries == ()
        # -- the c:legendEntry is removed entirely, not left as an
        # -- idx-only stub, keeping the XML minimal --
        assert chart.legend._element.findall(qn("c:legendEntry")) == []

    def it_ignores_include_for_a_non_hidden_idx(self):
        prs = Presentation()
        chart = _add_chart_with_four_series(prs)

        chart.legend.include_entry(0)  # nothing to include

        assert chart.legend.hidden_entries == ()

    def it_survives_save_and_reopen_round_trip(self):
        prs = Presentation()
        chart = _add_chart_with_four_series(prs)
        chart.legend.exclude_entry(1)
        chart.legend.exclude_entry(3)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        prs2 = Presentation(buf)
        slide = prs2.slides[-1]
        chart2 = next(s for s in slide.shapes if getattr(s, "has_chart", False)).chart
        assert chart2.legend.hidden_entries == (1, 3)
        # -- series are all still there --
        assert [s.name for s in chart2.series] == ["S0", "S1", "S2", "S3"]

    def it_writes_c_delete_with_explicit_val_1(self):
        # PowerPoint always writes val="1" on c:delete; assert we do too so
        # readers that interpret a missing val as False don't silently
        # un-hide the entry.
        prs = Presentation()
        chart = _add_chart_with_four_series(prs)

        chart.legend.exclude_entry(2)

        legend_elm = chart.legend._element
        matches = legend_elm.xpath("c:legendEntry/c:delete/@val")
        assert matches == ["1"]
