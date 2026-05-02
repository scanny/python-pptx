# pyright: reportPrivateUsage=false

"""Regression test for issue #650 — per-point color hides category names.

Issue #650 (https://github.com/scanny/python-pptx/issues/650) reports
that setting a per-point data-label color via
``DataLabel.font.color.rgb = ...`` causes the category name on that
point's data label to disappear.

Root cause: the ``CT_DLbl.new_dLbl`` factory hard-coded
``c:showCatName val="0"`` (and the other ``c:show*`` children) into
every newly-created per-point ``c:dLbl``. A child element under
``c:dLbl`` *overrides* the corresponding setting on the parent
series-level ``c:dLbls`` — so a point whose label container was created
merely as a vehicle for an ``spPr``/``txPr`` override would suppress the
category name that the series-level setting had enabled.

The fix: ``CT_DLbl.new_dLbl`` no longer emits the ``c:show*`` children;
inheritance from series/plot-level ``c:dLbls`` is allowed to take
effect. The series-level ``CT_DLbls.new_dLbls`` factory is unchanged.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.chart.datalabel import CT_DLbl
from pptx.oxml.ns import qn
from pptx.util import Inches


def _category_chart():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    cd = CategoryChartData()
    cd.categories = ["Alpha", "Beta", "Gamma"]
    cd.add_series("s1", (1.0, 2.0, 3.0))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        cd,
    ).chart
    return prs, chart


def _roundtrip(prs):
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)


class DescribeIssue650DataLabelCategoryVisibility(object):
    """Verify #650 — point color no longer hides categories."""

    def it_does_not_emit_show_children_from_new_dLbl_factory(self):
        """Unit-level pin on the fix.

        ``CT_DLbl.new_dLbl()`` must not emit any ``c:show*`` children;
        otherwise those children would override series-level
        inheritance.
        """
        dLbl = CT_DLbl.new_dLbl()

        for name in (
            "showLegendKey",
            "showVal",
            "showCatName",
            "showSerName",
            "showPercent",
            "showBubbleSize",
        ):
            assert dLbl.find(qn("c:%s" % name)) is None, (
                "new_dLbl() must not hard-code c:%s — it would override"
                " series-level defaults (issue #650)." % name
            )

    def it_keeps_categories_visible_when_point_font_color_is_set(self):
        """The reporter's scenario.

        Set series-level ``show_category_name`` to True, then author a
        per-point font color override. Category visibility must survive.
        """
        prs, chart = _category_chart()
        plot = chart.plots[0]
        plot.has_data_labels = True
        series = plot.series[0]

        # ---author wants category names on all labels---
        series.data_labels.show_category_name = True
        series.data_labels.show_value = True

        # ---author picks a per-point color for point 1---
        series.points[1].data_label.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

        # The per-point c:dLbl must NOT hard-code show* children that
        # would override the series-level defaults.
        point_dLbl = series.points[1].data_label._dLbl
        assert point_dLbl is not None
        for name in (
            "showLegendKey",
            "showVal",
            "showCatName",
            "showSerName",
            "showPercent",
            "showBubbleSize",
        ):
            assert point_dLbl.find(qn("c:%s" % name)) is None

        # Round-trip confirms the series-level settings still apply.
        reloaded = _roundtrip(prs)
        reloaded_chart = None
        for shp in reloaded.slides[0].shapes:
            if shp.has_chart:
                reloaded_chart = shp.chart
                break
        assert reloaded_chart is not None

        reloaded_series = reloaded_chart.plots[0].series[0]
        assert reloaded_series.data_labels.show_category_name is True
        assert reloaded_series.data_labels.show_value is True
        assert reloaded_series.points[1].data_label.font.color.rgb == RGBColor(0xFF, 0x00, 0x00)
