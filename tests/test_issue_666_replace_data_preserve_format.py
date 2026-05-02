# pyright: reportPrivateUsage=false

"""Regression test for issue #666 — Chart.replace_data preserves number format.

Issue #666 (https://github.com/scanny/python-pptx/issues/666) reports that
calling :meth:`Chart.replace_data` on a chart whose series had a non-default
``c:formatCode`` (e.g. ``"0.00%"`` or ``"#,##0"``) silently resets every
``c:numCache/c:formatCode`` (and the axis-displayed number format on
series-linked axes) back to the hard-coded default ``"General"``. The
reporter had to reach into the oxml layer after every ``replace_data`` call
to restore the formatting they had authored.

The fix (on branch ``fix/issue-666-replace-data-preserve-format``) teaches
each ``_BaseSeriesXmlRewriter._rewrite_ser_data`` subclass to capture the
existing ``c:formatCode`` on ``c:val`` / ``c:xVal`` / ``c:yVal`` /
``c:bubbleSize`` / numeric ``c:cat`` before the element is removed, then
re-apply it to the freshly generated replacement — but only when the caller
left ``chart_data.number_format`` at its ``"General"`` default, so an
explicit ``number_format=`` on ``CategoryChartData`` / ``ChartData`` /
``XyChartData`` / ``BubbleChartData`` still wins.

This test locks down the end-to-end flow:

1. Author a column chart with ``chart_data.number_format = "0.00%"``.
2. Reload the package, assert the authored format landed in the `c:val`
   ``c:formatCode``.
3. Build a fresh ``CategoryChartData`` without setting
   ``number_format`` and call ``chart.replace_data`` with it.
4. Save + reload and assert the ``"0.00%"`` format is still present — i.e.
   the bug is fixed.
5. Repeat but this time pass an explicit ``number_format="#,##0"`` on the
   replacement chart_data and assert the new value overrides the preserved
   one (confirming the explicit-override escape hatch still works).
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def _chart_in(slide):
    for shp in slide.shapes:
        if shp.has_chart:
            return shp.chart
    raise AssertionError("slide contains no chart")


def _val_format_codes(chart):
    """Return the list of ``c:val//c:formatCode`` strings for each c:ser."""
    chartSpace = chart._chartSpace
    return [fc.text for fc in chartSpace.xpath(".//c:ser/c:val//c:formatCode")]


def _build_chart_with_format(fmt):
    """Author a column chart whose series use *fmt* as their number format."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_data = CategoryChartData(number_format=fmt)
    chart_data.categories = ["A", "B", "C"]
    chart_data.add_series("s1", (0.1, 0.2, 0.3))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    )
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)


class DescribeIssue666ReplaceDataPreservesNumberFormat:
    """End-to-end coverage for #666."""

    def it_preserves_the_val_formatCode_across_replace_data(self):
        # -- author a chart with the percent format the reporter used --
        prs = _build_chart_with_format("0.00%")
        chart = _chart_in(prs.slides[0])
        assert _val_format_codes(chart) == ["0.00%"]

        # -- call replace_data with a fresh chart_data carrying **no**
        # -- explicit number_format — pre-fix this would reset the
        # -- preserved format back to "General". --
        new_cd = CategoryChartData()
        new_cd.categories = ["A", "B", "C"]
        new_cd.add_series("s1", (0.25, 0.5, 0.75))
        chart.replace_data(new_cd)

        # -- save + reload to confirm the preserved format survives the
        # -- zip round-trip, not just the in-memory mutation. --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)
        reloaded_chart = _chart_in(reloaded.slides[0])

        assert _val_format_codes(reloaded_chart) == ["0.00%"]

    def it_lets_an_explicit_number_format_override_the_preserved_one(
        self
    ):
        # -- start from a chart with "0.00%" authored on every series --
        prs = _build_chart_with_format("0.00%")
        chart = _chart_in(prs.slides[0])

        # -- explicit number_format on the replacement chart_data should
        # -- win over the preserved value (escape hatch for callers who
        # -- really do want to change the format). --
        new_cd = CategoryChartData(number_format="#,##0")
        new_cd.categories = ["A", "B", "C"]
        new_cd.add_series("s1", (1.0, 2.0, 3.0))
        chart.replace_data(new_cd)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)
        reloaded_chart = _chart_in(reloaded.slides[0])

        assert _val_format_codes(reloaded_chart) == ["#,##0"]

    def it_preserves_format_across_multiple_existing_series(self):
        """Every series's format is preserved, not just the first one."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        chart_data = CategoryChartData(number_format="0.00%")
        chart_data.categories = ["A", "B"]
        chart_data.add_series("s1", (0.1, 0.2))
        chart_data.add_series("s2", (0.3, 0.4))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(6),
            Inches(4),
            chart_data,
        )
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs = Presentation(buf)
        chart = _chart_in(prs.slides[0])
        assert _val_format_codes(chart) == ["0.00%", "0.00%"]

        # -- replace_data with default-formatted chart_data --
        new_cd = CategoryChartData()
        new_cd.categories = ["A", "B"]
        new_cd.add_series("s1", (0.5, 0.6))
        new_cd.add_series("s2", (0.7, 0.8))
        chart.replace_data(new_cd)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)
        reloaded_chart = _chart_in(reloaded.slides[0])

        assert _val_format_codes(reloaded_chart) == ["0.00%", "0.00%"]
