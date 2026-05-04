# pyright: reportPrivateUsage=false

"""Regression test for issue #719 — XML writer for THREE_D_PIE not implemented.

Issue #719 (https://github.com/scanny/python-pptx/issues/719) reported
``NotImplementedError: XML writer for chart type THREE_D_PIE (-4102) not
yet implemented`` when attempting to author or round-trip a 3-D pie
chart. The original error came from ``ChartXmlWriter`` having no
builder mapped to :attr:`XL_CHART_TYPE.THREE_D_PIE`, but issue #719
also surfaces as a ``ValueError: unsupported plot type c:pie3DChart``
when :meth:`Chart.replace_data` looks up ``self.chart_type`` — that
lookup iterates ``chart.plots``, which in turn dispatches through
:class:`PlotFactory`.

Two waves of fork work closed the gap:

- Wave 7 #266 added :class:`_Pie3DChartXmlWriter` so
  :meth:`Shapes.add_chart` accepts
  :attr:`XL_CHART_TYPE.THREE_D_PIE` / ``THREE_D_PIE_EXPLODED`` and
  emits a schema-valid ``c:pie3DChart`` wrapper plus a ``c:view3D``
  sibling on ``c:chart``.
- Wave 13 #321 added :class:`Pie3DPlot` and registered
  ``c:pie3DChart`` in both :class:`PlotFactory` and
  :class:`PlotTypeInspector`, which unblocks the
  ``chart.plots[0]`` / ``chart.chart_type`` / ``chart.replace_data``
  read path.

This suite pins the #719 resolution so the issue can be closed. Each
test exercises the real API against a fresh ``Presentation()`` — not
a mock and not the raw ``ChartXmlWriter`` — covering the authoring
path, the ``replace_data`` round-trip that originally surfaced the
``ValueError``, and save + reopen for both the plain and exploded
variants.

Cross-reference: issue #321 (``feat/issue-321-pie3d-plot``) is the
immediate parent fix for the ``Pie3DPlot`` read-side gap; #266 is the
earlier fix for the writer-side ``NotImplementedError``.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.chart import Chart
from pptx.chart.data import CategoryChartData
from pptx.chart.plot import Pie3DPlot
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

# -- fixtures -----------------------------------------------------------


@pytest.fixture
def _restore_part_factory():
    """Guard against pollution from tests that mutate `PartFactory.part_type_for`.

    ``tests/opc/test_package.py::DescribePartFactory`` overwrites
    ``PartFactory.part_type_for[CT.PML_SLIDE]`` with a Mock and does
    not restore it. Round-trip tests that call ``Presentation.save``
    + ``Presentation()`` reopen depend on the real registrations being
    in place. Mirrors the identical guard in
    ``tests/test_issue_235_3d_pie_verify.py``.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


def _pie_chart_data(values=(25, 35, 40)):
    """Return a three-slice ``CategoryChartData`` for pie-chart scenarios."""
    data = CategoryChartData()
    data.categories = ["Apples", "Bananas", "Cherries"]
    data.add_series("Fruit", values)
    return data


def _add_three_d_pie(prs, chart_type, values=(25, 35, 40)):
    """Author a ``chart_type`` chart on a new slide and return the ``Chart``."""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    gf = slide.shapes.add_chart(
        chart_type,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        _pie_chart_data(values),
    )
    return gf.chart


# -- the verification suite ---------------------------------------------


class DescribeIssue719ThreeDPieWriterVerify(object):
    """#719 verify-and-close: THREE_D_PIE end-to-end including ``replace_data``.

    Pins that ``add_chart`` + ``replace_data`` + save/reopen all work
    for both plain and exploded 3-D pie charts. The pre-fork error
    path raised ``NotImplementedError`` from ``ChartXmlWriter`` (fixed
    by #266) and ``ValueError`` from ``PlotFactory`` when
    ``replace_data`` consulted ``chart.chart_type`` (fixed by #321).
    """

    # -- authoring path -----------------------------------------------------

    def it_accepts_THREE_D_PIE_via_add_chart(self):
        """``add_chart`` emits a real ``c:pie3DChart`` element."""
        prs = Presentation()

        chart = _add_three_d_pie(prs, XL_CHART_TYPE.THREE_D_PIE)

        assert isinstance(chart, Chart)
        assert chart._chartSpace.find(f".//{qn('c:pie3DChart')}") is not None

    def it_resolves_THREE_D_PIE_through_chart_plots(self):
        """``chart.plots[0]`` returns a :class:`Pie3DPlot` (the #321 fix)."""
        prs = Presentation()
        chart = _add_three_d_pie(prs, XL_CHART_TYPE.THREE_D_PIE)

        # -- this is the call that raised ``ValueError: unsupported plot
        # -- type c:pie3DChart`` before #321 shipped Pie3DPlot. --
        first_plot = chart.plots[0]

        assert isinstance(first_plot, Pie3DPlot)

    def it_resolves_chart_type_to_THREE_D_PIE(self):
        """``chart.chart_type`` returns ``XL_CHART_TYPE.THREE_D_PIE``."""
        prs = Presentation()
        chart = _add_three_d_pie(prs, XL_CHART_TYPE.THREE_D_PIE)

        # -- ``chart_type`` invokes PlotTypeInspector on the first plot
        # -- — the other lookup path #321 had to fix. --
        assert chart.chart_type is XL_CHART_TYPE.THREE_D_PIE

    def it_resolves_chart_type_to_THREE_D_PIE_EXPLODED_for_the_exploded_variant(self):
        """Exploded variant lights up via the ``c:explosion`` detector."""
        prs = Presentation()
        chart = _add_three_d_pie(prs, XL_CHART_TYPE.THREE_D_PIE_EXPLODED)

        assert chart.chart_type is XL_CHART_TYPE.THREE_D_PIE_EXPLODED

    # -- replace_data (the literal #719 failure mode) -----------------------

    def it_supports_replace_data_on_THREE_D_PIE(self):
        """``chart.replace_data`` succeeds — the pre-fork ValueError is gone.

        Before #321, this call raised ``ValueError: unsupported plot
        type c:pie3DChart`` from ``PlotFactory`` via
        ``Chart._validate_chart_data_type`` -> ``self.chart_type`` ->
        ``self.plots[0]``.
        """
        prs = Presentation()
        chart = _add_three_d_pie(prs, XL_CHART_TYPE.THREE_D_PIE)

        new_data = CategoryChartData()
        new_data.categories = ["Red", "Green", "Blue", "Yellow"]
        new_data.add_series("Colors", (10, 20, 30, 40))

        # -- act: this is the exact call that blew up in #719 --
        chart.replace_data(new_data)

        # -- assert: the chart now carries the new series data --
        chartSpace = chart._chartSpace
        cat_pts = chartSpace.xpath(".//c:pie3DChart/c:ser/c:cat//c:pt")
        assert len(cat_pts) == 4
        val_pts = chartSpace.xpath(".//c:pie3DChart/c:ser/c:val//c:pt")
        assert len(val_pts) == 4

    def it_supports_replace_data_on_THREE_D_PIE_EXPLODED(self):
        """``replace_data`` also works for the exploded variant."""
        prs = Presentation()
        chart = _add_three_d_pie(prs, XL_CHART_TYPE.THREE_D_PIE_EXPLODED)

        new_data = CategoryChartData()
        new_data.categories = ["One", "Two"]
        new_data.add_series("Nums", (42, 58))

        chart.replace_data(new_data)

        chartSpace = chart._chartSpace
        assert len(chartSpace.xpath(".//c:pie3DChart/c:ser/c:cat//c:pt")) == 2
        assert len(chartSpace.xpath(".//c:pie3DChart/c:ser/c:val//c:pt")) == 2
        # -- the exploded-variant detector still fires post replace_data --
        assert chart.chart_type is XL_CHART_TYPE.THREE_D_PIE_EXPLODED

    # -- round-trip through save + reopen -----------------------------------

    def it_round_trips_a_THREE_D_PIE_chart(self, _restore_part_factory):
        """Save + reopen preserves the ``c:pie3DChart`` wrapper and Pie3DPlot."""
        prs = Presentation()
        _add_three_d_pie(prs, XL_CHART_TYPE.THREE_D_PIE)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        gf = next(s for s in prs2.slides[0].shapes if s.has_chart)
        reloaded = gf.chart
        assert reloaded._chartSpace.find(f".//{qn('c:pie3DChart')}") is not None
        assert isinstance(reloaded.plots[0], Pie3DPlot)
        assert reloaded.chart_type is XL_CHART_TYPE.THREE_D_PIE

    def it_round_trips_a_THREE_D_PIE_EXPLODED_chart(self, _restore_part_factory):
        """Save + reopen preserves the exploded variant."""
        prs = Presentation()
        _add_three_d_pie(prs, XL_CHART_TYPE.THREE_D_PIE_EXPLODED)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        gf = next(s for s in prs2.slides[0].shapes if s.has_chart)
        reloaded = gf.chart
        assert reloaded.chart_type is XL_CHART_TYPE.THREE_D_PIE_EXPLODED
        # -- c:explosion is preserved through the round-trip --
        explosion_elems = reloaded._chartSpace.xpath(".//c:pie3DChart/c:ser/c:explosion")
        assert len(explosion_elems) == 1
        assert explosion_elems[0].get("val") == "25"

    def it_round_trips_a_replace_data_rewrite_through_save_reopen(self, _restore_part_factory):
        """End-to-end: author, replace_data, save, reopen — new data survives."""
        prs = Presentation()
        chart = _add_three_d_pie(prs, XL_CHART_TYPE.THREE_D_PIE)

        new_data = CategoryChartData()
        new_data.categories = ["East", "West", "North", "South", "Center"]
        new_data.add_series("Regions", (11, 22, 33, 44, 55))
        chart.replace_data(new_data)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        gf = next(s for s in prs2.slides[0].shapes if s.has_chart)
        reloaded = gf.chart
        assert reloaded.chart_type is XL_CHART_TYPE.THREE_D_PIE
        assert len(reloaded._chartSpace.xpath(".//c:pie3DChart/c:ser/c:cat//c:pt")) == 5
        assert len(reloaded._chartSpace.xpath(".//c:pie3DChart/c:ser/c:val//c:pt")) == 5
