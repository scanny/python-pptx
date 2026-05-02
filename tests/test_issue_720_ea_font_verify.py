# pyright: reportPrivateUsage=false

"""Regression test for issue #720 — ``.font.name`` differs between WPS and PowerPoint for CJK text.

Issue #720 (https://github.com/scanny/python-pptx/issues/720) reports that a
caller who assigned ``run.font.name = "SimSun"`` (a Chinese typeface) on a CJK
run observed inconsistent rendering: WPS Office applied the requested CJK
typeface while Microsoft PowerPoint ignored it and rendered the text in the
theme's Latin typeface. The symptom is the OOXML font-slot design — a single
``a:rPr`` carries up to three typeface hints (``a:latin`` / ``a:ea`` / ``a:cs``)
and PowerPoint dispatches each run of text to the slot that matches its Unicode
script (CJK glyphs route to ``a:ea``). ``Font.name`` on its own only writes
``a:latin``, so CJK text is at the mercy of whatever ``a:ea`` the theme or
master supplies.

Wave 1 #337 (``feat/issue-337-font-ea-cs-slots``) shipped the two missing
accessors — :attr:`.Font.name_ea` (East-Asian) and :attr:`.Font.name_cs`
(complex-script) — alongside the existing :attr:`.Font.name` (Latin), making
all three OOXML font slots directly addressable from Python. With those in
place the #720 scenario is correctly authored by setting
``font.name = "Arial"`` *and* ``font.name_ea = "SimSun"`` on the same run;
PowerPoint then honours ``a:ea`` for CJK code points and both editors render
identically.

This regression suite pins the three-slot contract that resolves #720:

  * ``Font.name``, ``Font.name_ea``, and ``Font.name_cs`` are independent
    and can be set together on the same run without clobbering each other.
  * The emitted ``a:rPr`` carries all three child elements
    (``a:latin``, ``a:ea``, ``a:cs``) with the expected ``@typeface`` values.
  * The assignment survives a full ``Presentation.save`` + reopen cycle,
    so a ``.pptx`` written with the 3-slot pattern round-trips without
    losing the CJK or complex-script overrides that are the whole point
    of the fix.

Any future change that regresses to the single-slot behaviour (e.g., a
refactor that routes all three properties through ``a:latin``, or drops
``a:ea`` / ``a:cs`` from ``CT_TextCharacterProperties._tag_seq``) will
reproduce the #720 symptom and be caught here.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of ``PartFactory.part_type_for``.

    Mirrors the guard used by ``tests/test_issue_285_replace_text_preserve_format.py``
    and friends — ``Presentation(stream)`` reopen dispatches through
    ``PartFactory.part_type_for``, which earlier test modules overwrite with
    mocks and do not restore.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue720EaFontVerify:
    """Pins the three-slot (a:latin / a:ea / a:cs) font authoring contract."""

    def it_exposes_all_three_font_slot_properties_on_Font(self):
        """Guard: the public API surface that resolves #720 is present."""
        from pptx.text.text import Font

        # -- the two accessors added by #337 sit alongside the original ``name`` --
        assert hasattr(Font, "name")
        assert hasattr(Font, "name_ea")
        assert hasattr(Font, "name_cs")

    def it_lets_latin_and_east_asian_typefaces_coexist_on_one_run(self):
        """The core #720 pattern: CJK deck authored with the 3-slot approach.

        Setting ``font.name = "Arial"`` *and* ``font.name_ea = "SimSun"`` on
        the same run yields a run that renders consistently in both WPS and
        PowerPoint — PowerPoint's per-script dispatcher sends CJK code points
        to ``a:ea`` (SimSun) and Latin code points to ``a:latin`` (Arial).
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = textbox.text_frame.paragraphs[0].add_run()
        run.text = "Hello 世界"  # -- "Hello 世界" (mixed Latin + CJK) --

        run.font.name = "Arial"
        run.font.name_ea = "SimSun"

        # -- each getter reads back its own slot; neither clobbers the other --
        assert run.font.name == "Arial"
        assert run.font.name_ea == "SimSun"
        assert run.font.name_cs is None

    def it_lets_all_three_slots_be_set_independently_on_one_run(self):
        """The full 3-slot scenario — Latin + East-Asian + complex-script.

        A run that has to serve Latin, CJK, *and* e.g. Arabic or Hebrew text
        uses all three slots. Verifies they're truly independent (setting
        one never touches the other two).
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = textbox.text_frame.paragraphs[0].add_run()
        run.text = "Hello"

        run.font.name = "Calibri"
        run.font.name_ea = "MS Mincho"
        run.font.name_cs = "Arial Unicode MS"

        assert run.font.name == "Calibri"
        assert run.font.name_ea == "MS Mincho"
        assert run.font.name_cs == "Arial Unicode MS"

    def it_writes_a_latin_a_ea_and_a_cs_children_under_rPr(self):
        """Pin the XML shape PowerPoint reads.

        For the 3-slot pattern to resolve #720, the emitted ``a:rPr`` must
        carry three distinct child elements — ``a:latin``, ``a:ea``, and
        ``a:cs`` — each with its own ``@typeface``. This test inspects the
        live XML (not just the Python getters) so a future refactor that,
        say, routed both Latin and EA through ``a:latin`` would still be
        caught here.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = textbox.text_frame.paragraphs[0].add_run()
        run.text = "Hello 世界"
        run.font.name = "Arial"
        run.font.name_ea = "SimSun"
        run.font.name_cs = "Arial Unicode MS"

        rPr = run._r.get_or_add_rPr()
        latin_children = rPr.findall(qn("a:latin"))
        ea_children = rPr.findall(qn("a:ea"))
        cs_children = rPr.findall(qn("a:cs"))

        assert len(latin_children) == 1
        assert latin_children[0].get("typeface") == "Arial"
        assert len(ea_children) == 1
        assert ea_children[0].get("typeface") == "SimSun"
        assert len(cs_children) == 1
        assert cs_children[0].get("typeface") == "Arial Unicode MS"

    def it_preserves_ea_and_cs_when_name_is_reassigned(self):
        """Regression guard: reassigning ``font.name`` must not disturb
        ``a:ea`` / ``a:cs``. The #720 scenario would regress if setting
        ``font.name`` on a re-edit removed the previously set CJK override.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = textbox.text_frame.paragraphs[0].add_run()
        run.text = "Hello 世界"
        run.font.name = "Arial"
        run.font.name_ea = "SimSun"
        run.font.name_cs = "Mangal"

        # -- caller later swaps the Latin face; EA/CS must stay put --
        run.font.name = "Times New Roman"

        assert run.font.name == "Times New Roman"
        assert run.font.name_ea == "SimSun"
        assert run.font.name_cs == "Mangal"

    def it_round_trips_all_three_slots_through_save_and_reload(self, _restore_part_factory):
        """End-to-end: author a CJK run with the 3-slot pattern, save to a
        ``.pptx``, reopen, and verify every slot survives.

        This is the test #720 really needs — a ``.pptx`` written this way
        must be interpretable by both WPS and PowerPoint, which means the
        three ``a:latin`` / ``a:ea`` / ``a:cs`` children must survive the
        save and be re-read identically on load.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = textbox.text_frame.paragraphs[0].add_run()
        run.text = "Hello 世界"
        run.font.name = "Arial"
        run.font.name_ea = "SimSun"
        run.font.name_cs = "Arial Unicode MS"
        run.font.size = Pt(24)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        reloaded = prs2.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]

        assert reloaded.text == "Hello 世界"
        assert reloaded.font.name == "Arial"
        assert reloaded.font.name_ea == "SimSun"
        assert reloaded.font.name_cs == "Arial Unicode MS"
        assert reloaded.font.size == Pt(24)

    def it_clears_only_the_targeted_slot_when_assigning_None(self):
        """Assigning ``None`` to one slot removes just that child element —
        the other two slots are untouched. Mirrors the behaviour documented
        for ``font.name = None`` and extends it to the two new accessors.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = textbox.text_frame.paragraphs[0].add_run()
        run.text = "Hello"
        run.font.name = "Arial"
        run.font.name_ea = "SimSun"
        run.font.name_cs = "Mangal"

        # -- clear EA only --
        run.font.name_ea = None

        assert run.font.name == "Arial"
        assert run.font.name_ea is None
        assert run.font.name_cs == "Mangal"

        rPr = run._r.get_or_add_rPr()
        assert rPr.find(qn("a:latin")) is not None
        assert rPr.find(qn("a:ea")) is None
        assert rPr.find(qn("a:cs")) is not None
