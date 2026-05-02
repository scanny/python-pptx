# pyright: reportPrivateUsage=false

"""Regression test for issue #740 — paragraph font size ignored when text starts with newline.

Issue #740 (https://github.com/scanny/python-pptx/issues/740) reports that
assigning ``paragraph.text = "\\nHello"`` followed by
``paragraph.font.size = Pt(24)`` produces a paragraph whose run is rendered
at 18pt instead of the requested 24pt.

Root cause: the paragraph text setter inserts an ``a:br`` line-break element
(and then an ``a:r`` for ``"Hello"``) before the caller sets ``font.size``.
When the subsequent font-size assignment calls ``get_or_add_pPr()``, the
``a:pPr`` insertion logic in ``BaseOxmlElement.insert_element_before`` was
searching for the *first successor tag name listed* rather than the
*earliest-positioned existing child* among the successor set.

With the paragraph already holding ``[a:br, a:r]`` the old search walked
``("a:r", "a:br", "a:fld", "a:endParaRPr")`` and found ``a:r`` first, so
``a:pPr`` was inserted before ``a:r`` — after ``a:br``. The resulting
``<a:p><a:br/><a:pPr>…</a:pPr><a:r/></a:p>`` violates the schema's strict
ordering (``a:pPr`` must precede every other child), and PowerPoint drops
the out-of-order ``a:pPr`` on load, discarding the caller's font size.

The fix sets ``insert_element_before`` to pick the earliest existing child
whose tag belongs to the successor set, which preserves schema order no
matter which successor types are already present.

This regression test covers the reporter's workflow end-to-end: assign a
paragraph's text starting with a newline, set the paragraph's font size,
round-trip the file, and confirm the paragraph-default run properties
record the requested 24pt size in the correct XML position.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


class DescribeIssue740LeadingNewlineFontSize:
    """End-to-end regression for paragraph font-size lost behind leading newline."""

    def it_applies_paragraph_font_size_after_leading_newline(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        paragraph = textbox.text_frame.paragraphs[0]

        paragraph.text = "\nHello"
        paragraph.font.size = Pt(24)

        assert paragraph.font.size == Pt(24)

    def it_places_pPr_before_br_in_schema_order(self):
        """`a:pPr` must be the first child of `a:p` to round-trip through PowerPoint."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        paragraph = textbox.text_frame.paragraphs[0]

        paragraph.text = "\nHello"
        paragraph.font.size = Pt(24)

        children = list(paragraph._p)
        tags = [c.tag for c in children]
        # -- a:pPr must appear, and must come first --
        assert tags[0] == qn("a:pPr"), f"expected a:pPr first, got {tags}"
        # -- a:br and a:r must still follow in their original positions --
        assert qn("a:br") in tags
        assert qn("a:r") in tags
        assert tags.index(qn("a:br")) < tags.index(qn("a:r"))

    def it_survives_round_trip_through_save_and_reopen(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        paragraph = textbox.text_frame.paragraphs[0]
        paragraph.text = "\nHello"
        paragraph.font.size = Pt(24)

        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)

        prs2 = Presentation(buffer)
        reopened_paragraph = prs2.slides[0].shapes[-1].text_frame.paragraphs[0]
        assert reopened_paragraph.font.size == Pt(24)

    def it_preserves_existing_runs_when_font_is_set_after_content(self):
        """Other content (line-breaks, runs) must not be reordered by the pPr insertion."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        paragraph = textbox.text_frame.paragraphs[0]
        paragraph.text = "\nHello\nWorld"
        paragraph.font.size = Pt(24)

        # -- paragraph text uses vertical-tab for a:br boundaries; see the
        # -- _Paragraph.text getter docstring for the rationale --
        assert paragraph.text == "\vHello\vWorld"

        # -- the relative order of a:br and a:r elements is preserved --
        tags = [c.tag for c in paragraph._p]
        assert tags[0] == qn("a:pPr")
        content_tags = [t for t in tags if t != qn("a:pPr")]
        assert content_tags == [qn("a:br"), qn("a:r"), qn("a:br"), qn("a:r")]
