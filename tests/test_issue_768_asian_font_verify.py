# pyright: reportPrivateUsage=false

"""Regression test for issue #768 — Change font name not working for asian characters.

Issue #768 (https://github.com/scanny/python-pptx/issues/768) reports that
assigning ``run.font.name`` to a CJK typeface does not change how East-Asian
characters render in PowerPoint. The upstream tracker marks #768 as a
**duplicate of #720** — the root cause is the OOXML font-slot design: an
``a:rPr`` carries three independent typeface children (``a:latin`` /
``a:ea`` / ``a:cs``) and PowerPoint dispatches each code point to the slot
that matches its Unicode script. ``Font.name`` only writes ``a:latin``, so
Asian (CJK) glyphs remain bound to whatever ``a:ea`` the theme supplies.

The foundational fix shipped in Wave 1 #337
(``feat/issue-337-font-ea-cs-slots``), adding :attr:`.Font.name_ea`
(East-Asian) and :attr:`.Font.name_cs` (complex-script) alongside the
original :attr:`.Font.name` (Latin). With all three slots addressable, the
#768 scenario is correctly authored by setting ``font.name_ea`` to the
Asian typeface. See ``tests/test_issue_720_ea_font_verify.py`` for the
foundational three-slot regression suite; this module re-verifies the
resolution from the #768 reporter's perspective.

The suite pins:

  * ``font.name = "Calibri"`` alone emits ``a:latin`` and leaves ``a:ea``
    / ``a:cs`` absent — the pre-#337 behaviour the #768 reporter hit.
  * ``font.name_ea = "MS Mincho"`` emits an ``a:ea`` child with the
    requested typeface, which is what PowerPoint needs to render CJK glyphs
    in the caller's chosen face.
  * ``font.name_cs = "Arial Unicode MS"`` emits an ``a:cs`` child — the
    complex-script variant of the same fix.
  * All three slots can coexist on a single run, independently.
  * Every slot survives a ``Presentation.save`` + reopen round-trip, so a
    ``.pptx`` written with Asian typefaces loads cleanly in PowerPoint and
    in python-pptx itself.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of ``PartFactory.part_type_for``.

    Mirrors the guard used by ``tests/test_issue_720_ea_font_verify.py`` and
    friends — ``Presentation(stream)`` reopen dispatches through
    ``PartFactory.part_type_for``, which earlier test modules overwrite with
    mocks and do not restore.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


def _new_run_on_blank_slide():
    """Return a fresh ``(prs, run)`` — a blank-layout slide with an empty run."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    run = textbox.text_frame.paragraphs[0].add_run()
    return prs, run


class DescribeIssue768AsianFont:
    """Re-verifies #768 (duplicate of #720) — Asian typeface via ``a:ea`` slot.

    Cross-reference: ``tests/test_issue_720_ea_font_verify.py`` holds the
    foundational three-slot contract. Per the upstream tracker, #768 is a
    duplicate of #720.
    """

    def it_sets_only_a_latin_when_font_name_is_assigned_alone(self):
        """The pre-#337 behaviour the #768 reporter originally observed.

        Setting ``font.name`` alone writes ``a:latin`` and leaves ``a:ea`` /
        ``a:cs`` absent — which is why Asian characters in the run continued
        to render in the theme's East-Asian face (the #768 symptom). This
        test pins that single-slot shape so the EA/CS absence — and thus the
        need for ``name_ea`` / ``name_cs`` — is guarded.
        """
        _prs, run = _new_run_on_blank_slide()
        run.text = "Hello"

        run.font.name = "Calibri"

        rPr = run._r.get_or_add_rPr()
        latin_children = rPr.findall(qn("a:latin"))
        ea_children = rPr.findall(qn("a:ea"))
        cs_children = rPr.findall(qn("a:cs"))

        assert len(latin_children) == 1
        assert latin_children[0].get("typeface") == "Calibri"
        assert len(ea_children) == 0
        assert len(cs_children) == 0
        # -- the Python getters agree with the XML --
        assert run.font.name == "Calibri"
        assert run.font.name_ea is None
        assert run.font.name_cs is None

    def it_writes_a_ea_with_the_requested_typeface_when_name_ea_is_set(self):
        """The #768 fix: ``font.name_ea = "MS Mincho"`` emits ``a:ea``.

        That is the typeface PowerPoint reads for East-Asian glyphs — the
        slot the #768 reporter needed to address but couldn't with
        ``font.name`` alone.
        """
        _prs, run = _new_run_on_blank_slide()
        run.text = "日本語"  # -- Japanese text --

        run.font.name_ea = "MS Mincho"

        rPr = run._r.get_or_add_rPr()
        ea_children = rPr.findall(qn("a:ea"))

        assert len(ea_children) == 1
        assert ea_children[0].get("typeface") == "MS Mincho"
        assert run.font.name_ea == "MS Mincho"

    def it_writes_a_cs_with_the_requested_typeface_when_name_cs_is_set(self):
        """The complex-script companion: ``font.name_cs`` emits ``a:cs``.

        Shipped alongside ``name_ea`` in #337 so complex-script callers
        (Arabic / Hebrew / Thai / Devanagari) have the same escape hatch.
        """
        _prs, run = _new_run_on_blank_slide()
        run.text = "مرحبا"  # -- Arabic text --

        run.font.name_cs = "Arial Unicode MS"

        rPr = run._r.get_or_add_rPr()
        cs_children = rPr.findall(qn("a:cs"))

        assert len(cs_children) == 1
        assert cs_children[0].get("typeface") == "Arial Unicode MS"
        assert run.font.name_cs == "Arial Unicode MS"

    def it_sets_all_three_slots_independently_on_a_single_run(self):
        """The full multi-script pattern — Latin + East-Asian + complex-script.

        A run that has to serve Latin, CJK, *and* complex-script text uses
        all three slots together. Verifies they are truly independent (no
        slot clobbers another) both in the Python API and the live XML.
        """
        _prs, run = _new_run_on_blank_slide()
        run.text = "Hello 世界"

        run.font.name = "Calibri"
        run.font.name_ea = "MS Mincho"
        run.font.name_cs = "Arial Unicode MS"

        assert run.font.name == "Calibri"
        assert run.font.name_ea == "MS Mincho"
        assert run.font.name_cs == "Arial Unicode MS"

        rPr = run._r.get_or_add_rPr()
        latin_children = rPr.findall(qn("a:latin"))
        ea_children = rPr.findall(qn("a:ea"))
        cs_children = rPr.findall(qn("a:cs"))
        assert len(latin_children) == 1
        assert latin_children[0].get("typeface") == "Calibri"
        assert len(ea_children) == 1
        assert ea_children[0].get("typeface") == "MS Mincho"
        assert len(cs_children) == 1
        assert cs_children[0].get("typeface") == "Arial Unicode MS"

    def it_round_trips_all_three_slots_through_save_and_reload(
        self, _restore_part_factory
    ):
        """End-to-end: author a mixed-script run, save, reopen, verify.

        This is the test #768 really needs — a ``.pptx`` written with all
        three typeface overrides must re-read identically on load so the
        authored Asian/complex-script faces survive the trip into PowerPoint.
        """
        prs, run = _new_run_on_blank_slide()
        run.text = "Hello 世界"
        run.font.name = "Calibri"
        run.font.name_ea = "MS Mincho"
        run.font.name_cs = "Arial Unicode MS"
        run.font.size = Pt(24)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        reloaded = prs2.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]

        assert reloaded.text == "Hello 世界"
        assert reloaded.font.name == "Calibri"
        assert reloaded.font.name_ea == "MS Mincho"
        assert reloaded.font.name_cs == "Arial Unicode MS"
        assert reloaded.font.size == Pt(24)

    def it_is_the_same_three_slot_contract_that_resolves_issue_720(self):
        """Cross-reference guard.

        #768 is marked in the upstream tracker as a duplicate of #720. The
        resolution is the ``Font.name_ea`` / ``Font.name_cs`` accessors
        shipped by Wave 1 #337 and pinned by
        ``tests/test_issue_720_ea_font_verify.py``. This test asserts the
        three-slot surface area is live on the public ``Font`` class, so a
        future refactor that removes (or hides) ``name_ea`` / ``name_cs``
        regresses both #720 and #768 and is caught here.
        """
        from pptx.text.text import Font

        # -- all three slots are first-class on the public Font class --
        assert hasattr(Font, "name")
        assert hasattr(Font, "name_ea")
        assert hasattr(Font, "name_cs")

        # -- and each is an independently read-writable property --
        assert isinstance(Font.__dict__["name"], property)
        assert isinstance(Font.__dict__["name_ea"], property)
        assert isinstance(Font.__dict__["name_cs"], property)
        assert Font.__dict__["name"].fset is not None
        assert Font.__dict__["name_ea"].fset is not None
        assert Font.__dict__["name_cs"].fset is not None
