# pyright: reportPrivateUsage=false

"""Regression test for issue #784 — Movie.replace_media.

Issue #784 (https://github.com/scanny/python-pptx/issues/784) asks for an
API that swaps the underlying audio (or video) binary of an existing
movie/audio shape without disturbing its position, size, cropping, or
timing. The feature mirrors ``Picture.replace_image`` (issue #116) and
uses the same relationship-rewiring pattern delivered by Foundation F1
(``PartRelationshipCloner``) and Foundation F5 (embedded-workbook
handling).

This end-to-end test builds a minimal presentation in memory with a
single audio shape created via ``SlideShapes.add_movie``, exercises
``Movie.replace_media`` to swap the audio blob, and then round-trips the
package (save → reload) to confirm:

  1. ``Movie.replace_media`` adds a new media part whose blob matches
     the replacement bytes and rewires the shape's ``a:audioFile`` and
     ``p14:media`` rIds at it.
  2. The shape's position and size are preserved across the swap.
  3. After save + reload, the reloaded shape still yields the new
     blob via its media part.

Deferred (documented as issue #784 follow-up): modality switching
(swapping an audio shape for a video shape, or vice versa) is not
supported by ``replace_media`` and requires
``SlideShapes.add_movie`` + ``Shape.delete``.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.picture import Movie
from pptx.util import Inches


class DescribeIssue784RegressionReplaceAudio:
    """End-to-end coverage for :meth:`Movie.replace_media`."""

    def it_swaps_the_media_blob_in_place(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
        audio_bytes_v1 = b"\x52\x49\x46\x46" + b"\x00" * 36  # fake "RIFF" wav-ish
        audio_v1 = io.BytesIO(audio_bytes_v1)
        slide.shapes.add_movie(
            audio_v1,
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(2),
            mime_type="audio/wav",
        )
        # -- pull the Movie proxy from the slide's shapes iterator (the
        #    rehydrated shape is the real Movie subclass, even when
        #    add_movie's public return-type annotation is GraphicFrame) --
        media_shape = next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.MEDIA)
        assert isinstance(media_shape, Movie)

        # -- snapshot position / size before replace --
        original_left = media_shape.left
        original_top = media_shape.top
        original_width = media_shape.width
        original_height = media_shape.height

        # -- swap the media payload --
        audio_bytes_v2 = b"\x52\x49\x46\x46" + b"\x01" * 36  # different payload
        media_shape.replace_media(io.BytesIO(audio_bytes_v2), mime_type="audio/wav")

        # -- position/size preserved on the in-memory shape --
        assert media_shape.left == original_left
        assert media_shape.top == original_top
        assert media_shape.width == original_width
        assert media_shape.height == original_height

        # -- the new media part's blob matches the replacement bytes --
        new_video_rId = media_shape._pic.media_video_rId
        assert new_video_rId is not None
        media_part = media_shape.part.related_part(new_video_rId)
        assert media_part.blob == audio_bytes_v2

        # -- round-trip through save + reload and verify persistence --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)
        reloaded_slide = reloaded.slides[0]
        reloaded_movie = next(
            sh for sh in reloaded_slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.MEDIA
        )
        reloaded_rId = reloaded_movie._pic.media_video_rId
        assert reloaded_rId is not None
        reloaded_media_part = reloaded_movie.part.related_part(reloaded_rId)
        assert reloaded_media_part.blob == audio_bytes_v2
        # -- position/size preserved across round-trip --
        assert reloaded_movie.left == original_left
        assert reloaded_movie.top == original_top
        assert reloaded_movie.width == original_width
        assert reloaded_movie.height == original_height

    def it_preserves_the_audio_tag_after_replace(self):
        """Audio-shape semantics survive a replace — the ``a:audioFile`` tag stays."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        audio_v1 = io.BytesIO(b"first audio payload")
        slide.shapes.add_movie(
            audio_v1,
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(2),
            mime_type="audio/mpeg",
        )
        movie = next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.MEDIA)

        # -- confirm the pic was authored with an a:audioFile child (not videoFile) --
        pre_audioFiles = movie._pic.xpath("./p:nvPicPr/p:nvPr/a:audioFile")
        assert len(pre_audioFiles) == 1

        movie.replace_media(io.BytesIO(b"second audio payload"), mime_type="audio/mpeg")

        # -- still an audioFile, not a videoFile --
        post_audioFiles = movie._pic.xpath("./p:nvPicPr/p:nvPr/a:audioFile")
        post_videoFiles = movie._pic.xpath("./p:nvPicPr/p:nvPr/a:videoFile")
        assert len(post_audioFiles) == 1
        assert len(post_videoFiles) == 0
