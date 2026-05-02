# pyright: reportPrivateUsage=false

"""Regression test for issue #789 — data-label position corruption.

Issue #789 (https://github.com/scanny/python-pptx/issues/789): setting
``plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END`` on a
stacked column/bar chart produces a file PowerPoint refuses to open.
PowerPoint's UI restricts the set of legal ``c:dLblPos`` values per
chart type (see ECMA-376 §21.2.2.45). python-pptx previously let the
caller write any combination, silently producing corrupted files.

Resolution: ``DataLabels.position`` (both plot-level and series-level)
now validates against a per-chart-type whitelist and raises
``ValueError`` with a clear message when an incompatible value is
assigned. The XML is not mutated on rejection. Assigning ``None`` or
re-reading ``position`` is unaffected.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.util import Inches


def _add_category_chart(prs, chart_type):
    """Create a slide + chart of *chart_type* with two series, two categories."""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    cd = CategoryChartData()
    cd.categories = ["A", "B"]
    cd.add_series("S1", (1.0, 2.0))
    cd.add_series("S2", (3.0, 4.0))
    gf = slide.shapes.add_chart(
        chart_type,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        cd,
    )
    chart = gf.chart
    chart.plots[0].has_data_labels = True
    return chart


def _add_scatter_chart(prs):
    """Create a slide + scatter chart."""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    cd = XyChartData()
    s = cd.add_series("S1")
    s.add_data_point(1.0, 2.0)
    s.add_data_point(3.0, 4.0)
    gf = slide.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        cd,
    )
    chart = gf.chart
    chart.plots[0].has_data_labels = True
    return chart


class DescribeIssue789DataLabelPositionValidation:
    """``DataLabels.position`` setter validates per-chart-type compatibility."""

    def it_rejects_OUTSIDE_END_on_a_stacked_column_chart(self):
        """The exact case from the reporter — this was the corruption trigger."""
        prs = Presentation()
        chart = _add_category_chart(prs, XL_CHART_TYPE.COLUMN_STACKED)

        with pytest.raises(ValueError, match="COLUMN_STACKED"):
            chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END

    def it_rejects_OUTSIDE_END_on_a_stacked_bar_chart(self):
        prs = Presentation()
        chart = _add_category_chart(prs, XL_CHART_TYPE.BAR_STACKED)

        with pytest.raises(ValueError, match="BAR_STACKED"):
            chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END

    def it_accepts_OUTSIDE_END_on_a_clustered_column_chart(self):
        prs = Presentation()
        chart = _add_category_chart(prs, XL_CHART_TYPE.COLUMN_CLUSTERED)

        chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        assert chart.plots[0].data_labels.position == XL_LABEL_POSITION.OUTSIDE_END

    def it_accepts_INSIDE_END_on_a_stacked_column_chart(self):
        """The canonical legal position for stacked columns."""
        prs = Presentation()
        chart = _add_category_chart(prs, XL_CHART_TYPE.COLUMN_STACKED)

        chart.plots[0].data_labels.position = XL_LABEL_POSITION.INSIDE_END
        assert chart.plots[0].data_labels.position == XL_LABEL_POSITION.INSIDE_END

    def it_rejects_bar_positions_on_a_line_chart(self):
        prs = Presentation()
        chart = _add_category_chart(prs, XL_CHART_TYPE.LINE)

        with pytest.raises(ValueError, match="LINE"):
            chart.plots[0].data_labels.position = XL_LABEL_POSITION.INSIDE_BASE

    def it_accepts_ABOVE_on_a_line_chart(self):
        prs = Presentation()
        chart = _add_category_chart(prs, XL_CHART_TYPE.LINE)

        chart.plots[0].data_labels.position = XL_LABEL_POSITION.ABOVE
        assert chart.plots[0].data_labels.position == XL_LABEL_POSITION.ABOVE

    def it_rejects_bar_positions_on_a_scatter_chart(self):
        prs = Presentation()
        chart = _add_scatter_chart(prs)

        with pytest.raises(ValueError, match="XY_SCATTER"):
            chart.plots[0].data_labels.position = XL_LABEL_POSITION.INSIDE_END

    def it_rejects_on_a_series_level_data_labels_too(self):
        """Series-level ``DataLabels`` also validates (same corruption path)."""
        prs = Presentation()
        chart = _add_category_chart(prs, XL_CHART_TYPE.COLUMN_STACKED)
        series = chart.plots[0].series[0]

        with pytest.raises(ValueError, match="COLUMN_STACKED"):
            series.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END

    def it_does_not_mutate_XML_on_rejection(self):
        """A rejected assignment must leave the XML untouched."""
        prs = Presentation()
        chart = _add_category_chart(prs, XL_CHART_TYPE.COLUMN_STACKED)
        data_labels = chart.plots[0].data_labels
        before = data_labels._element.xml

        with pytest.raises(ValueError):
            data_labels.position = XL_LABEL_POSITION.OUTSIDE_END

        assert data_labels._element.xml == before

    def it_allows_assigning_None_to_clear_position(self):
        """``None`` is always legal — it removes c:dLblPos entirely."""
        prs = Presentation()
        chart = _add_category_chart(prs, XL_CHART_TYPE.COLUMN_STACKED)
        chart.plots[0].data_labels.position = XL_LABEL_POSITION.INSIDE_END
        assert chart.plots[0].data_labels.position == XL_LABEL_POSITION.INSIDE_END

        chart.plots[0].data_labels.position = None

        assert chart.plots[0].data_labels.position is None

    def it_round_trips_a_legal_position_across_save_and_reopen(self):
        """Sanity check: a legal position written by python-pptx survives a
        Presentation.save / reload — i.e. the file is still readable."""
        prs = Presentation()
        chart = _add_category_chart(prs, XL_CHART_TYPE.COLUMN_STACKED)
        chart.plots[0].data_labels.position = XL_LABEL_POSITION.INSIDE_END

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        prs2 = Presentation(buf)
        slide2 = list(prs2.slides)[-1]
        chart2 = next(shp.chart for shp in slide2.shapes if shp.has_chart)
        assert chart2.plots[0].data_labels.position == XL_LABEL_POSITION.INSIDE_END
