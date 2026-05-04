# pyright: reportUnknownMemberType=false, reportUnknownVariableType=false, reportUnknownParameterType=false, reportMissingParameterType=false, reportAttributeAccessIssue=false, reportOptionalMemberAccess=false

"""Regression tests for issue #793 - watermark recipe.

The reporter asked how to add a watermark (a transparent picture or a
lightly-colored piece of text) so that it appears on every slide. In
PowerPoint that is accomplished by adding the watermark shape to the
slide master, where it is inherited by every slide.

These tests lock in the documented recipe in
``docs/user/slides.rst`` - "Adding a watermark" - so that the code
paths it relies on stay working:

* ``prs.slide_masters[0].shapes.add_picture(...)`` on |MasterShapes|
* ``Picture.transparency = <percentage>`` (see issue #165)
* ``master.shapes.add_textbox(...)`` for a text-only watermark
* Round-trip of the master-level watermark shapes through save + reload
* ``slide.show_master_shapes = False`` opt-out for a single slide
"""

from __future__ import annotations

import io
import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.picture import Picture
from pptx.util import Emu, Inches, Pt

TEST_IMAGE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "test_files",
    "python-powered.png",
)


class DescribeIssue793WatermarkRecipe(object):
    """Regression suite for the docs recipe covering issue #793."""

    def it_adds_a_picture_watermark_to_the_slide_master(self):
        # -- The central recipe: add a picture to the master, set a
        # -- non-zero transparency so it reads as a watermark, and
        # -- confirm the picture is present on the master.
        prs = Presentation()
        master = prs.slide_masters[0]

        pic = master.shapes.add_picture(TEST_IMAGE, Inches(2), Inches(2), Inches(4), Inches(3))
        pic.transparency = 80.0

        assert pic.shape_type == MSO_SHAPE_TYPE.PICTURE
        assert pic.transparency == 80.0
        # -- the default template master has 5 latent placeholders; the
        # -- watermark is the 6th shape --
        assert pic in list(master.shapes)

    def it_round_trips_the_picture_watermark_through_save_and_reload(self):
        # -- The master-level picture and its transparency survive a
        # -- save + reopen; this is what guarantees PowerPoint will
        # -- render the watermark on every slide.
        prs = Presentation()
        master = prs.slide_masters[0]

        pic = master.shapes.add_picture(TEST_IMAGE, Inches(2), Inches(2), Inches(4), Inches(3))
        pic.transparency = 75.0

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        reloaded_pics = [
            s for s in reloaded.slide_masters[0].shapes if isinstance(s, Picture)
        ]
        assert len(reloaded_pics) == 1
        assert abs(reloaded_pics[0].transparency - 75.0) < 1e-6

    def it_makes_the_watermark_visible_on_every_inheriting_slide(self):
        # -- The watermark is authored once on the master but shows up
        # -- on every slide that inherits master shapes. python-pptx
        # -- does not render, but the inheritance contract is that
        # -- ``slide.show_master_shapes`` defaults to True and the
        # -- shape lives on ``master.shapes``, not on any one slide.
        prs = Presentation()
        master = prs.slide_masters[0]
        master.shapes.add_picture(
            TEST_IMAGE, Inches(2), Inches(2), Inches(4), Inches(3)
        ).transparency = 80.0

        # -- add a few slides from different layouts --
        slide_a = prs.slides.add_slide(prs.slide_layouts[0])
        slide_b = prs.slides.add_slide(prs.slide_layouts[1])
        slide_c = prs.slides.add_slide(prs.slide_layouts[5])

        # -- none of the slides carry the watermark shape themselves --
        for slide in (slide_a, slide_b, slide_c):
            slide_pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
            assert slide_pics == []
            # -- but every slide inherits master shapes by default --
            assert slide.show_master_shapes is True

    def it_supports_a_text_watermark_via_master_shapes_add_textbox(self):
        # -- Text-only variant of the recipe: a lightly-colored textbox
        # -- on the master. Exercises ``master.shapes.add_textbox`` and
        # -- the font sizing hook used in the documented snippet.
        prs = Presentation()
        master = prs.slide_masters[0]

        tb = master.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(1.5))
        tf = tb.text_frame
        tf.text = "DRAFT"
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(96)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        reloaded_master = reloaded.slide_masters[0]
        watermark_texts = [s.text_frame.text for s in reloaded_master.shapes if s.has_text_frame]
        assert "DRAFT" in watermark_texts

    def it_lets_a_single_slide_opt_out_via_show_master_shapes(self):
        # -- The watermark is inherited from the master, so hiding it on
        # -- an individual slide is a matter of assigning
        # -- ``slide.show_master_shapes = False`` - documented in the
        # -- recipe as the escape hatch (e.g. for a title slide that
        # -- should not display the watermark).
        prs = Presentation()
        master = prs.slide_masters[0]
        master.shapes.add_picture(
            TEST_IMAGE, Inches(2), Inches(2), Inches(4), Inches(3)
        ).transparency = 80.0

        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.show_master_shapes = False

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        # -- the master still carries the watermark picture --
        reloaded_master_pics = [
            s for s in reloaded.slide_masters[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        assert len(reloaded_master_pics) == 1

        # -- but the title slide has opted out of master shapes --
        assert reloaded.slides[0].show_master_shapes is False

    def it_places_the_watermark_where_the_caller_specified(self):
        # -- The recipe passes explicit left/top/width/height to
        # -- ``add_picture`` so the watermark lands where the author
        # -- intends; verify those geometry arguments round-trip.
        prs = Presentation()
        master = prs.slide_masters[0]

        left, top, width, height = Inches(2), Inches(2), Inches(4), Inches(3)
        pic = master.shapes.add_picture(TEST_IMAGE, left, top, width, height)

        assert pic.left == Emu(left)
        assert pic.top == Emu(top)
        assert pic.width == Emu(width)
        assert pic.height == Emu(height)
