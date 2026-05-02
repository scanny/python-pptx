# pyright: reportPrivateUsage=false

"""Regression test for issue #805 — replace audio / video in an existing slide.

Issue #805 (https://github.com/scanny/python-pptx/issues/805) asked for an
API that swaps the underlying audio (or video) binary behind an already-
authored media shape without dropping the shape and re-adding it — in
other words, ``Picture.replace_image`` (#116) but for the audio / video
case. The reporter's scenario: a slide deck template contains placeholder
movie / audio shapes and a downstream workflow wants to fill in the
actual media files without disturbing each shape's position, size, or
poster-frame.

This ask is resolved by Wave 6 #784 (``feat/issue-784-replace-audio``)
which shipped
:meth:`Movie.replace_media(new_path_or_file, mime_type=None)`. That
method loads the replacement bytes into a fresh |MediaPart|, rewires the
shape's ``a:videoFile`` / ``a:audioFile`` ``@r:link`` and ``p14:media``
``@r:embed`` rIds, and drops the old relationships so the previous media
part is eligible for garbage-collection on save. The shape's position,
size, poster frame, hyperlink, and ``p:timing`` entries are untouched —
only the media bytes behind the shape change.

This suite verifies that ``Movie.replace_media`` satisfies the #805 use
cases so the issue can be closed as *resolved by #784*. It exercises the
real API against a fresh ``Presentation()`` (no mocks), across both
audio and video payloads, and round-trips through ``Presentation.save``
+ reopen to pin that the rewired relationships survive packaging.

Note: modality switching — swapping an audio shape for a video shape or
vice versa — is a documented follow-up; it requires
``SlideShapes.add_movie`` + ``Shape.delete`` rather than
``replace_media``. A test below pins that the audio/video tag is
preserved across the swap (the guarantee that closes #805's "don't
disturb the shape" ask).
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.picture import Movie
from pptx.util import Inches


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of ``PartFactory.part_type_for``.

    ``Presentation(stream)`` reopen dispatches through
    ``PartFactory.part_type_for``, which other test modules overwrite
    with mocks and do not restore. Mirrors the guard used by
    ``tests/test_issue_720_ea_font_verify.py`` and friends.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue805ReplaceMediaVerify(object):
    """#805 replace audio/video in an existing slide — verified via #784.

    Pins that ``Movie.replace_media(path_or_file, mime_type=None)``
    swaps the media bytes behind an existing audio or video shape in
    place, preserves shape geometry and the media-tag modality, and
    survives a ``Presentation.save`` + reopen cycle.
    """

    # -- public API surface --------------------------------------------------

    def it_exposes_replace_media_on_Movie(self):
        """Guard: the public API surface that resolves #805 is present.

        A single ``hasattr`` check is enough — if a future refactor ever
        removed the method, this test fires before any of the behavioral
        tests below, making the regression easy to diagnose.
        """
        assert hasattr(Movie, "replace_media")

    # -- audio: the #805 reporter's primary scenario -------------------------

    def it_swaps_audio_bytes_in_place_without_moving_the_shape(self):
        """Replace the audio blob; position and size are untouched.

        This is the #805 reporter's canonical use case: a deck template
        contains an audio placeholder at a fixed position and the
        downstream workflow fills in the real audio file without
        disturbing layout.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        original_bytes = b"original-audio-payload"
        slide.shapes.add_movie(
            io.BytesIO(original_bytes),
            left=Inches(2),
            top=Inches(3),
            width=Inches(4),
            height=Inches(1),
            mime_type="audio/mpeg",
        )
        movie = next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.MEDIA)
        assert isinstance(movie, Movie)

        # -- snapshot geometry before the swap --
        left_before = movie.left
        top_before = movie.top
        width_before = movie.width
        height_before = movie.height

        # -- act: swap the media payload --
        new_bytes = b"replacement-audio-payload"
        movie.replace_media(io.BytesIO(new_bytes), mime_type="audio/mpeg")

        # -- geometry unchanged --
        assert movie.left == left_before
        assert movie.top == top_before
        assert movie.width == width_before
        assert movie.height == height_before

        # -- the new media part's blob matches the replacement bytes --
        new_rId = movie._pic.media_video_rId
        assert new_rId is not None
        media_part = movie.part.related_part(new_rId)
        assert media_part.blob == new_bytes

    def it_preserves_the_audio_modality_on_replace(self):
        """The a:audioFile tag stays an a:audioFile after replace.

        #805's "don't disturb the shape" guarantee requires that an
        audio shape remains an audio shape — replace_media swaps the
        bytes, not the modality.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_movie(
            io.BytesIO(b"audio-v1"),
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(2),
            mime_type="audio/mpeg",
        )
        movie = next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.MEDIA)
        # -- pre-condition: add_movie authored an a:audioFile child --
        assert len(movie._pic.xpath("./p:nvPicPr/p:nvPr/a:audioFile")) == 1

        movie.replace_media(io.BytesIO(b"audio-v2"), mime_type="audio/mpeg")

        # -- still an a:audioFile, never promoted to a:videoFile --
        assert len(movie._pic.xpath("./p:nvPicPr/p:nvPr/a:audioFile")) == 1
        assert len(movie._pic.xpath("./p:nvPicPr/p:nvPr/a:videoFile")) == 0

    # -- video: the second half of #805's ask --------------------------------

    def it_swaps_video_bytes_in_place_without_moving_the_shape(self):
        """Replace the video blob of a video shape in place.

        Same contract as the audio test, but with the ``a:videoFile``
        branch of the pic — #805 covers both modalities.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_movie(
            io.BytesIO(b"original-video-payload"),
            left=Inches(2),
            top=Inches(2),
            width=Inches(5),
            height=Inches(3),
            mime_type="video/mp4",
        )
        movie = next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.MEDIA)
        left_before, top_before = movie.left, movie.top
        width_before, height_before = movie.width, movie.height

        new_bytes = b"replacement-video-payload"
        movie.replace_media(io.BytesIO(new_bytes), mime_type="video/mp4")

        # -- geometry unchanged --
        assert movie.left == left_before
        assert movie.top == top_before
        assert movie.width == width_before
        assert movie.height == height_before

        # -- a:videoFile tag preserved --
        assert len(movie._pic.xpath("./p:nvPicPr/p:nvPr/a:videoFile")) == 1
        assert len(movie._pic.xpath("./p:nvPicPr/p:nvPr/a:audioFile")) == 0

        # -- the new media part's blob matches the replacement bytes --
        new_rId = movie._pic.media_video_rId
        assert new_rId is not None
        assert movie.part.related_part(new_rId).blob == new_bytes

    # -- round-trip through save + reopen ------------------------------------

    def it_round_trips_a_replaced_media_blob_through_save_and_reopen(self, _restore_part_factory):
        """End-to-end: author audio, replace it, save, reopen, verify.

        This is the test #805 really needs — a ``.pptx`` that's had its
        media swapped must open cleanly and the reloaded shape must
        still read back the *replacement* bytes (not the original).
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_movie(
            io.BytesIO(b"placeholder-audio"),
            left=Inches(1),
            top=Inches(1),
            width=Inches(3),
            height=Inches(1),
            mime_type="audio/mpeg",
        )
        movie = next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.MEDIA)

        # -- snapshot geometry for post-roundtrip verification --
        left_before = movie.left
        top_before = movie.top
        width_before = movie.width
        height_before = movie.height

        final_bytes = b"final-audio-payload-after-roundtrip"
        movie.replace_media(io.BytesIO(final_bytes), mime_type="audio/mpeg")

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        reloaded = next(sh for sh in prs2.slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.MEDIA)

        # -- geometry survives the round-trip --
        assert reloaded.left == left_before
        assert reloaded.top == top_before
        assert reloaded.width == width_before
        assert reloaded.height == height_before

        # -- the reloaded shape still points at the replacement bytes --
        reloaded_rId = reloaded._pic.media_video_rId
        assert reloaded_rId is not None
        assert reloaded.part.related_part(reloaded_rId).blob == final_bytes

        # -- a:audioFile modality preserved across packaging --
        assert len(reloaded._pic.xpath("./p:nvPicPr/p:nvPr/a:audioFile")) == 1

    # -- error handling ------------------------------------------------------

    def it_rejects_replace_media_on_a_non_media_pic(self):
        """A plain picture (no media rels) can't be replace_media'd.

        ``Picture.replace_image`` is the right API for image shapes;
        ``Movie.replace_media`` is specifically for audio/video shapes.
        Calling replace_media on a shape that has neither an
        ``a:videoFile`` nor an ``a:audioFile`` must raise — otherwise a
        caller who confuses the two would silently produce a broken
        ``.pptx``.
        """
        # -- fabricate a Movie proxy around a plain p:pic (no media child);
        #    add_picture does not return a Movie, so this is the only way
        #    to exercise the not-a-media-pic guard against real XML. --
        from pptx.oxml import parse_xml

        pic_xml = (
            '<p:pic xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
            '        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
            '        xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            "  <p:nvPicPr>"
            '    <p:cNvPr id="5" name="plain"/>'
            "    <p:cNvPicPr/>"
            "    <p:nvPr/>"
            "  </p:nvPicPr>"
            '  <p:blipFill><a:blip r:embed="rId1"/><a:stretch/></p:blipFill>'
            "  <p:spPr/>"
            "</p:pic>"
        )
        pic = parse_xml(pic_xml)
        movie = Movie(pic, None)

        with pytest.raises(ValueError, match="not a media pic"):
            movie.replace_media(io.BytesIO(b"bytes"), mime_type="audio/mpeg")
