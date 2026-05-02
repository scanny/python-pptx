# pyright: reportPrivateUsage=false

"""Regression test for issue #806 — link to image instead of embedding.

Issue #806 (https://github.com/scanny/python-pptx/issues/806) asks for an API
that places a picture shape on a slide whose underlying ``a:blip`` references
an *external* image via ``r:link="rIdX"`` rather than embedding the image
bytes via ``r:embed="rIdX"``. PowerPoint allows either form; the linked form
defers image retrieval to render time and keeps the package small.

The ``SlideShapes.add_picture_link(url, left, top, width, height)`` API
creates an external relationship of type
``http://schemas.openxmlformats.org/officeDocument/2006/relationships/image``
with ``Target=url`` / ``TargetMode="External"`` and emits ``<a:blip r:link=
"rIdX"/>`` referencing it. This regression test exercises that flow through
a full save / reopen round-trip to ensure the external relationship and the
``r:link`` attribute are preserved.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches


@pytest.fixture
def _restore_part_factory():
    """Guard against pollution from other tests that mutate `PartFactory.part_type_for`.

    See the identical fixture in ``tests/test_issue_400_animation_umbrella.py``.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.comments import CommentAuthorsPart, CommentsPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_COMMENTS: CommentsPart,
        CT.PML_COMMENT_AUTHORS: CommentAuthorsPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue806LinkedPicture(object):
    """The #806 flow: insert a picture that links to an external image URL."""

    def it_adds_a_picture_whose_blip_uses_r_link(self, _restore_part_factory):
        """Core ask of #806: the ``a:blip`` must carry ``r:link`` referencing an
        external relationship whose Target is the supplied URL and whose
        TargetMode is "External".
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        url = "https://example.com/images/logo.png"

        picture = slide.shapes.add_picture_link(
            url, Inches(1), Inches(1), Inches(2), Inches(2)
        )

        # -- a:blip on the returned picture uses r:link, not r:embed --
        blip = picture._element.xpath("p:blipFill/a:blip")[0]
        r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        rId = blip.get("{%s}link" % r_ns)
        assert rId is not None
        assert blip.get("{%s}embed" % r_ns) is None

        # -- the external image relationship on the slide part points to the URL --
        slide_part = slide.part
        rel = slide_part.rels[rId]
        assert rel.is_external is True
        assert rel.reltype == RT.IMAGE
        assert rel.target_ref == url

    def it_survives_save_and_reopen_round_trip(self, _restore_part_factory):
        """The external relationship and `r:link` attribute must round-trip
        through save + reopen, so the resulting .pptx can be opened by
        PowerPoint and the image will be fetched from the URL at render time.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        url = "https://example.com/images/logo.png"

        slide.shapes.add_picture_link(url, 0, 0, Inches(3), Inches(2))

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        prs2 = Presentation(buf)
        slide2 = prs2.slides[0]

        # -- last shape in the tree is the linked picture --
        picture2 = slide2.shapes[-1]
        blip2 = picture2._element.xpath("p:blipFill/a:blip")[0]
        r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        rId2 = blip2.get("{%s}link" % r_ns)
        assert rId2 is not None
        assert blip2.get("{%s}embed" % r_ns) is None

        # -- the relationship resolves to the URL with TargetMode="External" --
        rel = slide2.part.rels[rId2]
        assert rel.is_external is True
        assert rel.reltype == RT.IMAGE
        assert rel.target_ref == url

    def it_defaults_missing_dimensions_to_one_inch(self, _restore_part_factory):
        """When `width` and `height` are omitted the picture is sized to 1"x1".

        Unlike ``add_picture``, no image bytes are available so no native-size
        or aspect-ratio computation is possible; the defaults keep the shape
        visible while callers can always supply explicit dimensions.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        picture = slide.shapes.add_picture_link(
            "https://example.com/image.png", 0, 0
        )

        assert picture.width == Inches(1)
        assert picture.height == Inches(1)

    def it_does_not_create_an_image_part_for_the_linked_image(self, _restore_part_factory):
        """A linked picture must NOT cause image bytes to appear in the package
        for the linked URL itself — that is the entire point of linking rather
        than embedding. (The default-template may include other unrelated
        image parts; we verify no *new* image part is added.)
        """
        from pptx.parts.image import ImagePart

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        before = {
            id(p) for p in prs.part.package.iter_parts() if isinstance(p, ImagePart)
        }

        slide.shapes.add_picture_link(
            "https://example.com/foo.png", 0, 0, Inches(1), Inches(1)
        )

        after = {id(p) for p in prs.part.package.iter_parts() if isinstance(p, ImagePart)}
        assert after == before, "linked picture must not create a new image part"

        # -- and the slide must not hold an RT.IMAGE relationship to any
        # -- internal ImagePart; its only RT.IMAGE relationship must be
        # -- external (the URL) --
        rels = list(slide.part.rels.values())
        image_rels = [r for r in rels if r.reltype == RT.IMAGE]
        assert len(image_rels) == 1
        assert image_rels[0].is_external is True
