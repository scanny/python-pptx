# pyright: reportPrivateUsage=false

"""Regression test for issue #809 — resolve inherited slide background.

Issue #809 (https://github.com/scanny/python-pptx/issues/809) reported
that a slide whose layout carries a solid-color ``p:bg`` could not be
queried for its rendered background: reading
``slide.background.fill.fore_color.rgb`` either raised (the fill type
resolves to ``MSO_FILL.BACKGROUND`` / ``_NoFill``) or silently clobbered
the inheritance link — ``Slide.background.fill`` materializes a
``p:bgPr/a:noFill`` subtree on first read, overwriting the inherited
layout color with an empty override.

The fix adds a side-effect-free ``effective_background`` accessor on
``Slide``, ``SlideLayout``, and ``SlideMaster`` that walks
*slide → layout → master* and returns a read-only
:class:`pptx.slide._EffectiveBackground` view of the first ancestor
carrying an explicit ``p:bg``. The returned proxy exposes ``.source``
(which ancestor supplied the background), ``.owner`` (the slide-like
object), ``.bg_element`` (the resolved ``p:bg`` lxml element), and
``.fill`` (a non-destructive :class:`.FillFormat`, or |None| when the
resolved background is a theme-keyed ``p:bgRef``).

This suite pins:

1. Slide-level override wins — when the slide carries its own ``p:bg``
   ``effective_background.source`` is ``"slide"`` and ``.fill`` reads
   the slide's own color.
2. Layout-level inheritance — a slide with no ``p:bg`` whose layout
   carries a solid color reports ``source == "layout"`` and
   ``fill.fore_color.rgb`` equal to the layout color. (The #809 ask.)
3. Master-level inheritance — a slide whose layout *also* has no
   ``p:bg`` falls through to the master.
4. ``p:bgRef``-backed master — the default python-pptx template has a
   master ``p:bg`` that wraps a ``p:bgRef``; ``.fill`` returns |None|
   without mutating the XML.
5. Reading ``effective_background`` — and reading ``.fill`` on it — is
   side-effect free: the slide's ``p:bg`` stays absent, the layout's
   stays untouched.
6. ``effective_background`` on |SlideLayout| walks ``layout → master``.
7. ``effective_background`` on |SlideMaster| returns a wrapper of the
   master's own ``p:bg`` (or |None| when the master has none).
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL
from pptx.oxml.ns import qn
from pptx.slide import _EffectiveBackground


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of ``PartFactory.part_type_for``.

    Mirrors the guard used by ``tests/test_issue_366_slide_background_verify.py``
    and friends — ``Presentation(stream)`` reopen dispatches through
    ``PartFactory.part_type_for``, which earlier test modules overwrite
    with mocks and do not restore.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue809EffectiveBackgroundVerify:
    """#809 effective-background verify-and-close.

    Pins the read-only slide / layout / master background resolver so the
    #809 report can be closed.
    """

    # -- 1. Slide-level override ------------------------------------------

    def it_reports_slide_as_source_when_the_slide_has_its_own_p_bg(self):
        """A slide with an explicit `p:bg` wins over layout and master.

        ``effective_background.source`` is ``"slide"`` and ``.fill`` reads
        the slide's own color (not the layout / master color).
        """
        prs = Presentation()
        layout = prs.slide_layouts[0]
        # -- layout has a different color so we can tell them apart --
        layout.background.fill.solid()
        layout.background.fill.fore_color.rgb = RGBColor(0xAA, 0xBB, 0xCC)

        slide = prs.slides.add_slide(layout)
        # -- apply a distinct slide-level override --
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0x11, 0x22, 0x33)

        eff = slide.effective_background
        assert isinstance(eff, _EffectiveBackground)
        assert eff.source == "slide"
        assert eff.owner is slide
        assert eff.fill is not None
        assert eff.fill.type == MSO_FILL.SOLID
        assert eff.fill.fore_color.rgb == RGBColor(0x11, 0x22, 0x33)

    # -- 2. Layout-level inheritance (the #809 ask) -----------------------

    def it_resolves_layout_background_when_the_slide_has_none(self):
        """A slide with no `p:bg` inherits from its layout.

        The reporter's concrete case: layout defines a solid color,
        slide has no override; ``effective_background.fill.fore_color.rgb``
        must return the layout's color.
        """
        prs = Presentation()
        layout = prs.slide_layouts[0]
        layout.background.fill.solid()
        layout.background.fill.fore_color.rgb = RGBColor(0xFF, 0x00, 0x00)

        slide = prs.slides.add_slide(layout)

        # -- sanity: slide itself has no p:bg --
        assert slide._element.cSld.bg is None

        eff = slide.effective_background
        assert isinstance(eff, _EffectiveBackground)
        assert eff.source == "layout"
        assert eff.owner is slide.slide_layout
        assert eff.fill is not None
        assert eff.fill.type == MSO_FILL.SOLID
        assert eff.fill.fore_color.rgb == RGBColor(0xFF, 0x00, 0x00)

    # -- 3. Master-level inheritance (layout has no p:bg) -----------------

    def it_resolves_master_background_when_slide_and_layout_have_none(self):
        """A slide whose layout *also* lacks a `p:bg` falls through to the master.

        The default python-pptx template has a master-level ``p:bg``
        that wraps a theme-keyed ``p:bgRef`` — accessing
        ``effective_background`` must locate it and report
        ``source == "master"``. (``.fill`` is |None| for a ``p:bgRef``;
        see the dedicated bgRef scenario.)
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        # -- default template: neither slide nor layout has a p:bg --
        assert slide._element.cSld.bg is None
        assert slide.slide_layout._element.cSld.bg is None

        eff = slide.effective_background
        assert isinstance(eff, _EffectiveBackground)
        assert eff.source == "master"
        assert eff.owner is slide.slide_layout.slide_master
        # -- the default master carries a p:bgRef, not a p:bgPr --
        assert eff.bg_element.find(qn("p:bgRef")) is not None

    # -- 4. p:bgRef-backed master: .fill is None --------------------------

    def it_returns_None_fill_when_the_resolved_bg_is_a_style_reference(self):
        """``p:bgRef`` is not a :class:`FillFormat` surface.

        A theme-keyed background-style reference (``p:bg/p:bgRef``) carries
        no directly-readable color — the actual color is resolved against
        the theme's ``bgFillStyleLst`` at render time. ``.fill`` returns
        |None| for this shape rather than raising or clobbering the XML.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])

        eff = slide.effective_background
        assert eff is not None
        assert eff.source == "master"
        # -- bgRef resolved; .fill is None --
        assert eff.fill is None
        # -- but .bg_element still exposes the raw element --
        assert eff.bg_element is not None
        assert eff.bg_element.find(qn("p:bgRef")) is not None

    # -- 5. Read is side-effect free --------------------------------------

    def it_does_not_mutate_the_slide_xml_when_effective_background_is_read(self):
        """Reading ``effective_background`` (and ``.fill``) leaves inheritance intact.

        The whole point of ``effective_background`` is to **read**
        without side effects — in particular without materializing a
        ``p:bgPr/a:noFill`` on the slide (the side effect of
        ``Slide.background.fill``). Reading the proxy *and* its ``.fill``
        must leave the slide's ``p:cSld`` free of any ``p:bg`` child.
        """
        prs = Presentation()
        layout = prs.slide_layouts[0]
        layout.background.fill.solid()
        layout.background.fill.fore_color.rgb = RGBColor(0x33, 0x66, 0x99)
        slide = prs.slides.add_slide(layout)

        # -- snapshot the layout XML for later comparison --
        layout_xml_before = layout._element.cSld.xml

        # -- act: read the effective background and its fill --
        eff = slide.effective_background
        assert eff is not None
        _ = eff.fill
        _ = eff.fill.fore_color.rgb  # pyright: ignore[reportOptionalMemberAccess]

        # -- assert: the slide's p:cSld is still p:bg-free --
        assert slide._element.cSld.bg is None
        # -- and the layout XML is unchanged --
        assert layout._element.cSld.xml == layout_xml_before

    def and_it_does_not_mutate_the_master_when_reading_a_bgRef_effective_background(self):
        """Reading through a ``p:bgRef`` master background is side-effect free.

        Pins the fix for the bgRef-on-master case: the default python-pptx
        template carries a ``p:bg/p:bgRef`` on the master, and reading
        ``effective_background.fill`` must not swap it out for a
        ``p:bgPr/a:noFill`` stub.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        master = slide.slide_layout.slide_master
        master_xml_before = master._element.cSld.xml

        # -- act --
        eff = slide.effective_background
        assert eff is not None
        _ = eff.fill  # None for bgRef; must not mutate

        # -- assert: the master's p:cSld is unchanged --
        assert master._element.cSld.xml == master_xml_before

    # -- 6. SlideLayout.effective_background ------------------------------

    def it_walks_layout_to_master_on_SlideLayout_effective_background(self):
        """|SlideLayout| walks layout → master (no slide step)."""
        prs = Presentation()

        # -- a layout without its own p:bg resolves through to the master --
        layout = prs.slide_layouts[0]
        eff = layout.effective_background
        assert eff is not None
        assert eff.source == "master"

        # -- a layout with its own p:bg resolves at itself --
        layout.background.fill.solid()
        layout.background.fill.fore_color.rgb = RGBColor(0x00, 0xFF, 0x00)
        eff2 = layout.effective_background
        assert eff2 is not None
        assert eff2.source == "layout"
        assert eff2.owner is layout
        assert eff2.fill is not None
        assert eff2.fill.fore_color.rgb == RGBColor(0x00, 0xFF, 0x00)

    # -- 7. SlideMaster.effective_background ------------------------------

    def it_wraps_its_own_p_bg_on_SlideMaster_effective_background(self):
        """|SlideMaster| wraps its own ``p:bg`` when present."""
        prs = Presentation()
        master = prs.slide_masters[0]
        # -- the default master carries a p:bgRef --
        eff = master.effective_background
        assert eff is not None
        assert eff.source == "master"
        assert eff.owner is master
        assert eff.bg_element.find(qn("p:bgRef")) is not None

    # -- 8. Round-trip --------------------------------------------------

    def it_resolves_inherited_background_after_save_and_reopen(self, _restore_part_factory):
        """The #809 read path survives save + reopen.

        After saving the presentation with a layout-level color and
        reopening, reading ``effective_background.fill.fore_color.rgb``
        on the slide still returns the layout color.
        """
        prs = Presentation()
        layout = prs.slide_layouts[0]
        layout.background.fill.solid()
        layout.background.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
        prs.slides.add_slide(layout)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        slide2 = prs2.slides[0]

        eff = slide2.effective_background
        assert eff is not None
        assert eff.source == "layout"
        assert eff.fill is not None
        assert eff.fill.fore_color.rgb == RGBColor(0x12, 0x34, 0x56)
        # -- slide part remains bg-free (inheritance preserved) --
        assert slide2._element.cSld.bg is None
