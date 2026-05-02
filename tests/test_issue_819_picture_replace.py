# pyright: reportPrivateUsage=false

"""Regression test for issue #819 - image replacing.

Issue #819 (https://github.com/scanny/python-pptx/issues/819) asked for
a supported way to swap the pixel bytes of an existing picture shape in
place, without touching its position, size or identifying attributes.
The reporter's workflow:

    "I make a presentation in python-pptx and then manually go in and
     add some slides and then remember that I need to change one or
     more of the pictures. Then I want to be able to run the program
     again and have it replace the picture with the same name."

The reporter's workaround looked up pictures by their ``cNvPr/@descr``
alt-text, then performed a ``deleteShape`` + ``add_picture`` pair at
the captured ``top/left/width/height``. That approach loses the
reporter's identifying ``descr`` on the round-trip because
:meth:`.SlideShapes.add_picture` generates a fresh ``descr`` from the
source image's filename (e.g. ``image.png``), so the next run can no
longer find the same shape — which is exactly the misbehaviour the
reporter described ("more often than not, it is replaced with
image.png and because of this I am unable to find that picture the
next time I try to replace").

The reporter's second attempt poked at :class:`.ImagePart` internals
via ``shape.part.relatedpart(imgRID)`` (a method that does not exist)
and monkey-patched ``_blob`` — a path that cannot be made to work
safely because multiple pictures can share one image part through
content-hash deduplication.

Wave 6 #834 (``feat/issue-834-picture-replace-image``) shipped the
supported API that answers #819 directly:
:meth:`.Picture.replace_image` swaps the embedded image while keeping
every shape-level attribute intact — position, size, rotation, crop,
masking shape, outline, *and* the ``cNvPr/@name`` / ``cNvPr/@descr``
identity the reporter was keying off. Internally it adds (or reuses
via content-hash) an image part through
:meth:`.SlidePart.get_or_add_image_part`, rebinds
``p:pic/p:blipFill/a:blip/@r:embed`` to the new rId, and drops the
previous image relationship so the now-orphaned image part is
garbage-collected on save.

#819 is therefore a duplicate of #834 / #116 and is verified-and-closed
by this suite. The scenarios below pin the reporter's exact workflow
(find a picture by ``alt_text``, swap its pixels, confirm the
identifying descr + geometry still locate it on the second pass)
against the kind of silent regression that would reopen the issue.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.util import Inches


class DescribeIssue819ImageReplacing:
    """Round-trip regression for the #819 "find-by-descr, then replace" workflow."""

    def it_swaps_pixels_without_disturbing_the_descr_alt_text(self):
        # -- the reporter keyed off cNvPr/@descr to locate the shape on the
        # -- next program run; replace_image must not touch that identifier --
        prs = _prs_with_tagged_picture(
            "tests/test_files/python-powered.png",
            alt_text="company-logo",
        )
        picture = _find_picture_by_alt_text(prs, "company-logo")
        assert picture is not None

        original_descr = picture._element.nvPicPr.cNvPr.get("descr")
        original_name = picture._element.nvPicPr.cNvPr.get("name")
        original_blob = picture.image.blob

        picture.replace_image("tests/test_files/monty-truth.png")

        # -- the identifying attributes the reporter's loop matched against
        # -- are untouched, so the second run still finds the same shape --
        assert picture._element.nvPicPr.cNvPr.get("descr") == original_descr
        assert picture._element.nvPicPr.cNvPr.get("name") == original_name
        assert picture.alt_text == "company-logo"

        # -- but the pixel bytes did swap --
        assert picture.image.blob != original_blob
        with open("tests/test_files/monty-truth.png", "rb") as f:
            assert picture.image.blob == f.read()

    def it_preserves_position_and_size_across_the_replace(self):
        # -- the reporter's workaround captured top/left/width/height
        # -- specifically because add_picture would otherwise reset them;
        # -- replace_image must preserve them verbatim --
        prs = _prs_with_tagged_picture(
            "tests/test_files/python-powered.png",
            alt_text="hero-image",
            x=Inches(1),
            y=Inches(2),
            cx=Inches(3),
            cy=Inches(4),
        )
        picture = _find_picture_by_alt_text(prs, "hero-image")
        assert picture is not None

        picture.replace_image("tests/test_files/monty-truth.png")

        assert picture.left == Inches(1)
        assert picture.top == Inches(2)
        assert picture.width == Inches(3)
        assert picture.height == Inches(4)

    def it_supports_the_reporters_end_to_end_find_replace_roundtrip(self):
        # -- this is the reporter's full scenario: open a presentation, find
        # -- a picture by alt_text, replace it, save, reopen, and confirm
        # -- the same alt_text still locates a shape with the new pixels.
        # -- A second round trip exercises the "run the program again" part
        # -- of the reporter's workflow.
        prs = _prs_with_tagged_pictures(
            [
                ("tests/test_files/python-powered.png", "logo-primary"),
                ("tests/test_files/python-icon.jpeg", "logo-secondary"),
            ]
        )

        # -- first run: swap only the primary logo --
        target = _find_picture_by_alt_text(prs, "logo-primary")
        assert target is not None
        target.replace_image("tests/test_files/monty-truth.png")

        buf = io.BytesIO()
        prs.save(buf)

        # -- second run: reopen and verify the descr still locates the
        # -- same shape, now carrying the replacement bytes --
        buf.seek(0)
        prs2 = Presentation(buf)
        relocated = _find_picture_by_alt_text(prs2, "logo-primary")
        assert relocated is not None, (
            "cNvPr/@descr was lost across replace_image — the reporter's "
            "find-by-descr workflow would break on the second run"
        )
        with open("tests/test_files/monty-truth.png", "rb") as f:
            assert relocated.image.blob == f.read()

        # -- and the untouched sibling is still addressable by its descr --
        other = _find_picture_by_alt_text(prs2, "logo-secondary")
        assert other is not None
        with open("tests/test_files/python-icon.jpeg", "rb") as f:
            assert other.image.blob == f.read()

        # -- third run: swap again through a second save/reopen cycle to
        # -- pin the "run the program again" idempotency the reporter needs --
        relocated.replace_image("tests/test_files/python-powered.png")
        buf2 = io.BytesIO()
        prs2.save(buf2)
        buf2.seek(0)
        prs3 = Presentation(buf2)
        final = _find_picture_by_alt_text(prs3, "logo-primary")
        assert final is not None
        with open("tests/test_files/python-powered.png", "rb") as f:
            assert final.image.blob == f.read()

    def it_accepts_a_file_like_replacement_like_the_reporters_open_rb(self):
        # -- the reporter's second attempt read the replacement bytes with
        # -- `open(path, 'rb')`; replace_image supports that directly --
        prs = _prs_with_tagged_picture(
            "tests/test_files/python-powered.png",
            alt_text="chart-screenshot",
        )
        picture = _find_picture_by_alt_text(prs, "chart-screenshot")
        assert picture is not None

        with open("tests/test_files/monty-truth.png", "rb") as f:
            picture.replace_image(f)

        assert picture.alt_text == "chart-screenshot"
        with open("tests/test_files/monty-truth.png", "rb") as f:
            assert picture.image.blob == f.read()


# -- helpers -----------------------------------------------------------


def _prs_with_tagged_picture(
    image_path,
    alt_text,
    x=Inches(1),
    y=Inches(1),
    cx=Inches(2),
    cy=Inches(2),
):
    """Return a fresh |Presentation| holding one picture stamped with `alt_text`."""
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    picture = slide.shapes.add_picture(image_path, x, y, cx, cy)
    picture.alt_text = alt_text
    return prs


def _prs_with_tagged_pictures(tagged_paths):
    """Return a fresh |Presentation| with one slide holding the tagged pictures."""
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    for idx, (image_path, alt_text) in enumerate(tagged_paths):
        picture = slide.shapes.add_picture(
            image_path, Inches(1 + idx * 3), Inches(1), Inches(2), Inches(2)
        )
        picture.alt_text = alt_text
    return prs


def _find_picture_by_alt_text(prs, alt_text):
    """Return the first picture shape whose ``alt_text`` matches, or |None|.

    Mirrors the reporter's ``cNvPr/@descr`` lookup in issue #819 using the
    supported :attr:`.BaseShape.alt_text` accessor.
    """
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "alt_text", "") == alt_text:
                return shape
    return None
