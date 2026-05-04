# pyright: reportPrivateUsage=false

"""Regression test for issue #822 — embed an .xlsx file as an OLE object.

Issue #822 (https://github.com/scanny/python-pptx/issues/822) reports that
``SlideShapes.add_ole_object()`` could not be used to embed an Excel (.xlsx)
file on v0.6.21. The reporter expected an OLE icon on the slide that
launches Excel when double-clicked.

The fork-era generic OLE-object embedding path (issue #752) added
``PROG_ID.XLSX`` as a registered convenience member, and
``EmbeddedXlsxPart`` emits the xlsx content-type and
``/ppt/embeddings/Microsoft_Excel_Sheet%d.xlsx`` part-name. The arbitrary
``prog_id`` + ``extension`` overload additionally allows the raw progId
string ``"Excel.Sheet.12"`` to work with a caller-provided ``.xlsx``
extension hint.

This module exercises the #822 flow end-to-end to pin the behavior:

* author the embed with ``PROG_ID.XLSX`` from a str path,
* round-trip the package through save + reopen,
* confirm the reopened shape is detected as ``EMBEDDED_OLE_OBJECT``,
* confirm ``ole_format.prog_id == "Excel.Sheet.12"`` and the workbook
  bytes survive byte-for-byte,
* confirm the xlsx content-type and Office-package part-name are emitted,
* confirm a PNG icon provided via ``icon_file`` is embedded as an image
  part (the fork accepts PNG as well as the default EMF icon), and
* confirm the raw-string ``prog_id="Excel.Sheet.12"`` + ``extension="xlsx"``
  form also round-trips, via the generic OLE-object content-type path.
"""

from __future__ import annotations

import io
import zipfile

import xlsxwriter
from PIL import Image

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PROG_ID
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.util import Inches


def _make_xlsx_bytes() -> bytes:
    """Return a minimal valid .xlsx file as bytes via XlsxWriter."""
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    ws = wb.add_worksheet("Sheet1")
    ws.write(0, 0, "Issue")
    ws.write(0, 1, 822)
    ws.write(1, 0, "Hello, embedded xlsx!")
    wb.close()
    return buf.getvalue()


def _make_png_bytes() -> bytes:
    """Return a minimal PNG image as bytes (32x32 solid red)."""
    buf = io.BytesIO()
    Image.new("RGB", (32, 32), (220, 20, 60)).save(buf, format="PNG")
    return buf.getvalue()


class DescribeIssue822EmbedXlsxOleObject(object):
    """The #822 flow: embed an .xlsx file via ``PROG_ID.XLSX``."""

    def it_round_trips_an_xlsx_file_as_an_Excel_OLE_object(self, tmp_path):
        """``add_ole_object(xlsx_path, PROG_ID.XLSX, ...)`` embeds the
        workbook, produces an Office-package part, and survives save /
        reopen byte-for-byte — pinning the capability the v0.6.21 reporter
        was unable to reach.
        """
        xlsx_bytes = _make_xlsx_bytes()
        xlsx_path = tmp_path / "workbook.xlsx"
        xlsx_path.write_bytes(xlsx_bytes)

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        shape = slide.shapes.add_ole_object(
            str(xlsx_path),
            PROG_ID.XLSX,
            left=Inches(1),
            top=Inches(1),
        )

        # -- authored shape exposes the OLE metadata up front --
        assert shape.shape_type == MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT
        assert shape.ole_format.prog_id == "Excel.Sheet.12"
        assert shape.ole_format.blob == xlsx_bytes
        assert shape.ole_format.show_as_icon is True

        # -- save + reopen round-trip preserves progId and bytes --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        prs2 = Presentation(buf)
        slide2 = prs2.slides[0]
        ole_shapes = [
            s for s in slide2.shapes if s.shape_type == MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT
        ]
        assert len(ole_shapes) == 1
        reopened = ole_shapes[0]
        assert reopened.ole_format.prog_id == "Excel.Sheet.12"
        assert reopened.ole_format.blob == xlsx_bytes
        assert reopened.ole_format.show_as_icon is True

    def it_emits_the_xlsx_package_content_type_and_part_name(self, tmp_path):
        """``PROG_ID.XLSX`` selects the ``EmbeddedXlsxPart`` subclass, so the
        package emits the xlsx content-type and the
        ``Microsoft_Excel_Sheet%d.xlsx`` part-name — the Office-package
        path, distinct from the generic OLE-object path used by raw-string
        progIds.
        """
        xlsx_bytes = _make_xlsx_bytes()
        xlsx_path = tmp_path / "workbook.xlsx"
        xlsx_path.write_bytes(xlsx_bytes)

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.add_ole_object(
            str(xlsx_path), PROG_ID.XLSX, left=Inches(1), top=Inches(1)
        )

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        with zipfile.ZipFile(buf) as zf:
            names = zf.namelist()
            embed_names = [
                n
                for n in names
                if n.startswith("ppt/embeddings/Microsoft_Excel_Sheet")
                and n.endswith(".xlsx")
            ]
            assert len(embed_names) == 1
            # -- embedded part bytes match the input file exactly --
            assert zf.read(embed_names[0]) == xlsx_bytes
            # -- [Content_Types].xml declares the xlsx content-type --
            content_types_xml = zf.read("[Content_Types].xml").decode("utf-8")
            assert CT.SML_SHEET in content_types_xml

    def it_embeds_the_icon_PNG_when_one_is_supplied(self, tmp_path):
        """A caller-supplied ``icon_file`` PNG is embedded as an image part
        and survives round-trip. This pins that the icon is wired in
        alongside the xlsx payload — not silently dropped.
        """
        xlsx_bytes = _make_xlsx_bytes()
        xlsx_path = tmp_path / "workbook.xlsx"
        xlsx_path.write_bytes(xlsx_bytes)

        png_bytes = _make_png_bytes()
        png_path = tmp_path / "icon.png"
        png_path.write_bytes(png_bytes)

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.add_ole_object(
            str(xlsx_path),
            PROG_ID.XLSX,
            left=Inches(1),
            top=Inches(1),
            icon_file=str(png_path),
        )

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        with zipfile.ZipFile(buf) as zf:
            names = zf.namelist()
            png_names = [n for n in names if n.startswith("ppt/media/") and n.endswith(".png")]
            assert len(png_names) == 1
            # -- embedded icon bytes match the input PNG exactly --
            assert zf.read(png_names[0]) == png_bytes
            # -- and the xlsx part is still present alongside the icon --
            embed_names = [
                n
                for n in names
                if n.startswith("ppt/embeddings/Microsoft_Excel_Sheet")
                and n.endswith(".xlsx")
            ]
            assert len(embed_names) == 1

    def it_also_accepts_the_raw_string_Excel_Sheet_12_progId_with_extension(
        self, tmp_path
    ):
        """The arbitrary-``prog_id`` overload accepts the raw progId string
        ``"Excel.Sheet.12"`` and embeds the file through the generic
        OLE-object content-type path when ``extension="xlsx"`` is also
        supplied. This pins the alternate form the #822 reporter might
        reach for if ``PROG_ID.XLSX`` is unfamiliar.
        """
        xlsx_bytes = _make_xlsx_bytes()
        xlsx_path = tmp_path / "workbook.xlsx"
        xlsx_path.write_bytes(xlsx_bytes)

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        shape = slide.shapes.add_ole_object(
            str(xlsx_path),
            prog_id="Excel.Sheet.12",
            left=Inches(1),
            top=Inches(1),
            width=Inches(1),
            height=Inches(1),
            extension="xlsx",
        )

        assert shape.ole_format.prog_id == "Excel.Sheet.12"
        assert shape.ole_format.blob == xlsx_bytes

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        # -- raw-string progId uses the generic OLE-object embedding path:
        #    the part-name is ``oleObject*.xlsx`` with the generic OLE
        #    content-type, not the Microsoft_Excel_Sheet template. --
        with zipfile.ZipFile(buf) as zf:
            names = zf.namelist()
            ole_names = [
                n
                for n in names
                if n.startswith("ppt/embeddings/oleObject") and n.endswith(".xlsx")
            ]
            assert len(ole_names) == 1
            assert zf.read(ole_names[0]) == xlsx_bytes
            content_types_xml = zf.read("[Content_Types].xml").decode("utf-8")
            assert CT.OFC_OLE_OBJECT in content_types_xml

        # -- round-trip preserves the progId and payload --
        prs2 = Presentation(buf)
        slide2 = prs2.slides[0]
        ole_shapes = [
            s for s in slide2.shapes if s.shape_type == MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT
        ]
        assert len(ole_shapes) == 1
        assert ole_shapes[0].ole_format.prog_id == "Excel.Sheet.12"
        assert ole_shapes[0].ole_format.blob == xlsx_bytes
