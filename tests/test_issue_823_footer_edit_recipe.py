# pyright: reportUnknownMemberType=false, reportUnknownVariableType=false, reportUnknownParameterType=false, reportMissingParameterType=false, reportAttributeAccessIssue=false, reportOptionalMemberAccess=false

"""Regression tests for issue #823 — footer / slide-number / date placeholder edit recipe.

The reporter asked how to edit "the character in the lower-left corner" of
their slides. That character is almost always one of the *latent*
placeholders (footer, date, slide number) authored on the slide master
and/or overridden on a slide layout.

These tests lock in the documented recipe in
``docs/user/slides.rst`` — "Editing the footer, slide-number, or date
placeholder text" — so that the code paths it relies on stay working:

* ``SlideMaster.placeholders`` iteration by ``placeholder_format.type``
* ``SlideLayout.placeholders`` iteration by ``placeholder_format.type``
* ``.text`` assignment round-tripping through save + reload
* ``SlideMaster.header_footer`` visibility toggles
* Master-shapes fallback for decorative "watermark" shapes
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches


def _find_ph_by_type(container, ph_type):
    """Return the first placeholder on *container* with matching type, or None."""
    for ph in container.placeholders:
        if ph.placeholder_format.type == ph_type:
            return ph
    return None


class DescribeIssue823FooterEditRecipe(object):
    """Regression suite for the docs recipe covering issue #823."""

    def it_finds_the_footer_placeholder_on_the_master_by_type(self):
        # -- The master carries a footer placeholder; locating it by
        # -- PP_PLACEHOLDER.FOOTER (not by idx) is portable across
        # -- masters whose placeholder idx values differ.
        prs = Presentation()

        footer = _find_ph_by_type(prs.slide_master, PP_PLACEHOLDER.FOOTER)

        assert footer is not None
        assert footer.placeholder_format.type == PP_PLACEHOLDER.FOOTER

    def it_finds_the_slide_number_placeholder_on_the_master_by_type(self):
        prs = Presentation()

        sld_num = _find_ph_by_type(prs.slide_master, PP_PLACEHOLDER.SLIDE_NUMBER)

        assert sld_num is not None
        assert sld_num.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER

    def it_finds_the_date_placeholder_on_the_master_by_type(self):
        prs = Presentation()

        date_ph = _find_ph_by_type(prs.slide_master, PP_PLACEHOLDER.DATE)

        assert date_ph is not None
        assert date_ph.placeholder_format.type == PP_PLACEHOLDER.DATE

    def it_edits_the_master_footer_text_and_round_trips(self):
        # -- The central recipe: assign ``ph.text`` on the master's
        # -- footer placeholder; the edit survives save + reload.
        prs = Presentation()

        footer = _find_ph_by_type(prs.slide_master, PP_PLACEHOLDER.FOOTER)
        assert footer is not None
        footer.text = "Confidential - Q3 2026"

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        reloaded_footer = _find_ph_by_type(reloaded.slide_master, PP_PLACEHOLDER.FOOTER)
        assert reloaded_footer is not None
        assert reloaded_footer.text == "Confidential - Q3 2026"

    def it_edits_a_layout_level_footer_and_round_trips(self):
        # -- Layout-level footer edits override the master for slides
        # -- that use that layout; the edit also round-trips.
        prs = Presentation()
        layout = prs.slide_layouts[1]  # "Title and Content"

        footer = _find_ph_by_type(layout, PP_PLACEHOLDER.FOOTER)
        assert footer is not None
        footer.text = "Layout-specific footer"

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        reloaded_footer = _find_ph_by_type(reloaded.slide_layouts[1], PP_PLACEHOLDER.FOOTER)
        assert reloaded_footer is not None
        assert reloaded_footer.text == "Layout-specific footer"

    def it_independently_edits_master_and_layout_footers(self):
        # -- Master and layout footers are distinct XML nodes; editing
        # -- one does not mutate the other. This is load-bearing for
        # -- the "override on a single layout" variant of the recipe.
        prs = Presentation()

        master_footer = _find_ph_by_type(prs.slide_master, PP_PLACEHOLDER.FOOTER)
        layout_footer = _find_ph_by_type(prs.slide_layouts[0], PP_PLACEHOLDER.FOOTER)
        assert master_footer is not None
        assert layout_footer is not None

        master_footer.text = "MASTER"
        layout_footer.text = "LAYOUT"

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        assert _find_ph_by_type(reloaded.slide_master, PP_PLACEHOLDER.FOOTER).text == "MASTER"
        assert _find_ph_by_type(reloaded.slide_layouts[0], PP_PLACEHOLDER.FOOTER).text == "LAYOUT"

    def it_toggles_footer_visibility_via_header_footer(self):
        # -- Editing the text is meaningless if the ``<p:hf>`` toggle
        # -- has hidden the footer. The recipe pairs text-edits with a
        # -- visibility check on ``header_footer``.
        prs = Presentation()

        hf = prs.slide_master.header_footer
        hf.footer_visible = False
        assert hf.footer_visible is False

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)
        assert reloaded.slide_master.header_footer.footer_visible is False

    def it_toggles_slide_number_visibility_via_header_footer(self):
        prs = Presentation()

        hf = prs.slide_master.header_footer
        hf.slide_number_visible = False
        assert hf.slide_number_visible is False

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)
        assert reloaded.slide_master.header_footer.slide_number_visible is False

    def it_does_not_clone_latent_placeholders_onto_new_slides(self):
        # -- Pins the contract documented in the recipe: latent
        # -- placeholders (date, footer, slide number) are not cloned
        # -- onto a slide when it is created from its layout. This is
        # -- why the recipe edits the master / layout, not the slide.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])

        phs_by_type = {ph.placeholder_format.type for ph in slide.placeholders}

        assert PP_PLACEHOLDER.FOOTER not in phs_by_type
        assert PP_PLACEHOLDER.SLIDE_NUMBER not in phs_by_type
        assert PP_PLACEHOLDER.DATE not in phs_by_type

    def it_supports_the_watermark_shape_fallback_on_master(self):
        # -- When the "character in the lower-left corner" is actually
        # -- a decorative (non-placeholder) shape baked into the
        # -- master, iterating ``master.shapes`` and matching by text
        # -- or by name still works. Exercising that path here keeps
        # -- the fallback recipe from bit-rotting.
        prs = Presentation()
        master = prs.slide_master

        # -- add a textbox on the master with a recognisable watermark --
        textbox = master.shapes.add_textbox(Inches(0.2), Inches(6.8), Inches(2), Inches(0.3))
        textbox.text_frame.text = "Watermark"

        # -- the fallback: find it by text, edit it, round-trip --
        for shape in master.shapes:
            if shape.has_text_frame and "watermark" in shape.text_frame.text.lower():
                shape.text_frame.text = ""

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        # -- the watermark has been blanked out --
        watermark_texts = [
            s.text_frame.text for s in reloaded.slide_master.shapes if s.has_text_frame
        ]
        assert "Watermark" not in watermark_texts
