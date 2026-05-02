# pyright: reportPrivateUsage=false

"""Regression test for issue #830 — embed custom ``.otf`` font with python-pptx.

Issue #830 (https://github.com/scanny/python-pptx/issues/830) asks for a way
to embed an OpenType (``.otf``) font file into a presentation so that callers
can ship a custom typeface with the deck and PowerPoint will render text in
that font even on machines where the font is not installed.

Wave 3 #355 shipped :meth:`.Presentation.embed_font` accepting TrueType
(``.ttf``). This regression test verifies that the same API accepts OpenType
(``.otf``) files identically — OpenType is a valid font container and the
OOXML embedded-font part (``application/x-fontdata``) is the same envelope
for both. The library treats the file bytes as an opaque blob so the .otf
case works out of the box; this test pins that behavior.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.parts.font import FontPart


# -- A minimal ``.otf``-magic-header byte blob. Real OpenType files begin with
# -- the four-byte sniff tag ``OTTO`` (0x4F54544F) — distinct from TrueType's
# -- 0x00010000 scaler type. ``embed_font`` reads the file as an opaque blob so
# -- this stub is sufficient to exercise the .otf path end-to-end; PowerPoint
# -- itself parses the font table on load (not python-pptx's concern). --
_OTF_MAGIC_HEADER = b"OTTO"
_FAKE_OTF_BYTES = _OTF_MAGIC_HEADER + b"\x00\x0b\x00\x80\x00\x03\x00\x60" + b"\x00" * 248


class DescribeIssue830OtfFont(object):
    """The #830 flow: embed a custom ``.otf`` font and round-trip the package."""

    def it_accepts_an_otf_file_from_a_filesystem_path(self, tmp_path):
        """``embed_font(otf_path, typeface)`` embeds the ``.otf`` bytes and
        registers the typeface, identically to the ``.ttf`` path.
        """
        otf_path = tmp_path / "CustomFont.otf"
        otf_path.write_bytes(_FAKE_OTF_BYTES)

        prs = Presentation()
        prs.embed_font(str(otf_path), "CustomFont")

        assert prs.embedded_fonts == ("CustomFont",)

    def it_accepts_an_otf_file_from_a_file_like_object(self):
        """``embed_font`` also accepts an ``IO[bytes]`` carrying ``.otf`` bytes
        (matches the #355 contract for ``.ttf`` streams).
        """
        prs = Presentation()

        prs.embed_font(io.BytesIO(_FAKE_OTF_BYTES), "CustomFont")

        assert prs.embedded_fonts == ("CustomFont",)

    def it_creates_a_font_part_with_x_fontdata_content_type(self):
        """Both TrueType and OpenType embed into the same OOXML part kind —
        ``application/x-fontdata`` — because the fntdata envelope is
        container-agnostic. Pin that here so a future refactor doesn't split
        OTF into a separate content-type by mistake.
        """
        prs = Presentation()
        prs.embed_font(io.BytesIO(_FAKE_OTF_BYTES), "CustomFont")

        font_parts = [
            p for p in prs.part.package.iter_parts() if isinstance(p, FontPart)
        ]
        assert len(font_parts) == 1
        assert font_parts[0].content_type == CT.X_FONTDATA
        assert font_parts[0].blob == _FAKE_OTF_BYTES

    def it_survives_save_and_reopen_round_trip(self, tmp_path):
        """The embedded ``.otf`` font, its typeface name, and its byte payload
        must all survive a full save + reopen cycle, so the resulting .pptx
        can be opened by PowerPoint and the custom OpenType font is available
        to render slide text.
        """
        otf_path = tmp_path / "CustomFont.otf"
        otf_path.write_bytes(_FAKE_OTF_BYTES)

        prs = Presentation()
        prs.embed_font(str(otf_path), "CustomFont", style="regular")

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        # -- typeface survives --
        assert prs2.embedded_fonts == ("CustomFont",)

        # -- font part survives with the original OTF bytes intact --
        font_parts = [
            p for p in prs2.part.package.iter_parts() if isinstance(p, FontPart)
        ]
        assert len(font_parts) == 1
        assert font_parts[0].content_type == CT.X_FONTDATA
        assert font_parts[0].blob == _FAKE_OTF_BYTES

    def it_supports_multiple_otf_styles_under_one_typeface(self):
        """Callers can embed regular + bold (+ italic + boldItalic) ``.otf``
        files under a single typeface, matching PowerPoint's four style slots.
        """
        prs = Presentation()
        regular_bytes = _FAKE_OTF_BYTES
        bold_bytes = _OTF_MAGIC_HEADER + b"\xff" * 252

        prs.embed_font(io.BytesIO(regular_bytes), "CustomFont", style="regular")
        prs.embed_font(io.BytesIO(bold_bytes), "CustomFont", style="bold")

        assert prs.embedded_fonts == ("CustomFont",)
        font_parts = [
            p for p in prs.part.package.iter_parts() if isinstance(p, FontPart)
        ]
        assert len(font_parts) == 2
        blobs = {p.blob for p in font_parts}
        assert blobs == {regular_bytes, bold_bytes}
