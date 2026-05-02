# pyright: reportPrivateUsage=false

"""Regression test for issue #834 and #116 — ``Picture.replace_image``.

Issue #834 (https://github.com/scanny/python-pptx/issues/834) and its
earlier twin #116 ("feature: Picture.replace_image()") both request a
way to swap the pixel bytes of an already-placed picture shape while
keeping its position, size, rotation, cropping, masking shape, and any
other shape-level formatting intact. #116 was deferred in Wave 2
pending Foundation F1 (``PartRelationshipCloner``); F1 has since landed
on master and the replacement turned out to be simpler than cloning —
a new image part is added (or looked up) and the
``p:pic/p:blipFill/a:blip/@r:embed`` relationship is rebound in place.

This test exercises the end-to-end flow:

  1. Load a minimal presentation and add a picture with explicit
     size, rotation and a crop-rectangle.
  2. Call ``picture.replace_image(new_image_file)``.
  3. Round-trip the presentation through save / reopen.
  4. Verify the reloaded picture points at the new image bytes and that
     every geometry / crop property survived verbatim.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.util import Emu, Inches


class DescribeIssue834PictureReplaceImage:
    """Round-trip regression for Picture.replace_image (issues #834, #116)."""

    def it_swaps_the_embedded_image_preserving_geometry(
        self
    ):
        prs = Prs_with_picture_at(
            "tests/test_files/python-powered.png",
            x=Inches(1),
            y=Inches(2),
            cx=Inches(3),
            cy=Inches(4),
        )
        picture = prs.slides[0].shapes[0]

        # -- arrange: apply rotation and crop so we can verify preservation --
        picture.rotation = 45.0
        picture.crop_left = 0.1
        picture.crop_top = 0.2
        picture.crop_right = 0.3
        picture.crop_bottom = 0.4

        # -- original bytes for sanity comparison after replace --
        original_blob = picture.image.blob

        with open("tests/test_files/monty-truth.png", "rb") as f:
            new_bytes = f.read()
        assert new_bytes != original_blob, (
            "test precondition: fixture images must differ"
        )

        # -- act: replace the image --
        picture.replace_image("tests/test_files/monty-truth.png")

        # -- round-trip through save/reload to exercise the rel plumbing --
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        # -- assert: pixel bytes swapped --
        reloaded_picture = reloaded.slides[0].shapes[0]
        assert reloaded_picture.image.blob == new_bytes
        assert reloaded_picture.image.blob != original_blob

        # -- assert: geometry preserved verbatim --
        assert reloaded_picture.left == Inches(1)
        assert reloaded_picture.top == Inches(2)
        assert reloaded_picture.width == Inches(3)
        assert reloaded_picture.height == Inches(4)
        assert reloaded_picture.rotation == 45.0

        # -- assert: crop preserved (within float tolerance) --
        assert abs(reloaded_picture.crop_left - 0.1) < 1e-6
        assert abs(reloaded_picture.crop_top - 0.2) < 1e-6
        assert abs(reloaded_picture.crop_right - 0.3) < 1e-6
        assert abs(reloaded_picture.crop_bottom - 0.4) < 1e-6

    def it_reuses_the_image_part_when_the_replacement_is_the_existing_image(
        self
    ):
        # -- Replacing with the same file content should reuse the image part
        # -- and leave the blip/@r:embed pointing at the same rId.
        prs = Prs_with_picture_at(
            "tests/test_files/python-powered.png",
            x=Emu(0),
            y=Emu(0),
            cx=Inches(2),
            cy=Inches(2),
        )
        picture = prs.slides[0].shapes[0]
        original_rId = picture._pic.blip_rId
        original_image_blob = picture.image.blob

        picture.replace_image("tests/test_files/python-powered.png")

        # -- same content -> same rId, blob unchanged --
        assert picture._pic.blip_rId == original_rId
        assert picture.image.blob == original_image_blob

    def it_adds_a_new_image_part_and_drops_the_old_rel_when_replaced(
        self
    ):
        # -- Count image parts before/after a genuine replacement to confirm
        # -- the old rel is dropped and a new image part is added.
        prs = Prs_with_picture_at(
            "tests/test_files/python-powered.png",
            x=Emu(0),
            y=Emu(0),
            cx=Inches(2),
            cy=Inches(2),
        )
        slide_part = prs.slides[0].part
        picture = prs.slides[0].shapes[0]

        original_rId = picture._pic.blip_rId
        image_rel_rIds_before = {
            rId
            for rId, rel in slide_part.rels.items()
            if rel.reltype.endswith("/image") and not rel.is_external
        }
        assert original_rId in image_rel_rIds_before

        picture.replace_image("tests/test_files/monty-truth.png")

        new_rId = picture._pic.blip_rId
        assert new_rId != original_rId
        image_rel_rIds_after = {
            rId
            for rId, rel in slide_part.rels.items()
            if rel.reltype.endswith("/image") and not rel.is_external
        }
        # -- old rel is gone, new rel is present --
        assert original_rId not in image_rel_rIds_after
        assert new_rId in image_rel_rIds_after

    def it_raises_when_replacing_a_picture_with_no_embedded_image(self):
        # -- unit-level check; no need to restore the PartFactory since we
        # -- construct Picture directly from a cxml fragment.
        from pptx.shapes.picture import Picture

        from .unitutil.cxml import element

        pic = element("p:pic/p:blipFill/a:blip")
        picture = Picture(pic, None)

        with pytest.raises(ValueError, match="no embedded image to replace"):
            picture.replace_image("tests/test_files/monty-truth.png")


# -- helpers -----------------------------------------------------------


def Prs_with_picture_at(image_path, x, y, cx, cy):
    """Return a fresh |Presentation| with a single slide holding one picture."""
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # -- blank layout --
    slide = prs.slides.add_slide(blank_layout)
    slide.shapes.add_picture(image_path, x, y, cx, cy)
    return prs
