# pyright: reportPrivateUsage=false

"""Regression test for issue #835 — merge charts from different slides into 1.

Issue #835 (https://github.com/scanny/python-pptx/issues/835) asks:

    "Is it possible to merge two charts from different slides into 1 slide?
     Chart-1 is placed in slide-1 and chart-2 is placed in slide-2 and I
     want to combine chart-1 and chart-2 in slide-3 in another deck."

The underlying capability the reporter was missing — **moving a chart-bearing
slide from one presentation into another with the chart data intact** — is
served by two APIs shipped in earlier waves:

  * :meth:`Slides.add_slide_from_external` — per-slide control for pulling
    one chart-bearing slide from a source deck into a target deck.
  * :meth:`Presentation.merge` — bulk path that appends every slide of a
    source deck (including every chart-bearing slide) to a target deck.

Both go through the Wave 7 #934 cross-package cloner so chart parts get a
*distinct* embedded workbook in the target, and both round-trip through
``save`` + reopen with the cached chart values intact.

This regression suite pins both paths for the #835 scenario so the issue
can be closed as *resolved by the cross-presentation slide-copy wave*:

  * A source deck carrying three chart slides (with distinct values) can be
    copied one slide at a time into an empty target via
    :meth:`Slides.add_slide_from_external` — the chart's
    ``plots[0].series[0].values`` match the source.
  * :meth:`Presentation.merge` appends every chart slide from the source to
    the target in one call, preserving each chart's series values.
  * Save + reopen round-trip preserves the cached values of every merged
    chart.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

import pytest

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

if TYPE_CHECKING:
    from pptx.chart.chart import Chart
    from pptx.presentation import Presentation as PresentationClass
    from pptx.slide import Slide


# -- three distinct chart-data fingerprints so we can identify which source
# -- slide's chart ended up in which target position ---------------------
_CHART_DATA = [
    (("Q1", "Q2", "Q3"), (10.0, 20.0, 30.0), "Alpha"),
    (("Q1", "Q2", "Q3"), (40.0, 50.0, 60.0), "Bravo"),
    (("Q1", "Q2", "Q3"), (70.0, 80.0, 90.0), "Charlie"),
]


@pytest.fixture
def _restore_part_factory():
    """Guard against ``PartFactory.part_type_for`` pollution from other tests.

    Mirrors the identical guard used by ``tests/test_issue_175_*.py`` and
    ``tests/test_issue_403_*.py``; earlier test modules overwrite slide-part
    registrations with mocks and do not restore them, which breaks
    ``Presentation(stream)`` reopen in save/reopen round-trip assertions.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.embeddedpackage import EmbeddedXlsxPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
        CT.SML_SHEET: EmbeddedXlsxPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


def _build_source_deck() -> PresentationClass:
    """Return a presentation with three slides, each carrying one chart.

    Each chart carries a distinct ``(categories, values, series name)``
    fingerprint so per-slide identity can be verified after copying into
    another presentation.
    """
    prs = Presentation()
    for categories, values, series_name in _CHART_DATA:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        data = CategoryChartData()
        data.categories = list(categories)
        data.add_series(series_name, values)
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(3),
            data,
        )
    return prs


def _first_chart_on(slide: Slide) -> Chart:
    for shape in slide.shapes:
        if getattr(shape, "has_chart", False):
            return shape.chart
    raise AssertionError("no chart on slide")


def _values_of(slide: Slide) -> tuple[float, ...]:
    chart = _first_chart_on(slide)
    return tuple(chart.plots[0].series[0].values)


def _save_and_reopen(prs: PresentationClass) -> PresentationClass:
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)


class DescribeIssue835MergeChartsFromDifferentSlides(object):
    """#835 — close as resolved by the cross-presentation slide-copy wave.

    Pins the "pull chart-bearing slides from deck A into deck B" workflow
    the #835 reporter asked for, through both
    :meth:`Slides.add_slide_from_external` (per-slide) and
    :meth:`Presentation.merge` (bulk), including save+reopen round-trip
    and exact chart-data preservation.
    """

    # -- (1) per-slide path: add_slide_from_external preserves chart data --

    def it_copies_a_chart_bearing_slide_via_add_slide_from_external(self, _restore_part_factory):
        """The headline #835 scenario via the per-slide API.

        Build a source deck with three distinct chart slides, then pull
        one chart-bearing slide into an empty target deck. The appended
        slide must carry a chart whose ``series[0].values`` match the
        source slide's chart.
        """
        source = _build_source_deck()
        expected_values = _values_of(source.slides[1])  # middle chart

        target = Presentation()
        target_layout = target.slide_masters[0].slide_layouts[5]
        appended = target.slides.add_slide_from_external(source.slides[1], target_layout)

        assert _values_of(appended) == expected_values

    def it_preserves_chart_values_through_save_and_reopen_after_external_copy(
        self, _restore_part_factory
    ):
        """The per-slide #835 path survives the save+reopen round-trip.

        Pull all three chart-bearing slides into an empty target one at a
        time, save, reopen, and confirm every chart's series values match
        the source in the right order.
        """
        source = _build_source_deck()
        expected = [_values_of(sl) for sl in source.slides]

        target = Presentation()
        target_layout = target.slide_masters[0].slide_layouts[5]
        for src_slide in source.slides:
            target.slides.add_slide_from_external(src_slide, target_layout)

        reopened = _save_and_reopen(target)

        assert len(reopened.slides) == 3
        assert [_values_of(sl) for sl in reopened.slides] == expected

    # -- (2) bulk path: Presentation.merge preserves every chart's data ----

    def it_copies_every_chart_slide_via_presentation_merge(self, _restore_part_factory):
        """The bulk #835 scenario via :meth:`Presentation.merge`.

        Merging a source deck with three chart slides into an empty target
        appends all three chart slides, each carrying a chart whose
        ``series[0].values`` match the source's in source order.
        """
        source = _build_source_deck()
        expected = [_values_of(sl) for sl in source.slides]

        target = Presentation()
        appended = target.merge(source)

        assert len(appended) == 3
        assert len(target.slides) == 3
        assert [_values_of(sl) for sl in target.slides] == expected

    def it_preserves_every_merged_chart_through_save_and_reopen(self, _restore_part_factory):
        """The bulk #835 path survives the save+reopen round-trip.

        After ``merge`` + ``save`` + reopen, every chart slide in the
        target deck carries a chart whose series values still match the
        source deck's corresponding chart.
        """
        source = _build_source_deck()
        expected = [_values_of(sl) for sl in source.slides]

        target = Presentation()
        target.merge(source)

        reopened = _save_and_reopen(target)

        assert len(reopened.slides) == 3
        assert [_values_of(sl) for sl in reopened.slides] == expected

    # -- (3) cross-package isolation: merged charts own their workbook -----

    def it_gives_each_merged_chart_a_distinct_embedded_workbook(self, _restore_part_factory):
        """Wave 7 #934's full-fidelity guarantee, reaffirmed for #835.

        The cross-package cloner must install a fresh
        :class:`EmbeddedXlsxPart` in the target for every chart it copies,
        otherwise PowerPoint's "Edit Data" dialog on the merged deck would
        mutate the source deck's workbook.
        """
        from pptx.parts.embeddedpackage import EmbeddedXlsxPart

        source = _build_source_deck()
        source_xlsx_ids = {
            id(p) for p in source.part.package.iter_parts() if isinstance(p, EmbeddedXlsxPart)
        }

        target = Presentation()
        target.merge(source)

        target_xlsx_parts = [
            p for p in target.part.package.iter_parts() if isinstance(p, EmbeddedXlsxPart)
        ]
        assert target_xlsx_parts  # at least one xlsx materialised in target
        target_xlsx_ids = {id(p) for p in target_xlsx_parts}
        assert not target_xlsx_ids.intersection(source_xlsx_ids)
