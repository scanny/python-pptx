# pyright: reportPrivateUsage=false

"""Regression test for issue #838 — bounding box parsing for shapes inside groups.

Issue #838 (https://github.com/scanny/python-pptx/issues/838) reported that
``shape.left`` / ``shape.top`` / ``shape.width`` / ``shape.height`` returned
"wrong" values for shapes nested inside a :class:`.GroupShape` that had been
resized in PowerPoint. The root cause was not a parsing bug but rather the
OOXML semantics of ``p:grpSp``: when a group is resized, PowerPoint keeps the
child's raw ``a:off`` / ``a:ext`` unchanged and instead scales the group's
own ``a:ext`` relative to its ``a:chExt`` (child extent). The raw child
numbers are therefore expressed in the group's *local* coordinate system and
no longer match the slide-relative rectangle the shape actually renders at.

The fix for #925 (``feat/issue-925-group-shape-transform``) added four
read-only companion properties on :class:`.BaseShape` that composite each
enclosing group's ``a:chOff`` / ``a:chExt`` → ``a:off`` / ``a:ext`` linear
transform and return slide-relative values:

* :attr:`.BaseShape.effective_left`
* :attr:`.BaseShape.effective_top`
* :attr:`.BaseShape.effective_width`
* :attr:`.BaseShape.effective_height`

These are exactly the API surface #838 asked for. This suite verifies the
#925 properties resolve #838's reported scenarios and can be closed as
*resolved by #925*. The raw ``left``/``top``/``width``/``height`` still
return the group-local numbers for backwards compatibility.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

import pytest

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

if TYPE_CHECKING:
    from pptx.shapes.group import GroupShape


def _resize_group_on_slide(grp: GroupShape, new_cx: int, new_cy: int) -> None:
    """Scale `grp`'s slide-frame ``a:ext`` to (new_cx, new_cy) WITHOUT touching
    ``a:chExt`` — mimicking how PowerPoint records a user-resized group.

    After ``add_group_shape`` + ``recalculate_extents`` the group's
    ``a:ext`` equals its ``a:chExt``; children then render at their raw
    local size. To simulate the post-resize state that #838 reported we
    must diverge ``a:ext`` from ``a:chExt`` directly on the ``a:xfrm``
    element.
    """
    xfrm = grp._element.xfrm
    assert xfrm is not None
    ext = xfrm.find(qn("a:ext"))
    assert ext is not None
    ext.set("cx", str(new_cx))
    ext.set("cy", str(new_cy))


@pytest.fixture
def group_with_two_children_resized_half():
    """Return ``(prs, slide, grp, s1, s2)`` for a group that's been resized
    to half its original width on the slide.

    The raw child ``a:off``/``a:ext`` reflect authoring-time positions
    (``Inches(1)`` / ``Inches(3)`` etc); the group's slide-frame
    ``a:ext.cx`` has been halved relative to its ``a:chExt.cx``, so each
    child should render at half its raw width on the slide.
    """
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    s1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    s2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4), Inches(2), Inches(2), Inches(1))
    grp = slide.shapes.add_group_shape([s1, s2])

    # -- group now spans Inches(1..6) x Inches(1..3) with chExt == ext     --
    raw_cx = grp.width
    raw_cy = grp.height
    # -- Simulate the user resizing the group 2:1 horizontally in PowerPoint --
    _resize_group_on_slide(grp, raw_cx // 2, raw_cy)
    return prs, slide, grp, s1, s2


class DescribeIssue838GroupChildBoundingBox(object):
    """#838 group-child bounding-box parsing verify-and-close via #925.

    Pins that ``effective_left`` / ``effective_top`` / ``effective_width``
    / ``effective_height`` resolve the slide-relative rectangle the
    reporter expected, while the raw ``left`` / ``top`` / ``width`` /
    ``height`` continue to return the group-local values.
    """

    # -- the reported problem: raw values are group-local, not slide-relative --

    def it_returns_group_local_raw_left_for_child_of_resized_group(
        self, group_with_two_children_resized_half
    ):
        _, _, _, s1, _ = group_with_two_children_resized_half
        # -- raw left equals the authoring-time offset, unchanged by the    --
        # -- group resize. This is the "wrong-looking" value #838 reported. --
        assert s1.left == Inches(1)

    def it_returns_group_local_raw_width_for_child_of_resized_group(
        self, group_with_two_children_resized_half
    ):
        _, _, _, s1, _ = group_with_two_children_resized_half
        assert s1.width == Inches(2)

    # -- the #925 API the #838 reporter asked for ---------------------------

    def it_composites_the_group_transform_into_effective_left(
        self, group_with_two_children_resized_half
    ):
        # -- group was resized 2:1 in x: sx = ext.cx / chExt.cx = 0.5; the  --
        # -- group's a:off.x is unchanged at Inches(1), its chOff is       --
        # -- Inches(1). So child-local Inches(1) maps back to slide        --
        # -- Inches(1) (the anchor of the group itself).                    --
        _, _, grp, s1, _ = group_with_two_children_resized_half
        assert s1.effective_left == grp.left  # == Inches(1)

    def it_composites_the_group_transform_into_effective_width(
        self, group_with_two_children_resized_half
    ):
        _, _, _, s1, _ = group_with_two_children_resized_half
        # -- sx = 0.5 applied to raw cx (Inches(2)) -> Inches(1) on slide --
        assert s1.effective_width == Inches(1)

    def it_leaves_vertical_axis_unscaled_when_only_x_was_resized(
        self, group_with_two_children_resized_half
    ):
        # -- group chExt.cy == ext.cy so sy = 1.0 — vertical values pass --
        # -- through unchanged.                                           --
        _, _, _, s1, _ = group_with_two_children_resized_half
        assert s1.effective_top == Inches(1)
        assert s1.effective_height == Inches(1)

    def it_applies_the_transform_independently_to_each_child(
        self, group_with_two_children_resized_half
    ):
        # -- second child is further right in the group's local frame; its --
        # -- slide position must also be composited (not simply equal to   --
        # -- the group anchor).                                             --
        _, _, grp, _, s2 = group_with_two_children_resized_half
        # -- s2 local: off.x = Inches(4), chOff.x = Inches(1), sx = 0.5    --
        # --   effective_left = Inches(1) + (Inches(4) - Inches(1)) * 0.5  --
        # --                  = Inches(1) + Inches(1.5) = Inches(2.5)      --
        expected_left = grp.left + (Inches(4) - Inches(1)) // 2
        assert s2.effective_left == expected_left
        # -- effective_width = Inches(2) * 0.5 = Inches(1) --
        assert s2.effective_width == Inches(1)

    # -- reporter's expectation: a top-level shape round-trips unchanged ----

    def it_returns_raw_values_for_a_top_level_shape(self):
        # -- without a group ancestor, effective_* == raw (no compositing). --
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(2), Inches(3), Inches(4), Inches(5)
        )

        assert shape.effective_left == Inches(2)
        assert shape.effective_top == Inches(3)
        assert shape.effective_width == Inches(4)
        assert shape.effective_height == Inches(5)

    # -- the reported `GroupShape` itself also reports correct geometry ----

    def it_reports_effective_geometry_for_a_nested_group_shape(self):
        # -- A GroupShape nested inside another resized group also needs    --
        # -- its own position composited. This covers the #838 follow-up   --
        # -- where the reporter asked about nested groups.                 --
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        # -- inner group's children placed away from the top-left so the    --
        # -- outer's linear transform moves the inner group's anchor too.   --
        inner_s1 = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(3), Inches(1), Inches(2), Inches(2)
        )
        inner_s2 = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(6), Inches(1), Inches(2), Inches(2)
        )
        inner_grp = slide.shapes.add_group_shape([inner_s1, inner_s2])
        # -- Now wrap that inner group in an outer group, then resize the  --
        # -- outer group 2:1 horizontally on the slide.                    --
        outer_grp = slide.shapes.add_group_shape([inner_grp])
        outer_raw_cx = outer_grp.width
        outer_raw_cy = outer_grp.height
        _resize_group_on_slide(outer_grp, outer_raw_cx // 2, outer_raw_cy)

        # -- inner group's own raw left is preserved by the outer resize;  --
        # -- its effective_left is composited through the outer's 0.5 x.   --
        # -- outer: off.x = chOff.x = Inches(3), sx = 0.5                  --
        # -- inner local: off.x = Inches(3)                                --
        # -- effective = Inches(3) + (Inches(3) - Inches(3)) * 0.5         --
        # --   => Inches(3) (at the outer's origin)                        --
        assert inner_grp.left == Inches(3)
        assert inner_grp.effective_left == Inches(3)
        # -- inner group's effective width is halved by the outer scale    --
        assert inner_grp.width == Inches(5)  # spans Inches(3..8)
        assert inner_grp.effective_width == Inches(5) // 2

        # -- the innermost shape composites through both groups. inner_s2 --
        # -- is at local off.x = Inches(6); the inner group's ext equals  --
        # -- its chExt so its scale is 1.0, then the outer scales 0.5:    --
        # --   inner stage: Inches(6) (unchanged, sy=sx=1.0)                --
        # --   outer stage: Inches(3) + (Inches(6) - Inches(3)) * 0.5      --
        # --              = Inches(3) + Inches(1.5) = Inches(4.5)          --
        expected = Inches(3) + (Inches(6) - Inches(3)) // 2
        assert inner_s2.effective_left == expected
        # -- effective_width = raw Inches(2) * 0.5 = Inches(1) --
        assert inner_s2.effective_width == Inches(1)

    # -- round-trip: values survive save/reopen ----------------------------

    def it_round_trips_effective_geometry_through_save_reopen(
        self, group_with_two_children_resized_half
    ):
        prs, _, _, s1_orig, _ = group_with_two_children_resized_half
        orig_eff_left = s1_orig.effective_left
        orig_eff_width = s1_orig.effective_width
        s1_id = s1_orig.shape_id

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        slide2 = prs2.slides[0]

        # -- find the group and the reopened child by shape id --
        grp2 = next(s for s in slide2.shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP)
        s1_reopened = next(s for s in grp2.shapes if s.shape_id == s1_id)

        # -- raw values are preserved unchanged --
        assert s1_reopened.left == Inches(1)
        assert s1_reopened.width == Inches(2)
        # -- composited effective values are preserved unchanged --
        assert s1_reopened.effective_left == orig_eff_left
        assert s1_reopened.effective_width == orig_eff_width
