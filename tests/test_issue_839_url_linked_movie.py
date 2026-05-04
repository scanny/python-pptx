# pyright: reportPrivateUsage=false

"""Regression test for issue #839 — "Add youtube video to presentation (via url)".

Issue #839 (https://github.com/scanny/python-pptx/issues/839) reports that
``SlideShapes.add_movie`` unconditionally embeds the media bytes as a
``MediaPart``, with no code path for a URL-linked (online) video. PowerPoint
supports this via its "Insert Online Video" command, which writes a
``p:pic`` whose ``a:videoFile`` and ``p14:media`` descriptors both carry
``r:link`` pointing at an *external* relationship (``TargetMode="External"``),
instead of a media part inside the package.

:meth:`.SlideShapes.add_movie_link` closes this gap. The caller supplies:

  * a URL — for YouTube, the ``https://www.youtube.com/embed/<ID>`` form is
    what PowerPoint's player loads; ordinary ``watch?v=…`` URLs don't play
    in-slide;
  * a poster-frame image — required, because there's no media part from
    which to derive a default loudspeaker graphic;
  * left/top/width/height — position and size in EMU.

The emitted XML matches what PowerPoint writes for an online-video shape:
both ``a:videoFile`` and the ``p14:media`` descriptor reference the same
external-URL relationship via ``r:link``. No media bytes are copied into the
``.pptx``.

This suite pins the round-trip behavior end-to-end against a fresh
``Presentation()``.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.shapes.picture import Movie
from pptx.util import Inches

# --- Minimal 1x1 transparent PNG used as a poster frame. ------------------
_PNG_BYTES = (
    b"\x89PNG\r\n\x1a\n"
    b"\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
    b"\x00\x00\x00\rIDATx\x9cc\xf8\xcf\xc0\x00\x00\x00\x03\x00\x01\xde\xfb\xc4f"
    b"\x00\x00\x00\x00IEND\xaeB\x60\x82"
)

_YT_URL = "https://www.youtube.com/embed/dQw4w9WgXcQ"


class DescribeIssue839UrlLinkedMovie:
    """Acceptance criteria for URL-linked (online) video insertion."""

    def it_returns_a_Movie(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        movie = slide.shapes.add_movie_link(
            _YT_URL, io.BytesIO(_PNG_BYTES), Inches(1), Inches(1), Inches(4), Inches(3)
        )

        assert isinstance(movie, Movie)

    def it_writes_videoFile_and_p14_media_with_r_link(self):
        """Both descriptors point at the same external-URL rId via r:link."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        movie = slide.shapes.add_movie_link(
            _YT_URL, io.BytesIO(_PNG_BYTES), Inches(1), Inches(1), Inches(4), Inches(3)
        )

        pic = movie._element
        video_links = pic.xpath("./p:nvPicPr/p:nvPr/a:videoFile/@r:link")
        media_links = pic.xpath("./p:nvPicPr/p:nvPr/p:extLst/p:ext/p14:media/@r:link")
        assert len(video_links) == 1
        assert len(media_links) == 1
        assert video_links[0] == media_links[0]
        # -- embedded-media r:embed must NOT be emitted for a link-only shape --
        assert pic.xpath("./p:nvPicPr/p:nvPr/p:extLst/p:ext/p14:media/@r:embed") == []

    def it_creates_a_NoRels_external_VIDEO_relationship(self):
        """The url becomes a VIDEO-typed external-mode rel on the slide part."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        movie = slide.shapes.add_movie_link(
            _YT_URL, io.BytesIO(_PNG_BYTES), Inches(1), Inches(1), Inches(4), Inches(3)
        )

        rId = movie._element.xpath("./p:nvPicPr/p:nvPr/a:videoFile/@r:link")[0]
        rel = slide.part.rels[rId]
        assert rel.reltype == RT.VIDEO
        assert rel.is_external
        assert rel.target_ref == _YT_URL

    def it_embeds_the_poster_frame(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        movie = slide.shapes.add_movie_link(
            _YT_URL, io.BytesIO(_PNG_BYTES), Inches(1), Inches(1), Inches(4), Inches(3)
        )

        poster = movie.poster_frame
        assert poster is not None
        assert poster.blob == _PNG_BYTES

    def it_does_not_embed_any_media_bytes(self):
        """No MediaPart is created — a `blob` fetch therefore comes back None."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        movie = slide.shapes.add_movie_link(
            _YT_URL, io.BytesIO(_PNG_BYTES), Inches(1), Inches(1), Inches(4), Inches(3)
        )

        # -- blob reads the MediaPart; there is none for a link-only movie --
        assert movie.blob is None
        assert movie.content_type is None

    def it_survives_a_save_load_roundtrip(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.add_movie_link(
            _YT_URL, io.BytesIO(_PNG_BYTES), Inches(1), Inches(1), Inches(4), Inches(3)
        )

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        slide2 = prs2.slides[0]
        movies = [sh for sh in slide2.shapes if isinstance(sh, Movie)]
        assert len(movies) == 1
        movie = movies[0]
        rId = movie._element.xpath("./p:nvPicPr/p:nvPr/a:videoFile/@r:link")[0]
        rel = slide2.part.rels[rId]
        assert rel.is_external
        assert rel.target_ref == _YT_URL

    def it_positions_the_shape_where_requested(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        movie = slide.shapes.add_movie_link(
            _YT_URL,
            io.BytesIO(_PNG_BYTES),
            Inches(2),
            Inches(3),
            Inches(5),
            Inches(4),
        )

        assert movie.left == Inches(2)
        assert movie.top == Inches(3)
        assert movie.width == Inches(5)
        assert movie.height == Inches(4)

    @pytest.mark.parametrize(
        "url",
        [
            "https://www.youtube.com/embed/dQw4w9WgXcQ",
            "https://player.vimeo.com/video/123456789",
            "https://example.com/videos/clip.mp4",
        ],
    )
    def it_accepts_any_URL_the_caller_supplies(self, url):
        """`add_movie_link` doesn't validate the URL — the caller picks a playable form."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        movie = slide.shapes.add_movie_link(
            url, io.BytesIO(_PNG_BYTES), Inches(1), Inches(1), Inches(4), Inches(3)
        )

        rId = movie._element.xpath("./p:nvPicPr/p:nvPr/a:videoFile/@r:link")[0]
        assert slide.part.rels[rId].target_ref == url
