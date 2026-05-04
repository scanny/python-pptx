# pyright: reportPrivateUsage=false, reportUnknownVariableType=false
# pyright: reportUnknownMemberType=false, reportAttributeAccessIssue=false
# pyright: reportOptionalMemberAccess=false

"""Verify-close regression test for issue #848 — ``_font_scale`` not
available for ``text_frame``.

Issue #848 (https://github.com/scanny/python-pptx/issues/848) asked for
programmatic access to the font-scale-reduction and line-spacing-
reduction values that PowerPoint writes into the ``a:normAutofit``
element under an ``a:bodyPr`` when autofit shrinks overflowing text.
Until the fork-era resolutions landed, callers had no public surface
to read or emit those attributes, even though the underlying XML is
what PowerPoint uses to render the shrink-on-overflow effect.

This is a verify-close pin: the feature shipped in two parts —

* **#715** (Wave 12) — ``TextFrame.font_scale`` read/write on the
  ``fontScale`` attribute of ``a:normAutofit``.
* **#969** (Wave 14) — ``TextFrame.line_space_reduction`` read/write
  on the ``lnSpcReduction`` attribute of ``a:normAutofit``, plus the
  shared "write to the ``a:normAutofit`` choice sibling, creating
  one if needed and replacing any ``a:noAutofit`` / ``a:spAutoFit``
  sibling" behavior.

This suite exercises the real public API on a fresh ``Presentation()``
— no mocks — and pins the complete user-visible contract: defaults,
XML emission with the correct ``parts-per-hundred-thousand``
encoding, round-trip through save and reopen, and the "clear to
default" behavior that removes the attribute.

If the feature is ever regressed so callers again cannot read or set
``font_scale`` / ``line_space_reduction`` on a ``TextFrame``, the
tests here will fail loudly and point at this issue.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

import pytest

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

if TYPE_CHECKING:
    from pptx.presentation import Presentation as PresentationType

# -- helpers --------------------------------------------------------------


def _fresh_textbox():
    """Return ``(prs, text_frame)`` — a new textbox on a blank slide."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    return prs, tb.text_frame


def _roundtrip(prs: "PresentationType"):
    """Serialize `prs` to BytesIO and reopen. Return the reloaded Presentation."""
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)


class DescribeIssue848FontScaleVerify(object):
    """Verify-close regression suite for issue #848.

    Pins ``TextFrame.font_scale`` and ``TextFrame.line_space_reduction``
    against the resolutions delivered by #715 and #969.
    """

    # -- 1: defaults on an untouched TextFrame ------------------------------

    def it_reads_the_default_font_scale_as_100_percent(self):
        # -- Pin: a fresh textbox has no a:normAutofit, so font_scale
        # -- reads as the PowerPoint default of 100.0% (no reduction).
        # -- Resolves the read half of #848 (delivered via #715).
        _, tf = _fresh_textbox()

        assert tf.font_scale == 100.0

    def it_reads_the_default_line_space_reduction_as_0_percent(self):
        # -- Pin: a fresh textbox has no a:normAutofit, so
        # -- line_space_reduction reads as the PowerPoint default of
        # -- 0.0% (no reduction).
        # -- Resolves the read half of #848 (delivered via #969).
        _, tf = _fresh_textbox()

        assert tf.line_space_reduction == 0.0

    # -- 2: setter emits the correct parts-per-hundred-thousand XML --------

    def it_emits_normAutofit_fontScale_in_parts_per_hundred_thousand(self):
        # -- Pin: setting font_scale = 50.0 (a 50% reduction) emits
        # -- <a:normAutofit fontScale="50000"/> under <a:bodyPr>. The
        # -- OOXML integer encoding is parts per hundred thousand:
        # -- 50% -> 50000. Resolves the write half of #848 via #715.
        _, tf = _fresh_textbox()

        tf.font_scale = 50.0

        bodyPr = tf._txBody.bodyPr
        normAutofit = bodyPr.find(qn("a:normAutofit"))
        assert normAutofit is not None
        assert normAutofit.get("fontScale") == "50000"
        # -- readback returns the percent float --
        assert tf.font_scale == 50.0

    def it_emits_normAutofit_lnSpcReduction_in_parts_per_hundred_thousand(self):
        # -- Pin: setting line_space_reduction = 20.0 (a 20% reduction)
        # -- emits <a:normAutofit lnSpcReduction="20000"/> under
        # -- <a:bodyPr>. Same parts-per-hundred-thousand encoding.
        # -- Resolves the write half of #848 via #969.
        _, tf = _fresh_textbox()

        tf.line_space_reduction = 20.0

        bodyPr = tf._txBody.bodyPr
        normAutofit = bodyPr.find(qn("a:normAutofit"))
        assert normAutofit is not None
        assert normAutofit.get("lnSpcReduction") == "20000"
        # -- readback returns the percent float --
        assert tf.line_space_reduction == 20.0

    def it_supports_setting_both_attributes_on_the_same_normAutofit(self):
        # -- Pin: the two setters share a single a:normAutofit child —
        # -- setting one does not clobber the other's attribute. This
        # -- is the behavior PowerPoint emits when both reductions
        # -- apply to the same shape.
        _, tf = _fresh_textbox()

        tf.font_scale = 85.0
        tf.line_space_reduction = 10.0

        bodyPr = tf._txBody.bodyPr
        normAutofit = bodyPr.find(qn("a:normAutofit"))
        assert normAutofit is not None
        assert normAutofit.get("fontScale") == "85000"
        assert normAutofit.get("lnSpcReduction") == "10000"
        assert tf.font_scale == 85.0
        assert tf.line_space_reduction == 10.0

    # -- 3: round-trip through save and reopen -----------------------------

    def it_round_trips_font_scale_through_save_and_reopen(self):
        # -- Pin: font_scale survives serialize + reopen. This catches
        # -- a serializer regression that silently drops the
        # -- a:normAutofit child or its fontScale attribute.
        prs, tf = _fresh_textbox()
        tf.text = "shrink-on-overflow"
        tf.font_scale = 75.0

        reloaded = _roundtrip(prs)

        reloaded_shapes = [s for s in reloaded.slides[0].shapes if s.has_text_frame]
        reloaded_tf = reloaded_shapes[0].text_frame
        assert reloaded_tf.font_scale == 75.0

    def it_round_trips_line_space_reduction_through_save_and_reopen(self):
        # -- Pin: line_space_reduction survives serialize + reopen.
        prs, tf = _fresh_textbox()
        tf.text = "shrink-on-overflow"
        tf.line_space_reduction = 15.0

        reloaded = _roundtrip(prs)

        reloaded_shapes = [s for s in reloaded.slides[0].shapes if s.has_text_frame]
        reloaded_tf = reloaded_shapes[0].text_frame
        assert reloaded_tf.line_space_reduction == 15.0

    def it_round_trips_both_attributes_on_the_same_normAutofit(self):
        # -- Pin: when both attributes are set, both survive the
        # -- round-trip on a single a:normAutofit child.
        prs, tf = _fresh_textbox()
        tf.text = "double-reduction"
        tf.font_scale = 60.0
        tf.line_space_reduction = 25.0

        reloaded = _roundtrip(prs)

        reloaded_shapes = [s for s in reloaded.slides[0].shapes if s.has_text_frame]
        reloaded_tf = reloaded_shapes[0].text_frame
        assert reloaded_tf.font_scale == 60.0
        assert reloaded_tf.line_space_reduction == 25.0

    # -- 4: "clear to default" behavior ------------------------------------

    def it_clears_fontScale_when_assigned_the_default_of_100(self):
        # -- Pin: assigning the default value of 100.0 removes the
        # -- fontScale attribute from a:normAutofit (while leaving
        # -- the element itself in place, matching the existing unit
        # -- test in tests/text/test_text.py). Resolves the "how do I
        # -- clear it again" question from #848.
        _, tf = _fresh_textbox()
        tf.font_scale = 85.0
        # -- sanity: the attribute exists after the initial set --
        assert tf._txBody.bodyPr.find(qn("a:normAutofit")).get("fontScale") == "85000"

        tf.font_scale = 100.0

        normAutofit = tf._txBody.bodyPr.find(qn("a:normAutofit"))
        assert normAutofit is not None
        assert normAutofit.get("fontScale") is None
        assert tf.font_scale == 100.0

    def it_clears_lnSpcReduction_when_assigned_the_default_of_0(self):
        # -- Pin: assigning 0.0 removes the lnSpcReduction attribute
        # -- from a:normAutofit. Symmetric "clear" behavior to the
        # -- fontScale case above.
        _, tf = _fresh_textbox()
        tf.line_space_reduction = 20.0
        assert tf._txBody.bodyPr.find(qn("a:normAutofit")).get("lnSpcReduction") == "20000"

        tf.line_space_reduction = 0.0

        normAutofit = tf._txBody.bodyPr.find(qn("a:normAutofit"))
        assert normAutofit is not None
        assert normAutofit.get("lnSpcReduction") is None
        assert tf.line_space_reduction == 0.0

    # -- 5: input-validation contract --------------------------------------

    def it_rejects_font_scale_values_outside_the_valid_range(self):
        # -- Pin: font_scale only accepts values in 1.0..100.0
        # -- (inclusive); anything else raises ValueError rather than
        # -- silently producing an out-of-spec @fontScale that
        # -- PowerPoint rejects at open time.
        _, tf = _fresh_textbox()

        with pytest.raises(ValueError):
            tf.font_scale = 0.5
        with pytest.raises(ValueError):
            tf.font_scale = 150.0

    def it_rejects_line_space_reduction_values_outside_the_valid_range(self):
        # -- Pin: line_space_reduction only accepts values in
        # -- 0.0..100.0 (inclusive); anything else raises ValueError.
        _, tf = _fresh_textbox()

        with pytest.raises(ValueError):
            tf.line_space_reduction = -1.0
        with pytest.raises(ValueError):
            tf.line_space_reduction = 150.0
