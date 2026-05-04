# pyright: reportPrivateUsage=false

"""Regression test for issue #872 — marker border color on LINE_MARKERS charts.

Issue #872 (https://github.com/scanny/python-pptx/issues/872) asked for a way to
set the *border* color of data-point markers on a line chart with markers. The
target OOXML shape is::

    c:ser
      c:marker
        c:spPr
          a:ln
            a:solidFill
              a:srgbClr @val

The capability is already reachable through the existing public API:
``series.marker`` (``_MarkerMixin.marker``) returns a :class:`~pptx.chart.marker.Marker`
proxy whose ``.format`` property returns a :class:`~pptx.dml.chtfmt.ChartFormat`
wrapping the ``c:marker`` element's ``c:spPr``. From there, ``format.line``
provides the standard :class:`~pptx.dml.line.LineFormat` surface — ``.color.rgb``
writes the ``a:ln/a:solidFill/a:srgbClr`` tuple, and ``.width`` writes
``a:ln/@w``.

This suite pins the resolution so #872 can be closed. Each test exercises the
real public API on a freshly-built ``Presentation()`` (no mocks), writes the
expected XML, and — where applicable — confirms the value round-trips through
save + reopen.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.chart.marker import Marker
from pptx.dml.chtfmt import ChartFormat
from pptx.dml.color import RGBColor
from pptx.dml.line import LineFormat
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of `PartFactory.part_type_for`.

    Mirrors the guard used by other ``test_issue_*_verify.py`` suites that
    exercise save + reopen round-trips.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.embeddedpackage import EmbeddedXlsxPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.SML_SHEET: EmbeddedXlsxPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


def _make_line_markers_chart(prs):
    """Return the chart object of a fresh LINE_MARKERS chart on a new slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    data = CategoryChartData()
    data.categories = ["A", "B", "C"]
    data.add_series("Ser1", (1.0, 2.0, 3.0))
    return slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS, Inches(1), Inches(1), Inches(5), Inches(3), data
    ).chart


class DescribeIssue872MarkerBorderColorVerify(object):
    """#872 verify-and-close: marker border color via series.marker.format.line."""

    # -- 1: the public API chain is wired through ---------------------------

    def it_exposes_marker_format_line_on_a_line_markers_series(self):
        """Pins the proxy chain asked for by the reporter.

        ``series.marker`` returns a ``Marker``; ``marker.format`` returns a
        ``ChartFormat``; ``format.line`` returns a ``LineFormat``. These are
        the steps a caller takes to reach the marker-border line.
        """
        prs = Presentation()
        chart = _make_line_markers_chart(prs)
        series = chart.plots[0].series[0]

        assert isinstance(series.marker, Marker)
        assert isinstance(series.marker.format, ChartFormat)
        assert isinstance(series.marker.format.line, LineFormat)

    # -- 2: the authoring recipe emits the expected XML ---------------------

    def it_writes_marker_border_color_to_c_ser_c_marker_c_spPr_a_ln(self):
        """Pins that ``series.marker.format.line.color.rgb = RGBColor(...)``
        emits ``c:ser/c:marker/c:spPr/a:ln/a:solidFill/a:srgbClr @val``.
        """
        prs = Presentation()
        chart = _make_line_markers_chart(prs)
        series = chart.plots[0].series[0]

        # -- act: the issue-#872 recipe --
        series.marker.format.line.color.rgb = RGBColor(0xFF, 0x00, 0x00)

        # -- assert: the exact element path the reporter needs --
        ser = series._element
        ln = ser.xpath("c:marker/c:spPr/a:ln")
        assert len(ln) == 1
        srgb = ser.xpath("c:marker/c:spPr/a:ln/a:solidFill/a:srgbClr")
        assert len(srgb) == 1
        assert srgb[0].get("val") == "FF0000"

    def and_it_also_supports_setting_the_marker_border_width(self):
        """Border width (the companion to color) also writes under the same
        ``c:marker/c:spPr/a:ln`` element via ``LineFormat.width``.
        """
        prs = Presentation()
        chart = _make_line_markers_chart(prs)
        series = chart.plots[0].series[0]

        series.marker.format.line.color.rgb = RGBColor(0x33, 0x66, 0x99)
        series.marker.format.line.width = Pt(2)

        ser = series._element
        ln = ser.xpath("c:marker/c:spPr/a:ln")
        assert len(ln) == 1
        # -- Pt(2) -> 2 * 12700 EMU = 25400 --
        assert ln[0].get("w") == "25400"

    # -- 3: the value round-trips through save + reopen ---------------------

    def it_round_trips_marker_border_color_through_save_and_reopen(self, _restore_part_factory):
        """Save + reopen preserves the authored marker-border color.

        This is the load-bearing end-to-end check — the reporter needs the
        color to still be present when the file is opened by downstream
        consumers (PowerPoint, or python-pptx itself).
        """
        prs = Presentation()
        chart = _make_line_markers_chart(prs)
        series = chart.plots[0].series[0]
        series.marker.format.line.color.rgb = RGBColor(0xFF, 0x66, 0x00)
        series.marker.format.line.width = Pt(1.5)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        # -- slide_layouts[5] is "Title Only" — the chart is the first
        # -- non-placeholder shape, so pick the one carrying `has_chart`. --
        reloaded_gf = next(s for s in prs2.slides[0].shapes if s.has_chart)
        reloaded_series = reloaded_gf.chart.plots[0].series[0]
        assert reloaded_series.marker.format.line.color.rgb == RGBColor(0xFF, 0x66, 0x00)
        # -- Pt(1.5) -> 1.5 * 12700 EMU = 19050 --
        assert reloaded_series.marker.format.line.width == 19050

    # -- 4: the same recipe applies to XY (scatter) and Radar series -------

    def it_works_on_XY_scatter_series_too(self):
        """``XySeries`` also gets the ``_MarkerMixin.marker`` surface."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        data = XyChartData()
        series_data = data.add_series("Scatter1")
        series_data.add_data_point(1.0, 2.0)
        series_data.add_data_point(2.0, 4.0)
        series_data.add_data_point(3.0, 6.0)
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.XY_SCATTER,
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(3),
            data,
        ).chart
        series = chart.plots[0].series[0]

        series.marker.format.line.color.rgb = RGBColor(0x00, 0xCC, 0x00)

        srgb = series._element.xpath("c:marker/c:spPr/a:ln/a:solidFill/a:srgbClr")
        assert len(srgb) == 1
        assert srgb[0].get("val") == "00CC00"

    def it_works_on_radar_chart_series_too(self):
        """``RadarSeries`` inherits ``_MarkerMixin`` as well."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        data = CategoryChartData()
        data.categories = ["A", "B", "C"]
        data.add_series("RadarSer", (1.0, 2.0, 3.0))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.RADAR_MARKERS,
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(3),
            data,
        ).chart
        series = chart.plots[0].series[0]

        series.marker.format.line.color.rgb = RGBColor(0x11, 0x22, 0x33)

        srgb = series._element.xpath("c:marker/c:spPr/a:ln/a:solidFill/a:srgbClr")
        assert len(srgb) == 1
        assert srgb[0].get("val") == "112233"
