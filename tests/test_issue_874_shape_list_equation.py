# pyright: reportPrivateUsage=false

"""Regression test for issue #874 — shape.shapes skips equation-bearing shapes.

Issue #874 (https://github.com/scanny/python-pptx/issues/874) reported
that a shape whose text contained ``"500-7,000 m^3  per day."`` was
missing from ``slide.shapes`` — i.e. ``len(slide.shapes)`` returned one
less than the shape count the user saw in PowerPoint, and the shape
itself could not be discovered via iteration.

Root cause: when the user typed ``m^3``, PowerPoint auto-formatted the
superscript as an OMML equation (``m:oMath``) and wrapped the whole
enclosing ``p:sp`` in an ``mc:AlternateContent`` element — a
``mc:Choice`` subtree containing the equation-bearing shape, plus a
``mc:Fallback`` subtree with a plain-text rendering. Prior to Foundation
F3, ``CT_GroupShape.iter_shape_elms`` only looked at top-level
``p:sp`` / ``p:grpSp`` / ``p:graphicFrame`` / ``p:cxnSp`` / ``p:pic``
children of ``p:spTree`` and silently skipped ``mc:AlternateContent``
wrappers, which made the wrapped shape invisible to ``slide.shapes``.

Foundation F3 (``feat/foundation-f3-mc-alternate-content``, at
``b3a22344`` on master) taught ``iter_shape_elms`` to walk into the
first ``mc:Choice`` child transparently while leaving the
``mc:Fallback`` intact for round-trip fidelity. Issue #126
(``feat/issue-126-math-equation-read``, at ``b28845cd``) layered on
read-only access to the OMML subtree via ``BaseShape.has_math_equation``
and ``BaseShape.math_equation_xml``.

Together those two changes make #874's use case work end-to-end: the
equation-bearing shape now appears in ``slide.shapes`` iteration, and
callers can detect / introspect the embedded equation.

This test loads the exact ``Presentation2.pptx`` that the reporter
attached to the issue thread (the title text ``"500-7,000 m^3  per
day."`` is the very string that triggered the bug) and pins the flow so
future refactors of the shape-tree iterator or the ``mc:AlternateContent``
handler can't regress it.
"""

from __future__ import annotations

from os.path import abspath, dirname, join

from pptx import Presentation


def _fixture_path() -> str:
    return abspath(
        join(
            dirname(__file__),
            "..",
            "features",
            "steps",
            "test_files",
            "issue-874-equation-in-text.pptx",
        )
    )


class DescribeIssue874EquationInShapeText(object):
    """The #874 flow: open the reporter's real .pptx, iterate
    ``slide.shapes``, and confirm the equation-bearing shape is not
    dropped.
    """

    def it_surfaces_equation_wrapped_shapes_in_slide_shapes(self):
        """Core repro: the shape whose text contains ``m^3`` (which
        PowerPoint auto-formatted into an OMML equation and wrapped in
        ``mc:AlternateContent``) must appear in ``slide.shapes`` just
        like any other shape.

        Pre-F3, this assertion would have failed with
        ``len(shapes) == 5`` — the equation-bearing ``Rectangle 5`` was
        silently skipped because its ``p:sp`` was wrapped in
        ``mc:AlternateContent`` rather than sitting as a direct child
        of ``p:spTree``.
        """
        prs = Presentation(_fixture_path())
        slide = prs.slides[0]

        # -- the reporter's slide has six shapes total: five plain
        # -- ``p:sp`` direct children of ``p:spTree`` plus one that
        # -- lives inside a ``mc:AlternateContent`` wrapper --
        shapes = list(slide.shapes)
        assert len(shapes) == 6
        assert len(slide.shapes) == 6

        # -- and exactly one of those shapes carries the OMML equation
        # -- that triggered the original wrapping --
        equation_shapes = [s for s in shapes if s.has_math_equation]
        assert len(equation_shapes) == 1

        # -- the wrapped shape round-trips its identity (name from
        # -- ``p:cNvPr/@name``) so callers can target it by name just
        # -- like an unwrapped shape --
        eq_shape = equation_shapes[0]
        assert eq_shape.name == "Rectangle 5"

        # -- the reporter's text fragment is still reachable through
        # -- the normal ``.text_frame`` accessor on the wrapped shape;
        # -- the "500-7,000 " prefix that precedes the ``m^3`` OMML
        # -- run survives as plain text --
        assert eq_shape.has_text_frame is True
        assert "500-7,000" in eq_shape.text_frame.text

    def it_exposes_the_embedded_OMML_subtree_for_the_wrapped_shape(
        self
    ):
        """Follow-up from the issue thread: the reporter asked to be
        able to *extract data* from the equation-bearing shape rather
        than silently drop it. #126 shipped the read-only OMML API the
        caller needs — ``has_math_equation`` + ``math_equation_xml`` —
        which must work for the ``mc:AlternateContent``-wrapped shape
        the same way it does for an un-wrapped one.
        """
        prs = Presentation(_fixture_path())
        slide = prs.slides[0]

        eq_shapes = [s for s in slide.shapes if s.has_math_equation]
        assert len(eq_shapes) == 1
        eq_shape = eq_shapes[0]

        oMath_xml = eq_shape.math_equation_xml
        assert oMath_xml is not None
        # -- it's a well-formed OMML fragment the caller can hand to an
        # -- external converter (pandoc, omml.xsl, etc.) --
        assert oMath_xml.startswith("<m:oMath")
        assert (
            'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
            in oMath_xml
        )
        # -- the ``m^3`` superscript that triggered the wrapping is
        # -- preserved: OMML emits the base as ``<m:e>`` and the
        # -- exponent as ``<m:sup>`` inside ``<m:sSup>`` --
        assert "<m:sSup" in oMath_xml or "m:sSup" in oMath_xml

    def it_preserves_mc_Fallback_round_trip_for_wrapped_equation_shape(
        self
    ):
        """Iterating ``slide.shapes`` must not destructively strip the
        ``mc:Fallback`` subtree — the equation-bearing shape surfaces
        via ``mc:Choice`` but ``mc:Fallback`` stays intact for
        round-trip fidelity (PowerPoint falls back to it on viewers
        that don't support the required extension namespace).
        """
        prs = Presentation(_fixture_path())
        slide = prs.slides[0]

        spTree = slide.shapes._spTree
        # -- baseline: the fixture has exactly one ``mc:AlternateContent``
        # -- wrapper at the top level of ``p:spTree`` --
        mc_ns = "http://schemas.openxmlformats.org/markup-compatibility/2006"
        ac_nodes_before = spTree.findall("{%s}AlternateContent" % mc_ns)
        assert len(ac_nodes_before) == 1
        fallback_before = ac_nodes_before[0].find("{%s}Fallback" % mc_ns)
        assert fallback_before is not None

        # -- force full iteration; this is what used to skip the
        # -- wrapped shape pre-F3 --
        _ = list(slide.shapes)

        # -- iteration is non-destructive: the wrapper is still there
        # -- with its Fallback subtree intact --
        ac_nodes_after = spTree.findall("{%s}AlternateContent" % mc_ns)
        assert len(ac_nodes_after) == 1
        fallback_after = ac_nodes_after[0].find("{%s}Fallback" % mc_ns)
        assert fallback_after is not None

    def it_reports_the_reporter_title_text_on_the_wrapped_shape(self):
        """The issue's title pins the exact string the user typed:
        ``"500-7,000 m^3  per day."`` (note the double-space before
        ``per``). The ``m^3`` part became an OMML run so it doesn't
        round-trip through ``.text_frame.text``, but the surrounding
        literal text does — and the fact that we can read *any* text
        from the shape at all is the #874 fix (pre-F3 the shape was
        invisible so even ``shape.name`` was unreachable).
        """
        prs = Presentation(_fixture_path())
        slide = prs.slides[0]

        # -- find the equation-bearing shape --
        (eq_shape,) = [s for s in slide.shapes if s.has_math_equation]

        # -- the literal text runs around the equation survive and are
        # -- discoverable via ``.text_frame.text`` --
        text = eq_shape.text_frame.text
        assert "500-7,000" in text
        # -- and the OMML element is a child of the text frame's
        # -- paragraph, not a sibling of the shape --
        txBody = eq_shape.text_frame._txBody
        oMaths = txBody.findall(
            ".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"
        )
        assert len(oMaths) >= 1


