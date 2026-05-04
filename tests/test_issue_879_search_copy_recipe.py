# pyright: reportPrivateUsage=false

"""Regression tests for issue #879 — search + cross-presentation copy recipe.

Issue #879 (https://github.com/scanny/python-pptx/issues/879) asks how to
find a slide in one presentation (by title text, by stable id, by layout
name, or by any shape's text) and copy it into another presentation.
Both capabilities already ship:

  * A short Python helper iterating ``prs.slides`` and its shapes covers
    any practical search predicate (:attr:`.SlideShapes.title`,
    ``shape.text_frame.text``, ``slide.slide_layout.name``).
  * :meth:`.Slides.get_by_slide_id` (Wave 1/2) resolves a stashed
    presentation-stable id back to the |Slide|.
  * :meth:`.Slides.add_slide_from_external` (Wave 1 #1036) clones the
    located slide into the target deck at full fidelity, re-establishing
    image, chart, media, OLE-object, and hyperlink parts in the target
    package.

This suite pins the documented recipe under
``docs/user/slides.rst`` → "Searching and copying slides between
presentations" so the code paths it composes stay working:

  * Title-text search (``slide.shapes.title.text_frame.text``) across a
    source deck, skipping Blank-layout slides without a title.
  * Cross-presentation copy via
    :meth:`.Slides.add_slide_from_external`, bound to a target layout.
  * Save + reopen round-trip: the grafted slide's title text survives.
  * The ``find_slide_containing(prs, phrase)`` variation for any-shape
    text matching.
  * Stable-``slide_id`` lookup via :meth:`.Slides.get_by_slide_id` as an
    alternative keying.
  * Layout-name keying as another alternative predicate.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.package import PartFactory
from pptx.parts.chart import ChartPart
from pptx.parts.coreprops import CorePropertiesPart
from pptx.parts.embeddedpackage import EmbeddedXlsxPart
from pptx.parts.image import ImagePart
from pptx.parts.media import MediaPart
from pptx.parts.presentation import PresentationPart
from pptx.parts.slide import (
    NotesMasterPart,
    NotesSlidePart,
    SlideLayoutPart,
    SlideMasterPart,
    SlidePart,
)
from pptx.util import Inches, Pt


@pytest.fixture
def _restore_part_factory():
    """Guard against ``PartFactory.part_type_for`` pollution from other tests.

    Mirrors the identical guard used by ``tests/test_issue_175_*.py``,
    ``tests/test_issue_696_*.py``, and siblings. Earlier modules may
    overwrite slide-part registrations with mocks and not restore them,
    which breaks ``Presentation(stream)`` on reopen.
    """
    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
        CT.SML_SHEET: EmbeddedXlsxPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


# -- documented recipe helpers: mirror the snippets in docs/user/slides.rst --


def find_slide_by_title(prs, title):
    """Return the first slide in *prs* whose title text equals *title*."""
    for slide in prs.slides:
        title_shape = slide.shapes.title
        if title_shape is None:
            continue
        if not title_shape.has_text_frame:
            continue
        if title_shape.text_frame.text == title:
            return slide
    return None


def find_slide_containing(prs, phrase):
    """Return the first slide in *prs* containing *phrase* in any shape."""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and phrase in shape.text_frame.text:
                return slide
    return None


def find_slide_by_layout_name(prs, layout_name):
    """Return the first slide in *prs* whose layout.name matches."""
    for slide in prs.slides:
        if slide.slide_layout.name == layout_name:
            return slide
    return None


def _add_titled_slide(prs, title_text, layout_index=5):
    """Append a slide with a title placeholder carrying *title_text*."""
    slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
    title_shape = slide.shapes.title
    assert title_shape is not None, "layout %r has no title placeholder" % layout_index
    title_shape.text_frame.text = title_text
    return slide


def _add_blank_slide(prs):
    """Append a slide built from the Blank layout (index 6 — no title)."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def _slide_title(slide):
    """Return the slide's title text, or None when no title placeholder."""
    t = slide.shapes.title
    if t is None or not t.has_text_frame:
        return None
    return t.text_frame.text


class DescribeIssue879SearchCopyRecipe(object):
    """End-to-end coverage of the #879 search-and-copy recipe."""

    # -- (1) primary recipe: find slide by title, copy to target, round-trip --

    def it_finds_a_slide_by_title_in_the_source_deck(self):
        # -- Pins the helper in the docs: title-text lookup returns the
        # -- matching slide and None when no match exists. Non-title
        # -- (Blank-layout) slides must be skipped without raising. --
        source = Presentation()
        _add_titled_slide(source, "Intro")
        wanted = _add_titled_slide(source, "Q3 Highlights")
        _add_blank_slide(source)  # no title — must be skipped
        _add_titled_slide(source, "Outro")

        assert find_slide_by_title(source, "Q3 Highlights") is wanted
        assert find_slide_by_title(source, "Nope") is None

    def it_copies_the_found_slide_into_a_target_presentation(self, _restore_part_factory):
        # -- Headline recipe: locate by title, graft via
        # -- add_slide_from_external; title text rides along on the clone. --
        source = Presentation()
        _add_titled_slide(source, "Intro")
        _add_titled_slide(source, "Q3 Highlights")
        _add_titled_slide(source, "Outro")

        target = Presentation()
        target_pre_count = len(target.slides)

        wanted = find_slide_by_title(source, "Q3 Highlights")
        assert wanted is not None
        cloned = target.slides.add_slide_from_external(wanted, target.slide_layouts[5])

        assert len(target.slides) == target_pre_count + 1
        assert cloned in list(target.slides)
        assert _slide_title(cloned) == "Q3 Highlights"

    def it_round_trips_the_title_through_save_and_reopen(self, _restore_part_factory):
        # -- The grafted slide's title text must survive a save/reopen
        # -- cycle, otherwise the recipe is useless for real-world decks. --
        source = Presentation()
        _add_titled_slide(source, "Intro")
        _add_titled_slide(source, "Q3 Highlights")

        target = Presentation()
        wanted = find_slide_by_title(source, "Q3 Highlights")
        assert wanted is not None
        target.slides.add_slide_from_external(wanted, target.slide_layouts[5])

        buf = io.BytesIO()
        target.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        titles = [_slide_title(s) for s in reloaded.slides]
        assert "Q3 Highlights" in titles

    def it_does_not_mutate_the_source_presentation(self, _restore_part_factory):
        # -- The "library deck" idiom in the docs relies on the source
        # -- being a read-only source of slides. The copy must leave its
        # -- slide count, title text, and slide_ids alone. --
        source = Presentation()
        _add_titled_slide(source, "Intro")
        _add_titled_slide(source, "Q3 Highlights")
        _add_titled_slide(source, "Outro")
        pre_count = len(source.slides)
        pre_titles = [_slide_title(s) for s in source.slides]
        pre_ids = [s.slide_id for s in source.slides]

        target = Presentation()
        wanted = find_slide_by_title(source, "Q3 Highlights")
        assert wanted is not None
        target.slides.add_slide_from_external(wanted, target.slide_layouts[5])

        assert len(source.slides) == pre_count
        assert [_slide_title(s) for s in source.slides] == pre_titles
        assert [s.slide_id for s in source.slides] == pre_ids

    # -- (2) variations on the search predicate ------------------------------

    def it_supports_the_any_shape_text_search_variant(self):
        # -- The docs' "find_slide_containing" helper matches a phrase in
        # -- any text-bearing shape, not just the title. Useful when the
        # -- target phrase lives in a body placeholder or a textbox. --
        source = Presentation()
        _add_titled_slide(source, "Intro")
        carrier = source.slides.add_slide(source.slide_layouts[5])
        tb = carrier.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = "Project Bluebell milestones"
        run.font.size = Pt(18)
        _add_titled_slide(source, "Outro")

        found = find_slide_containing(source, "Bluebell")
        assert found is carrier
        assert find_slide_containing(source, "not-in-deck") is None

    def it_supports_stable_slide_id_lookup_after_reordering(self):
        # -- Once a slide_id is stashed, the slide remains reachable via
        # -- Slides.get_by_slide_id even if the deck is reordered. This
        # -- is the "durable pointer" variant of the docs recipe. --
        source = Presentation()
        _add_titled_slide(source, "Intro")
        wanted = _add_titled_slide(source, "Q3 Highlights")
        _add_titled_slide(source, "Outro")
        sid = wanted.slide_id

        # -- reorder the deck; the id-based lookup still finds the slide --
        source.slides.move_slide(wanted, 0)

        assert source.slides.get_by_slide_id(sid) is wanted
        assert source.slides.get_by_slide_id(9_999_999) is None
        assert source.slides.get_by_slide_id(9_999_999, default="fallback") == "fallback"

    def it_supports_layout_name_keying(self):
        # -- The "find by layout name" variant in the docs recipe --
        # -- surfaces slides that use a distinctive layout (for example,
        # -- the only slide built from "Title Only"). --
        source = Presentation()
        _add_titled_slide(source, "Intro", layout_index=1)  # Title and Content
        unique = _add_titled_slide(source, "Cover", layout_index=5)  # Title Only
        _add_titled_slide(source, "Outro", layout_index=1)

        assert find_slide_by_layout_name(source, "Title Only") is unique
        assert find_slide_by_layout_name(source, "Nonexistent Layout") is None

    # -- (3) composite: predicate-agnostic end-to-end copy -------------------

    def it_copies_a_slide_located_by_any_shape_text_and_round_trips(self, _restore_part_factory):
        # -- Ties the any-shape search to the cross-presentation copy:
        # -- find by phrase, clone to target, confirm the phrase is on
        # -- the cloned slide after a save/reopen round-trip. --
        source = Presentation()
        _add_titled_slide(source, "Intro")
        slide = source.slides.add_slide(source.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = "Nuestra experiencia en su industria"
        run.font.size = Pt(18)

        target = Presentation()
        found = find_slide_containing(source, "experiencia")
        assert found is slide
        target.slides.add_slide_from_external(found, target.slide_layouts[5])

        buf = io.BytesIO()
        target.save(buf)
        buf.seek(0)
        reloaded = Presentation(buf)

        assert find_slide_containing(reloaded, "experiencia") is not None
