# pyright: reportPrivateUsage=false

"""Regression test for issue #933 — ``Slide.follow_master_background`` settable.

Issue #933 (https://github.com/scanny/python-pptx/issues/933) asked for a
write-side hook paired with the read-only
:attr:`.Slide.follow_master_background` flag, so a caller who has painted
a custom background onto a slide can flip it back to master inheritance
without reaching into the underlying ``p:cSld/p:bg`` element.

The capability was delivered in the Wave 13 follow-up to issue #366,
which turned :attr:`.Slide.follow_master_background` into a dual
bool-like/callable proxy (:class:`pptx.slide._FollowMasterBackground`).
Reading the attribute still returns a bool-like object — ``True`` when
the slide has no ``p:bg`` override and ``False`` once an explicit
background is authored. *Calling* the returned object
(``slide.follow_master_background()``) drops any ``p:bg`` child on the
slide's ``p:cSld``, restoring master inheritance — PowerPoint's *Reset
Background* button equivalent. The call returns the owning slide so it
can be chained.

This suite pins the #933 reporter's perspective:

1. A fresh slide inherits master background — ``bool(slide.follow_master_background)``
   is ``True`` and no ``p:bg`` child exists.
2. After authoring a custom background,
   ``bool(slide.follow_master_background)`` is ``False`` and the
   ``p:bg`` child is present.
3. ``slide.follow_master_background()`` removes the ``p:bg`` child,
   flips ``bool(slide.follow_master_background)`` back to ``True``, and
   returns the slide.
4. The reverted state round-trips through ``Presentation.save`` + reopen.
5. Calling ``follow_master_background()`` on a slide that already
   inherits from the master is idempotent — still no ``p:bg``, still
   ``True``, still returns the slide.

Cross-reference:

* Wave 13 follow-up commit — ``feat: add Slide.follow_master_background()
  method to revert a custom slide background to master inheritance``
  (HISTORY.rst).
* Companion issue #366 — the authoring half of the contract — is pinned
  by ``tests/test_issue_366_slide_background_verify.py``.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.slide import Slide, _FollowMasterBackground


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of ``PartFactory.part_type_for``.

    Mirrors the guard in ``tests/test_issue_366_slide_background_verify.py``:
    ``Presentation(stream)`` reopen dispatches through
    ``PartFactory.part_type_for``, which earlier test modules overwrite
    with mocks and do not restore.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue933FollowMasterBackgroundVerify:
    """#933 follow_master_background is now settable via callable proxy.

    Pins the dual-form contract — ``bool(...)`` to read, ``(...)`` to
    reset — end-to-end, including save + reopen.
    """

    # -- 1. Default fresh slide: follows master ---------------------------

    def it_reads_true_on_a_fresh_slide_with_no_bg_override(self):
        """A freshly-added slide inherits its background.

        The ``follow_master_background`` descriptor returns a dual
        bool-like/callable proxy. On a fresh slide with no ``p:bg``
        child its bool value is ``True``.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        follows = slide.follow_master_background

        # -- proxy object, not a raw bool --
        assert isinstance(follows, _FollowMasterBackground)
        assert bool(follows) is True
        assert follows == True  # noqa: E712
        # -- no p:bg child under p:cSld --
        assert slide._element.cSld.bg is None

    # -- 2. After authoring: reads False ----------------------------------

    def it_reads_false_after_an_explicit_background_is_applied(self):
        """Authoring a solid background flips the read state.

        Painting ``slide.background.fill.solid()`` stamps a ``p:bg``
        child, so ``bool(slide.follow_master_background)`` is ``False``.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0x00, 0x00)

        assert bool(slide.follow_master_background) is False
        assert slide._element.cSld.bg is not None

    # -- 3. Calling the proxy resets to master inheritance ----------------

    def it_reverts_to_master_inheritance_when_called(self):
        """The #933 fix: calling the proxy drops ``p:bg``.

        After authoring an explicit background, calling
        ``slide.follow_master_background()`` removes the ``p:bg`` child
        from the slide's ``p:cSld`` element and returns the slide for
        chaining. ``bool(slide.follow_master_background)`` flips back
        to ``True``.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        # -- arrange: apply a custom background so p:bg exists --
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0x33, 0x66, 0x99)
        assert bool(slide.follow_master_background) is False
        assert slide._element.cSld.bg is not None

        # -- act: call the proxy to reset --
        result = slide.follow_master_background()

        # -- returns the slide for chaining --
        assert result is slide
        # -- p:bg is gone; master inheritance restored --
        assert slide._element.cSld.bg is None
        assert slide._element.cSld.find(qn("p:bg")) is None
        assert bool(slide.follow_master_background) is True

    # -- 4. Reverted state round-trips through save + reopen --------------

    def it_round_trips_the_reset_through_save_and_reload(self, _restore_part_factory):
        """The reverted state survives a save + reopen cycle.

        Authoring a background, resetting via the callable proxy, then
        saving and reopening must produce a deck whose slide has no
        ``p:bg`` override — regressions that cached the proxy's
        pre-reset state or skipped the element removal on write would
        be caught here.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0x00, 0x80, 0x00)
        # -- reset via the callable proxy before saving --
        slide.follow_master_background()

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        slide2 = prs2.slides[0]

        assert bool(slide2.follow_master_background) is True
        assert slide2._element.cSld.bg is None
        assert slide2._element.cSld.find(qn("p:bg")) is None

    # -- 5. Idempotent on a slide that already follows master -------------

    def it_is_idempotent_on_a_slide_that_already_follows_master(self):
        """Calling the proxy on an inheriting slide is a no-op.

        No ``p:bg`` child exists to drop, so the method must not raise,
        must leave the element tree unchanged, and must still return
        the slide.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        assert bool(slide.follow_master_background) is True
        assert slide._element.cSld.bg is None
        pre_xml = slide._element.xml

        result = slide.follow_master_background()

        assert result is slide
        assert bool(slide.follow_master_background) is True
        assert slide._element.cSld.bg is None
        # -- the ``p:sld`` XML is byte-for-byte unchanged --
        assert slide._element.xml == pre_xml

    # -- 6. The proxy is a fresh object on each access --------------------

    def it_returns_a_fresh_proxy_on_each_read(self):
        """Each access returns a distinct :class:`_FollowMasterBackground`.

        The proxy is documented as "produced by
        :attr:`Slide.follow_master_background` on each access" — two
        reads yield two independent instances, both agreeing with the
        current slide state.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        follows_a = slide.follow_master_background
        follows_b = slide.follow_master_background

        assert isinstance(follows_a, _FollowMasterBackground)
        assert isinstance(follows_b, _FollowMasterBackground)
        # -- separate instances --
        assert follows_a is not follows_b
        # -- but both agree on the slide state --
        assert bool(follows_a) is bool(follows_b) is True
        assert isinstance(slide, Slide)
