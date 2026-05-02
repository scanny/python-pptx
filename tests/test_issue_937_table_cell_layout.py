"""Regression test for issue #937 — arrange tickers side-by-side in table cells.

Issue #937 (https://github.com/scanny/python-pptx/issues/937) asks how to
arrange multiple short strings (the reporter used stock tickers) *side by
side* within a table cell, so that the layout reads as a row of tickers
rather than a column. PowerPoint has no native "multi-column text frame"
inside a single ``a:tc``; the two ways PowerPoint itself produces that
visual layout are:

  * **Multiple adjacent cells with the internal borders hidden.** This is
    the approach Martin Packer recommended in the issue thread (from his
    external ``md2pptx`` project). PowerPoint draws one cell per ticker
    but the hidden internal borders make the group read as one visual
    "super-cell". Merging a header cell across the group keeps the row
    above intact.

  * **Multiple runs (or runs separated by whitespace/tabs) in a single
    cell paragraph.** Each ticker is an ``a:r`` run inside a single
    ``a:p``; the runs lay out horizontally on one line because a
    paragraph always flows left-to-right. This is the approach suggested
    by scanny on the companion Stack Overflow question (fixed-pitch font
    + spaces).

Every API needed to realise *either* approach is already shipped in
python-pptx as of Wave 3:

  * ``_Cell.margin_left`` / ``.margin_right`` / ``.margin_top`` /
    ``.margin_bottom`` — tighten the cell padding so tickers fit.
  * ``_Cell.text_frame`` + ``paragraphs[0].alignment`` — centre or
    justify the ticker text inside its cell.
  * ``_Cell.merge`` — span a header cell ("Bottom Fishing") across the
    group of per-ticker sub-cells.
  * ``_Cell.border_left`` / ``.border_right`` / ``.border_top`` /
    ``.border_bottom`` — shipped by ``feat/issue-71-cell-borders``
    (Wave 3). Calling ``cell.border_<side>.fill.background()`` hides
    that edge so the group of sub-cells reads as a single visual cell.
  * ``_Paragraph.add_run`` — pre-dates the issue; each ticker becomes
    its own ``a:r`` with independent font / bold / colour, and the runs
    flow side-by-side on one line.

This regression test pins *both* recipes through
``Presentation.save`` + reopen so a future refactor that drops any of
the supporting APIs (cell borders, merge, margins, per-run font, or
paragraph alignment) fails loudly here rather than at user-report
time.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


@pytest.fixture
def _restore_part_factory():
    """Guard against pollution from tests that mutate `PartFactory.part_type_for`.

    See the identical fixture in ``tests/test_issue_705_shape_shadows.py``,
    ``tests/test_issue_1033_line_arrow.py``, and
    ``tests/test_issue_560_data_label_colors.py`` for the rationale — other
    regression suites mock out the part registry and, if their teardown
    leaks, this suite's ``Presentation.save`` + reopen round-trips
    encounter a ``Mock`` in place of ``SlidePart`` and explode with
    ``AttributeError: Mock object has no attribute 'slide'``.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.comments import CommentAuthorsPart, CommentsPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_COMMENTS: CommentsPart,
        CT.PML_COMMENT_AUTHORS: CommentAuthorsPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue937TableCellLayout(object):
    """End-to-end regression coverage for issue #937."""

    def it_arranges_tickers_side_by_side_via_adjacent_cells_with_hidden_borders(
        self, _restore_part_factory
    ):
        # -- Recipe A: one sub-cell per ticker, merged header row above,
        # -- internal vertical borders hidden so the group reads as a
        # -- single visual cell. This is the approach the reporter's
        # -- screenshot ("Basic Materials / Bottom Fishing") actually
        # -- shows in PowerPoint.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        tickers = ["CGAU", "NEM", "GOLD", "VALE"]
        shape = slide.shapes.add_table(
            rows=2,
            cols=len(tickers),
            left=Inches(1),
            top=Inches(1),
            width=Inches(6),
            height=Inches(1),
        )
        table = shape.table

        # -- header row: merge all columns into one "Bottom Fishing" cell
        header_origin = table.cell(0, 0)
        header_origin.merge(table.cell(0, len(tickers) - 1))
        header_origin.text = "Bottom Fishing"
        header_origin.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        # -- ticker row: one cell per ticker, centred, with internal
        # -- vertical borders hidden via `fill.background()`
        for j, ticker in enumerate(tickers):
            cell = table.cell(1, j)
            cell.text = ticker
            cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            # -- tighten padding so short tickers pack tight
            cell.margin_left = Inches(0.02)
            cell.margin_right = Inches(0.02)
            # -- hide the interior vertical borders between adjacent
            # -- ticker cells so the row reads as one visual "super-cell"
            if j > 0:
                cell.border_left.fill.background()
            if j < len(tickers) - 1:
                cell.border_right.fill.background()

        # -- round-trip through save + reopen
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reopened = Presentation(buf)

        rt_table = reopened.slides[0].shapes[0].table

        # -- header cell survives the merge + keeps its centred alignment
        rt_header = rt_table.cell(0, 0)
        assert rt_header.is_merge_origin is True
        assert rt_header.span_width == len(tickers)
        assert rt_header.text == "Bottom Fishing"
        assert rt_header.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER

        # -- each ticker lands in its own cell, in order, with the
        # -- expected alignment + margin
        for j, ticker in enumerate(tickers):
            rt_cell = rt_table.cell(1, j)
            assert rt_cell.text == ticker
            assert rt_cell.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
            assert rt_cell.margin_left == Inches(0.02)
            assert rt_cell.margin_right == Inches(0.02)

        # -- interior vertical borders are "no-fill" (hidden); the
        # -- outermost ticker cells keep their inherited borders on the
        # -- outer edges.
        for j in range(1, len(tickers)):
            assert (
                rt_table.cell(1, j).border_left.fill.type
                == MSO_FILL_TYPE.BACKGROUND
            )
        for j in range(len(tickers) - 1):
            assert (
                rt_table.cell(1, j).border_right.fill.type
                == MSO_FILL_TYPE.BACKGROUND
            )

    def it_arranges_tickers_side_by_side_via_multiple_runs_in_one_cell(
        self, _restore_part_factory
    ):
        # -- Recipe B: everything lives in a single cell; each ticker is
        # -- its own `a:r` run inside the cell's one `a:p`, so the runs
        # -- flow horizontally on one line. Per-run `.font` lets the
        # -- caller colour / weight each ticker independently — which
        # -- plain `cell.text = "\n".join(...)` can never do since that
        # -- writes one paragraph *per* ticker (the stacked-column
        # -- output the reporter wanted to avoid).
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        tickers = ["CGAU", "NEM", "GOLD", "VALE"]
        shape = slide.shapes.add_table(
            rows=1, cols=1, left=Inches(1), top=Inches(1), width=Inches(4), height=Inches(1)
        )
        cell = shape.table.cell(0, 0)

        paragraph = cell.text_frame.paragraphs[0]
        for i, ticker in enumerate(tickers):
            if i > 0:
                # -- whitespace-only run so PowerPoint wraps tickers on
                # -- natural word boundaries when the cell is narrow
                spacer = paragraph.add_run()
                spacer.text = "  "
            run = paragraph.add_run()
            run.text = ticker
            run.font.size = Pt(9)
            run.font.bold = i % 2 == 0
            # -- demonstrate per-run colour (the whole point of runs)
            run.font.color.rgb = RGBColor(0x33, 0x66, 0x99)

        # -- round-trip through save + reopen
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reopened = Presentation(buf)

        rt_paragraph = (
            reopened.slides[0].shapes[0].table.cell(0, 0).text_frame.paragraphs[0]
        )

        # -- the cell content is a single paragraph (the tickers sit
        # -- side-by-side on one line, not stacked)
        assert (
            len(reopened.slides[0].shapes[0].table.cell(0, 0).text_frame.paragraphs)
            == 1
        )

        # -- each ticker + spacer is its own run, so per-run formatting
        # -- survives and tickers stay independently styleable
        rt_tickers = [r.text for r in rt_paragraph.runs if r.text.strip()]
        assert rt_tickers == tickers

        # -- per-run font settings survived the round-trip
        ticker_runs = [r for r in rt_paragraph.runs if r.text.strip()]
        assert all(r.font.size == Pt(9) for r in ticker_runs)
        assert [r.font.bold for r in ticker_runs] == [True, False, True, False]
        assert all(
            r.font.color.rgb == RGBColor(0x33, 0x66, 0x99) for r in ticker_runs
        )

    def it_shrinks_cell_padding_so_small_text_fits_in_narrow_cells(
        self, _restore_part_factory
    ):
        # -- the reporter's secondary complaint was "the output is
        # -- currently overflowing the PowerPoint slide". The fix is
        # -- not a new API — it is the existing `_Cell.margin_*`
        # -- setters, which accept `Length` values (inc. zero) and
        # -- round-trip through `a:tcPr/@mar[LRTB]`. This test pins
        # -- that contract so the knob the reporter actually needs
        # -- keeps working.
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_table(
            rows=1, cols=1, left=Inches(1), top=Inches(1), width=Inches(1), height=Inches(0.3)
        )
        cell = shape.table.cell(0, 0)

        cell.margin_left = Inches(0.02)
        cell.margin_right = Inches(0.02)
        cell.margin_top = Inches(0.01)
        cell.margin_bottom = Inches(0.01)
        cell.text = "CGAU"

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        reopened = Presentation(buf)

        rt_cell = reopened.slides[0].shapes[0].table.cell(0, 0)
        assert rt_cell.margin_left == Inches(0.02)
        assert rt_cell.margin_right == Inches(0.02)
        assert rt_cell.margin_top == Inches(0.01)
        assert rt_cell.margin_bottom == Inches(0.01)
        assert rt_cell.text == "CGAU"
