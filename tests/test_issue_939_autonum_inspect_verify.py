# pyright: reportPrivateUsage=false

"""Regression test for issue #939 — read-access to paragraph bullet settings.

Issue #939 (https://github.com/scanny/python-pptx/issues/939) asked for a
way to *inspect* whether a text-frame paragraph has auto-numbered bullets
enabled. The companion authoring surface was added by #100 as
``_BulletFormat.auto_number(scheme, start_at=None)``; the reporter wanted
the symmetric reader.

This fork already ships the full read surface on ``_BulletFormat``:

* ``.type``  -> ``"char"`` | ``"autonum"`` | ``"none"`` | ``None``
* ``.char``  -> ``str | None`` (the literal bullet character)
* ``.number_scheme``  -> ``PP_AUTO_NUMBER_SCHEME | None``
* ``.start_at``  -> ``int | None`` (ECMA-376 default of 1 when ``@startAt``
  is absent on an ``a:buAutoNum`` child)

plus the inherit-semantics contract that a paragraph with no explicit
bullet returns ``None`` from every reader (the effective bullet is
inherited from the style hierarchy).

This suite pins those contracts end-to-end against the real public API
(no mocks) so #939 can be closed. Each assertion is phrased from the
reporter's perspective: open a ``.pptx``, walk its paragraphs, and be
able to tell whether each one has an auto-numbered bullet, a character
bullet, an explicit "no bullet", or inherits.
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.enum.text import PP_AUTO_NUMBER, PP_AUTO_NUMBER_SCHEME
from pptx.util import Inches


@pytest.fixture
def _restore_part_factory():
    """Guard against test-module pollution of ``PartFactory.part_type_for``.

    Mirrors the guard used by other round-trip regression tests in
    this directory; required because some unit-test modules overwrite
    the real slide-part registrations with mocks and don't restore
    them.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


class DescribeIssue939AutonumInspectVerify(object):
    """#939 verify-and-close: bullet-type / scheme / start_at / char readers."""

    # -- type reader --------------------------------------------------

    def it_reads_autonum_as_the_bullet_type(self, _restore_part_factory):
        """``_BulletFormat.type`` returns ``"autonum"`` after ``auto_number``."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        p = tb.text_frame.paragraphs[0]

        p.bullet.auto_number(PP_AUTO_NUMBER_SCHEME.ARABIC_PERIOD)

        assert p.bullet.type == "autonum"

    def it_reads_char_as_the_bullet_type_after_character(self, _restore_part_factory):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        p = tb.text_frame.paragraphs[0]

        p.bullet.character("*")

        assert p.bullet.type == "char"

    def it_reads_none_as_the_bullet_type_after_none(self, _restore_part_factory):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        p = tb.text_frame.paragraphs[0]

        p.bullet.none()

        assert p.bullet.type == "none"

    def it_reads_None_bullet_type_when_no_explicit_bullet_set(self, _restore_part_factory):
        """A paragraph that inherits its bullet returns |None| from every reader."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        p = tb.text_frame.paragraphs[0]

        assert p.bullet.type is None
        assert p.bullet.char is None
        assert p.bullet.number_scheme is None
        assert p.bullet.start_at is None

    # -- number_scheme / start_at readers -----------------------------

    def it_reads_the_auto_number_scheme_written_by_setter(self, _restore_part_factory):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        p = tb.text_frame.paragraphs[0]

        p.bullet.auto_number(PP_AUTO_NUMBER_SCHEME.ROMAN_UC_PERIOD, start_at=5)

        assert p.bullet.number_scheme is PP_AUTO_NUMBER_SCHEME.ROMAN_UC_PERIOD
        assert p.bullet.start_at == 5

    def it_defaults_start_at_to_1_when_startAt_attribute_is_absent(self, _restore_part_factory):
        """ECMA-376 default: an ``a:buAutoNum`` with no ``@startAt`` begins at 1."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        p = tb.text_frame.paragraphs[0]

        # -- call without start_at so no @startAt attribute is written --
        p.bullet.auto_number(PP_AUTO_NUMBER_SCHEME.ARABIC_PERIOD)

        assert p.bullet.number_scheme is PP_AUTO_NUMBER_SCHEME.ARABIC_PERIOD
        assert p.bullet.start_at == 1

    def it_returns_None_for_scheme_and_start_at_on_a_char_bullet(self, _restore_part_factory):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        p = tb.text_frame.paragraphs[0]

        p.bullet.character("•")

        assert p.bullet.number_scheme is None
        assert p.bullet.start_at is None

    # -- char reader --------------------------------------------------

    def it_reads_the_character_written_by_setter(self, _restore_part_factory):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        p = tb.text_frame.paragraphs[0]

        p.bullet.character("§")

        assert p.bullet.char == "§"
        assert p.bullet.type == "char"

    def it_returns_None_for_char_on_an_autonum_bullet(self, _restore_part_factory):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        p = tb.text_frame.paragraphs[0]

        p.bullet.auto_number(PP_AUTO_NUMBER_SCHEME.ARABIC_PERIOD)

        assert p.bullet.char is None

    # -- enum alias -------------------------------------------------

    def it_accepts_the_PP_AUTO_NUMBER_alias(self, _restore_part_factory):
        """The docs use ``PP_AUTO_NUMBER``; assert it's the same class."""
        assert PP_AUTO_NUMBER is PP_AUTO_NUMBER_SCHEME

    # -- round-trip through save + reload -----------------------------

    def it_round_trips_the_full_autonum_read_surface_through_reload(self, _restore_part_factory):
        """The reporter's scenario: open a ``.pptx``, inspect paragraph bullets.

        This test authors three paragraphs (autonum / char / inherit),
        saves to bytes, reopens into a fresh ``Presentation``, and
        asserts every reader returns the expected value.
        """
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(3))
        tf = tb.text_frame

        p0 = tf.paragraphs[0]
        p0.text = "numbered"
        p0.bullet.auto_number(PP_AUTO_NUMBER_SCHEME.ALPHA_LC_PAREN_R, start_at=4)

        p1 = tf.add_paragraph()
        p1.text = "starred"
        p1.bullet.character("*")

        p2 = tf.add_paragraph()
        p2.text = "inherited"  # -- no explicit bullet --

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)
        tf2 = prs2.slides[0].shapes[-1].text_frame

        q0, q1, q2 = tf2.paragraphs[0], tf2.paragraphs[1], tf2.paragraphs[2]

        # -- autonum paragraph: scheme + start_at preserved --
        assert q0.bullet.type == "autonum"
        assert q0.bullet.number_scheme is PP_AUTO_NUMBER_SCHEME.ALPHA_LC_PAREN_R
        assert q0.bullet.start_at == 4
        assert q0.bullet.char is None

        # -- character paragraph --
        assert q1.bullet.type == "char"
        assert q1.bullet.char == "*"
        assert q1.bullet.number_scheme is None
        assert q1.bullet.start_at is None

        # -- inherited paragraph: every reader returns None --
        assert q2.bullet.type is None
        assert q2.bullet.char is None
        assert q2.bullet.number_scheme is None
        assert q2.bullet.start_at is None
