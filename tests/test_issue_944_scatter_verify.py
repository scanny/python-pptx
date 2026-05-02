# pyright: reportPrivateUsage=false

"""Regression test for issue #944 — TreeMap and ScatterPlot chart support.

Issue #944 (https://github.com/scanny/python-pptx/issues/944) asked for
support for two chart kinds: *TreeMap* and *ScatterPlot*.

The **scatter** half is already fully supported: ``Shapes.add_chart``
accepts all five ``XL_CHART_TYPE.XY_SCATTER*`` members and dispatches
to :class:`_XyChartXmlWriter` (see ``src/pptx/chart/xmlwriter.py``),
emitting a schema-valid ``c:scatterChart`` whose ``c:scatterStyle/@val``
("lineMarker" / "smoothMarker") and optional ``c:marker/c:symbol="none"``
select the visual variant. The read path is covered by :class:`XyPlot`
and :class:`XySeries` (``src/pptx/chart/plot.py``,
``src/pptx/chart/series.py``). This file is the #944-specific
verify-and-close pin for the scatter half: it exercises all five enum
members through the real :meth:`Shapes.add_chart` API against a fresh
``Presentation()``, round-trips each through ``Presentation.save`` +
reopen, and asserts the ``scatter_style`` distinctions that Excel
itself uses to tell the variants apart at read time.

The **treemap** half of #944 is a duplicate of `#371`_ ("Treemap
Charts") and is blocked on the F4 chartex-foundation work described
in ``docs/dev/analysis/chartex-treemap.rst`` (Wave 8). Treemap is an
Office 2016+ "extended" chart kind in the ``cx:`` namespace, not the
legacy ``c:`` namespace scatter lives in, so it requires a wholly new
part class and oxml tree. Nothing in #944 changes the treemap picture;
see the design note for specifics.

.. _#371: https://github.com/scanny/python-pptx/issues/371
"""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

# -- fixtures -----------------------------------------------------------


@pytest.fixture
def _restore_part_factory():
    """Guard against pollution from tests that mutate `PartFactory.part_type_for`.

    ``tests/opc/test_package.py::DescribePartFactory`` overwrites
    ``PartFactory.part_type_for[CT.PML_SLIDE]`` with a Mock and does
    not restore it. Round-trip tests that call ``Presentation.save``
    + ``Presentation()`` reopen depend on the real registrations being
    in place. Mirrors the identical guard elsewhere in the
    regression-test suite.
    """
    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.package import PartFactory
    from pptx.parts.chart import ChartPart
    from pptx.parts.coreprops import CorePropertiesPart
    from pptx.parts.image import ImagePart
    from pptx.parts.media import MediaPart
    from pptx.parts.presentation import PresentationPart
    from pptx.parts.slide import (
        NotesMasterPart,
        NotesSlidePart,
        SlideLayoutPart,
        SlideMasterPart,
        SlidePart,
    )

    saved = dict(PartFactory.part_type_for)
    expected = {
        CT.PML_PRESENTATION_MAIN: PresentationPart,
        CT.PML_PRES_MACRO_MAIN: PresentationPart,
        CT.PML_TEMPLATE_MAIN: PresentationPart,
        CT.PML_SLIDESHOW_MAIN: PresentationPart,
        CT.OPC_CORE_PROPERTIES: CorePropertiesPart,
        CT.PML_NOTES_MASTER: NotesMasterPart,
        CT.PML_NOTES_SLIDE: NotesSlidePart,
        CT.PML_SLIDE: SlidePart,
        CT.PML_SLIDE_LAYOUT: SlideLayoutPart,
        CT.PML_SLIDE_MASTER: SlideMasterPart,
        CT.DML_CHART: ChartPart,
        CT.JPEG: ImagePart,
        CT.PNG: ImagePart,
        CT.MP4: MediaPart,
    }
    PartFactory.part_type_for.update(expected)
    try:
        yield
    finally:
        PartFactory.part_type_for.clear()
        PartFactory.part_type_for.update(saved)


def _xy_chart_data():
    """Return a two-series ``XyChartData`` fixture."""
    data = XyChartData()
    s1 = data.add_series("Alpha")
    s1.add_data_point(1.0, 2.0)
    s1.add_data_point(2.0, 3.0)
    s1.add_data_point(3.0, 1.5)
    s2 = data.add_series("Beta")
    s2.add_data_point(0.5, 1.2)
    s2.add_data_point(1.5, 2.7)
    s2.add_data_point(2.5, 3.1)
    return data


def _add_scatter_chart(chart_type):
    """Author a slide with a scatter chart of ``chart_type`` and return the chart."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    graphic_frame = slide.shapes.add_chart(
        chart_type,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        _xy_chart_data(),
    )
    return prs, graphic_frame.chart


# -- the verification suite ---------------------------------------------


class DescribeIssue944ScatterVerify(object):
    """#944 (scatter half) verify-and-close.

    Pins that ``Shapes.add_chart`` accepts every
    ``XL_CHART_TYPE.XY_SCATTER*`` member, emits a ``c:scatterChart``
    with the expected ``scatterStyle``/``marker`` shape, and round-trips
    through ``Presentation.save`` + reopen.
    """

    # -- authoring: every XY_SCATTER* variant is accepted -----------------

    @pytest.mark.parametrize(
        "chart_type",
        [
            XL_CHART_TYPE.XY_SCATTER,
            XL_CHART_TYPE.XY_SCATTER_LINES,
            XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
            XL_CHART_TYPE.XY_SCATTER_SMOOTH,
            XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS,
        ],
    )
    def it_accepts_every_XY_SCATTER_variant_in_add_chart(self, chart_type):
        """Every five-member ``XY_SCATTER*`` enum is usable with ``add_chart``."""
        _prs, chart = _add_scatter_chart(chart_type)

        # -- assert: a c:scatterChart wrapper is present --
        assert chart._chartSpace.find(f".//{qn('c:scatterChart')}") is not None
        # -- and chart_type is reported back accurately --
        assert chart.chart_type == chart_type

    # -- scatter_style distinguishes the five variants --------------------

    def it_emits_lineMarker_style_for_XY_SCATTER(self):
        """Plain ``XY_SCATTER`` emits ``c:scatterStyle val="lineMarker"``."""
        _prs, chart = _add_scatter_chart(XL_CHART_TYPE.XY_SCATTER)
        style = chart._chartSpace.xpath(".//c:scatterChart/c:scatterStyle")[0]
        assert style.get("val") == "lineMarker"

    def it_emits_lineMarker_style_for_XY_SCATTER_LINES(self):
        """``XY_SCATTER_LINES`` also uses ``lineMarker`` (lines + markers)."""
        _prs, chart = _add_scatter_chart(XL_CHART_TYPE.XY_SCATTER_LINES)
        style = chart._chartSpace.xpath(".//c:scatterChart/c:scatterStyle")[0]
        assert style.get("val") == "lineMarker"

    def it_emits_smoothMarker_style_for_XY_SCATTER_SMOOTH(self):
        """``XY_SCATTER_SMOOTH`` uses the ``smoothMarker`` spline variant."""
        _prs, chart = _add_scatter_chart(XL_CHART_TYPE.XY_SCATTER_SMOOTH)
        style = chart._chartSpace.xpath(".//c:scatterChart/c:scatterStyle")[0]
        assert style.get("val") == "smoothMarker"

    def it_suppresses_markers_for_the_NO_MARKERS_variants(self):
        """The *_NO_MARKERS* variants inject ``<c:marker><c:symbol val="none"/></c:marker>``."""
        for chart_type in (
            XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
            XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS,
        ):
            _prs, chart = _add_scatter_chart(chart_type)
            symbols = chart._chartSpace.xpath(".//c:scatterChart/c:ser/c:marker/c:symbol")
            assert len(symbols) >= 1
            assert all(s.get("val") == "none" for s in symbols)

    def it_omits_none_marker_for_the_marker_variants(self):
        """The marker-bearing variants do not emit ``c:symbol val="none"``."""
        for chart_type in (
            XL_CHART_TYPE.XY_SCATTER,
            XL_CHART_TYPE.XY_SCATTER_LINES,
            XL_CHART_TYPE.XY_SCATTER_SMOOTH,
        ):
            _prs, chart = _add_scatter_chart(chart_type)
            none_symbols = chart._chartSpace.xpath(
                ".//c:scatterChart/c:ser/c:marker/c:symbol[@val='none']"
            )
            assert none_symbols == []

    # -- data carriage ----------------------------------------------------

    def it_carries_xVal_and_yVal_pairs_into_each_series(self):
        """Every ``c:ser`` carries ``c:xVal`` + ``c:yVal`` with the authored point count."""
        _prs, chart = _add_scatter_chart(XL_CHART_TYPE.XY_SCATTER)
        sers = chart._chartSpace.xpath(".//c:scatterChart/c:ser")
        assert len(sers) == 2
        for ser in sers:
            xvals = ser.xpath("./c:xVal//c:pt")
            yvals = ser.xpath("./c:yVal//c:pt")
            assert len(xvals) == 3
            assert len(yvals) == 3

    # -- round-trip through save + reopen ---------------------------------

    @pytest.mark.parametrize(
        "chart_type",
        [
            XL_CHART_TYPE.XY_SCATTER,
            XL_CHART_TYPE.XY_SCATTER_LINES,
            XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
            XL_CHART_TYPE.XY_SCATTER_SMOOTH,
            XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS,
        ],
    )
    def it_round_trips_each_variant_through_save_reopen(self, chart_type, _restore_part_factory):
        """Each ``XY_SCATTER*`` survives save + reopen with the same ``chart_type``."""
        prs, _chart = _add_scatter_chart(chart_type)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        graphic_frame = next(s for s in prs2.slides[0].shapes if s.has_chart)
        reloaded = graphic_frame.chart
        # -- c:scatterChart still present --
        assert reloaded._chartSpace.find(f".//{qn('c:scatterChart')}") is not None
        # -- the read path reports the same chart_type back --
        assert reloaded.chart_type == chart_type

    # -- regression: the reporter's exact scenario ------------------------

    def it_reproduces_the_944_reporter_scenario(self, _restore_part_factory):
        """The reporter's "add a scatter chart to a slide" workflow completes.

        #944 asked for "TreeMap and ScatterPlot" support. The scatter
        path is exercised end-to-end here: create a fresh presentation,
        add a scatter chart with real XY data, save, reopen, and verify
        the chart reads back as a scatter chart with the same series
        count and data-point counts.
        """
        prs, chart = _add_scatter_chart(XL_CHART_TYPE.XY_SCATTER_LINES)
        assert chart.chart_type == XL_CHART_TYPE.XY_SCATTER_LINES

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        graphic_frame = next(s for s in prs2.slides[0].shapes if s.has_chart)
        chart2 = graphic_frame.chart
        assert chart2.chart_type == XL_CHART_TYPE.XY_SCATTER_LINES
        # -- two series, three points each, survived the round-trip --
        sers = chart2._chartSpace.xpath(".//c:scatterChart/c:ser")
        assert len(sers) == 2
        for ser in sers:
            assert len(ser.xpath("./c:xVal//c:pt")) == 3
            assert len(ser.xpath("./c:yVal//c:pt")) == 3
