# pyright: reportPrivateUsage=false

"""Regression test for issue #956 — "Deleting slides leaves unused parts behind".

Issue #956 (https://github.com/scanny/python-pptx/issues/956) asked whether
:meth:`Slides.delete` actually garbage-collects the parts that only the
deleted slide referenced. The reporter worried that images, charts, and
other descendant parts might linger inside the ``.pptx`` zip after a slide
was removed — making a heavily-edited deck grow forever.

The answer is that the package-save path already implements reachability-
based part garbage collection. :meth:`.OpcPackage.save` passes the result
of :meth:`.OpcPackage.iter_parts` to :class:`.PackageWriter`; that iterator
walks the rels graph depth-first from the package root, so any part no
longer reachable from the presentation part (the slide and its
uniquely-referenced descendants — notes slide, uniquely-referenced image,
chart, embedded xlsx, media, tags parts) is silently dropped on save.

:meth:`Slides.delete` correctly makes the slide unreachable:

  * removes the ``p:sldId`` entry from ``p:sldIdLst`` (so nothing in
    ``presentation.xml`` references the slide any more);
  * then calls :meth:`.XmlPart.drop_rel` on the presentation-part, which
    removes the presentation -> slide relationship because its reference
    count in ``presentation.xml`` is now zero.

#956 is therefore a verify-close: the scenarios below exercise the
reporter's exact concern — delete a slide, save, reopen, and confirm the
orphaned slide / image / chart / embedded xlsx / notes parts are gone from
the zip — plus the complementary invariant that *shared* parts (an image
referenced by another surviving slide) are preserved.
"""

from __future__ import annotations

import io
import zipfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.parts.chart import ChartPart
from pptx.parts.embeddedpackage import EmbeddedXlsxPart
from pptx.parts.image import ImagePart
from pptx.parts.slide import NotesSlidePart, SlidePart
from pptx.util import Inches


def _image_shas(prs):
    """Return the set of image-part SHA1s reachable from ``prs``."""
    return {p.sha1 for p in prs.part.package.iter_parts() if isinstance(p, ImagePart)}


def _parts_of_type(prs, cls):
    """Return list of parts of ``cls`` reachable from ``prs`` (in iter order)."""
    return [p for p in prs.part.package.iter_parts() if isinstance(p, cls)]


def _zip_entries(buf):
    """Return the list of zip-member names in the saved package ``buf``."""
    with zipfile.ZipFile(buf) as zf:
        return zf.namelist()


class DescribeIssue956DeleteSlideCleanup:
    """#956 slide-delete part-cleanup verify-and-close."""

    # -- image-part cleanup ---------------------------------------------

    def it_drops_the_deleted_slides_unique_image_from_the_package(self):
        # -- The reporter's headline concern: an image only referenced by
        # -- the deleted slide must be gone from the saved zip. We sample
        # -- the image by SHA1 so any accidental retention by partname
        # -- renumbering is caught too.
        prs = Presentation()
        layout = prs.slide_layouts[6]
        shas = []
        for img in ("python-powered.png", "monty-truth.png", "python-icon.jpeg"):
            slide = prs.slides.add_slide(layout)
            pic = slide.shapes.add_picture("tests/test_files/%s" % img, Inches(1), Inches(1))
            shas.append(pic.image.sha1)
        deleted_sha, kept_sha_a, kept_sha_b = shas

        prs.slides.delete(prs.slides[0])

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        present = _image_shas(prs2)
        assert deleted_sha not in present
        assert kept_sha_a in present
        assert kept_sha_b in present

    def it_preserves_an_image_shared_with_another_surviving_slide(self):
        # -- Pin the inverse invariant: a part that *another* slide still
        # -- points at (via a shared relationship — Slides.duplicate
        # -- produces exactly this shape) must *not* be dropped when one
        # -- of its referencing slides is deleted.
        prs = Presentation()
        layout = prs.slide_layouts[6]
        s_source = prs.slides.add_slide(layout)
        s_source.shapes.add_picture("tests/test_files/python-powered.png", Inches(1), Inches(1))
        s_other = prs.slides.add_slide(layout)
        s_other.shapes.add_picture("tests/test_files/monty-truth.png", Inches(1), Inches(1))
        # -- duplicate the source slide — the duplicate shares the source's
        # -- ImagePart by package reference. --
        prs.slides.duplicate(s_source)
        shared_sha = prs.slides[0].shapes[0].image.sha1

        prs.slides.delete(prs.slides[0])

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        # -- the shared image survives because the duplicate slide still
        # -- holds a rel to it --
        assert shared_sha in _image_shas(prs2)

    # -- slide and notes-slide cleanup ----------------------------------

    def it_drops_the_deleted_slide_part_itself(self):
        prs = Presentation()
        layout = prs.slide_layouts[6]
        prs.slides.add_slide(layout)
        prs.slides.add_slide(layout)
        prs.slides.add_slide(layout)

        prs.slides.delete(prs.slides[1])

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        slide_parts = _parts_of_type(prs2, SlidePart)
        assert len(slide_parts) == 2
        # -- and nothing in the zip points at a dangling slide part --
        zip_slide_entries = [n for n in _zip_entries(buf) if n.startswith("ppt/slides/slide")]
        assert len(zip_slide_entries) == 2

    def it_drops_the_deleted_slides_notes_slide_part(self):
        # -- A notes slide is owned by exactly one slide (it carries a
        # -- back-reference to its owner). Deleting the owning slide must
        # -- also drop the notes part.
        prs = Presentation()
        layout = prs.slide_layouts[6]
        s_with_notes = prs.slides.add_slide(layout)
        s_with_notes.notes_slide.notes_text_frame.text = "orphan these notes"
        s_keep = prs.slides.add_slide(layout)
        s_keep.notes_slide.notes_text_frame.text = "keep these notes"

        prs.slides.delete(s_with_notes)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        notes_parts = _parts_of_type(prs2, NotesSlidePart)
        assert len(notes_parts) == 1
        # -- and the surviving notes' text proves it belongs to the kept slide --
        assert "keep these notes" in notes_parts[0].blob.decode("utf-8")
        # -- belt-and-braces: exactly one notesSlideN.xml (no orphan) --
        notes_xml_entries = [
            n
            for n in _zip_entries(buf)
            if n.startswith("ppt/notesSlides/notesSlide") and n.endswith(".xml")
        ]
        assert len(notes_xml_entries) == 1

    # -- chart + embedded xlsx cleanup ----------------------------------

    def it_drops_the_deleted_slides_chart_and_embedded_xlsx(self):
        prs = Presentation()
        layout = prs.slide_layouts[6]
        s_with_chart = prs.slides.add_slide(layout)
        chart_data = CategoryChartData()
        chart_data.categories = ("a", "b", "c")
        chart_data.add_series("series", (1, 2, 3))
        s_with_chart.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(4),
            chart_data,
        )
        prs.slides.add_slide(layout)

        prs.slides.delete(s_with_chart)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        prs2 = Presentation(buf)

        # -- chart + its embedded xlsx descendant both gone --
        assert _parts_of_type(prs2, ChartPart) == []
        assert _parts_of_type(prs2, EmbeddedXlsxPart) == []
        # -- and nothing in the zip either --
        entries = _zip_entries(buf)
        assert not any("ppt/charts/chart" in n for n in entries)
        assert not any("ppt/embeddings/" in n for n in entries)

    # -- whole-package shrink -------------------------------------------

    def it_shrinks_the_saved_package_after_a_delete(self):
        # -- The reporter's bottom-line: the zip gets smaller. Build a
        # -- deck with images, delete one slide, save twice (before and
        # -- after) and compare byte-lengths.
        prs_before = Presentation()
        layout = prs_before.slide_layouts[6]
        for img in ("python-powered.png", "monty-truth.png", "python-icon.jpeg"):
            slide = prs_before.slides.add_slide(layout)
            slide.shapes.add_picture("tests/test_files/%s" % img, Inches(1), Inches(1))
        before_buf = io.BytesIO()
        prs_before.save(before_buf)
        before_size = len(before_buf.getvalue())

        # -- now build the same deck, delete a slide, save --
        prs_after = Presentation()
        layout = prs_after.slide_layouts[6]
        for img in ("python-powered.png", "monty-truth.png", "python-icon.jpeg"):
            slide = prs_after.slides.add_slide(layout)
            slide.shapes.add_picture("tests/test_files/%s" % img, Inches(1), Inches(1))
        prs_after.slides.delete(prs_after.slides[0])
        after_buf = io.BytesIO()
        prs_after.save(after_buf)
        after_size = len(after_buf.getvalue())

        # -- the kept deck should be meaningfully smaller (the deleted
        # -- slide carried its own image part). Use an inequality so the
        # -- test is robust to small serialization-overhead changes. --
        assert after_size < before_size

    # -- round-trip stability (the "heavily edited" scenario) -----------

    def it_gc_is_idempotent_across_repeated_add_delete_cycles(self):
        # -- Pin the reporter's worst-case: a deck that goes through
        # -- many add-then-delete cycles should not accumulate orphan
        # -- parts at save time. After 10 add/delete cycles of a slide
        # -- with a unique image, the saved zip should contain exactly
        # -- the parts a one-slide deck needs — not 10 stale images.
        prs = Presentation()
        layout = prs.slide_layouts[6]
        kept = prs.slides.add_slide(layout)
        kept.shapes.add_picture("tests/test_files/python-powered.png", Inches(1), Inches(1))
        for _ in range(10):
            transient = prs.slides.add_slide(layout)
            transient.shapes.add_picture("tests/test_files/monty-truth.png", Inches(1), Inches(1))
            prs.slides.delete(transient)

        buf = io.BytesIO()
        prs.save(buf)

        # -- zip contains exactly one slide and (slide-image-wise) exactly
        # -- one ppt/media/ image — the kept slide's. The ten transient
        # -- monty-truth.png parts have all been GC'd. --
        entries = _zip_entries(buf)
        slide_entries = [n for n in entries if n.startswith("ppt/slides/slide")]
        assert len([n for n in slide_entries if n.endswith(".xml")]) == 1
        media_entries = [n for n in entries if n.startswith("ppt/media/")]
        assert len(media_entries) == 1
