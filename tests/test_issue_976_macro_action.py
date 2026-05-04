"""Regression tests for issue #976 — `ActionSetting.macro` and .pptm save.

Issue #976 asked for a way to attach a ``ppaction://macro`` action-verb
click-handler to a shape so that clicking the shape during a slide show
invokes a VBA macro in a ``.pptm`` package.

These tests pin the XML-level behavior of the new
:attr:`pptx.action.ActionSetting.macro` read/write property and the
automatic content-type swap that :meth:`pptx.presentation.Presentation.save`
performs when the caller supplies a ``.pptm`` / ``.ppsm`` path.
"""

from __future__ import annotations

import io
import zipfile

from pptx import Presentation
from pptx.enum.action import PP_ACTION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches


class DescribeIssue976MacroAction(object):
    """Regression suite for issue #976 — macro click action."""

    def it_writes_a_macro_action_attribute(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )

        shape.click_action.macro = "Module1.MySub"

        hlinkClick = shape._element.nvSpPr.cNvPr.hlinkClick
        assert hlinkClick is not None
        assert hlinkClick.action == "ppaction://macro?name=Module1.MySub"
        # -- no relationship needed; macro action is wholly in the action URL
        assert not hlinkClick.rId
        # -- round-trips through the ActionSetting API
        assert shape.click_action.action is PP_ACTION.RUN_MACRO
        assert shape.click_action.macro == "Module1.MySub"

    def it_reads_the_macro_name_from_an_existing_pptm(self):
        # -- exercise the canonical sample shipped with the behave tests:
        # -- slide 3, shape index 13 carries ppaction://macro?name=Dummy_Macro
        prs = Presentation("/home/ben/code/python-pptx/features/steps/test_files/act-props.pptm")
        shape = prs.slides[2].shapes[13]

        assert shape.click_action.action is PP_ACTION.RUN_MACRO
        assert shape.click_action.macro == "Dummy_Macro"

    def it_replaces_an_existing_macro_action(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )
        shape.click_action.macro = "Old.Sub"

        shape.click_action.macro = "New.Sub"

        assert shape.click_action.macro == "New.Sub"
        hlinkClick = shape._element.nvSpPr.cNvPr.hlinkClick
        assert hlinkClick.action == "ppaction://macro?name=New.Sub"

    def it_replaces_a_prior_URL_hyperlink_with_a_macro(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )
        shape.click_action.hyperlink.address = "https://example.com/"
        # -- confirm the URL is wired before we overwrite it
        assert shape.click_action.action is PP_ACTION.HYPERLINK

        shape.click_action.macro = "TakeOver"

        assert shape.click_action.action is PP_ACTION.RUN_MACRO
        assert shape.click_action.macro == "TakeOver"
        # -- the external hyperlink rel must be gone (else PowerPoint
        # -- would still show the URL)
        hlinkClick = shape._element.nvSpPr.cNvPr.hlinkClick
        assert not hlinkClick.rId

    def it_removes_the_macro_via_None_assignment(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )
        shape.click_action.macro = "Doomed"

        shape.click_action.macro = None

        assert shape.click_action.action is PP_ACTION.NONE
        assert shape.click_action.macro is None
        assert shape._element.nvSpPr.cNvPr.hlinkClick is None

    def it_works_on_hover_action_too(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )

        shape.hover_action.macro = "OnHover"

        assert shape.hover_action.macro == "OnHover"
        assert shape.hover_action.action is PP_ACTION.RUN_MACRO
        # -- click_action remains empty; hover lives on a different element
        assert shape.click_action.action is PP_ACTION.NONE

    def it_round_trips_a_macro_action_through_save_and_reopen(self, tmp_path):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )
        shape.click_action.macro = "Module1.Do"
        out = tmp_path / "out.pptm"
        prs.save(str(out))

        reopened = Presentation(str(out))
        shape2 = reopened.slides[0].shapes[0]
        assert shape2.click_action.macro == "Module1.Do"
        assert shape2.click_action.action is PP_ACTION.RUN_MACRO

    # -- content-type auto-swap on save -------------------------------

    def it_writes_macro_enabled_content_type_when_saving_to_pptm(self, tmp_path):
        prs = Presentation()
        out = tmp_path / "deck.pptm"
        prs.save(str(out))

        with zipfile.ZipFile(str(out)) as zf:
            content_types = zf.read("[Content_Types].xml").decode("utf-8")

        # -- the presentation-part override carries the macro-enabled type
        assert "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml" in content_types
        # -- and the non-macro variant is not written
        assert (
            "vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
            not in content_types
        )

    def it_writes_slideshow_macro_content_type_when_saving_to_ppsm(self, tmp_path):
        prs = Presentation()
        out = tmp_path / "deck.ppsm"
        prs.save(str(out))

        with zipfile.ZipFile(str(out)) as zf:
            content_types = zf.read("[Content_Types].xml").decode("utf-8")

        assert "application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml" in content_types

    def it_keeps_regular_content_type_when_saving_to_pptx(self, tmp_path):
        prs = Presentation()
        out = tmp_path / "deck.pptx"
        prs.save(str(out))

        with zipfile.ZipFile(str(out)) as zf:
            content_types = zf.read("[Content_Types].xml").decode("utf-8")

        assert (
            "vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
            in content_types
        )
        assert "macroEnabled" not in content_types

    def it_does_not_swap_when_writing_to_a_stream(self):
        prs = Presentation()
        stream = io.BytesIO()
        prs.save(stream)

        stream.seek(0)
        with zipfile.ZipFile(stream) as zf:
            content_types = zf.read("[Content_Types].xml").decode("utf-8")

        # -- stream saves keep whatever content type the part currently
        # -- has; for a fresh Presentation() that's the regular variant
        assert "macroEnabled" not in content_types

    def it_opens_a_macro_enabled_pptm_successfully(self):
        # -- opening a package whose main content type is PML_PRES_MACRO_MAIN
        # -- must not raise, even though the file extension happens to be .pptm
        prs = Presentation("/home/ben/code/python-pptx/features/steps/test_files/act-props.pptm")
        assert len(prs.slides) >= 1
