"""Integration tests for password-protected .pptx read/write.

These exercise the real msoffcrypto-tool integration end-to-end without mocks, so they
complement the more granular tests in :mod:`tests.opc.test__crypto` and
:mod:`tests.opc.test_serialized`.
"""

from __future__ import annotations

import importlib.util
import io
import os
import zipfile

import pytest

from pptx import Presentation
from pptx.exc import EncryptedPackageError
from pptx.opc._crypto import decrypt_stream

from .unitutil.file import absjoin, test_file_dir

_MINIMAL_PPTX = absjoin(test_file_dir, "minimal.pptx")

# -- Issue #327: skip the entire module when msoffcrypto-tool is absent;
#    every test here requires the optional dependency. --
pytestmark = pytest.mark.skipif(
    importlib.util.find_spec("msoffcrypto") is None,
    reason="msoffcrypto-tool is not installed (optional test dependency for issue #327)",
)


class DescribePasswordRoundTrip:
    """Save-with-password then open-with-password returns the same content."""

    def it_round_trips_a_password_protected_pptx_through_a_stream(self):
        prs = Presentation(_MINIMAL_PPTX)

        encrypted = io.BytesIO()
        prs.save(encrypted, password="s3cret")

        encrypted.seek(0)
        prs2 = Presentation(encrypted, password="s3cret")

        # -- the encrypted bytes start with the OLE2 compound-document magic --
        encrypted.seek(0)
        assert encrypted.read(8) == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
        # -- and the re-opened presentation behaves like the original --
        assert len(prs2.slides) == len(prs.slides)

    def it_round_trips_a_password_protected_pptx_through_a_file_path(self, tmp_path: object):
        prs = Presentation(_MINIMAL_PPTX)
        out_path = os.path.join(str(tmp_path), "encrypted.pptx")

        prs.save(out_path, password="s3cret")
        prs2 = Presentation(out_path, password="s3cret")

        assert os.path.getsize(out_path) > 0
        assert len(prs2.slides) == len(prs.slides)

    def it_combines_zip_date_time_and_password_on_save(self):
        """Passing both kwargs together produces an encrypted .pptx whose inner
        (plaintext) zip carries the fixed zip_date_time stamp on every member.

        Exercises the orthogonal reconciliation of #702 (zip_date_time) and #668
        (password): the plaintext zip is first built with fixed member timestamps,
        and then wrapped in the ECMA-376 Agile Encryption CFBF container.
        """
        prs = Presentation(_MINIMAL_PPTX)

        encrypted = io.BytesIO()
        fixed_dt = (2020, 1, 1, 0, 0, 0)
        prs.save(encrypted, zip_date_time=fixed_dt, password="s3cret")

        # -- the outer bytes are the OLE2 / CFBF envelope --
        outer = encrypted.getvalue()
        assert outer[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

        # -- decrypt and inspect the inner zip; every member must carry the fixed stamp --
        plain = decrypt_stream(io.BytesIO(outer), "s3cret")
        with zipfile.ZipFile(io.BytesIO(plain), "r") as z:
            dates = {info.filename: info.date_time for info in z.infolist()}
        assert dates, "decrypted zip should contain at least one member"
        for name, d in dates.items():
            assert d == fixed_dt, f"member {name!r} has date_time {d!r}"

        # -- and the decrypted package opens round-trip --
        prs2 = Presentation(io.BytesIO(outer), password="s3cret")
        assert len(prs2.slides) == len(prs.slides)


class DescribeOpenPasswordProtectedPptx:
    """Error handling when opening an encrypted file."""

    def it_raises_when_no_password_is_supplied(self, encrypted_pptx_stream: io.BytesIO):
        with pytest.raises(EncryptedPackageError, match="password-protected"):
            Presentation(encrypted_pptx_stream)

    def it_raises_on_a_wrong_password(self, encrypted_pptx_stream: io.BytesIO):
        with pytest.raises(EncryptedPackageError, match="does not match"):
            Presentation(encrypted_pptx_stream, password="WRONG")

    # -- fixtures ------------------------------------------------------

    @pytest.fixture
    def encrypted_pptx_stream(self) -> io.BytesIO:
        prs = Presentation(_MINIMAL_PPTX)
        out = io.BytesIO()
        prs.save(out, password="right")
        out.seek(0)
        return out
