# pyright: reportPrivateUsage=false

"""Unit-test suite for the slide-level ``Slide.tags`` accessor (issue #578).

Covers both the :class:`pptx.slide.SlideTags` dict-like proxy and the end-to-end
round-trip through save/reopen that exercises the |TagsPart| wiring.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from pptx import Presentation
from pptx.slide import SlideTags


@pytest.fixture
def prs_with_slide():
    prs = Presentation()
    slide_layout = prs.slide_layouts[5]
    prs.slides.add_slide(slide_layout)
    return prs


class DescribeSlideTags:
    """Unit tests for the dict-like `SlideTags` proxy."""

    def it_is_empty_for_a_slide_with_no_tags_part(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        tags = slide.tags

        assert isinstance(tags, SlideTags)
        assert len(tags) == 0
        assert list(tags) == []
        assert "missing" not in tags
        assert tags.get("missing") is None
        assert tags.get("missing", "default") == "default"
        # -- read-only iteration should NOT create the tags part --
        assert slide.has_tags is False

    def it_raises_KeyError_on_indexed_access_to_missing_tag(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        with pytest.raises(KeyError, match="absent"):
            _ = slide.tags["absent"]

    def it_creates_the_tags_part_on_first_write(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        slide.tags["priority"] = "high"

        assert slide.has_tags is True
        assert slide.tags["priority"] == "high"
        assert len(slide.tags) == 1

    def it_stores_multiple_distinct_tags_in_document_order(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        slide.tags["a"] = "1"
        slide.tags["b"] = "2"
        slide.tags["c"] = "3"

        assert list(slide.tags) == ["a", "b", "c"]
        assert slide.tags.keys() == ["a", "b", "c"]
        assert slide.tags.values() == ["1", "2", "3"]
        assert slide.tags.items() == [("a", "1"), ("b", "2"), ("c", "3")]

    def it_overwrites_an_existing_tag_without_adding_a_duplicate(self, prs_with_slide):
        slide = prs_with_slide.slides[0]
        slide.tags["k"] = "v1"

        slide.tags["k"] = "v2"

        assert len(slide.tags) == 1
        assert slide.tags["k"] == "v2"

    def it_removes_a_tag_by_name(self, prs_with_slide):
        slide = prs_with_slide.slides[0]
        slide.tags["k"] = "v"
        slide.tags["j"] = "w"

        del slide.tags["k"]

        assert "k" not in slide.tags
        assert len(slide.tags) == 1

    def but_raises_KeyError_when_deleting_a_missing_tag(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        with pytest.raises(KeyError):
            del slide.tags["nope"]

    def and_raises_KeyError_when_deleting_before_any_tags_exist(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        # -- no tags part yet; delete must still raise without side effect --
        with pytest.raises(KeyError):
            del slide.tags["nope"]
        assert slide.has_tags is False

    def it_rejects_non_string_names(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        with pytest.raises(TypeError):
            slide.tags[42] = "x"  # type: ignore[index]

    def it_rejects_non_string_values(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        with pytest.raises(TypeError):
            slide.tags["k"] = 42  # type: ignore[assignment]

    def it_returns_False_for_non_string_contains_query(self, prs_with_slide):
        slide = prs_with_slide.slides[0]
        slide.tags["k"] = "v"

        assert (42 in slide.tags) is False

    def it_caches_the_SlideTags_proxy(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        t1 = slide.tags
        t2 = slide.tags

        assert t1 is t2


class DescribeSlide_tags_integration:
    def it_writes_the_expected_xml_payloads_when_saved(self, prs_with_slide):
        slide = prs_with_slide.slides[0]
        slide.tags["priority"] = "high"
        slide.tags["owner"] = "alice"

        buf = io.BytesIO()
        prs_with_slide.save(buf)
        buf.seek(0)

        with zipfile.ZipFile(buf, "r") as z:
            tag_xml = z.read("ppt/tags/tag1.xml").decode()
            slide_xml = z.read("ppt/slides/slide1.xml").decode()
            slide_rels = z.read("ppt/slides/_rels/slide1.xml.rels").decode()
            content_types = z.read("[Content_Types].xml").decode()

        assert '<p:tag name="priority" val="high"/>' in tag_xml
        assert '<p:tag name="owner" val="alice"/>' in tag_xml
        # -- slide XML carries the custDataLst/tags reference --
        assert "custDataLst" in slide_xml
        assert "<p:tags " in slide_xml
        # -- rel to the tags part is present --
        assert "relationships/tags" in slide_rels
        assert "tag1.xml" in slide_rels
        # -- [Content_Types].xml registers the tags part type --
        assert "tag1.xml" in content_types
        assert "presentationml.tags+xml" in content_types

    def it_round_trips_through_save_and_reopen(self, prs_with_slide):
        slide = prs_with_slide.slides[0]
        slide.tags["priority"] = "high"
        slide.tags["owner"] = "alice"

        buf = io.BytesIO()
        prs_with_slide.save(buf)
        buf.seek(0)

        prs2 = Presentation(buf)
        slide2 = prs2.slides[0]

        assert slide2.has_tags is True
        assert dict(slide2.tags.items()) == {"priority": "high", "owner": "alice"}
        # -- mutating round-tripped tags continues to work --
        slide2.tags["owner"] = "bob"
        del slide2.tags["priority"]
        assert slide2.tags.items() == [("owner", "bob")]

    def it_does_not_create_a_tags_part_on_read_only_access(self, prs_with_slide):
        slide = prs_with_slide.slides[0]

        # -- side-effect-free reads --
        _ = len(slide.tags)
        _ = list(slide.tags)
        _ = "any" in slide.tags
        _ = slide.tags.get("any")

        buf = io.BytesIO()
        prs_with_slide.save(buf)
        buf.seek(0)

        with zipfile.ZipFile(buf, "r") as z:
            assert all("tag" not in n.lower() for n in z.namelist()), (
                "no tags part should have been materialized on read-only access"
            )
